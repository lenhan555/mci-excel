"""Dựng deck Buổi 5: Slide_B5.pptx từ Slide_Script_B5.md, theo style của Slide_B1.pptx (cùng khung với build_slides_b4.py).

- Master, layout, theme và slide bìa lấy từ Slide_B1.pptx (không sửa file B1; chỉ đọc làm khuôn).
- Chữ trên slide khai báo trong script này, khớp phần "Trên slide" của kịch bản.
- Speaker notes lấy NGUYÊN VĂN từ kịch bản (Lời giảng, Ghi chú giảng viên), bỏ định dạng markdown,
  nên sửa lời giảng thì sửa Slide_Script_B5.md rồi chạy lại script này.
- Quiz: shape đáp án tên "Answer …", lớp mờ "Dim …", hiện khi bấm (build_student_slides.py sẽ xóa ở bản _HocVien).
- Ảnh chụp thật lấy từ Materials/B5/Images (tên S<số slide>_…png, danh sách: _DANH_SACH_ANH.md); chưa có thì để khung [CHỤP].
- Số liệu khớp build_b5_meta.json (build_b5_meta.py).

    python3 _generator/build_slides_b5.py     # cần python-pptx
    python3 _generator/build_student_slides.py B5
"""
import copy, os, re, sys
for s_ in (sys.stdout, sys.stderr):
    if hasattr(s_, "reconfigure"): s_.reconfigure(encoding="utf-8", errors="replace")
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

ROOT = os.environ.get("MCI_ROOT") or os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
TEMPLATE = os.path.join(ROOT, "1_Sessions", "B1_Nen_tang_Excel_Du_lieu", "Slide_B1.pptx")
B5 = os.path.join(ROOT, "1_Sessions", "B5_Data_Model_DAX")
SCRIPT = os.path.join(B5, "Slide_Script_B5.md")
OUT = os.path.join(B5, "Slide_B5.pptx")
IMG = os.path.join(ROOT, "Materials", "B5", "Images")        # ảnh chụp thật: S<số slide>_<chủ đề>.png
MISSING = []

NAVY, INK, INK2, MUTED = "203864", "1B2230", "585858", "6B6F76"
RED, GOLD, ALT, HAIR, WHITE = "C00000", "FFC000", "F3F5F9", "D9DEE8", "FFFFFF"
GOOD, GOOD_BG, BAD_TXT, BAD_BG, ACCENT = "2E7D32", "E8F5E9", "9C0006", "FFC7CE", "3398DA"
L, R, W = 1.67, 12.74, 11.07          # lề trái, lề phải, bề rộng vùng nội dung (in)

# ---------------------------------------------------------------------
# Khuôn: mở B1, giữ slide bìa, xóa các slide còn lại
# ---------------------------------------------------------------------
prs = Presentation(TEMPLATE)
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst)[1:]:
    prs.part.drop_rel(sldId.rId); sldIdLst.remove(sldId)
LAYOUT = {l.name: l for l in prs.slide_layouts}

def rgb(h): return RGBColor.from_string(h)

def _runs(p, parts, size, color, bold):
    for part in parts:
        if isinstance(part, str): part = (part, {})
        text, f = part
        r = p.add_run(); r.text = text
        r.font.size = Pt(f.get("size", size)); r.font.bold = f.get("b", bold)
        r.font.color.rgb = rgb(f.get("color", color)); r.font.name = "Arial"
        if f.get("i"): r.font.italic = True

def tb(s, x, y, w, h, paras, size=18, color=INK, bold=False, anchor="t", align=None, name="TextBox", bullet=False, space=6):
    """paras: str | list[para]; para: str | list[run]; run: str | (text, {b, color, size, i})."""
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); sh.name = name
    tf = sh.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "ctr": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str): paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align: p.alignment = {"l": PP_ALIGN.LEFT, "ctr": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        if i: p.space_before = Pt(space)
        if isinstance(para, str): para = [para]
        if bullet: para = [("•  ", {"color": NAVY, "b": True})] + list(para)
        _runs(p, para, size, color, bold)
    return sh

def box(s, x, y, w, h, fill=None, line=None, lw=1, shape=MSO_SHAPE.RECTANGLE, radius=None, name="Box", alpha=None):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h)); sh.name = name
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0] = radius
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    else: sh.fill.background()
    if line: sh.line.color.rgb = rgb(line); sh.line.width = Pt(lw)
    else: sh.line.fill.background()
    sh.shadow.inherit = False
    if alpha is not None:
        clr = sh.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        a = etree.SubElement(clr, qn("a:alpha")); a.set("val", str(alpha))
    return sh

def chip(s, x, y, w, text, fill=None, line=NAVY, color=NAVY, size=14, name="Chip", h=0.4):
    sh = box(s, x, y, w, h, fill=fill, line=line, lw=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2, name=name)
    tf = sh.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; _runs(p, [text], size, color, True)
    return sh

def vline(s, x, y, h):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x), Inches(y + h))
    ln.line.color.rgb = rgb(HAIR); ln.line.width = Pt(1); return ln

def content(title, rail):
    s = prs.slides.add_slide(LAYOUT["Title Only"])
    s.shapes.title.text = title
    tb(s, 0, 1.5, 1.07, 5.25, rail, size=14, color=WHITE, bold=True, anchor="ctr", align="ctr", name="Rail label")
    s.shapes[-1].text_frame._bodyPr.set("vert", "vert270")
    return s

def section(title, body=None):
    s = prs.slides.add_slide(LAYOUT["Section Header"])
    s.shapes.title.text = title
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx == 1:
            if body: ph.text = body
            else: ph._element.getparent().remove(ph._element)
    return s

def tip(s, x, y, w, h, text, pro=False):
    box(s, x, y, w, h, fill=ALT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, name=("PRO TIP" if pro else "TIP") + " callout")
    if pro:
        chip(s, x + 0.12, y + 0.12, 1.1, "PRO TIP", fill=GOLD, line=None, color=INK, size=12, name="PRO TIP chip", h=0.35)
        tb(s, x + 1.32, y + 0.1, w - 1.44, h - 0.2, text, size=16, name="PRO TIP text")
    else:
        tb(s, x + 0.12, y + 0.12, 0.6, 0.35, "TIP", size=14, color=NAVY, bold=True, anchor="ctr")
        tb(s, x + 0.75, y + 0.1, w - 0.87, h - 0.2, text, size=16, name="TIP text")

def tool(s, x, y, w, name, purpose, h_purpose=0.7, size=24):
    tb(s, x, y, w, 0.5, name, size=size, color=RED, bold=True, name="Tool name")
    tb(s, x, y + 0.5, w, h_purpose, purpose, size=16, color=INK2, name="Tool purpose")

def table(s, x, y, colw, rows, size=16, head_size=None, rowh=0.45, bold_first=False, name="Table",
          header=True, fills=None, colors=None, align=None):
    """rows[0] là tiêu đề (nền navy). fills/colors: {(r,c): hex}. Chữ trong ô có thể là list run."""
    nr, nc = len(rows), len(rows[0])
    gf = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(sum(colw)), Inches(rowh * nr)); gf.name = name
    tbl = gf.table
    tbl._tbl.tblPr.set("firstRow", "1"); tbl._tbl.tblPr.set("bandRow", "0")
    style = tbl._tbl.tblPr.find(qn("a:tableStyleId"))
    if style is not None: tbl._tbl.tblPr.remove(style)
    for j, cw in enumerate(colw): tbl.columns[j].width = Inches(cw)
    for i in range(nr):
        tbl.rows[i].height = Inches(rowh)
        for j in range(nc):
            c = tbl.cell(i, j); v = rows[i][j]
            c.margin_left = c.margin_right = Inches(0.08); c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]
            if align and align[j]: p.alignment = {"l": PP_ALIGN.LEFT, "ctr": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align[j]]
            head = header and i == 0
            col = WHITE if head else (colors or {}).get((i, j), INK)
            parts = v if isinstance(v, list) else [str(v)]
            _runs(p, parts, head_size or size, col, head or (bold_first and j == 0))
            tcPr = c._tc.get_or_add_tcPr()
            for side in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                ln = etree.SubElement(tcPr, qn(side))
                if side in ("a:lnL", "a:lnR"):
                    ln.set("w", "0"); etree.SubElement(ln, qn("a:noFill"))
                else:
                    ln.set("w", "12700"); sf = etree.SubElement(ln, qn("a:solidFill"))
                    etree.SubElement(sf, qn("a:srgbClr")).set("val", HAIR)
            f = NAVY if head else (fills or {}).get((i, j))
            if f:
                sf = etree.SubElement(tcPr, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr")).set("val", f)
            else:
                etree.SubElement(tcPr, qn("a:noFill"))
    return gf

def fx(s, x, y, w, formula, size=16, name="Formula bar"):
    """Thanh công thức mô phỏng Excel: ô 'fx' + công thức."""
    box(s, x, y, w, 0.45, fill=WHITE, line=HAIR, name=name)
    tb(s, x + 0.1, y, 0.45, 0.45, [[("fx", {"i": True})]], size=14, color=MUTED, anchor="ctr")
    tb(s, x + 0.6, y, w - 0.7, 0.45, formula, size=size, color=INK, anchor="ctr", name="Formula")

def stat(s, x, y, w, h, label, value, sub=None, name="Stat card", vcolor=NAVY, vsize=32):
    box(s, x, y, w, h, fill=WHITE, line=HAIR, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, name=name)
    paras = [[(label, {"size": 14, "color": INK2})], [(value, {"size": vsize, "b": True, "color": vcolor})]]
    if sub: paras.append([(sub, {"size": 14, "color": INK})])
    tb(s, x + 0.2, y + 0.1, w - 0.4, h - 0.2, paras, anchor="ctr", space=2)

def shot(s, x, y, w, h, what, where, img, name="Image placeholder [CHỤP]"):
    """Ảnh chụp thật (Excel 365 Windows, zoom 100%), xem CLAUDE.md, mục Ảnh minh họa.
    Có file Materials/B5/Images/<img> thì chèn ảnh vào khung (giữ tỉ lệ, canh giữa, viền hairline);
    chưa có thì để khung chờ [CHỤP] ghi nội dung, file · sheet · vùng và tên file ảnh cần lưu."""
    path = os.path.join(IMG, img)
    if os.path.exists(path):
        pic = s.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))
        if pic.height > Inches(h):
            ratio = Inches(h) / pic.height; pic.height = Inches(h); pic.width = int(pic.width * ratio)
        pic.left = Inches(x) + (Inches(w) - pic.width) // 2; pic.top = Inches(y) + (Inches(h) - pic.height) // 2
        pic.line.color.rgb = rgb(HAIR); pic.line.width = Pt(0.75); pic.name = f"Screenshot {img[:-4]}"
        return pic
    sh = box(s, x, y, w, h, fill=ALT, line="A6A6A6", lw=1.25, name=name)
    sh.line.dash_style = MSO_LINE.DASH
    tb(s, x + 0.2, y + 0.1, w - 0.4, h - 0.2, [[("[CHỤP] ", {"b": True, "color": NAVY}), (what, {"color": INK2})],
                                               [(f"{where} → {img}", {"color": MUTED, "size": 12})]], size=14, anchor="ctr", align="ctr", space=4)
    MISSING.append(img)
    return sh

def footnote(s, text):
    tb(s, L, 6.8, W, 0.35, text, size=12, color=MUTED, anchor="b", name="Footnote")

# ---------------------------------------------------------------------
# Speaker notes từ kịch bản
# ---------------------------------------------------------------------
md = open(SCRIPT, encoding="utf-8").read()
blocks = re.split(r"^### Slide (\d+) · ", md, flags=re.M)
NOTES = {}

def clean(t):
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t); t = re.sub(r"(?<![\w*])\*(?=\S)(.+?)(?<=\S)\*(?![\w*])", r"\1", t)
    return t.replace("`", "").strip()

def fmt_min(m):
    return (f"{m:g}").replace(".", ",")

for k in range(1, len(blocks), 2):
    n = int(blocks[k]); body = blocks[k + 1]
    head, rest = body.split("\n", 1)
    body = re.split(r"^(?:---|## )", rest, flags=re.M)[0]
    m = re.search(r"· ([\d.]+)'\s*$", head) or re.search(r"\((\d+) phút\)", head)
    out = []
    ans = re.search(r"\*\*([A-D])\)", body)
    if head.startswith("QUESTION"):
        out.append(f"Khoảng 35 giây. Đáp án: {ans.group(1)}.")
    elif m:
        out.append(f"Thời lượng: {fmt_min(float(m.group(1)))} phút.")
    for sec in re.finditer(r"^\*\*(Lời giảng[^*]*|Ghi chú giảng viên[^*]*):\*\*\n(.*?)(?=^\*\*(?:Lời giảng|Ghi chú giảng viên)|\Z)", body, flags=re.M | re.S):
        label, txt = sec.group(1), sec.group(2)
        head_, _, tail_ = label.partition(" (")
        out.append(head_.upper() + (f" ({tail_}" if tail_ else ""))
        for line in txt.splitlines():
            line = line.rstrip()
            if line.startswith(">"): line = line[1:].strip()
            elif line.lstrip().startswith("- "): line = line.lstrip()[2:]
            else: line = line.strip()
            if line: out.append(clean(line))
    NOTES[n] = "\n".join(out)
assert len(NOTES) == 39, sorted(NOTES)

def notes(s, n):
    s.notes_slide.notes_text_frame.text = NOTES[n]

# ---------------------------------------------------------------------
# Quiz: 4 phương án, đáp án hiện khi bấm (Q1 = slide 34 → nhãn "Question 1")
# ---------------------------------------------------------------------
def quiz(n, topic, question, options, answer, rail="PHẦN 9 · QUIZ & TỔNG KẾT", qsize=22, osize=18):
    s = content(f"Question {n - 33}: {topic}", rail)
    tb(s, L, 1.65, W, 1.0, question, size=qsize, color=INK)
    pos = {"A": (L, 2.85), "B": (7.33, 2.85), "C": (L, 4.35), "D": (7.33, 4.35)}
    ow, oh = 5.41, 1.3
    for letter, (x, y) in pos.items():
        box(s, x, y, ow, oh, fill=WHITE, line=HAIR, name=f"Option {letter}")
        if letter == answer:
            box(s, x, y, ow, oh, fill=GOOD_BG, line=GOOD, lw=3, name=f"Answer highlight {letter}")
        lt = box(s, x + 0.2, y + (oh - 0.6) / 2, 0.6, 0.6, fill=NAVY, shape=MSO_SHAPE.OVAL, name=f"Letter {letter}")
        _center_text(lt, letter)
        if letter == answer:
            al = box(s, x + 0.2, y + (oh - 0.6) / 2, 0.6, 0.6, fill=GOOD, shape=MSO_SHAPE.OVAL, name=f"Answer letter {letter}")
            _center_text(al, letter)
        tb(s, x + 1.0, y, ow - 1.1, oh, options[letter], size=osize, anchor="ctr")
    for letter, (x, y) in pos.items():
        if letter != answer: box(s, x, y, ow, oh, fill=WHITE, name=f"Dim {letter}", alpha=60000)
    ax, ay = pos[answer]
    chip(s, ax + ow - 1.6, ay - 0.2, 1.45, "✓ Đáp án", fill=GOOD, line=None, color=WHITE, size=14, name="Answer chip", h=0.32)
    _reveal(s, [sh for sh in s.shapes if sh.name.startswith(("Answer", "Dim"))])
    notes(s, n)

def _center_text(sh, text):
    tf = sh.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; _runs(p, [text], 20, WHITE, True)

def _reveal(s, shapes):
    """Một lần bấm: mọi shape đáp án/làm mờ hiện dần (fade), giống deck B1."""
    ids = [sh.shape_id for sh in shapes]
    cid = [3]
    def nid(): cid[0] += 1; return cid[0]
    pars = []
    for k, spid in enumerate(ids):
        a, b = nid(), nid(); c = nid()
        node = "clickEffect" if k == 0 else "withEffect"
        pars.append(f'<p:par><p:cTn id="{a}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="{node}">'
                    f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:set><p:cBhvr><p:cTn id="{b}" dur="1" fill="hold">'
                    f'<p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
                    f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
                    f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{c}" dur="400"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
                    f'</p:childTnLst></p:cTn></p:par>')
    bld = "".join(f'<p:bldP spid="{i}" grpId="0" animBg="1"/>' for i in ids)
    xml = ('<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par>'
           '<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek">'
           '<p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst><p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>'
           f'<p:childTnLst><p:par><p:cTn id="{nid() + 900}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{"".join(pars)}</p:childTnLst></p:cTn></p:par>'
           '</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>'
           f'<p:bldLst>{bld}</p:bldLst></p:timing>')
    s._element.append(etree.fromstring(xml))

# =====================================================================
# SLIDES
# =====================================================================
BLUE, ORANGE = "2A78D6", "EB6834"
R0 = "PHẦN 0 · MỞ ĐẦU"; R1 = "PHẦN 1 · DATA MODEL"; R2 = "PHẦN 2 · RELATIONSHIP"
R3 = "PHẦN 3 · COLUMN VS MEASURE"; R4 = "PHẦN 4 · FILTER CONTEXT"; R5 = "PHẦN 5 · CALCULATE & ALL"
R6 = "PHẦN 6 · TIME INTELLIGENCE"; R7 = "PHẦN 7 · KPI DICTIONARY"; R8 = "PHẦN 8 · THỰC HÀNH"; R9 = "PHẦN 9 · QUIZ & TỔNG KẾT"

def card(s, x, y, w, h, title, body, size=15, fill=WHITE, line=HAIR, title_color=NAVY, name="Card"):
    box(s, x, y, w, h, fill=fill, line=line, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05, name=name)
    paras = [[(title, {"b": True, "color": title_color, "size": size + 1})]] + [p if isinstance(p, list) else [p] for p in body]
    tb(s, x + 0.18, y + 0.12, w - 0.36, h - 0.24, paras, size=size, space=5)

def bricks(s, done, nxt):
    box(s, 9.07, 4.6, 3.67, 1.0, fill=NAVY, name=f"Brick {done[:2]}")
    tb(s, 9.17, 4.6, 3.47, 1.0, "✓  " + done, size=18, color=WHITE, bold=True, anchor="ctr")
    box(s, 9.07, 3.3, 3.67, 1.0, fill=WHITE, line=BLUE, lw=2.25, name=f"Brick {nxt[:2]}")
    tb(s, 9.17, 3.3, 3.47, 1.0, nxt, size=18, color=NAVY, bold=True, anchor="ctr")

def arrow(s, x1, y1, x2, y2, color=ACCENT, w=2.25, name="Arrow"):
    a = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    a.line.color.rgb = rgb(color); a.line.width = Pt(w); a.name = name
    ln = a.line._get_or_add_ln(); te = etree.SubElement(ln, qn("a:tailEnd")); te.set("type", "triangle")
    return a

def tbl_box(s, x, y, w, h, title, sub, fill=WHITE, line=NAVY, tcolor=NAVY, name="Model table"):
    """[VẼ] một bảng trong mô hình: tên bảng + số dòng / vai trò."""
    box(s, x, y, w, h, fill=fill, line=line, lw=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08, name=name)
    tb(s, x + 0.1, y + 0.05, w - 0.2, h - 0.1, [[(title, {"b": True, "color": tcolor, "size": 15})], [(sub, {"size": 12, "color": INK2})]],
       anchor="ctr", align="ctr", space=2)

def numbered(s, x, y, w, h, items, size=16, space=12):
    tb(s, x, y, w, h, [[(f"{k}. ", {"b": True, "color": NAVY}), t] if isinstance(t, str) else [(f"{k}. ", {"b": True, "color": NAVY})] + t
                       for k, t in enumerate(items, 1)], size=size, space=space)

# 1 · Bìa
s1 = prs.slides[0]
for sh in s1.shapes:
    if sh.name.startswith("subTitle"):
        sh.text_frame.paragraphs[0].runs[0].text = "Buổi 5: Mô hình dữ liệu & DAX"
        for r in sh.text_frame.paragraphs[0].runs[1:]: r.text = ""
    elif sh.has_text_frame and sh.text_frame.text.startswith("Yêu cầu máy"):
        p = sh.text_frame.paragraphs[0]
        p.runs[0].text = "Bắt buộc máy Windows có Power Pivot (Data → Manage Data Model)"
        for r in p.runs[1:]: r.text = ""
        p2 = copy.deepcopy(p._p); sh.text_frame._txBody.append(p2)
        p2r = sh.text_frame.paragraphs[-1].runs
        p2r[0].text = "Mở sẵn: file khởi đầu B4_MCI_Mart_Answer.xlsx (đáp án Buổi 4)"
        for r in p2r[1:]: r.text = ""
        sh.height = Inches(0.9)
    elif sh.has_text_frame and sh.text_frame.text.startswith("Mở sẵn"):
        p = sh.text_frame.paragraphs[0]
        p.runs[0].text = "Mở sẵn: file khởi đầu B4_MCI_Mart_Answer.xlsx (đáp án Buổi 4)"
        for r in p.runs[1:]: r.text = ""
s1.notes_slide.notes_text_frame.text = NOTES[1]

# 2 · Câu hỏi mở buổi
s = content("Câu hỏi mở buổi", R0)
tb(s, L, 1.75, W, 1.5, [[("\"Doanh thu tháng này đạt bao nhiêu % mục tiêu, và vì sao mỗi báo cáo cũ lại ra một con số khác nhau?\"", {"b": True, "color": NAVY})]], size=26)
chip(s, L, 3.3, 2.9, "Tầng 3: Diagnostic", fill=WHITE, line=NAVY, color=NAVY, size=14, name="Tier chip")
stat(s, L, 4.0, 4.6, 1.9, "Báo cáo A · Gross Margin % Quý 1/2025", "34,18%", "tổng lợi nhuận ÷ tổng doanh thu", name="Stat A", vsize=34)
tb(s, 6.37, 4.0, 1.0, 1.9, [[("≠", {"b": True, "color": BAD_TXT, "size": 48})]], anchor="ctr", align="ctr", name="Not equal")
stat(s, 7.47, 4.0, 4.6, 1.9, "Báo cáo B · Gross Margin % Quý 1/2025", "33,77%", "trung bình biên lợi nhuận từng dòng", name="Stat B", vsize=34, vcolor=BAD_TXT)
tb(s, L, 6.05, W, 0.5, "Cùng một dữ liệu. Không ai bịa số, chỉ là hai công thức khác nhau.", size=16, color=INK2)
notes(s, 2)

# 3 · Recap B4 & file khởi đầu
s = content("Recap Buổi 4 & file khởi đầu", R0)
tb(s, L, 1.75, 7.1, 4.6, [
    [("Viên gạch B4: ", {"b": True, "color": NAVY}), "6 bảng sạch trong Data Model, Sales 30.247 dòng (4.436.849.500), Product 40, Store 12, Customer 1.500, Target 144, Calendar 731. Refresh tự cập nhật."],
    [("Bài về nhà B4: ", {"b": True, "color": NAVY}), "đối soát trước khi tính % target, ST07 tháng 4 thiếu file bổ sung chỉ đạt 54,3%, đủ file đạt 97,4%; cả chuỗi 103,1%."],
    [("File khởi đầu hôm nay ", {"b": True, "color": NAVY}), "= đáp án B4 (B4_MCI_Mart_Answer.xlsx). Lưu thành B5_<HoTen>.xlsx. KHÔNG bấm Refresh All."],
    [("Giảng viên làm mẫu trên 2 file gọn: ", {"b": True, "color": NAVY}), "Demo/B5_Demo_Model.xlsx (Phần 1–2, chưa nối) · Demo/B5_Demo_DAX.xlsx (Phần 3–7, đã nối); mỗi file có bản _HocVien."],
    [("Cuối buổi: ", {"b": True, "color": NAVY}), "đối soát với KPI_SUMIFS (Thực hành 2), đáp án ở sheet DoiSoat_B2 của Practice/B5_MCI_Mart_Answer.xlsx; file này là file khởi đầu Buổi 6."],
], size=16, bullet=True, space=8)
bricks(s, "B4 · Pipeline Power Query", "B5 · Star schema + DAX")
footnote(s, "Demo: ../B4_Power_Query/Practice/B4_MCI_Mart_Answer.xlsx (sheet Kiem_tra_B4) · ../B4_Power_Query/Homework/B4_BTVN_MCI_Mart_T4_2025_Answer.xlsx (sheet KetQua_T4)")
notes(s, 3)

# 4 · Chia phần 1
notes(section("1. Data Model: nhiều bảng, một mô hình"), 4)

# 5 · Power Pivot window
s = content("Data View · Diagram View", R1)
table(s, L, 1.55, [4.6, 6.47], [
    ["Tên", "Mục đích"],
    [[("Data → Manage Data Model", {"b": True, "color": RED})], "Mở cửa sổ Power Pivot, nơi chứa Data Model, tách riêng khỏi worksheet (hoặc Power Pivot → Manage)"],
    [[("Data View", {"b": True, "color": RED})], "Mỗi bảng một tab; xem dữ liệu, thêm calculated column; dòng Record ghi số dòng"],
    [[("Diagram View", {"b": True, "color": RED})], "Mỗi bảng là một hộp; relationship là đường nối giữa các hộp"],
], size=14, rowh=0.5)
shot(s, L, 3.65, 5.43, 2.05, "Data View, tab Sales: 6 tab bảng ở dưới, Record 1 of 30,247", "Demo/B5_Demo_Model.xlsx · Data → Manage Data Model", "S05a_data_view.png")
shot(s, 7.31, 3.65, 5.43, 2.05, "Diagram View: 6 hộp bảng, CHƯA có đường nối nào", "Demo/B5_Demo_Model.xlsx · Home → Diagram View", "S05b_diagram_roi.png")
tip(s, L, 5.8, 6.9, 0.95, "Không thấy Manage Data Model: bật add-in Power Pivot (File → Options → COM Add-ins).")
tip(s, 8.75, 5.8, 3.99, 0.95, "Diagram View: kéo lookup bao quanh fact.", pro=True)
footnote(s, "Ảnh: Demo/B5_Demo_Model.xlsx, Data → Manage Data Model (Data View, Diagram View)")
notes(s, 5)

# 6 · Normalization · Fact vs Lookup
s = content("Normalization · Fact vs Lookup", R1)
tb(s, L, 1.55, W, 0.5, [[("Normalization: ", {"b": True, "color": NAVY}), "mỗi thông tin lưu đúng MỘT chỗ, tên ngành nằm ở Product, không lặp lại trên từng dòng bán hàng."]], size=16)
tb(s, L, 2.25, 4.9, 0.4, "Fact: ghi sự kiện, có số để cộng", size=15, color=NAVY, bold=True)
tbl_box(s, L, 2.7, 4.9, 1.15, "Sales · 30.247 dòng", "grain: 1 sản phẩm trong 1 đơn", fill=NAVY, line=NAVY, tcolor=WHITE, name="Fact Sales")
s.shapes[-1].text_frame.paragraphs[1].runs[0].font.color.rgb = rgb("E3ECF7")
tbl_box(s, L, 4.0, 4.9, 1.15, "Target · 144 dòng", "grain: tháng × cửa hàng, fact thứ hai", fill=NAVY, line=NAVY, tcolor=WHITE, name="Fact Target")
s.shapes[-1].text_frame.paragraphs[1].runs[0].font.color.rgb = rgb("E3ECF7")
tb(s, 6.97, 2.25, 5.77, 0.4, "Lookup: mô tả, mỗi đối tượng một dòng", size=15, color=NAVY, bold=True)
for k, (t, n_, d_) in enumerate([("Product", "40", "sản phẩm"), ("Store", "12", "cửa hàng"), ("Customer", "1.500", "khách hàng"), ("Calendar", "731", "ngày")]):
    tbl_box(s, 6.97 + (k % 2) * 2.95, 2.7 + (k // 2) * 1.3, 2.82, 1.15, f"{t} · {n_}", d_, name=f"Lookup {t}")
tip(s, L, 5.55, W, 1.0, "Bảng fact dài, có cột số (Quantity, Revenue). Bảng lookup ngắn, chủ yếu là chữ. Target là fact: nó ghi một con số kế hoạch, nhưng không có cột sản phẩm.")
footnote(s, "Số dòng: Data Model của file khởi đầu (Power Pivot → Data View)")
notes(s, 6)

# 7 · Primary key · Foreign key
s = content("Primary key · Foreign key", R1)
table(s, L, 1.55, [2.2, 3.4, 5.47], [
    ["Bảng lookup", "Primary key (duy nhất)", "Foreign key ở bảng fact (lặp lại)"],
    ["Product", "Product_Code: 40 mã", "Sales[Product_Code]"],
    ["Store", "Store_Code: 12 mã", "Sales[Store_Code] · Target[Store_Code]"],
    ["Customer", "Customer_ID: 1.500 mã", "Sales[Customer_ID]"],
    ["Calendar", "Date: 731 ngày", "Sales[Order_Date] · Target[Ngay_Dau_Thang]"],
], size=15, rowh=0.55, bold_first=True)
stat(s, L, 4.55, 5.43, 1.35, "Khóa chính phải DUY NHẤT", "P021 (B2) · P015 (B4)", "trùng thì không tạo được quan hệ", vsize=22, vcolor=BAD_TXT)
stat(s, 7.31, 4.55, 5.43, 1.35, "Khóa ngoại không có trong lookup", "0 dòng", "đã kiểm bằng Left Anti ở B4", vsize=26, vcolor=GOOD)
tip(s, L, 6.0, W, 0.75, "Ngay_Dau_Thang (B4 thêm vào Target) là cái móc để Target nối vào Calendar.")
notes(s, 7)

# 8 · Relationships vs Merged tables
s = content("Relationships vs Merged tables", R1)
card(s, L, 1.6, 5.43, 2.3, "Merge (Buổi 4)", ["Kéo cột của lookup vào bảng fact → một bảng to, thông tin lặp lại trên mọi dòng.",
                                            "Lookup trùng khóa → số dòng phình ra (28.328 → 29.408), không báo lỗi."], size=15, name="Card Merge")
card(s, 7.31, 1.6, 5.43, 2.3, "Relationship", ["Giữ các bảng riêng, nối bằng khóa; Pivot lấy cột từ bảng nào cũng được.",
                                              "Mỗi thông tin một chỗ → mô hình hình ngôi sao (star schema)."], size=15, title_color=GOOD, name="Card Relationship")
stat(s, L, 4.1, 5.43, 1.6, "Merge Product, Store, Customer vào Sales", "453.705 ô", "30.247 dòng × 15 cột mô tả, lặp lại", vsize=30, vcolor=BAD_TXT)
stat(s, 7.31, 4.1, 5.43, 1.6, "Giữ lookup riêng + relationship", "9.220 ô", "Product 40×4 · Store 12×5 · Customer 1.500×6", vsize=30, vcolor=GOOD)
tip(s, L, 5.9, W, 0.8, "Star schema: fact ở giữa, lookup xung quanh, mình dựng nó ở Phần 2.")
footnote(s, "Số ô tính từ Data Model của file khởi đầu (build_b5_meta.json)")
notes(s, 8)

# 9 · Chia phần 2
notes(section("2. Relationship: nối bảng bằng khóa, không bằng công thức"), 9)

# 10 · Create Relationship
s = content("Tạo và quản lý Relationship", R2)
table(s, L, 1.55, [4.6, 6.47], [
    ["Tên", "Mục đích"],
    [[("Kéo thả trong Diagram View", {"b": True, "color": RED})], "Kéo cột khóa ngoại của bảng fact, thả lên khóa chính của bảng lookup"],
    [[("Design → Create Relationship", {"b": True, "color": RED})], "Hộp thoại chọn bảng và cột hai bên, khi sơ đồ quá rối để kéo"],
    [[("Design → Manage Relationships", {"b": True, "color": RED})], "Xem, sửa, xóa từng relationship (ngoài Excel: Data → Relationships)"],
], size=14, rowh=0.48)
shot(s, L, 3.6, 4.6, 2.15, "Create Relationship: Sales / Order_Date ↔ Calendar / Date", "Power Pivot → Design → Create Relationship", "S10a_create_relationship.png")
shot(s, 6.47, 3.6, 6.27, 2.15, "Diagram View: Sales ở giữa, 4 đường nối có ký hiệu 1 (lookup) và * (Sales)", "Sau khi tạo 4 relationship của Sales", "S10b_diagram_sales.png")
tip(s, L, 5.85, W, 0.9, "Hai cột phải CÙNG KIỂU dữ liệu (Order_Date và Date đều là Date nhờ Using Locale ở B4). Báo trùng giá trị = khóa lookup bị trùng → Remove Duplicates trong Power Query.")
footnote(s, "Ảnh: Demo/B5_Demo_Model.xlsx, Power Pivot → Diagram View; Design → Create Relationship")
notes(s, 10)

# 11 · Cardinality · Filter direction
s = content("Cardinality 1-* · Filter direction", R2)
tbl_box(s, L + 0.6, 1.75, 3.0, 1.0, "Store · 12 dòng", "Region = Nam", name="Diagram Store")
tbl_box(s, L + 0.6, 4.35, 3.0, 1.0, "Sales · 30.247 dòng", "chỉ còn dòng cửa hàng miền Nam", fill=NAVY, line=NAVY, tcolor=WHITE, name="Diagram Sales")
s.shapes[-1].text_frame.paragraphs[1].runs[0].font.color.rgb = rgb("E3ECF7")
arrow(s, L + 2.1, 2.8, L + 2.1, 4.3, color=ACCENT, w=3, name="Filter arrow")
tb(s, L + 2.3, 2.85, 0.5, 0.4, [[("1", {"b": True, "color": NAVY, "size": 20})]], name="Card one")
tb(s, L + 2.3, 3.85, 0.5, 0.4, [[("*", {"b": True, "color": NAVY, "size": 24})]], name="Card many")
tb(s, 6.0, 1.7, 6.74, 4.2, [
    [("Cardinality: ", {"b": True, "color": RED}), "mức độ duy nhất của khóa ở hai đầu, 1 (lookup, duy nhất) – * (fact, lặp lại). Data Model chỉ nhận 1-1 và 1-*; khóa trùng ở cả hai đầu (nhiều–nhiều) không tạo được."],
    [("Filter direction: ", {"b": True, "color": RED}), "bộ lọc chảy TỪ LOOKUP XUỐNG FACT. Chọn Region = Nam ở Store → Sales chỉ còn dòng của cửa hàng miền Nam. Chiều ngược lại không có."],
    [("Quy tắc vàng: ", {"b": True, "color": NAVY}), "không bao giờ nối thẳng hai bảng fact. Sales và Target gặp nhau qua lookup chung."],
], size=16, space=14)
tip(s, L, 5.9, W, 0.8, "Bộ lọc chảy như nước, từ trên xuống: lookup ở trên, fact ở dưới.")
footnote(s, "Sơ đồ khái niệm: bộ lọc từ Store chảy xuống Sales qua relationship Store_Code")
notes(s, 11)

# 12 · Fact thứ hai: Target
s = content("Fact thứ hai: Target lọc qua lookup chung", R2)
tool(s, L, 1.55, W, "Insert → PivotTable ▾ → From Data Model  ·  relationship cho Target", "Pivot đọc thẳng Data Model. Target chỉ lọc được qua lookup CHUNG với Sales: Target[Store_Code] → Store · Target[Ngay_Dau_Thang] → Calendar.", 0.6, size=20)
shot(s, L, 2.8, 4.75, 2.85, "Category × Sum of Target: 6 dòng đều 3.665.000.000; thanh vàng \"Relationships between tables may be needed\"", "Sheet PV_Grain, TRƯỚC khi nối Target", "S12a_target_lap.png")
shot(s, 6.57, 2.8, 3.6, 2.85, "Diagram View: 6 bảng, 6 relationship; Sales và Target cùng nối Store, Calendar", "Sau Slide 12", "S12b_diagram_hai_fact.png")
stat(s, 10.32, 2.8, 2.42, 1.35, "Chưa nối · mọi ngành", "3.665 tr", "= target cả năm", vsize=22, vcolor=BAD_TXT)
stat(s, 10.32, 4.3, 2.42, 1.35, "Đã nối · theo miền", "1.227 · 1.685 · 753", "triệu (Bắc · Nam · Trung)", vsize=16, vcolor=GOOD)
tip(s, L, 5.8, W, 0.95, "Cùng một con số ở MỌI dòng = trường ở Rows không lọc được bảng của trường ở Values. Kiểm tra Diagram View. Target không có cột sản phẩm: ngành hàng lặp lại là ĐÚNG.")
footnote(s, "Ảnh: Demo/B5_Demo_Model.xlsx, sheet PV_Grain (Insert → PivotTable → From Data Model); Power Pivot → Diagram View")
notes(s, 12)

# 13 · Nghỉ
notes(section("Nghỉ 10 phút", "Trước khi quay lại: mở Diagram View, kiểm tra đủ 6 relationship"), 13)

# 14 · Chia phần 3
notes(section("3. Calculated Column và Measure: tính từng dòng hay tính từng ô"), 14)

# 15 · Calculated Column · RELATED
s = content("Calculated Column · RELATED", R3)
table(s, L, 1.55, [4.6, 6.47], [
    ["Tên", "Mục đích"],
    [[("Calculated Column", {"b": True, "color": RED}), (" (Data View → Add Column)", {"color": INK2})], "Công thức tính cho TỪNG DÒNG, kết quả LƯU thành một cột của bảng"],
    [[("RELATED(<cột ở bảng lookup>)", {"b": True, "color": RED})], "Lấy giá trị từ bảng lookup theo relationship, XLOOKUP của Data Model"],
], size=14, rowh=0.55)
fx(s, L, 3.35, W, "= Sales[Quantity] * RELATED(Product[Unit_Cost])        → đổi tên cột: COGS", size=16)
shot(s, L, 3.95, 6.9, 1.85, "Data View bảng Sales: cột COGS ở cuối, thanh công thức =Sales[Quantity]*RELATED(Product[Unit_Cost])", "Power Pivot → Data View → Sales", "S15_cot_cogs.png")
tip(s, 8.75, 3.95, 3.99, 1.85, "Không có địa chỉ ô kiểu A1: công thức viết theo tên cột, áp cho cả cột. Đổi tên: double-click tiêu đề.")
tip(s, L, 5.9, W, 0.8, "RELATED chỉ chạy khi ĐÃ CÓ relationship, và đi từ fact sang lookup (phía \"nhiều\" sang phía \"một\").", pro=True)
footnote(s, "Ảnh: Demo/B5_Demo_DAX.xlsx, Power Pivot → Data View, bảng Sales, cột COGS")
notes(s, 15)

# 16 · Vì sao KPI tỷ lệ phải là measure
s = content("Vì sao KPI tỷ lệ phải là measure", R3)
fx(s, L, 1.55, W, "GM_Dong = DIVIDE(Sales[Revenue] - Sales[COGS], Sales[Revenue])      (calculated column, thêm tạm)", size=15)
table(s, L, 2.15, [2.0, 2.45, 2.45], [
    ["", "Calculated Column", "Measure"],
    ["Tính theo", "từng dòng (row context)", "từng ô Pivot (filter context)"],
    ["Lưu trong mô hình", "có: tốn bộ nhớ", "không: tính khi cần"],
    ["Dùng ở", "Rows, Columns, Filters, Slicer", "chỉ ở Values"],
    ["Hợp với", "thuộc tính của một dòng", "KPI, tổng, TỶ LỆ"],
], size=13, rowh=0.45, bold_first=True)
stat(s, L, 4.6, 2.2, 1.4, "Average of GM_Dong", "33,77%", "trung bình của tỷ lệ", vsize=24, vcolor=BAD_TXT)
stat(s, L + 2.35, 4.6, 2.2, 1.4, "Sum of GM_Dong", "208.904%", "vô nghĩa", vsize=22, vcolor=BAD_TXT)
stat(s, L + 4.7, 4.6, 2.2, 1.4, "Đúng (B2, ô E30)", "34,18%", "tổng LN ÷ tổng DT", vsize=24, vcolor=GOOD)
shot(s, 8.75, 2.15, 3.99, 3.85, "Pivot PV_Loc lọc Nam = 2025, Quy = 1: Average of GM_Dong 33,77% và Sum of GM_Dong", "Sheet PV_Loc", "S16_gm_dong.png")
tip(s, L, 6.03, W, 0.74, "Pivot đặt ở sheet PV_Loc, ô A3 (Existing Worksheet). File demo có sẵn sheet này; file B5_<HoTen>.xlsx: bấm dấu + cạnh tab sheet, đặt tên PV_Loc.")
footnote(s, "Ảnh: Demo/B5_Demo_DAX.xlsx, sheet PV_Loc (cột GM_Dong thêm tạm, xóa sau demo). Đối chiếu: sheet KPI_SUMIFS, ô E30")
notes(s, 16)

# 17 · Implicit vs Explicit measure
s = content("Implicit vs Explicit measure", R3)
table(s, L, 1.55, [4.6, 6.47], [
    ["Tên", "Mục đích"],
    [[("Implicit measure", {"b": True, "color": RED})], "Tick cột số vào Values → \"Sum of …\" tự tạo: không tên, không định dạng, KHÔNG dùng lại được"],
    [[("Power Pivot → Measures → New Measure…", {"b": True, "color": RED})], "Explicit measure: tự đặt tên, viết DAX, chọn định dạng; dùng lại trong measure khác và CUBEVALUE (B6)"],
    [[("DIVIDE(tử, mẫu)", {"b": True, "color": RED})], "Phép chia an toàn: mẫu bằng 0 hoặc trống → ô trống thay vì lỗi"],
], size=13, rowh=0.5)
table(s, L, 3.65, [2.3, 4.6], [
    ["Measure (bảng Sales)", "DAX"],
    ["Total Revenue", "SUM(Sales[Revenue])"],
    ["Total COGS", "SUM(Sales[COGS])"],
    ["Gross Margin %", "DIVIDE([Total Revenue] - [Total COGS], [Total Revenue])"],
], size=12, rowh=0.42, bold_first=True, name="Measures")
shot(s, 8.75, 3.65, 3.99, 2.1, "Hộp Measure: Gross Margin %, Formula, Check formula, Category Number, Format Percentage", "Power Pivot → Measures → New Measure", "S17_measure_dialog.png")
tip(s, L, 5.9, W, 0.8, "Cột luôn kèm tên bảng: Sales[Revenue]. Measure ghi trần: [Total Revenue]. Quý 1/2025: 941.454.000 · 619.663.000 · 34,2%.")
footnote(s, "Ảnh: Demo/B5_Demo_DAX.xlsx, Power Pivot → Measures → New Measure (Gross Margin %)")
notes(s, 17)

# 18 · Chia phần 4
notes(section("4. Filter context: ô này đang bị lọc bởi những gì?"), 18)

# 19 · Filter context, Chuỗi C lần 3
s = content("Filter context: Chuỗi C lần 3", R4)
box(s, L, 1.7, 2.3, 1.2, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08, name="Spine cell")
tb(s, L, 1.7, 2.3, 1.2, [[("90.210.000", {"b": True, "color": WHITE, "size": 22})], [("một ô Pivot", {"color": "E3ECF7", "size": 12})]], anchor="ctr", align="ctr", space=2)
for k, t in enumerate(["Store[Region] = Bắc", "Calendar[Nam] = 2025", "Calendar[Quy] = 1", "Calendar[Thang] = 3"]):
    chip(s, L + 2.75, 1.55 + k * 0.41, 2.6, t, fill=WHITE, line=NAVY, color=NAVY, size=12, name=f"Filter {k+1}", h=0.36)
box(s, L + 5.8, 1.85, 2.5, 0.9, fill=ALT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1, name="Spine sum")
tb(s, L + 5.8, 1.85, 2.5, 0.9, [[("SUM(Sales[Revenue])", {"b": True, "color": NAVY, "size": 14})]], anchor="ctr", align="ctr")
arrow(s, L + 2.7, 2.3, L + 2.35, 2.3, name="Arrow 1"); arrow(s, L + 5.75, 2.3, L + 5.4, 2.3, name="Arrow 2")
table(s, L, 3.2, [1.3, 3.8], [["Buổi", "Cùng một ý"], ["B2", "tiêu chí SUMIFS (KPI_SUMIFS!D6)"], ["B3", "tọa độ ô Pivot"], ["B5", "filter context của measure"]],
      size=13, rowh=0.42, bold_first=True, name="Three times")
numbered(s, L, 5.05, 5.1, 1.6, ["Đọc tọa độ ô: hàng, cột, Filters, Slicer", "Bộ lọc chảy theo relationship, lookup → Sales", "Tính công thức trên các dòng còn lại"], size=14, space=4)
shot(s, 6.97, 3.2, 5.77, 3.45, "Pivot Rows Region · Columns Thang · Filters Nam = 2025, Quy = 1 · Values Total Revenue; ô Bắc × 3 = 90.210.000 đang chọn", "Sheet PV_Loc", "S19_pivot_mien_thang.png")
footnote(s, "Ảnh: bản sao Practice/B5_MCI_Mart_Answer.xlsx, sheet PV_Loc (trên lớp: Demo/B5_Demo_DAX.xlsx). Đối chiếu: sheet KPI_SUMIFS, ô D6 và E9")
notes(s, 19)

# 20 · Lookup vs fact · Hide from Client Tools
s = content("Lọc bằng cột của lookup", R4)
tool(s, L, 1.55, W, "Hide from Client Tools  (Power Pivot: chuột phải tiêu đề cột)", "Ẩn cột khỏi danh sách trường của Pivot, để người dùng không kéo nhầm cột của bảng fact vào Rows / Columns.", 0.55, size=20)
shot(s, L, 2.7, 6.9, 2.85, "Hai Pivot cạnh nhau (Nam = 2025): trái Rows Calendar[Thang]: đúng; phải Rows Sales[Thang], chỉ còn Sum of Target, 3.665.000.000 ở mọi dòng", "Sheet PV_Lookup_vs_Fact, trước khi ẩn cột", "S20_lookup_vs_fact.png")
card(s, 8.75, 2.7, 3.99, 2.85, "Quy tắc", ["Rows, Columns, Filters, Slicer: lấy cột của LOOKUP (Calendar, Store, Product, Customer).",
                                          [("Sales[Thang] ", {"b": True}), "chỉ lọc được Sales, không chảy sang Target."]], size=14, name="Rule card")
tip(s, L, 5.7, W, 1.0, "Ẩn: Sales[Nam], [Thang], [Order_Date], [Store_Code], [Product_Code], [Customer_ID] · Target[Store_Code], [Thang], [Ngay_Dau_Thang]. Cột ẩn chuyển màu xám trong Data View.")
footnote(s, "Ảnh: bản sao Demo/B5_Demo_DAX.xlsx, sheet PV_Lookup_vs_Fact (tạo trong demo)")
notes(s, 20)

# 21 · Chia phần 5
notes(section("5. CALCULATE và ALL: đổi bộ lọc có chủ đích"), 21)

# 22 · CALCULATE
s = content("CALCULATE", R5)
tool(s, L, 1.55, W, "CALCULATE(<biểu thức>, <bộ lọc 1>, <bộ lọc 2>, …)", "Tính biểu thức dưới bộ lọc mình đặt; bộ lọc trong CALCULATE GHI ĐÈ bộ lọc cùng cột của ô Pivot, các bộ lọc khác giữ nguyên.", 0.6, size=22)
fx(s, L, 2.8, W, 'Revenue Mien Nam := CALCULATE([Total Revenue], Store[Region] = "Nam")', size=16)
shot(s, L, 3.45, 6.9, 2.3, "PV_Loc (Q1/2025): Rows Region; Total Revenue và Revenue Mien Nam, 427.663.000 ở cả 3 dòng", "Sheet PV_Loc", "S22_calculate.png")
stat(s, 8.75, 3.45, 3.99, 2.3, "Revenue Mien Nam ở dòng Bắc, Trung, Nam", "427.663.000", "Region bị thay bằng Nam; bộ lọc Quý 1 vẫn giữ", vsize=26)
tip(s, L, 5.9, W, 0.8, "Demo xong, xóa measure này: Power Pivot → Measures → Manage Measures → Delete.")
footnote(s, "Ảnh: bản sao Practice/B5_MCI_Mart_Answer.xlsx, sheet PV_Loc (trên lớp: Demo/B5_Demo_DAX.xlsx)")
notes(s, 22)

# 23 · ALL → % of Total
s = content("ALL → % of Total", R5)
tool(s, L, 1.55, W, "ALL(<bảng hoặc cột>)", "Bỏ mọi bộ lọc trên bảng (hoặc cột) đó: dùng trong CALCULATE để lấy \"mẫu số chung\".", 0.5, size=22)
table(s, L, 2.65, [3.0, 8.07], [
    ["Measure", "DAX"],
    ["Revenue All Stores", "CALCULATE([Total Revenue], ALL(Store))"],
    ["% Revenue Share", "DIVIDE([Total Revenue], [Revenue All Stores])"],
], size=14, rowh=0.45, bold_first=True, name="Measures")
shot(s, L, 4.1, 6.9, 1.65, "PV_Loc (Q1): Total Revenue · Revenue All Stores 941.454.000 mọi dòng · % Revenue Share", "Sheet PV_Loc", "S23_pct_of_total.png")
for k, (r_, v_) in enumerate([("Bắc", "34,4%"), ("Nam", "45,4%"), ("Trung", "20,2%")]):
    stat(s, 8.75 + k * 1.36, 4.1, 1.28, 1.65, r_, v_, vsize=18, name=f"Share {r_}")
tip(s, L, 5.9, W, 0.8, "Đúng Pct_Share của Buổi 1. B1 khóa mẫu số bằng $B$5; DAX khóa bằng ALL. ALL(Store) bỏ lọc miền, cửa hàng, loại cửa hàng; bộ lọc Calendar vẫn giữ. Tỷ trọng theo ngành: ALL(Product).")
footnote(s, "Ảnh: bản sao Practice/B5_MCI_Mart_Answer.xlsx, sheet PV_Loc (trên lớp: Demo/B5_Demo_DAX.xlsx)")
notes(s, 23)

# 24 · Target measures
s = content("Measure cho Target", R5)
table(s, L, 1.55, [3.0, 8.07], [
    ["Measure", "DAX"],
    ["Total Target  (bảng Target)", "SUM(Target[Target])"],
    ["Revenue vs Target %", "DIVIDE([Total Revenue], [Total Target])"],
    ["Variance", [("IF(ISBLANK([Total Revenue]), BLANK(), ", {"b": True, "color": RED}), "[Total Revenue] - [Total Target])"]],
], size=14, rowh=0.45, bold_first=True, name="Measures")
shot(s, L, 3.5, 6.9, 2.25, "PV_Target (Nam = 2025): Total Revenue, Total Target, Revenue vs Target %, Variance theo Calendar[Thang]; T5–T12 chỉ có target", "Sheet PV_Target", "S24_target_thang.png")
for k, (m_, v_, c_) in enumerate([("T1", "120,6%", GOOD), ("T2", "82,5%", BAD_TXT), ("T3", "94,6%", BAD_TXT), ("T4", "103,1%", GOOD)]):
    stat(s, 8.75 + (k % 2) * 2.03, 3.5 + (k // 2) * 1.15, 1.96, 1.1, m_ + "/2025", v_, vsize=20, vcolor=c_, name=f"Month {m_}")
tip(s, L, 5.9, W, 0.8, "Không có IF(ISBLANK…): Variance tháng 5 = −288.000.000, \"hụt 288 triệu\" chỉ vì chưa có dữ liệu. T1 + T2 gộp: 103,2% (lệch mùa Tết). Dòng Grand Total 33,8% = 4 tháng ÷ target cả năm: không đọc.")
footnote(s, "Ảnh: bản sao Practice/B5_MCI_Mart_Answer.xlsx, sheet PV_Target (trên lớp: Demo/B5_Demo_DAX.xlsx)")
notes(s, 24)

# 25 · Revenue vs Target % chỉ theo chiều Target có
s = content("Target không lọc theo ngành", R5)
card(s, L, 1.6, 4.6, 4.15, "Rows = Product[Category] · Quý 1", [
    "Bánh kẹo 36,2% · Gia dụng 21,7% · Thực phẩm khô 18,0% · Đồ uống 11,8% · Chăm sóc cá nhân 11,2% · Sữa 1,8%",
    [("= doanh thu ngành ÷ target CẢ CHUỖI (936.000.000) → vô nghĩa", {"b": True, "color": BAD_TXT})],
    "Measure không báo lỗi, người dùng phải biết giới hạn."], size=14, title_color=BAD_TXT, name="Card category")
shot(s, 6.47, 1.6, 3.95, 4.15, "Rows Store_Code · Columns Thang 1–4 · Revenue vs Target %; ô ST07 × 3 = 44,9% đang chọn", "Sheet PV_Target", "S25_store_thang.png")
stat(s, 10.57, 1.6, 2.17, 2.0, "ST07 · T3/2025", "44,9%", "8.539.000 / 19.000.000", vsize=24, vcolor=BAD_TXT)
stat(s, 10.57, 3.75, 2.17, 2.0, "ST07 · T4/2025", "97,4%", "hồi phục sau cải tạo", vsize=24, vcolor=GOOD)
tip(s, L, 5.9, W, 0.8, "KPI dictionary ghi rõ: Revenue vs Target % lọc theo store, region, month, KHÔNG theo product/category.")
footnote(s, "Ảnh: bản sao Practice/B5_MCI_Mart_Answer.xlsx, sheet PV_Target_CuaHang (trên lớp: sheet PV_Target, Demo/B5_Demo_DAX.xlsx)")
notes(s, 25)

# 26 · Chia phần 6
notes(section("6. Time intelligence: lũy kế và so với cùng kỳ"), 26)

# 27 · YTD · cùng kỳ năm trước
s = content("YTD · cùng kỳ năm trước", R6)
table(s, L, 1.55, [3.3, 3.6], [
    ["Tên", "Mục đích"],
    [[("Design → Mark as Date Table", {"b": True, "color": RED})], "Khai báo Calendar là bảng lịch: mỗi ngày một dòng, liên tục (731 ngày ✔)"],
    [[("DATESYTD(Calendar[Date])", {"b": True, "color": RED})], "Từ đầu năm tới ngày cuối của ô → lũy kế"],
    [[("DATEADD(Calendar[Date], -1, YEAR)", {"b": True, "color": RED})], "Lùi các ngày của ô một năm → cùng kỳ"],
], size=12, rowh=0.48)
shot(s, 8.75, 1.55, 3.99, 1.9, "Hộp Mark as Date Table, Date column = Date", "Power Pivot, tab Calendar → Design", "S27a_mark_date_table.png")
table(s, L, 3.6, [1.75, 5.15], [
    ["Measure", "DAX"],
    ["Revenue YTD", "CALCULATE([Total Revenue], DATESYTD(Calendar[Date]))"],
    ["Revenue LY", "CALCULATE([Total Revenue], DATEADD(Calendar[Date], -1, YEAR))"],
    ["Revenue YoY %", "IF(ISBLANK([Total Revenue]), BLANK(), DIVIDE([Total Revenue] - [Revenue LY], [Revenue LY]))"],
], size=11, rowh=0.48, bold_first=True, name="Measures")
shot(s, 8.75, 3.6, 3.99, 2.0, "PV_YoY T1–T4/2025: Revenue, LY, YTD, YoY %; YTD T4 1.238.419.000", "Sheet PV_YoY", "S27b_yoy_ytd.png")
tip(s, L, 5.75, 5.6, 1.0, "YoY T1 +42,8% · T2 −0,9% · T3 +13,7% · T4 +24,0%; YTD T4 +21,0%. Không chặn ô trống: T5 = −100%.")
tip(s, 7.42, 5.75, 5.32, 1.0, "SAMEPERIODLASTYEAR cho cùng kết quả với DATEADD lùi 1 năm.", pro=True)
footnote(s, "Ảnh: Demo/B5_Demo_DAX.xlsx (Design → Mark as Date Table); bản sao Practice/B5_MCI_Mart_Answer.xlsx, sheet PV_YoY")
notes(s, 27)

# 28 · Like-for-like
s = content("Like-for-like: ST12 và ô YoY trống", R6)
shot(s, L, 1.6, 6.9, 3.55, "PV_YoY: Rows Store_Code, Filters Nam = 2025 và Thang 1–4; Total Revenue, Revenue LY, Revenue YoY %; dòng ST12 hai ô trống", "Sheet PV_YoY", "S28_yoy_cua_hang.png")
stat(s, 8.75, 1.6, 3.99, 1.15, "ST12 khai trương 01/06/2024", "79.399.000", "T1–T4/2025 · không có kỳ gốc", vsize=22)
stat(s, 8.75, 2.8, 3.99, 1.15, "YTD T4 · cả chuỗi 12 cửa hàng", "+21,0%", "gồm cả ST12", vsize=24, vcolor=BAD_TXT)
stat(s, 8.75, 4.0, 3.99, 1.15, "YTD T4 · like-for-like 11 cửa hàng", "+13,3%", "chỉ cửa hàng có đủ hai kỳ", vsize=24, vcolor=GOOD)
tip(s, L, 5.25, W, 0.72, "D2: thấy ô trống, hỏi vì sao; thấy tổng đẹp, hỏi tổng gồm những gì. ST10 +37,9% · ST07 −3,2% · ST04 −1,5%.")
tip(s, L, 6.02, W, 0.72, "Bài về nhà: like-for-like thành measure, không cần lọc tay: Revenue LFL := CALCULATE([Total Revenue], Store[Open_Date] < DATE(2024,1,1))", pro=True)
footnote(s, "Ảnh: bản sao Practice/B5_MCI_Mart_Answer.xlsx, sheet PV_YoY (trên lớp: Demo/B5_Demo_DAX.xlsx)")
notes(s, 28)

# 29 · Chia phần 7
notes(section("7. KPI dictionary: một định nghĩa, mọi báo cáo"), 29)

# 30 · KPI dictionary
s = content("KPI dictionary", R7)
table(s, L, 1.45, [1.75, 4.35, 2.3, 1.3, 1.37], [
    ["KPI", "Measure (DAX)", "Dimensions", "Owner", "Target"],
    ["Total Revenue", "SUM(Sales[Revenue])", "region, store, category, month", "Sales Director", "= Total Target"],
    ["Total Orders", "DISTINCTCOUNT(Sales[Order_ID])", "region, store, month", "Sales Director", "1.000 đơn/tháng"],
    ["Avg Order Value", "DIVIDE([Total Revenue], [Total Orders])", "region, store, month", "Sales Director", "300.000"],
    ["Gross Margin %", "DIVIDE([Total Revenue] - [Total COGS], [Total Revenue])", "category, region, month", "Finance", "33%"],
    ["Revenue YoY %", "IF(ISBLANK([Total Revenue]), BLANK(), DIVIDE([Total Revenue] - [Revenue LY], [Revenue LY]))", "region, category, month", "Finance", "+10%"],
    ["Active Customers", "DISTINCTCOUNT(Sales[Customer_ID])", "region, segment, month", "Marketing", "700 KH/tháng"],
    ["Revenue vs Target %", "DIVIDE([Total Revenue], [Total Target])", [("store, region, month: KHÔNG category", {"b": True, "color": BAD_TXT})], "Sales Director", "100%"],
], size=11, rowh=0.36, bold_first=True, name="KPI dictionary")
tb(s, L, 4.58, W, 0.5, [
    [("Measure = semantic layer: ", {"b": True, "color": NAVY}), "KPI định nghĩa một lần trong mô hình; mọi Pivot, mọi dashboard (Buổi 6) chỉ gọi tên measure."],
    [("Ô Description ", {"b": True, "color": NAVY}), "của hộp Measure: ghi định nghĩa và giới hạn ngay trong mô hình."],
], size=13, bullet=True, space=2, name="Semantic layer")
shot(s, L, 5.15, 3.3, 1.55, "Rows Calendar[Nam]: Gross Margin % 31,2% · 34,1%", "Sheet PV_YoY", "S30_gm_theo_nam.png")
card(s, 5.12, 5.15, 7.62, 1.55, "D2: kiểm giả định (câu chuyện #5)", [
    "GM% 2024: 31,2% → 2025: 34,1%. Giá vốn chỉ có MỘT bảng (Product[Unit_Cost]) cho cả hai năm, giá bán 2025 cao hơn 2024 khoảng 3,9% → ghi vào cột \"Giới hạn\"."], size=13, name="Card D2")
footnote(s, "Course Plan, Phụ lục C · Ảnh: bản sao Practice/B5_MCI_Mart_Answer.xlsx, sheet PV_GM_Nam (trên lớp: sheet PV_YoY, Demo/B5_Demo_DAX.xlsx)")
notes(s, 30)

# 31 · Thực hành 1
s = content("Thực hành 1: Hoàn thành bộ 6 KPI", R8)
chip(s, L, 1.55, 1.53, "10 phút", name="Time chip")
numbered(s, L, 2.1, 5.3, 3.0, [
    "Viết 3 measure (bảng Sales): Total Orders = DISTINCTCOUNT(Sales[Order_ID]) · Avg Order Value = DIVIDE([Total Revenue], [Total Orders]) · Active Customers = DISTINCTCOUNT(Sales[Customer_ID]).",
    "Sheet mới PV_KPI_Thang (From Data Model): Filters Calendar[Nam] = 2025 · Rows Calendar[Thang] (lọc 1–4) · Values: 6 KPI.",
    "Kiểm tay AOV: lọc Product[Category] = Bánh kẹo, Calendar[Quy] = 1 → đọc dòng Grand Total: Total Revenue ÷ Total Orders = AOV?",
], size=14, space=10)
table(s, 7.17, 1.55, [0.45, 1.3, 0.75, 0.9, 0.75, 0.8, 0.62], [
    ["T", "Revenue", "Orders", "AOV", "GM %", "YoY %", "KH"],
    ["1", "425.850.000", "1.398", "304.614", "35,2%", "+42,8%", "831"],
    ["2", "245.996.000", "823", "298.902", "33,7%", "−0,9%", "605"],
    ["3", "269.608.000", "923", "292.100", "33,1%", "+13,7%", "637"],
    ["4", "296.965.000", "966", "307.417", "33,8%", "+24,0%", "670"],
], size=11, rowh=0.45, name="Check table", align=["ctr", "r", "r", "r", "r", "r", "r"])
box(s, 7.17, 4.0, 5.57, 1.75, fill=ALT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05, name="AOV check")
tb(s, 7.37, 4.1, 5.17, 1.55, [[("Kiểm tay: Bánh kẹo, Quý 1", {"b": True, "color": NAVY})],
                              [("338.733.000 ÷ 1.739 đơn = ", {}), ("194.786", {"b": True, "color": GOOD})],
                              [("AVERAGE(Sales[Revenue]) = 153.830: trung bình một DÒNG, không phải một ĐƠN", {"size": 13, "color": INK2})]], size=15, space=4)
footnote(s, "Kiểm tra: Filters Calendar[Nam] = 2025, Rows Calendar[Thang] 1–4 · KH = Active Customers")
notes(s, 31)

# 32 · Thực hành 2
s = content("Thực hành 2: Đối soát với KPI_SUMIFS", R8)
chip(s, L, 1.55, 1.53, "10 phút", name="Time chip")
numbered(s, L, 2.1, 6.9, 2.4, [
    "Sheet mới PV_DoiSoat_B2 (From Data Model): Filters Nam = 2025, Quy = 1 · Rows Store[Region] · Columns Calendar[Thang] · Values: Total Revenue, Total Orders, Avg Order Value, Gross Margin %.",
    "GETPIVOTDATA (gõ =, bấm ô Pivot) trừ ô KPI_SUMIFS: Bắc × 3 (D6) · Doanh thu Quý 1 (E9) · Số đơn (E16) · GM% (E30) → phải bằng 0.",
], size=14, space=10)
items32 = ["Diagram View: hình ngôi sao, 6 bảng, 6 relationship; Sales và Target gặp nhau qua Store, Calendar",
           "6 KPI + Total Target, Revenue vs Target %, Variance là explicit measure, có định dạng",
           "Đối soát với KPI_SUMIFS: chênh lệch 0"]
for k, t in enumerate(items32):
    y = 4.55 + k * 0.62
    box(s, L, y + 0.13, 0.3, 0.3, fill=WHITE, line=NAVY, lw=1.5, name=f"Checkbox {k+1}")
    tb(s, L + 0.5, y, 6.4, 0.56, t, size=14, anchor="ctr")
box(s, 8.87, 1.55, 3.87, 5.05, fill=ALT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03, name="Check box")
tb(s, 9.07, 1.7, 3.47, 4.8, [[("Kiểm tra: Quý 1/2025", {"b": True, "color": NAVY, "size": 16})], [("941.454.000", {"b": True, "color": NAVY, "size": 24})],
                             "3.144 đơn · AOV 299.445", "Gross Margin 34,18%", "Bắc × 3: 90.210.000",
                             [("Khớp = B2 và B5 cùng nói một sự thật", {"b": True, "color": GOOD})]], size=15, space=8)
footnote(s, "Làm mẫu: Practice/B5_MCI_Mart_Answer.xlsx, sheet PV_DoiSoat_B2 (ô J11, N14) · đủ 64 ô Miền × Tháng × 4 KPI: sheet DoiSoat_B2 cùng file")
notes(s, 32)

# 33 · Quiz cover
notes(section("Excel Quiz Challenge", "4 câu hỏi nhanh · câu 3–4 là tình huống"), 33)
quiz(34, "Measure vs Column", "KPI Gross Margin % nên viết thế nào trong Data Model?",
     {"A": "Calculated column GM từng dòng, rồi lấy Average", "B": "Calculated column GM từng dòng, rồi lấy Sum",
      "C": "Measure: DIVIDE(tổng doanh thu − tổng giá vốn, tổng doanh thu)", "D": "Implicit measure Average of Revenue"}, "C", osize=15)
quiz(35, "Hai bảng fact", "Sales và Target đều có cột Store_Code. Nối thế nào để phân tích doanh thu so với target theo cửa hàng?",
     {"A": "Relationship thẳng giữa Sales[Store_Code] và Target[Store_Code]", "B": "Nối cả hai vào lookup chung Store (và Calendar), lọc qua lookup",
      "C": "Merge Target vào Sales bằng Power Query", "D": "Không nối được, phải dùng SUMIFS"}, "B", osize=15)
quiz(36, "Tình huống: Con số lặp lại", "Pivot Quý 1/2025: Rows = Product[Category], Values = Total Target. Sáu dòng đều ra 936.000.000. Vì sao?",
     {"A": "Thiếu Mark as Date Table", "B": "Measure Total Target viết sai",
      "C": "Target có grain tháng × cửa hàng, không có chiều sản phẩm; Product không lọc được Target", "D": "Phải Refresh All"}, "C", qsize=20, osize=14)
quiz(37, "Tình huống: Tổng lệch", "Active Customers Quý 1 theo miền: 492 + 657 + 307 = 1.456. Nhưng Grand Total ghi 1.208. Kết luận nào đúng?",
     {"A": "Pivot bị lỗi, phải Refresh", "B": "Thiếu relationship Sales – Customer",
      "C": "Không lỗi: khách mua ở nhiều miền chỉ đếm một lần ở dòng tổng; DISTINCTCOUNT không cộng dồn", "D": "Phải đổi sang SUM(Sales[Customer_ID])"}, "C", qsize=20, osize=14)

# 38 · Tổng kết + insight
s = content("Tổng kết & insight của bạn", R9)
tb(s, L, 1.65, 5.3, 0.45, "Điểm chính cần nhớ", size=18, color=NAVY, bold=True)
tb(s, L, 2.2, 5.3, 4.5, ["(1) Star schema: fact ở giữa, lookup xung quanh, 1-*; bộ lọc chảy lookup → fact. Hai fact gặp nhau qua lookup chung.",
                        "(2) Calculated column tính từng dòng; MEASURE tính từng ô. Tỷ lệ và đếm phân biệt luôn là measure.",
                        "(3) Filter context = tọa độ của ô. CALCULATE đổi bộ lọc; ALL bỏ bộ lọc.",
                        "(4) Time intelligence cần Calendar liên tục; hàng/cột lấy từ Calendar, không lấy từ fact.",
                        "(5) KPI dictionary: mỗi KPI một measure, một định nghĩa, kèm giới hạn, đối soát trước khi dùng."], size=15, space=9)
box(s, 6.97, 1.65, 5.77, 3.8, fill=ALT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03, name="Insight box")
tb(s, 7.17, 1.8, 5.37, 0.45, "Insight của bạn (1 câu)", size=18, color=NAVY, bold=True)
tb(s, 7.17, 2.35, 5.37, 2.2, "\"Theo Revenue vs Target %, ___ (con số), nguyên nhân khả dĩ là ___, nên ___ (hàm ý).\"", size=19)
for k, t in enumerate(["Con số", "Nguyên nhân", "Hàm ý"]):
    chip(s, 7.17 + k * 1.84, 4.8, 1.69, t, fill=WHITE, line=HAIR, name=f"Tag {t}", h=0.45)
tb(s, 6.97, 5.6, 5.77, 0.9, "Gợi ý: so tháng 1 với tháng 2, rồi mở Calendar xem Tết rơi vào ngày nào.", size=14, color=INK2)
notes(s, 38)

# 39 · Preview B6
s = content("Buổi sau: Dashboard", R9)
tb(s, L, 1.65, 7.0, 0.9, "\"Giám đốc bán hàng có 5 giây. Họ cần thấy gì đầu tiên?\"", size=20, color=NAVY, bold=True)
tb(s, L, 2.55, 7.0, 1.0, [[("Hôm nay số đã đúng và thống nhất. ", {"b": True}), "Buổi 6: trình bày cho người ra quyết định, KPI card bằng CUBEVALUE đọc thẳng các measure hôm nay."]], size=15)
tb(s, L, 3.6, 7.0, 3.1, [[("Về nhà", {"b": True, "color": NAVY})],
                          [("Bài tập về nhà (bắt buộc): ", {"b": True}), "B5_BTVN_MCI_Mart_KPI_T1_T4_2025.xlsx: báo cáo tháng 1–4 của phòng Kinh doanh lệch số Tài chính: đối soát, tìm 4 lỗi, sửa mô hình và measure, KPI dictionary, 3 câu hỏi cho Sales Director, 1 câu insight"],
                          [("Tập trước capstone: ", {"b": True}), "C3 Mô hình · C4 KPI"],
                          "Nộp trước Buổi 6 tại hocvien.mcivietnam.com · đọc handout B5 · tick \"Checklist trước Buổi 6\""], size=14, space=6)
bricks(s, "B5 · Star schema + DAX", "B6 · Dashboard")
tb(s, 9.07, 5.7, 3.67, 0.5, "File đáp án B5 (Practice/B5_MCI_Mart_Answer.xlsx) là file khởi đầu Buổi 6.", size=13, color=INK2)
notes(s, 39)

assert len(prs.slides) == 39, len(prs.slides)
prs.save(OUT)
print("ok", OUT, len(prs.slides), "slide;", f"còn {len(MISSING)} khung [CHỤP] chưa có ảnh:" if MISSING else "đủ ảnh.", " ".join(MISSING))
