"""Dựng handout tự học Buổi 4 (Power Query): Handout/B4_Handout.docx bằng python-docx.

Số liệu lấy từ build_b4_meta.json (build_b4_homework.py). Ảnh chụp thật lấy từ Materials/B4/Images (H<nn>_…png);
chưa có ảnh thì để khung [CHỤP]. Xuất PDF bằng Microsoft Word (xem 2_Dataset/README.md).
    python3 _generator/build_docs_b4.py
"""
import json, os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.image.image import Image as DocxImage

ROOT = os.path.normpath(os.environ.get("MCI_ROOT") or os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
GEN = os.path.dirname(os.path.abspath(__file__))
OUT = f"{ROOT}/1_Sessions/B4_Power_Query/Handout"
IMG = f"{ROOT}/Materials/B4/Images"   # ảnh chụp thật, danh sách: _DANH_SACH_ANH.md
MISSING = []
os.makedirs(OUT, exist_ok=True)
M = json.load(open(f"{GEN}/build_b4_meta.json", encoding="utf-8"))
H = M["hw"]
RED, NAVY, GOLD, GREY = RGBColor(0xC0, 0, 0), RGBColor(0x20, 0x38, 0x64), RGBColor(0x8A, 0x65, 0x00), RGBColor(0x58, 0x58, 0x58)
vn = lambda x: f"{round(x):,}".replace(",", ".")
pct = lambda x: f"{x * 100:.1f}%".replace(".", ",")


def base_doc(footer="MCI – Magic Code Institute · Excel for Business Intelligence Analyst · Buổi 4"):
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for m in ("left_margin", "right_margin"): setattr(sec, m, Cm(2))
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    for name, size, color in (("Heading 1", 16, NAVY), ("Heading 2", 13, NAVY), ("Heading 3", 11.5, RED)):
        h = d.styles[name]; h.font.name = "Arial"; h.font.size = Pt(size); h.font.color.rgb = color; h.font.bold = True
        rf = h.element.rPr.rFonts
        for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(k)) is not None: del rf.attrib[qn(k)]
        for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"): rf.set(qn(k), "Arial")
    fp = sec.footer.paragraphs[0]
    fp.text = footer
    fp.runs[0].font.size = Pt(8); fp.runs[0].font.color.rgb = GREY
    return d


def para(d, text="", bold=False, color=None, size=None, italic=False):
    p = d.add_paragraph()
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic
        if color: r.font.color.rgb = color
        if size: r.font.size = Pt(size)
    return p


def rich(d, parts):
    p = d.add_paragraph()
    for t, f in parts:
        r = p.add_run(t); r.bold = f.get("bold", False); r.italic = f.get("italic", False)
        if f.get("color"): r.font.color.rgb = f["color"]
        if f.get("mono"): r.font.name = "Consolas"; r.font.size = Pt(9.5)
    return p


def bullet(d, text):
    p = d.add_paragraph(style="List Bullet"); p.add_run(text); return p


def code(d, text):
    """Khối mã M: chữ Consolas trên nền xám nhạt."""
    for line in text.splitlines():
        p = d.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.left_indent = Cm(0.5)
        r = p.add_run(line); r.font.name = "Consolas"; r.font.size = Pt(9)
        pPr = p._p.get_or_add_pPr(); shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F3F5F9"); pPr.append(shd)
    d.add_paragraph().paragraph_format.space_after = Pt(0)


def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_)
    tcPr.append(shd)


def table(d, rows, widths, header=True, size=9.5, first_bold=False):
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for j, w in enumerate(widths): t.columns[j].width = Cm(w)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.width = Cm(widths[j]); c.text = ""
            r = c.paragraphs[0].add_run(str(v)); r.font.size = Pt(size)
            if i < len(rows) - 1: c.paragraphs[0].paragraph_format.keep_with_next = True   # không tách bảng qua trang
            if header and i == 0:
                r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); shade(c, "203864")
            elif first_bold and j == 0:
                r.bold = True
    d.add_paragraph()
    return t


def figure(d, what, where, img, width=15):
    """Chèn ảnh chụp thật từ Materials/B4/Images nếu đã có; chưa có thì để khung chờ [CHỤP] (CLAUDE.md, mục Ảnh minh họa)."""
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(f"{IMG}/{img}"):
        px = DocxImage.from_file(f"{IMG}/{img}")
        w = min(width, px.px_width * 2.54 / 96)
        if w * px.px_height / px.px_width > 10: w = 10 * px.px_width / px.px_height
        p.add_run().add_picture(f"{IMG}/{img}", width=Cm(w)); return p
    MISSING.append(img)
    r = p.add_run(f"[CHỤP] {what}"); r.italic = True; r.font.color.rgb = GREY; r.font.size = Pt(9)
    r = p.add_run(f"\n{where} → {img}"); r.font.color.rgb = GREY; r.font.size = Pt(8)
    pPr = p._p.get_or_add_pPr(); bdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "dashed"); e.set(qn("w:sz"), "6"); e.set(qn("w:color"), "A6A6A6"); e.set(qn("w:space"), "8"); bdr.append(e)
    pPr.append(bdr)
    return p


def tool(d, name, purpose, how=None, example=None, tip=None, protip=None, shot=None, m=None):
    """Khối công cụ: tên (đỏ) → mục đích → cách dùng → bảng ví dụ trên case → mã M → ảnh → TIP / PRO TIP."""
    d.add_heading(name, level=3).paragraph_format.keep_with_next = True
    para(d, purpose).paragraph_format.keep_with_next = True
    for h in how or []: bullet(d, h)
    if example: table(d, example[0], example[1])
    if m: code(d, m)
    if shot: figure(d, *shot)
    if tip: rich(d, [("TIP: ", {"bold": True, "color": NAVY}), (tip, {})])
    if protip: rich(d, [("PRO TIP: ", {"bold": True, "color": GOLD}), (protip, {})])


F1 = M["files"]["Sales_2025-01.csv"]
JS = M["jan_store"]
top_store = max(JS, key=lambda k: JS[k]["rev"])

# =====================================================================
d = base_doc()
para(d, "MCI – MAGIC CODE INSTITUTE", bold=True, color=GREY, size=9)
d.add_heading("Handout Buổi 4 — Tiền xử lý dữ liệu với Power Query", level=1)
para(d, "Tài liệu tự học đi kèm Buổi 4. Mục 1–3 tóm tắt những gì đã làm trên lớp; mục 4–9 là công cụ KHÔNG giảng chi tiết trên lớp, đọc sau buổi học rồi làm "
        "bài tập về nhà (bắt buộc) ở mục 10. Ví dụ viết trên các file Practice/ của Buổi 4 (Data_Sales/*.csv, B4_Lookups.xlsx, B4_Target_2025.xlsx) và "
        "file đáp án B4_MCI_Mart_Answer.xlsx. Tên lệnh giữ tiếng Anh như trên Ribbon của Power Query Editor.", italic=True, color=GREY)

d.add_heading("Mục lục", level=2)
for s in ["1. Tóm tắt: pipeline Power Query của Buổi 4 · lệnh chính", "2. Locale: đọc đúng ngày và số kiểu Việt Nam",
          "3. Đối soát ngay trong Power Query", "4. Index Column và công cụ cho cột số", "5. Pivot Column và Transpose (chiều ngược của Unpivot)",
          "6. Quản lý query trong workbook: Duplicate / Reference, Load To, Properties", "7. From Folder nâng cao: tên file (Source.Name), Transform Sample File",
          "8. PRO TIP: tham số hóa đường dẫn thư mục, Rolling Calendar", "9. Xử lý sự cố: Formula.Firewall, privacy levels, lỗi hay gặp",
          "10. Bài tập về nhà (bắt buộc)", "11. Checklist trước Buổi 5"]:
    bullet(d, s)

# 1 Tóm tắt
d.add_heading("1. Tóm tắt: pipeline Power Query của Buổi 4", level=2)
para(d, "Power Query ghi mỗi thao tác làm sạch thành một bước trong Applied Steps. Refresh là chạy lại toàn bộ các bước trên dữ liệu mới, "
        "nên việc làm sạch chỉ làm MỘT LẦN. Kết quả cuối buổi là 6 bảng sạch, nạp riêng vào Data Model — viên gạch cho Buổi 5.")
table(d, [["Query", "Nguồn", "Các bước chính", "Load", "Số dòng"],
          ["Sales", "Practice/Data_Sales/ (From Folder)", "Xóa Changed Type tự động · Trim Store_Code · Using Locale (vi-VN) cho Order_Date, Unit_Price · Remove Duplicates · Revenue, Nam, Thang",
           "Table Sales (sheet Sales_PQ) + Data Model", f"{vn(M['q_rows'])} → {vn(M['all_rows'])} sau khi thả tháng 4"],
          ["Product", "B4_Lookups.xlsx, bảng Product", "Remove Duplicates theo Product_Code (P015 trùng)", "Connection + Data Model", f"{M['product_raw']} → {M['product']}"],
          ["Store · Customer", "B4_Lookups.xlsx", "Giữ nguyên", "Connection + Data Model", f"{M['store']} · {vn(M['customer'])}"],
          ["Target", "B4_Target_2025.xlsx", "Remove Top Rows 2 · Use First Row as Headers · Unpivot Other Columns · Thang dạng số · Ngay_Dau_Thang", "Connection + Data Model", str(M["target_rows"])],
          ["Calendar", "Blank Query", "List.Dates từ 01/01/2024, 731 ngày · Nam, Thang, Quy, Thu", "Connection + Data Model", str(M["calendar_rows"])]],
      [2.2, 3.4, 6.0, 2.4, 3.0], size=8.5, first_bold=True)
table(d, [["Lệnh / phím tắt", "Tác dụng"],
          ["Data → Get Data → From File → From Text/CSV · From Workbook · From Folder", "Kết nối nguồn; bấm Transform Data để mở Query Editor trước khi nạp"],
          ["Data → Get Data → Launch Power Query Editor", "Mở lại Query Editor"],
          ["Dấu X cạnh một bước trong Applied Steps", "Xóa bước (Query Editor không có Undo)"],
          ["Home → Close & Load ▾ → Close & Load To…", "Chọn Table / Only Create Connection / Add to Data Model"],
          ["Data → Queries & Connections", "Mở khung quản lý query bên phải workbook"],
          ["Data → Refresh All · Ctrl+Alt+F5", "Chạy lại mọi query rồi cập nhật Pivot"]], [8.0, 8.0], first_bold=True)
figure(d, "Hộp From Text/CSV: File Origin 65001 (UTF-8), Delimiter Comma, Data Type Detection, các nút Load / Transform Data",
       "Practice/Data_Sales/Sales_2025-01.csv → Data → Get Data → From File → From Text/CSV", "H01_from_text_csv.png")

# 2 Locale
d.add_heading("2. Locale: đọc đúng ngày và số kiểu Việt Nam", level=2)
para(d, "Power Query đọc ngày và số theo định dạng vùng miền (locale). Mặc định là locale của máy đang chạy Excel, không phải của người tạo ra file. "
        "File CSV của MCI Mart viết ngày dd/mm/yyyy và giá có dấu chấm ngăn nghìn (\"19.000\").")
table(d, [["Đọc cột theo", "Order_Date \"13/01/2025\"", "Order_Date \"02/01/2025\"", "Unit_Price \"19.000\""],
          ["English (United States)", "Error (không có tháng 13)", "01/02/2025 — sai, không báo lỗi", "19 — sai 1.000 lần, không báo lỗi"],
          ["Vietnamese (Vietnam)", "13/01/2025 ✔", "02/01/2025 ✔", "19000 ✔"]], [4.0, 3.8, 4.4, 3.8], first_bold=True)
para(d, f"Trên file tháng 1: đọc kiểu Mỹ thì {vn(F1['day13'])} / {vn(F1['raw'])} dòng Order_Date báo lỗi, {vn(F1['day_le12'])} dòng còn lại không báo lỗi, trong đó {vn(F1['day_le12'] - F1['day_le12_same'])} dòng bị đảo ngày–tháng ({vn(F1['day_le12_same'])} dòng có ngày = tháng như 01/01 vẫn đúng); "
        f"Unit_Price chạy từ {vn(M['jan_price_us_min'])} đến {vn(M['jan_price_us_max'])} thay vì {vn(M['jan_price_min'])} đến {vn(M['jan_price_max'])}.")
tool(d, "Change Type → Using Locale…", "Đổi kiểu một cột theo locale của dữ liệu nguồn.",
     how=["Xóa bước Changed Type tự động (dấu X), vì bước đó đoán kiểu theo locale máy.",
          "Chuột phải tiêu đề cột → Change Type → Using Locale… → Data Type: Date (hoặc Whole Number / Decimal Number) → Locale: Vietnamese (Vietnam) → OK.",
          "Cột không có dấu ngăn nghìn, không phải ngày (Line_No, Quantity) đổi kiểu thường là đủ."],
     tip="Công thức của bước có thêm tham số locale: Table.TransformColumnTypes(…, {{\"Order_Date\", type date}}, \"vi-VN\"). File chạy đúng trên mọi máy, kể cả máy đồng nghiệp để tiếng Anh.")
tool(d, "Locale của cả workbook (Query Options)", "Đặt locale mặc định cho mọi query trong file, thay vì theo máy.",
     how=["Data → Get Data → Query Options → CURRENT WORKBOOK → Regional Settings → Locale: Vietnamese (Vietnam) → OK.",
          "Muốn tái hiện lỗi của máy tiếng Anh trên máy tiếng Việt (để kiểm thử): đặt tạm English (United States)."],
     shot=("Hộp Query Options, mục Current Workbook → Regional Settings, ô Locale = Vietnamese (Vietnam)", "Data → Get Data → Query Options (trên B4_<HoTen>.xlsx)", "H02_query_options_locale.png"),
     tip="Using Locale trên từng cột vẫn là cách chắc chắn nhất: bước ghi rõ locale, không phụ thuộc thiết lập của workbook hay của máy.")
tool(d, "Data Type Detection", "Hộp From Text/CSV có ô Data Type Detection: Based on first 200 rows (mặc định) · Based on entire dataset · Do not detect data types.",
     tip="Chọn Do not detect data types thì Power Query không tạo bước Changed Type tự động; bạn tự đổi kiểu bằng Using Locale ngay từ đầu.")

# 3 Đối soát
d.add_heading("3. Đối soát ngay trong Power Query", level=2)
para(d, "Bảng xem trước chỉ ghi \"999+ rows\". Muốn biết con số chính xác trước khi nạp, thêm tạm một bước tính rồi xóa bước đó.")
table(d, [["Cần kiểm tra", "Thao tác", "Kết quả trên case"],
          ["Số dòng", "Transform → Count Rows (xem xong bấm X xóa bước Counted Rows)", f"Sales_2025-01.csv sau Remove Duplicates: {vn(F1['rows'])}"],
          ["Tổng một cột", "Chọn cột → Transform → Statistics → Sum (xem xong xóa bước)", f"Revenue của Sales (2024 + Quý 1/2025): {vn(M['q_rev'])}"],
          ["Số giá trị khác nhau", "View → Column profile (entire data set) → Column statistics: Distinct", f"Store_Code tháng 1 trước Trim: {F1['distinct_store_raw']} (đúng ra 12)"],
          ["Theo nhóm", "Reference → Group By (Advanced): Count Rows, Sum", f"12 cửa hàng tháng 1; cao nhất {top_store}: {vn(JS[top_store]['rows'])} dòng, {vn(JS[top_store]['rev'])}"],
          ["Mã không có trong danh mục", "Merge Queries as New, Join Kind = Left Anti", f"Sales × Product: {M['left_anti_rows']} dòng"],
          ["Số dòng sau Merge", "Count Rows trên query đã Merge", f"{vn(M['q_rows'])} (nếu Product còn P015 trùng: {vn(M['merge_trap_rows'])})"]],
      [3.4, 7.0, 5.6], first_bold=True)
figure(d, "Tab Transform: nhóm Table (Count Rows) và nhóm Number Column (Statistics ▾ → Sum)", "Query Editor, query Sales (B4_<HoTen>.xlsx)", "H03_count_rows_statistics.png")
table(d, [["Tháng", "Số dòng", "Doanh thu", "Tháng", "Số dòng", "Doanh thu"]] +
      [[f"{a['thang']:02d}/{a['nam']}", vn(a["rows"]), vn(a["rev"]), f"{b['thang']:02d}/{b['nam']}", vn(b["rows"]), vn(b["rev"])]
       for a, b in zip(M["by_month"][:8], M["by_month"][8:])], [2.2, 2.2, 3.6, 2.2, 2.2, 3.6])
para(d, f"Bảng trên là Pivot PV_Refresh (Nam × Thang) sau khi thả file tháng 4: tổng {vn(M['all_rows'])} dòng, {vn(M['all_rev'])}. "
        "Dùng nó để đối soát mỗi khi có tháng mới.", italic=True, color=GREY)

# 4 Index, number tools
d.add_heading("4. Index Column và công cụ cho cột số", level=2)
tool(d, "Add Column → Index Column", "Thêm cột số thứ tự 0, 1, 2… (From 0) hoặc 1, 2, 3… (From 1), hoặc tùy chọn điểm bắt đầu và bước nhảy (Custom).",
     how=["Giữ thứ tự gốc của file trước khi Sort / Merge (Merge có thể làm đảo thứ tự dòng).",
          "Truy ngược về file nguồn: trên Sales_T1, Index From 2 (dòng 1 của CSV là tên cột) → số trong cột Index chính là số dòng trong file CSV.",
          "Tạo khóa dòng khi dữ liệu không có khóa duy nhất."],
     shot=("Hộp Add Index Column (Custom: Starting Index 2, Increment 1) và cột Index mới trên Sales_T1", "Query Sales_T1 (Practice/Data_Sales/Sales_2025-01.csv)", "H04_index_column.png"),
     tip="Thêm Index TRƯỚC khi Remove Duplicates thì số dòng vẫn khớp file gốc, tiện ghi lại vị trí dòng lỗi.")
tool(d, "Transform → Standard / Rounding / Information / Statistics", "Các phép tính trên cột số. Ở tab Transform thì sửa cột đang có; ở tab Add Column thì tạo cột mới, giữ cột gốc.",
     example=([["Nhóm lệnh", "Ví dụ", "Kết quả"],
               ["Standard → Divide", f"Doanh_thu của DoiSoat_T1 ÷ 1.000.000", f"{top_store}: {vn(JS[top_store]['rev'])} → " + f"{JS[top_store]['rev'] / 1e6:.3f}".replace(".", ",")],
               ["Rounding → Round… (1 chữ số)", "Làm tròn cột triệu đồng", f"{JS[top_store]['rev'] / 1e6:.1f}".replace(".", ",")],
               ["Information → Is Even", "Line_No chẵn / lẻ", "TRUE / FALSE"],
               ["Statistics → Sum / Average / Count Distinct Values", "Một con số cho cả cột (đối soát)", "Xem mục 3"]], [5.2, 6.0, 4.8]),
     tip="Đừng sửa giá sai Locale bằng Standard → Multiply × 1.000: file tháng sau có giá đúng thì bước nhân làm sai ngược lại. Sửa tận gốc bằng Using Locale.")

# 5 Pivot Column, Transpose
d.add_heading("5. Pivot Column và Transpose", level=2)
tool(d, "Transform → Pivot Column", "Ngược của Unpivot: lấy các giá trị của một cột làm tên cột mới. Dùng khi cần bảng ngang để in hoặc gửi đi, không dùng làm nguồn phân tích.",
     how=["Trên một Reference của query Target (Store_Code, Thang, Target, Ngay_Dau_Thang): bỏ cột Ngay_Dau_Thang, chọn cột Thang → Transform → Pivot Column.",
          "Values Column: Target; Advanced options → Aggregate Value Function: Don't Aggregate (mỗi ô chỉ có một giá trị) → OK."],
     shot=("Hộp Pivot Column: Values Column = Target, Aggregate = Don't Aggregate; kết quả 12 dòng × 13 cột", "Reference của query Target", "H05_pivot_column.png"),
     tip="Có nhiều dòng cùng tọa độ thì chọn Sum; chọn Don't Aggregate sẽ ra lỗi ở các ô đó.")
tool(d, "Transform → Transpose", "Lật cả bảng: dòng thành cột, cột thành dòng. Không nhận biết \"cột cùng loại\" như Unpivot.",
     example=([["Bảng Target_2025 (12 dòng × 14 cột)", "Unpivot Other Columns", "Transpose"],
               ["Kết quả", "144 dòng × 4 cột, dạng tidy", "14 dòng × 12 cột, mỗi cửa hàng thành một cột — vẫn không tidy"]], [5.4, 5.3, 5.3]),
     tip="Transpose hợp với bảng bị xoay nghiêng hoàn toàn (tên trường nằm ở cột A). Sau Transpose thường phải Use First Row as Headers.")

# 6 Workbook queries
d.add_heading("6. Quản lý query trong workbook", level=2)
tool(d, "Data → Queries & Connections", "Khung bên phải workbook liệt kê mọi query, số dòng đã nạp và lỗi nếu có. Chuột phải một query để:",
     example=([["Lệnh", "Dùng khi"],
               ["Edit", "Mở Query Editor tại query đó"],
               ["Duplicate", "Chép rời mọi bước sang query mới; sửa bản sao không ảnh hưởng bản gốc (thử nghiệm)"],
               ["Reference", "Query mới đọc KẾT QUẢ của query gốc; gốc đổi thì nó đổi theo (bảng đối soát, Group By, Merge)"],
               ["Load To…", "Đổi nơi nạp: Table / Only Create Connection / Add to Data Model"],
               ["Move to Group", "Gom query vào nhóm (ví dụ nhóm Nhap cho query tập)"],
               ["Properties", "Enable background refresh · Refresh data when opening the file · Refresh this connection on Refresh All"]], [3.5, 12.5]),
     shot=("Khung Queries & Connections với menu chuột phải của query Sales", "B4_<HoTen>.xlsx · Data → Queries & Connections", "H06_queries_pane_menu.png"),
     tip="Pivot dựng trên bảng do query nạp ra mà không đổi sau Refresh All: vào Properties của query, bỏ tick Enable background refresh để query chạy xong rồi Pivot mới cập nhật.",
     protip="Duplicate hay Reference? Muốn sửa thử mà không ảnh hưởng query gốc → Duplicate. Muốn một bảng phái sinh luôn đi theo query gốc → Reference.")
tool(d, "Data → Get Data → Data Source Settings", "Xem mọi nguồn dữ liệu mà workbook đang dùng; Change Source để trỏ sang đường dẫn mới khi chuyển thư mục hoặc chuyển máy; Edit Permissions để đặt privacy level (mục 9).")

# 7 From Folder nâng cao
d.add_heading("7. From Folder nâng cao: tên file và Transform Sample File", level=2)
tool(d, "Cột Source.Name → Extract → Text Before Delimiter", "Combine giữ lại cột Source.Name (tên file). Khi thông tin nằm trong tên file (mã cửa hàng, tháng), tách nó ra thành cột.",
     how=["Chọn Source.Name → Transform → Extract → Text Before Delimiter → Delimiter \"_\" → OK. \"ST07_2025-04.csv\" → \"ST07\".",
          "Đổi tên cột thành Store_Code. Cần tháng thì Text Between Delimiters \"_\" và \".\" → \"2025-04\"."],
     shot=("Query gộp từ thư mục Cua_hang: cột Source.Name và hộp Text Before Delimiter", "Homework/B4_BTVN_Du_lieu/Cua_hang (From Folder)", "H07_source_name.png"),
     tip="Bài tập về nhà cần đúng thao tác này: các file cửa hàng không có cột Store_Code.")
tool(d, "Transform Sample File", "Query trợ giúp mà Combine tạo ra. Bước thêm ở đây chạy trên TỪNG FILE trước khi nối — dùng cho việc phải làm theo vị trí trong mỗi file.",
     how=["Ví dụ mỗi file có một dòng \"Tổng cộng\" ở cuối: mở Transform Sample File → Home → Remove Rows → Remove Bottom Rows → 1.",
          "Việc làm theo tên cột (Trim, Using Locale, Remove Duplicates) làm trên query gộp là đủ, dễ xem hơn."],
     tip="Nhóm query trợ giúp (Sample File, Parameter, Transform File, Transform Sample File) đừng xóa. Thư mục chỉ nên chứa file dữ liệu: thêm bước lọc Extension = \".csv\" ở đầu query gộp nếu có người để lẫn file khác.")

# 8 PRO TIP
d.add_heading("8. PRO TIP: tham số hóa đường dẫn thư mục, Rolling Calendar", level=2)
tool(d, "Home → Manage Parameters → New Parameter", "Đường dẫn thư mục nằm ở một chỗ duy nhất. Chuyển máy hoặc chuyển thư mục: sửa tham số, không sửa từng query.",
     how=["Manage Parameters → New: Name DataFolder, Type Text, Current Value = đường dẫn thư mục Practice của bạn, có dấu \\ ở cuối → OK.",
          "Bước Source của CẢ HAI query Sales và Sample File (nhóm Helper Queries): thay chuỗi đường dẫn bằng tham số như dưới đây. Sample File cũng đọc thư mục — sót nó thì sang máy khác vẫn lỗi.",
          "Danh sách tham số có thêm Parameter1 (kiểu Binary) do Combine tạo: để nguyên.",
          "File đáp án B4 (file khởi đầu Buổi 5) dùng đúng tham số này: mở trên máy mới thì Manage Parameters → sửa DataFolder → Refresh All."],
     m='Source = Folder.Files(DataFolder & "Data_Sales")',
     shot=("Hộp Manage Parameters: tham số DataFolder (Type Text, Current Value) và Parameter1 do Combine tạo", "Practice/B4_MCI_Mart_Answer.xlsx · Launch Power Query Editor → Manage Parameters", "H08_manage_parameters.png"))
tool(d, "Rolling Calendar", "Calendar tự kéo dài theo thời gian: điểm cuối là 31/12 của năm hiện tại, lấy từ đồng hồ máy. Không phải sửa tay mỗi năm.",
     m=('let\n    StartDate = #date(2024, 1, 1),\n    EndDate = Date.EndOfYear(Date.From(DateTime.LocalNow())),\n'
        '    Dates = List.Dates(StartDate, Duration.Days(EndDate - StartDate) + 1, #duration(1, 0, 0, 0)),\n'
        '    ToTable = Table.FromList(Dates, Splitter.SplitByNothing(), {"Date"}),\n    Typed = Table.TransformColumnTypes(ToTable, {{"Date", type date}})\nin\n    Typed'),
     tip=f"Trên lớp dùng số cố định {M['calendar_rows']} ngày (01/01/2024–31/12/2025) để cả lớp cùng một con số. Calendar phải phủ trọn các năm của mọi bảng fact.")

# 9 Sự cố
d.add_heading("9. Xử lý sự cố: Formula.Firewall, privacy levels, lỗi hay gặp", level=2)
para(d, "Power Query có một \"tường lửa\" chặn việc gửi dữ liệu của nguồn này sang nguồn khác khi chưa biết mức riêng tư (privacy level) của chúng. "
        "Lỗi hay xuất hiện ở capstone, khi kết hợp file nội bộ với nguồn web, hoặc dựng đường dẫn file từ một ô trong workbook.")
table(d, [["Thông báo", "Nguyên nhân", "Xử lý"],
          ["Formula.Firewall: Query 'X' (step 'Y') references other queries or steps, so it may not directly access a data source. Please rebuild this data combination.",
           "Một bước vừa đọc nguồn dữ liệu vừa dùng kết quả của query khác (ví dụ Folder.Files lấy đường dẫn từ query đọc ô trong sheet)",
           "Tách: mỗi query chỉ đọc một nguồn; query kết hợp chỉ tham chiếu các query đó. Hoặc dùng Parameter (mục 8) thay cho query đọc ô"],
          ["Formula.Firewall: Query 'X' … has privacy levels that cannot be used together. Please rebuild this data combination.",
           "Hai nguồn có privacy level khác nhau (Public với Organizational / Private)",
           "Data → Get Data → Data Source Settings → chọn nguồn → Edit Permissions → Privacy Level: đặt cùng mức (thường Organizational)"],
          ["Information is required about data privacy", "Nguồn chưa được đặt privacy level", "Bấm Continue → chọn mức cho từng nguồn"]],
      [6.2, 4.8, 5.0], size=8.5)
figure(d, "Data Source Settings → Edit Permissions: ô Privacy Level", "B4_<HoTen>.xlsx · Data → Get Data → Data Source Settings", "H09_privacy_levels.png")
rich(d, [("Cách cuối cùng: ", {"bold": True, "color": NAVY}),
         ("Data → Get Data → Query Options → CURRENT WORKBOOK → Privacy → Ignore the Privacy Levels. Chỉ dùng khi mọi nguồn đều của bạn / công ty bạn; "
          "với dữ liệu khách hàng hoặc nguồn web lạ, đặt privacy level đúng thay vì tắt tường lửa.", {})])
table(d, [["Lỗi", "Dấu hiệu", "Xử lý"],
          ["Bước Changed Type tự động đoán theo locale máy", "Ngày ≥ 13 thành Error; giá 19.000 thành 19", "Xóa bước đó; Using Locale (mục 2)"],
          ["File mới thả vào thư mục (From Folder) đổi tên hoặc thêm cột", "Không báo lỗi: cột sai tên để trống ở các dòng của file đó, cột mới bị bỏ; bộ lọc/Pivot có mục (blank)", "Đối soát sau mỗi lần Refresh; nhờ nguồn giữ tên cột, hoặc Transform Sample File → Advanced Editor: Table.RenameColumns(…, {{\"StoreCode\", \"Store_Code\"}}, MissingField.Ignore)"],
          ["Tên cột của từng CSV lọt vào dữ liệu", "Dòng có Order_ID = \"Order_ID\"", "Dùng Combine & Transform; hoặc lọc bỏ dòng đó"],
          ["Khoảng trắng thừa trong mã", "Distinct nhiều hơn thực tế; Merge ra null", "Trim / Clean trước Merge"],
          ["Lookup có khóa trùng", f"Số dòng tăng sau Merge ({vn(M['q_rows'])} → {vn(M['merge_trap_rows'])})", "Remove Duplicates theo cột khóa ở lookup"],
          ["Chép file mới nhầm thư mục", "Refresh không đổi gì", "Di chuột lên query trong Queries & Connections để xem đường dẫn nguồn"],
          ["Pivot không đổi sau Refresh All", "Query đã có dữ liệu mới, Pivot vẫn số cũ", "Properties → bỏ tick Enable background refresh; Refresh All lại"],
          ["Mở file đáp án trên máy khác", "Could not find a part of the path …", "Manage Parameters → sửa DataFolder (mục 8) hoặc Data Source Settings → Change Source"]],
      [4.6, 5.4, 6.0], first_bold=True)

# 10 Homework
d.add_heading("10. Bài tập về nhà (bắt buộc)", level=2)
rich(d, [("File đề: ", {"bold": True}), (f"B4_BTVN_MCI_Mart_T4_2025.xlsx và thư mục B4_BTVN_Du_lieu (Cua_hang: {H['files_main']} file CSV tháng 4/2025, mỗi cửa hàng một file; "
                                         "Bo_sung: 1 file). Làm ngay trên file đề, nộp trước Buổi 5, tên file B4_BTVN_<HoTen>.xlsx. Sheet De_bai ghi đủ yêu cầu.", {})])
rich(d, [("Nộp bài: ", {"bold": True, "color": NAVY}), ("https://hocvien.mcivietnam.com/auth#/ — Tài khoản: email bạn đăng ký tại MCI · Mật khẩu: hv12345678", {})])
for s in ["From Folder thư mục Cua_hang → Combine & Transform; query Sales_T4. Store_Code tách từ tên file (mục 7).",
          "Làm sạch: Using Locale cho ngày và giá, bỏ dòng \"Tổng cộng\", Trim + UPPERCASE Product_Code, Remove Duplicates, cột Revenue.",
          "Kiểm tra mã sản phẩm bằng Merge Left Anti với sheet Product: ghi số dòng trước và sau khi sửa.",
          "Group By Store_Code (Reference của Sales_T4), so từng cửa hàng với sheet So_lieu_he_thong. Lệch ở đâu? Xử lý bằng thư mục Bo_sung, Refresh All đến khi khớp.",
          "Unpivot Target_2025, lọc tháng 4, Merge với bảng Group By → Pct_Dat = Doanh_thu / Target (query KetQua_T4).",
          "Load Sales_T4, KetQua_T4 ra sheet; các query khác Only Create Connection. Sheet Insight: lỗi đã sửa + 1 câu insight."]:
    bullet(d, s)
rich(d, [("Tự kiểm tra trước khi nộp: ", {"bold": True, "color": NAVY}),
         (f"Sales_T4 {vn(H['rows'])} dòng, doanh thu {vn(H['revenue'])} VNĐ — đúng bằng dòng tháng 4 của PV_Refresh trên lớp; Left Anti 0 dòng; "
          f"Group By khớp So_lieu_he_thong ở cả 12 cửa hàng; cả chuỗi đạt {pct(H['pct_all'])} target tháng 4.", {})])

# 11 Checklist
h = d.add_heading("11. Checklist trước Buổi 5", level=2); h.paragraph_format.keep_with_next = True
for s in [f"File B4 của bạn có 6 query Sales, Product, Store, Customer, Target, Calendar; Sales {vn(M['all_rows'])} dòng, {vn(M['all_rev'])} sau khi thả tháng 4.",
          f"Product {M['product']} dòng · Store {M['store']} · Customer {vn(M['customer'])} · Target {M['target_rows']} (có Thang dạng số và Ngay_Dau_Thang) · Calendar {M['calendar_rows']}.",
          "Product, Store, Customer, Target, Calendar: Only Create Connection + Data Model. Data → Manage Data Model mở cửa sổ Power Pivot thấy đủ 6 bảng.",
          "Pivot PV_Refresh (Nam × Thang) đổi số sau Refresh All mà không sửa bước nào.",
          "Đã viết câu insight cuối Buổi 4 (ST12, hoặc Target tháng 1 / tháng 2 và lịch Tết).",
          "Đã nộp bài tập về nhà (B4_BTVN_<HoTen>.xlsx) và tìm ra cửa hàng lệch ở bước đối soát.",
          "Nếu file chưa đạt: Buổi 5 lớp phát file khởi đầu (đáp án B4, B4_MCI_Mart_Answer.xlsx). Trên máy của bạn, sửa tham số DataFolder trước khi Refresh (mục 8).",
          "Power Pivot đã bật: File → Options → Add-ins → COM Add-ins → Microsoft Power Pivot for Excel (checklist Buổi 1)."]:
    bullet(d, "☐  " + s).paragraph_format.keep_with_next = True
d.paragraphs[-1].paragraph_format.keep_with_next = False
d.save(f"{OUT}/B4_Handout.docx")
print("ok;", f"còn {len(MISSING)} khung [CHỤP] chưa có ảnh:" if MISSING else "đủ ảnh.", " ".join(MISSING))
