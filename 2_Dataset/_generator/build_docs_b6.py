"""Dựng handout tự học Buổi 6 (Thiết kế và xây dựng dashboard): Handout/B6_Handout.docx bằng python-docx.

Số liệu lấy từ build_b6_meta.json (build_b6_meta.py). Ảnh chụp thật lấy từ Materials/B6/Images (H<nn>_…png);
chưa có ảnh thì để khung [CHỤP]. Xuất PDF xem 2_Dataset/README.md.
    python3 _generator/build_docs_b6.py
"""
import json, os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.image.image import Image as DocxImage

ROOT = os.path.normpath(os.environ.get("MCI_ROOT") or os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
GEN = os.path.dirname(os.path.abspath(__file__))
OUT = f"{ROOT}/1_Sessions/B6_Dashboard/Handout"
IMG = f"{ROOT}/Materials/B6/Images"   # ảnh chụp thật, danh sách: _DANH_SACH_ANH.md
MISSING = []
os.makedirs(OUT, exist_ok=True)
M = json.load(open(f"{GEN}/build_b6_meta.json", encoding="utf-8"))
C = M["card"]; T = M["targets"]; F = M["forecast"]; S = M["scr"]; HW = M["hw"]
RED, NAVY, GOLD, GREY = RGBColor(0xC0, 0, 0), RGBColor(0x20, 0x38, 0x64), RGBColor(0x8A, 0x65, 0x00), RGBColor(0x58, 0x58, 0x58)
vn = lambda x: f"{round(x):,}".replace(",", ".")
pct = lambda x: f"{x * 100:.1f}%".replace(".", ",")
pct0 = lambda x: f"{x * 100:.0f}%"
SIGN = {"ok": "✔ Đạt", "warn": "▲ Cảnh báo", "bad": "✘ Không đạt"}


def base_doc(footer="MCI – Magic Code Institute · Excel for Business Intelligence Analyst · Buổi 6"):
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for m in ("left_margin", "right_margin"): setattr(sec, m, Cm(2))
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    for name, size, color in (("Heading 1", 16, NAVY), ("Heading 2", 13, NAVY), ("Heading 3", 11.5, RED)):
        h = d.styles[name]; h.font.name = "Arial"; h.font.size = Pt(size); h.font.color.rgb = color; h.font.bold = True
        rf = h.element.rPr.rFonts
        for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(k)) is not None: del rf.attrib[qn(k)]
        for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"): rf.set(qn(k), "Arial")
    fp = sec.footer.paragraphs[0]
    fp.text = footer
    fp.runs[0].font.size = Pt(8); fp.runs[0].font.color.rgb = GREY
    return d


def para(d, text="", bold=False, color=None, size=None, italic=False):
    p = d.add_paragraph()
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic
        if color: r.font.color.rgb = color
        if size: r.font.size = Pt(size)
    return p


def rich(d, parts):
    p = d.add_paragraph()
    for t, f in parts:
        r = p.add_run(t); r.bold = f.get("bold", False); r.italic = f.get("italic", False)
        if f.get("color"): r.font.color.rgb = f["color"]
        if f.get("mono"): r.font.name = "Consolas"; r.font.size = Pt(9.5)
    return p


def bullet(d, text):
    p = d.add_paragraph(style="List Bullet"); p.add_run(text); return p


def code(d, text):
    """Khối mã: chữ Consolas trên nền xám nhạt."""
    for line in text.splitlines():
        p = d.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.left_indent = Cm(0.5)
        r = p.add_run(line); r.font.name = "Consolas"; r.font.size = Pt(9)
        pPr = p._p.get_or_add_pPr(); shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F3F5F9"); pPr.append(shd)
    d.add_paragraph().paragraph_format.space_after = Pt(0)


def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_)
    tcPr.append(shd)


def table(d, rows, widths, header=True, size=9.5, first_bold=False):
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for j, w in enumerate(widths): t.columns[j].width = Cm(w)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.width = Cm(widths[j]); c.text = ""
            r = c.paragraphs[0].add_run(str(v)); r.font.size = Pt(size)
            if i < len(rows) - 1: c.paragraphs[0].paragraph_format.keep_with_next = True   # không tách bảng qua trang
            if header and i == 0:
                r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); shade(c, "203864")
            elif first_bold and j == 0:
                r.bold = True
    d.add_paragraph()
    return t


def figure(d, what, where, img, width=15):
    """Chèn ảnh chụp thật từ Materials/B6/Images nếu đã có; chưa có thì để khung chờ [CHỤP] (CLAUDE.md, mục Ảnh minh họa)."""
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(f"{IMG}/{img}"):
        px = DocxImage.from_file(f"{IMG}/{img}")
        w = min(width, px.px_width * 2.54 / 96)
        if w * px.px_height / px.px_width > 10: w = 10 * px.px_width / px.px_height
        p.add_run().add_picture(f"{IMG}/{img}", width=Cm(w)); return p
    MISSING.append(img)
    r = p.add_run(f"[CHỤP] {what}"); r.italic = True; r.font.color.rgb = GREY; r.font.size = Pt(9)
    r = p.add_run(f"\n{where} → {img}"); r.font.color.rgb = GREY; r.font.size = Pt(8)
    pPr = p._p.get_or_add_pPr(); bdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "dashed"); e.set(qn("w:sz"), "6"); e.set(qn("w:color"), "A6A6A6"); e.set(qn("w:space"), "8"); bdr.append(e)
    pPr.append(bdr)
    return p


def tool(d, name, purpose, how=None, example=None, tip=None, protip=None, shot=None, m=None):
    """Khối công cụ: tên (đỏ) → mục đích → cách dùng → bảng ví dụ trên case → mã → ảnh → TIP / PRO TIP."""
    d.add_heading(name, level=3).paragraph_format.keep_with_next = True
    para(d, purpose).paragraph_format.keep_with_next = True
    for h in how or []: bullet(d, h)
    if example: table(d, example[0], example[1])
    if m: code(d, m)
    if shot: figure(d, *shot)
    if tip: rich(d, [("TIP: ", {"bold": True, "color": NAVY}), (tip, {})])
    if protip: rich(d, [("PRO TIP: ", {"bold": True, "color": GOLD}), (protip, {})])


# ---------------------------------------------------------------- số liệu dùng trong handout
c4 = C["4"]; c3 = C["3"]
reg = M["region_t4"]; reg_tot = sum(reg.values())
top5 = M["top5_t4"]; top5_share = M["top5_share_t4"]
cat = M["cat_ytd"]
dist = M["dist"]; sc = M["scatter"]
dow = F["dow_mean"]

# =====================================================================
d = base_doc()
para(d, "MCI – MAGIC CODE INSTITUTE", bold=True, color=GREY, size=9)
d.add_heading("Handout Buổi 6 — Thiết kế và xây dựng dashboard", level=1)
para(d, "Tài liệu tự học đi kèm Buổi 6. Phần Tóm tắt nhắc lại những gì đã làm trên lớp (brief, chart lõi, thẻ KPI, dự báo); mục 1–7 là công cụ KHÔNG giảng chi tiết "
        "trên lớp hoặc chỉ nói lướt, đọc sau buổi học; mục 8 là checklist chấm chéo, bài tập về nhà (bắt buộc) và checklist trước Buổi 7. "
        "Ví dụ viết trên file khởi đầu B5_MCI_Mart_Answer.xlsx và các file trong thư mục Demo/ của Buổi 6. Tên lệnh giữ tiếng Anh như trên Ribbon. "
        "Số trong ảnh chụp theo định dạng en-US: dấu phẩy ngăn nghìn, dấu chấm thập phân (296,965,000 trong ảnh là 296.965.000 trong bảng).",
     italic=True, color=GREY)

d.add_heading("Mục lục", level=2)
for s in ["Tóm tắt: brief, chart lõi, thẻ KPI, SCR, dự báo",
          "1. Chart ít dùng hơn: Area, Pie, Bubble, Radar, Surface & Contour, Stock, Funnel, Sunburst, Histogram và Pareto nâng cao",
          "2. Brief mẫu, 3 câu hỏi then chốt và bản in trắng",
          "3. CUBEVALUE, CUBEMEMBER và các hàm CUBE khác: cú pháp, Convert to Formulas, lỗi #N/A",
          "4. Conditional Formatting Rules Manager, Icon Sets, Theme cho dashboard",
          "5. Form Controls, dynamic range, custom number format",
          "6. Dự báo nâng cao: FORECAST.ETS.SEASONALITY, FORECAST.ETS.STAT, Data Completion, Aggregate Duplicates",
          "7. Sneak peek Power BI: khi nào nên chuyển công cụ",
          "8. Checklist chấm chéo, bài tập về nhà (bắt buộc) và checklist trước Buổi 7"]:
    bullet(d, s)

# ---------------------------------------------------------------- Tóm tắt
d.add_heading("Tóm tắt: brief, chart lõi, thẻ KPI, SCR, dự báo", level=2)
para(d, "Thứ tự thiết kế một dashboard: Audience → Câu hỏi và quyết định → KPI → Layout → Chart → Build → Story → Phát hành. "
        "Câu hỏi tốt có đủ bốn thành phần: chỉ số, chiều, thời gian, mốc so sánh. \"Xem doanh thu\" chỉ là chủ đề.")
table(d, [["Câu hỏi dữ liệu", "Chart lõi", "Thay thế"],
          ["Xu hướng theo thời gian", "Line", "Area"],
          ["So sánh giữa các nhóm", "Bar / Column", "Bullet"],
          ["Cơ cấu, tỷ trọng", "Stacked bar", "Donut (tối đa 5 phần), Treemap"],
          ["Hai thước đo khác đơn vị; Actual vs Target", "Combo 2 trục", "Column + đường target"],
          ["Phân phối, outlier", "Histogram", "Box & Whisker"],
          ["Quan hệ giữa hai biến", "Scatter (+ trendline)", "Bubble"],
          ["Địa lý", "Map (Filled Map)", "Bar theo vùng"],
          ["Chênh lệch cộng dồn (target → actual)", "Waterfall", "—"]],
      [7.0, 4.5, 4.5], first_bold=True, size=9)
para(d, "Quy tắc RAG của buổi: ✔ Đạt khi mức đạt ≥ 100%, ▲ Cảnh báo từ 90% đến dưới 100%, ✘ Không đạt dưới 90% (theo Warning 90% của KPI Dictionary). Mức đạt = giá trị ÷ target (riêng YoY: YoY ÷ 10%). "
        "Không bao giờ chỉ dùng màu để truyền nghĩa: luôn kèm biểu tượng và chữ. Ngưỡng nằm ở sheet Cai_dat để đổi một nơi.")
rows = [["Thẻ KPI (tháng 4/2025)", "Giá trị", "Target", "Mức đạt", "Trạng thái"]]
spec = [("Doanh thu", "rev", vn(c4["rev"]), vn(c4["target"])),
        ("Số đơn", "orders", vn(c4["orders"]), vn(T["orders"])),
        ("AOV", "aov", vn(c4["aov"]), vn(T["aov"])),
        ("Gross Margin %", "gm", pct(c4["gm"]), pct(T["gm"])),
        ("YoY %", "yoy", "+" + pct(c4["yoy"]), "+" + pct(T["yoy"])),
        ("Khách hoạt động", "cust", vn(c4["cust"]), vn(T["cust"]))]
for nm, k, val, tg in spec:
    rows.append([nm, val, tg, pct(c4["att"][k]), SIGN[c4["rag"][k]]])
table(d, rows, [4.2, 3.2, 3.2, 2.4, 3.0], first_bold=True, size=9)
para(d, f"Số tự đối chiếu: tháng 3/2025 doanh thu {vn(c3['rev'])} / target {vn(c3['target'])} = {pct(c3['att']['rev'])} (hụt {vn(-c3['variance'])}); "
        f"lũy kế T1–T4 {vn(S['ytd_actual'])} / {vn(S['ytd_target'])} = {pct(S['ytd_pct'])}. "
        "Insight theo SCR: Situation = con số có so sánh; Complication = vấn đề và nguyên nhân đã kiểm chứng bằng drill-down; Resolution = đề xuất cho đúng người xem trong brief. "
        "Dự báo = mẫu hình quá khứ + giả định, luôn kèm dải tin cậy.", italic=True, color=GREY)

# ---------------------------------------------------------------- 1 Chart ít dùng
d.add_heading("1. Chart ít dùng hơn", level=2)
para(d, "Tám loại lõi ở Tóm tắt trả lời phần lớn câu hỏi. Các chart dưới đây có chỗ dùng riêng nhưng dễ bị lạm dụng. "
        "Quy tắc chung: chỉ dùng khi nó trả lời một câu hỏi trong brief mà chart lõi không trả lời được. Nội dung chart theo tài liệu Excel Data Visualization (trang 7, 9, 12, 16–19).")
table(d, [["Chart", "Dùng cho", "Ví dụ trên MCI Mart", "Lưu ý"],
          ["Area", "Thay đổi cơ cấu theo thời gian",
           "Doanh thu theo ngành theo tháng T1–T4. Có sáu ngành, nên gộp Sữa (1,8%) vào nhóm Khác để còn năm",
           "Giữ dưới 6 nhóm. Dùng Data Validation và định dạng tùy chỉnh để làm nổi bật một nhóm"],
          ["Pie / Donut", "Tỷ trọng của một tổng, ít phần",
           f"Doanh thu tháng 4 theo miền: Nam {pct(reg['Nam']/reg_tot)} · Bắc {pct(reg['Bắc']/reg_tot)} · Trung {pct(reg['Trung']/reg_tot)}",
           "Tối đa 5 phần, không 3D. Sáu ngành T1–T4 (34,5% … 1,8%) thì dùng bar"],
          ["Bubble", "Scatter có thêm chiều thứ ba (kích thước)",
           "12 cửa hàng: tuổi cửa hàng (X), doanh thu (Y), số đơn (kích thước)",
           "Dùng màu làm chiều thứ tư (ví dụ miền). Kết hợp ô công thức và Form Control để làm bubble chạy theo tháng"],
          ["Radar", "So sánh từ ba biến định lượng quanh một tâm",
           "Mức đạt 6 KPI của tháng 2 và tháng 4 chồng lên nhau",
           "Chuẩn hóa mỗi biến về cùng thang (0–1, 0–100) trước khi vẽ; giữ ít nhóm để đỡ nhiễu"],
          ["Surface & Contour", "Dữ liệu ba chiều để tìm tổ hợp tối ưu",
           "MCI Mart chỉ có ngày, không có giờ, nên chưa có ví dụ tự nhiên (ví dụ chung: tai nạn theo giờ × thứ)",
           "Heat map (Conditional Formatting Color Scales) kể được cùng câu chuyện thì dùng heat map; tránh wireframe vì khó đọc"],
          ["Stock", "Giá mở, cao, thấp, đóng và khối lượng",
           "MCI Mart không có dữ liệu giá cổ phiếu; chart này dành cho dữ liệu tài chính thị trường",
           "Đặt tay trục min/max; đổi trục ngày thành trục văn bản để bỏ các khoảng trống ngày nghỉ"],
          ["Funnel", "Tiến trình qua các giai đoạn",
           "MCI Mart chưa có số lượt xem hay lượt nhấp; ví dụ chung: lượt xem → nhấp → đặt hàng của web bán hàng",
           "Hiển thị % tổng ở mỗi giai đoạn thay vì số tuyệt đối; tô màu nhấn về giai đoạn cuối"],
          ["Treemap / Sunburst", "Dữ liệu phân cấp có nhóm và nhóm con",
           "Doanh thu theo Ngành → Sản phẩm (Treemap), Miền → Thành phố → Cửa hàng (Sunburst)",
           "Treemap cho 1–2 cấp hoặc khi so kích thước; Sunburst cho nhiều cấp. Dữ liệu nguồn phải gom nhóm và sắp xếp trước"],
          ["Histogram & Pareto nâng cao", "Phân phối liên tục; Pareto thêm đường tích lũy theo thứ tự quan trọng",
           f"Giá trị đơn Q1: {vn(dist['n_orders'])} đơn, trung vị {vn(dist['median'])}, trung bình {vn(dist['mean'])}, lệch phải. Top 5 sản phẩm tháng 4 chiếm {pct(top5_share)} doanh thu",
           "Chỉnh bin size cho đến khi hình dạng ổn định; Pareto cho thấy ít mặt hàng gánh phần lớn doanh thu"]],
      [2.6, 3.4, 5.4, 4.6], first_bold=True, size=8.5)
tool(d, "Histogram: chỉnh bin và đọc đuôi", "Cùng một dữ liệu, bin khác nhau cho hình khác nhau. Bin quá rộng giấu đuôi phải, bin quá hẹp tạo nhiễu.",
     how=["Chọn cột giá trị → Insert → Insert Statistic Chart → Histogram.",
          "Chuột phải trục ngang → Format Axis → Axis Options → Bin width = 100000 (bin 100.000); thử thêm Overflow bin để gom đuôi phải vào một cột.",
          f"Đọc: trung vị {vn(dist['median'])} thấp hơn trung bình {vn(dist['mean'])} vì đuôi phải kéo trung bình lên; {vn(dist['n_outlier'])} đơn vượt ngưỡng outlier {vn(dist['fence'])} (Q3 + 1,5 × IQR)."],
     shot=("Menu Insert → All Charts liệt kê Histogram, Box & Whisker, Treemap, Sunburst, Waterfall, Funnel; bên cạnh là histogram giá trị đơn Q1 với bin 100.000",
           "Demo/B6_Demo_Chart_HocVien.xlsx · Demo_Phan_phoi · Insert → All Charts", "H01_chart_it_dung.png"),
     tip="Muốn Pareto thật sự theo sản phẩm: Pivot doanh thu theo sản phẩm, sắp giảm dần, thêm cột % lũy kế rồi dùng combo cột + đường (trục phụ 0–100%).")

# ---------------------------------------------------------------- 2 Brief
d.add_heading("2. Brief mẫu, 3 câu hỏi then chốt và bản in trắng", level=2)
para(d, "Brief là một trang, bốn mục. Viết xong brief mới chọn chart: mỗi visual trên dashboard phải trả lời đúng một câu hỏi trong brief; visual nào không trả lời câu nào thì bỏ.")
d.add_heading("Ba câu hỏi then chốt khi chọn chart", level=3)
table(d, [["#", "Câu hỏi", "Ví dụ"],
          ["1", "Dữ liệu của bạn thuộc loại nào?", "Số, danh mục, thời gian hay địa lý"],
          ["2", "Bạn muốn người xem nhận ra điều gì?", "So sánh, cơ cấu, xu hướng, phân phối hay quan hệ"],
          ["3", "Ai xem?", "Giám đốc cần ít chart và chữ to hơn analyst"]], [1.0, 7.0, 8.0], first_bold=True)
d.add_heading("Brief mẫu: dashboard cho Sales Director", level=3)
table(d, [["Mục", "Câu hỏi", "Dashboard MCI Mart"],
          ["Người xem", "Executive, operational hay self-service?", "Sales Director: executive, xem trong họp ban điều hành"],
          ["Câu hỏi (3–5, đo lường được)", "Chỉ số · chiều · thời gian · mốc so sánh",
           "Q1 Doanh thu tháng này đạt bao nhiêu % target, so với tháng trước và cùng kỳ? · Q2 Cửa hàng nào dưới 90% target? · "
           "Q3 Ngày nào trong tháng bất thường? · Q4 Năm sản phẩm đầu chiếm bao nhiêu % doanh thu? · Q5 Thành phố nào đóng góp nhiều nhất?"],
          ["Tần suất refresh", "Bao lâu một lần?", "Hằng tháng, khi có file tháng mới"],
          ["Quyết định", "Dashboard giúp quyết định gì?", "Hỗ trợ cửa hàng nào trước; điều chỉnh target"]],
      [3.6, 4.6, 7.8], first_bold=True, size=9)
para(d, "Câu hỏi yếu (chủ đề): \"Xem doanh thu\", \"Phân tích khách hàng\". Viết lại thành câu có đủ bốn thành phần: \"Cửa hàng nào có doanh thu tháng 4/2025 dưới 90% target?\" "
        f"(chỉ số = doanh thu, chiều = cửa hàng, thời gian = tháng 4/2025, mốc = 90% target). Đáp án tháng 4: {', '.join(M['stores_below_90']['4'])}.")
d.add_heading("Bản in trắng (in ra hoặc chép vào sheet Brief)", level=3)
table(d, [["Mục", "Điền vào đây"],
          ["Người xem (executive / operational / self-service)", "\n"],
          ["Câu hỏi 1 (chỉ số · chiều · thời gian · mốc so sánh)", "\n"],
          ["Câu hỏi 2", "\n"],
          ["Câu hỏi 3", "\n"],
          ["Câu hỏi 4 (tùy chọn)", "\n"],
          ["Câu hỏi 5 (tùy chọn)", "\n"],
          ["Tần suất refresh", "\n"],
          ["Quyết định dashboard hỗ trợ", "\n"]], [7.0, 9.0], first_bold=True)
para(d, "Tự kiểm tra: mỗi câu hỏi có đủ bốn thành phần; không câu nào bắt đầu bằng \"Xem…\" hoặc \"Phân tích…\"; mỗi câu đã gán cho một visual.", italic=True, color=GREY)

# ---------------------------------------------------------------- 3 CUBE
d.add_heading("3. CUBEVALUE, CUBEMEMBER và các hàm CUBE khác", level=2)
para(d, "Các hàm CUBE lấy số từ Data Model vào đúng một ô, không qua Pivot. Thẻ KPI, vùng dữ liệu cho waterfall và map cần đúng điều đó. "
        "Nguyên tắc: mọi con số trên dashboard là một measure của mô hình, gọi qua CUBEVALUE hoặc Pivot; không viết lại công thức KPI ở tầng hiển thị. "
        "CUBEVALUE chỉ chạy trong file có Data Model (Power Pivot) và Excel 365 Windows.")
tool(d, "CUBEVALUE", "Trả về một giá trị tổng hợp (measure) từ Data Model, theo các bộ lọc truyền vào.",
     m=('=CUBEVALUE("ThisWorkbookDataModel",\n   "[Measures].[Total Revenue]",\n   "[Calendar].[Nam].&[2025]",\n   Slicer_Thang, Slicer_Region)'),
     how=["Đối số 1: tên kết nối. File có Data Model dùng \"ThisWorkbookDataModel\".",
          "Đối số 2: measure, dạng \"[Measures].[Tên measure]\".",
          "Các đối số sau: bộ lọc. Có thể là chuỗi thành viên (\"[Calendar].[Nam].&[2025]\"), kết quả của CUBEMEMBER (tham chiếu ô), hoặc chính một Slicer (tên Slicer_Thang, Slicer_Region; xem tên đúng trong Slicer → Slicer Settings → Name to use in formulas).",
          f"Bấm Slicer sang tháng 3, ô trả {vn(c3['rev'])}; trở lại tháng 4, ô trả {vn(c4['rev'])}."],
     tip="Slicer làm đối số của CUBEVALUE là cách thẻ KPI nghe theo Slicer. Report Connections chỉ áp dụng cho Pivot, không áp dụng cho ô CUBEVALUE.")
tool(d, "CUBEMEMBER", "Trả về một thành viên (một cửa hàng, một sản phẩm, một tháng…) để dùng làm bộ lọc hoặc nhãn.",
     m=('=CUBEMEMBER("ThisWorkbookDataModel",\n   "[Store].[Store_Code].&[ST07]")'),
     how=["Ô chứa CUBEMEMBER hiển thị tên thành viên (ví dụ ST07); truyền ô đó vào CUBEVALUE như một bộ lọc.",
          "Waterfall: cột A là CUBEMEMBER của từng cửa hàng, cột B là CUBEVALUE đo measure Variance cho cửa hàng ở cột A cùng Slicer tháng."],
     example=([["Cột A (CUBEMEMBER)", "Cột B (CUBEVALUE, measure Variance, tháng 3/2025)"],
               ["ST07", f"{vn(M['store_month']['ST07']['3']['var'])}  ({pct(M['store_month']['ST07']['3']['pct'])} target)"],
               ["ST12", f"{vn(M['store_month']['ST12']['3']['var'])}  ({pct(M['store_month']['ST12']['3']['pct'])} target)"],
               ["Cả chuỗi (Target → Actual)", f"{vn(-c3['variance'])} hụt: {vn(c3['target'])} → {vn(c3['rev'])}"]], [6.0, 10.0]))
tool(d, "Convert to Formulas", "Chuyển cả một Pivot từ Data Model thành lưới ô CUBEVALUE / CUBEMEMBER chỉ bằng một nút; Excel tự viết hết cú pháp.",
     how=["Chọn một ô trong Pivot (From Data Model) → PivotTable Analyze → OLAP Tools → Convert to Formulas.",
          "Pivot biến mất; thay vào đó là các ô công thức. Từ đây bố cục không còn bị Pivot co giãn khi Refresh.",
          "Không hoàn tác được sau khi lưu: làm trên bản sao hoặc giữ một bản Pivot ở sheet khác."],
     shot=("Menu PivotTable Analyze → OLAP Tools đang mở, dòng Convert to Formulas được chọn; Pivot doanh thu theo tháng phía sau",
           "Bản sao Practice/B5_MCI_Mart_Answer.xlsx · một sheet Pivot từ Data Model · PivotTable Analyze → OLAP Tools", "H02_convert_to_formulas.png"))
tool(d, "Lỗi hay gặp với CUBEVALUE", "Hàm báo lỗi thường không phải do mô hình, mà do gõ sai chữ.",
     example=([["Dấu hiệu", "Nguyên nhân", "Xử lý"],
               ["#N/A", "Gõ sai tên measure hoặc thành viên (ví dụ [Measures].[Total Revenu])", "Dùng gợi ý khi gõ, hoặc Convert to Formulas để Excel viết hộ"],
               ["#NAME?", "Tên Slicer sai, hoặc kết nối ThisWorkbookDataModel không có", "Slicer → Slicer Settings để xem tên đúng; kiểm tra file có Data Model"],
               ["Ô hiện \"#GETTING_DATA\" rồi mới ra số", "Excel đang tải từ Data Model", "Chờ vài giây; nếu mãi không ra, kiểm tra Power Pivot đã bật"],
               ["Số khác Pivot", "Bộ lọc trong CUBEVALUE khác bộ lọc của Pivot (năm, tháng, miền)", "Đối chiếu từng bộ lọc; số của hai nơi phải khớp vì cùng một measure"]], [3.6, 6.8, 5.6]),
     shot=("Ô CUBEVALUE báo #N/A vì gõ sai tên measure; thanh công thức hiện công thức, bên dưới là gợi ý autocomplete đúng",
           "Bản sao Practice/B5_MCI_Mart_Answer.xlsx · sheet mới, gõ thử trong ô B3", "H03_cube_na.png"),
     tip="Gõ \"[Measures].\" rồi dừng: Excel hiện danh sách 14 measure để chọn, không cần nhớ tên.")
d.add_heading("Bảy hàm CUBE", level=3)
table(d, [["Hàm", "Dùng để", "Ghi chú"],
          ["CUBEVALUE", "Lấy một giá trị tổng hợp", "Hàm dùng nhiều nhất của dashboard"],
          ["CUBEMEMBER", "Lấy một thành viên của chiều", "Nhãn hàng, nhãn cột, bộ lọc"],
          ["CUBESET", "Tạo một tập thành viên, có thể sắp theo measure", "Bước đầu của Top N"],
          ["CUBERANKEDMEMBER", "Lấy thành viên thứ k trong một tập", "Top N: CUBESET sắp giảm dần theo measure, rồi k = 1, 2, 3…"],
          ["CUBESETCOUNT", "Đếm số thành viên của tập", "Kiểm tra Top N có đủ dòng chưa"],
          ["CUBEMEMBERPROPERTY", "Lấy thuộc tính của thành viên", "Ít dùng với Data Model"],
          ["CUBEKPIMEMBER", "Lấy KPI định nghĩa trong mô hình", "Dùng khi mô hình có KPI (Value, Goal, Status)"]], [4.2, 6.3, 5.5], first_bold=True, size=9)
rich(d, [("PRO TIP: ", {"bold": True, "color": GOLD}),
         ("Top 5 không cần Pivot (cú pháp tham khảo, thử trên bản sao file): ô A1 =CUBESET(\"ThisWorkbookDataModel\", \"[Product].[Product_Name].Children\", \"Sản phẩm\", 2, \"[Measures].[Total Revenue]\") (đối số 2 = sắp giảm dần theo measure), "
          "rồi các ô dưới =CUBERANKEDMEMBER(\"ThisWorkbookDataModel\", $A$1, 1) … 5 và CUBEVALUE cho từng sản phẩm. "
          f"Kết quả tháng 4/2025: {top5[0]['name']} {vn(top5[0]['rev'])}; năm sản phẩm đầu cộng lại chiếm {pct(top5_share)} doanh thu tháng.", {})])

# ---------------------------------------------------------------- 4 CF + Theme
d.add_heading("4. Conditional Formatting Rules Manager, Icon Sets, Theme cho dashboard", level=2)
para(d, "Trên lớp mỗi thẻ KPI có một Icon Set (✔ ▲ ✘) và nhãn chữ. Mục này giải thích cách quản lý các quy tắc đó và cách đổi giao diện cả dashboard bằng Theme.")
tool(d, "Rules Manager: xem, sửa và đổi thứ tự quy tắc", "Một ô có thể có nhiều quy tắc. Rules Manager cho thấy tất cả, theo thứ tự ưu tiên từ trên xuống.",
     how=["Home → Conditional Formatting → Manage Rules. Ô Show formatting rules for: chọn This Worksheet để thấy mọi quy tắc của sheet.",
          "Mỗi dòng có Applies to (vùng áp dụng). Quy tắc trùng vùng thì dòng trên thắng khi xung đột; tick Stop If True để dừng không xét tiếp quy tắc dưới.",
          "Edit Rule → với Icon Set: icon thứ nhất when value is ≥ 1 (Type Number), icon thứ hai ≥ 0,9, còn lại là icon thứ ba. Mức đạt phải là số (0,957), không phải chữ \"95,7%\".",
          "Muốn đổi ngưỡng một nơi: Type = Formula và trỏ tới ô trong Cai_dat (=Cai_dat!$C$13 cho ngưỡng Đạt, =Cai_dat!$C$14 cho ngưỡng Cảnh báo). Sau này sửa ô trong Cai_dat là mọi thẻ đổi theo."],
     example=([["Mức đạt", "Icon", "Nhãn chữ"], ["≥ 100%", "✔ xanh", "Đạt"], ["từ 90% đến dưới 100%", "▲ vàng", "Cảnh báo"], ["dưới 90%", "✘ đỏ", "Không đạt"]], [5.0, 4.0, 4.0]),
     shot=("Hộp Conditional Formatting Rules Manager (This Worksheet) liệt kê quy tắc Icon Set của sáu thẻ KPI; hộp Edit Rule mở với Type = Formula",
           "Practice/B6_MCI_Mart_Dashboard_Answer.xlsx · sheet Dashboard · Home → Conditional Formatting → Manage Rules", "H04_rules_manager.png"),
     tip="Không bao giờ chỉ dùng màu để truyền nghĩa: khoảng 8% nam giới không phân biệt được đỏ và xanh. Hình dạng khác nhau (tick, chấm than, chéo) cộng nhãn chữ thì ai cũng đọc được.",
     protip="Dùng thêm Color Scales làm heat map (ví dụ mức đạt target của 12 cửa hàng × 4 tháng) khi cần nhìn cả bảng cùng lúc; vẫn phải kèm số trong ô.")
tool(d, "Theme cho dashboard", "Theme gói màu, phông chữ và hiệu ứng của cả workbook. Đổi theme là đổi giao diện mọi chart, bảng và hình dạng dùng màu theme, không cần tô lại từng cái.",
     how=["Page Layout → Themes (chọn sẵn), hoặc Colors / Fonts để đổi riêng từng phần.",
          "Customize Colors → đặt sáu màu nhấn theo màu thương hiệu hoặc theo quy ước dashboard (một miền một màu, mốc target luôn xám nét đứt).",
          "Themes → Save Current Theme → lưu file .thmx để dùng lại cho mọi dashboard cùng bộ.",
          "Chart và hình vẽ phải chọn màu từ bảng màu của theme (không chọn More Colors) thì mới đổi theo khi đổi theme."],
     shot=("Page Layout → Themes mở, hộp Customize Theme Colors với sáu màu nhấn; dashboard phía sau đổi màu theo",
           "Practice/B6_MCI_Mart_Dashboard_Answer.xlsx · Page Layout → Themes", "H05_theme.png"),
     tip="Dùng số màu ít: một màu chính, một màu nhấn cho cảnh báo, xám cho mốc so sánh. Dashboard càng nhiều màu càng khó đọc nhanh.")

# ---------------------------------------------------------------- 5 Form controls
d.add_heading("5. Form Controls, dynamic range, custom number format", level=2)
tool(d, "Form Controls", "Điều khiển (hộp chọn, nút xoay, thanh cuộn) nằm trên sheet, gắn với một ô; người xem đổi lựa chọn, công thức và chart đổi theo mà không cần Slicer hay Pivot.",
     how=["Bật tab Developer: File → Options → Customize Ribbon → tick Developer.",
          "Developer → Insert → Form Controls (không phải ActiveX) → Combo Box → kéo vẽ lên sheet.",
          "Chuột phải → Format Control → Control: Input range = vùng danh sách (ví dụ tên 12 cửa hàng), Cell link = ô nhận số thứ tự chọn (ví dụ $H$1).",
          "Dùng =INDEX(vùng_danh_sách, $H$1) để lấy tên đang chọn, rồi đưa vào công thức của chart (ví dụ SUMIFS hoặc CUBEVALUE với CUBEMEMBER nối từ tên đó).",
          "Cell link trả về số thứ tự (1, 2, 3…), không phải tên: luôn cần INDEX để đổi sang tên."],
     shot=("Hộp Format Control của Combo Box: Input range, Cell link, Drop down lines; bên cạnh là chart đổi theo cửa hàng đang chọn",
           "Bản sao Demo/B6_Demo_Chart_HocVien.xlsx · tự dựng theo các bước trên", "H06_form_control.png"),
     tip="File có Form Control lưu được bình thường ở .xlsx (không cần .xlsm vì không có macro).")
tool(d, "Dynamic range (vùng dữ liệu tự co giãn)", "Chart gắn vào vùng cố định sẽ không có dòng mới khi thêm dữ liệu. Có hai cách để chart tự mở rộng.",
     how=["Cách đơn giản nhất: chuyển dữ liệu thành Excel Table (Ctrl+T). Chart và công thức dùng tên cột của Table (tblNgay[Doanh thu]) tự mở rộng khi thêm dòng.",
          "Cách dùng Name Manager: định nghĩa tên bằng OFFSET + COUNTA, ví dụ =OFFSET(Sheet1!$B$2,0,0,COUNTA(Sheet1!$B:$B)-1,1); trong Select Data của chart nhập =TênFile.xlsx!TênVùng.",
          "Ưu tiên Table; chỉ dùng OFFSET khi dữ liệu không thể thành Table. OFFSET là hàm volatile, tính lại mỗi lần sheet đổi."],
     tip="Bảng Ngày–Doanh thu dùng cho Forecast Sheet (Demo/B6_Demo_DuBao.xlsx, tblNgay, 486 ngày) là Table nên thêm ngày mới thì Forecast Sheet chọn lại vùng là xong.")
tool(d, "Custom number format", "Định dạng hiển thị, không đổi giá trị trong ô. Dùng để rút gọn số lớn và gắn biểu tượng trạng thái ngay trong định dạng.",
     example=([["Mục đích", "Mã định dạng (viết theo en-US)", "Hiển thị với 296965000 hoặc 0,957"],
               ["Triệu đồng", '#,##0.0,,"tr"', "297.0tr (en-US) / 297,0tr (vi-VN)"],
               ["Nghìn đồng", '#,##0,"k"', "296,965k"],
               ["Phần trăm một số lẻ", "0.0%", "95.7%"],
               ["Biểu tượng theo mức đạt", '[>=1]"✔ Đạt";[>=0.9]"▲ Cảnh báo";"✘ Không đạt"', "0,957 → ▲ Cảnh báo"],
               ["Mũi tên tăng giảm", '[Color10]▲ 0.0%;[Red]▼ 0.0%;0.0%', "+0,240 → ▲ 24,0%; −0,009 → ▼ 0,9%"]], [4.0, 6.8, 5.2]),
     how=["Chọn ô → Ctrl+1 → Number → Custom → gõ mã vào ô Type.",
          "Mỗi mã chia nhiều phần bằng dấu chấm phẩy: dương;âm;không;chữ. Điều kiện trong ngoặc vuông như [>=1] cho phép tối đa hai điều kiện, phần còn lại là mặc định.",
          "Dấu phẩy sau mã số (\"#,##0,,\") chia cho 1.000 cho mỗi dấu phẩy: hai dấu chia cho một triệu.",
          "Dấu thập phân và dấu ngăn nghìn trong mã theo cài đặt vùng của máy: máy vi-VN dùng dấu chấm ngăn nghìn và dấu phẩy thập phân, nên mã trên gõ theo dấu của máy mình; Excel hiển thị đúng khi mở lại."],
     tip="Ô vẫn giữ số gốc nên vẫn tính, vẫn tô Icon Set được. Đây là lý do nên dùng định dạng thay vì dùng công thức TEXT() cho nhãn hiển thị.")

# ---------------------------------------------------------------- 6 Forecast
d.add_heading("6. Dự báo nâng cao: ETS, mùa vụ, Data Completion, Aggregate Duplicates", level=2)
para(d, "Trên lớp: Trendline tuyến tính cho đường thẳng bỏ qua nhịp cuối tuần; Forecast Sheet (ETS) học được nhịp lặp đều. "
        f"Bảng Ngày–Doanh thu 01/01/2024–30/04/2025 có {F['n_days']} ngày; Trendline tuyến tính độ dốc {vn(F['linear_all']['slope_per_day'])} mỗi ngày, R² = {format(F['linear_all']['r2'], '.3f').replace('.', ',')}" +
        f" (gần như không giải thích gì). Trung bình mỗi ngày theo thứ: Hai {vn(dow['Thứ Hai'])} … Chủ nhật {vn(dow['Chủ nhật'])}; "
        f"Chủ nhật hơn Thứ Hai {pct(F['sunday_vs_monday'] - 1)}, cuối tuần hơn ngày thường {pct(F['weekend_vs_weekday'] - 1)}.")
table(d, [["Hàm", "Cú pháp", "Dùng để"],
          ["FORECAST.ETS", "(target_date, values, timeline, [seasonality], [data_completion], [aggregation])", "Giá trị dự báo cho một ngày"],
          ["FORECAST.ETS.CONFINT", "(target_date, values, timeline, [confidence_level], [seasonality], [data_completion], [aggregation])", "Nửa độ rộng của dải tin cậy tại ngày đó (mặc định 95%)"],
          ["FORECAST.ETS.SEASONALITY", "(values, timeline, [data_completion], [aggregation])", "Độ dài chu kỳ mùa vụ Excel tự phát hiện"],
          ["FORECAST.ETS.STAT", "(values, timeline, statistic_type, [seasonality], [data_completion], [aggregation])", "Thống kê của mô hình: 1 Alpha, 2 Beta, 3 Gamma, 4 MASE, 5 SMAPE, 6 MAE, 7 RMSE, 8 Step size"]],
      [4.2, 7.8, 4.0], first_bold=True, size=8.5)
tool(d, "FORECAST.ETS.SEASONALITY: để Excel tự tìm chu kỳ", "Hàm trả chu kỳ mùa vụ (số kỳ) mà Excel phát hiện. Với dữ liệu hằng ngày, bạn kỳ vọng 7 (tuần).",
     how=["Trên file Demo/B6_Demo_DuBao.xlsx, sheet Demo_ETS: =FORECAST.ETS.SEASONALITY(tblNgay[Doanh thu], tblNgay[Ngay]).",
          "Kết quả đo trên 486 ngày dữ liệu này là 121, không phải 7. Vì ta biết nhịp của nghiệp vụ là tuần, ta đặt tay Seasonality = 7 thay vì tin con số tự phát hiện.",
          "Quy tắc: biết rõ nhịp của nghiệp vụ (tuần, tháng, quý) thì đặt tay; chỉ để Detect Automatically khi chưa biết, và luôn kiểm tra kết quả."],
     protip="Dữ liệu theo tháng với seasonality 12 cần ít nhất hai năm để mô hình thấy hai chu kỳ. Với 16 tháng, dự báo theo ngày (seasonality 7) đáng tin hơn dự báo theo tháng.")
tool(d, "FORECAST.ETS.STAT: đo độ chính xác của mô hình", "Cho các chỉ số lỗi để so sánh nhiều cấu hình (seasonality 7 hay tự động, Data Completion nào).",
     how=["=FORECAST.ETS.STAT(tblNgay[Doanh thu], tblNgay[Ngay], 6, 7) trả MAE (sai số tuyệt đối trung bình); 5 là SMAPE, 7 là RMSE.",
          "Cùng dữ liệu, cùng thước đo, đổi một tham số mỗi lần rồi so sánh: cấu hình có sai số thấp hơn thì tốt hơn trên dữ liệu đã có.",
          "Kiểm chứng ngược: dùng dữ liệu đến hết tháng 3, dự báo tháng 4 (Forecast Start = 01/04/2025), rồi so với thực tế. Sai số cả tháng nhỏ không có nghĩa từng ngày chính xác: dải tin cậy 95% rộng hơn nhiều so với biến động từng ngày."])
tool(d, "Hai tùy chọn trong hộp Forecast Sheet", "Data → Forecast Sheet → Options mở rộng ra các tùy chọn dưới đây.",
     example=([["Tùy chọn", "Làm gì", "Khi nào dùng"],
               ["Fill Missing Points Using: Interpolation / Zeros", "Chuỗi ngày bị thiếu thì điền bằng nội suy hay bằng 0", "Ngày thiếu do không bán thì Zeros; ngày thiếu do lỗi dữ liệu thì Interpolation. Excel cho phép tới 30% điểm thiếu"],
               ["Aggregate Duplicates Using: Average, Sum, Count, Median, Max, Min…", "Một ngày có nhiều dòng thì gộp thế nào", "Bảng Ngày–Doanh thu nên SUM theo ngày từ gốc; nếu để Average, mỗi ngày bị lấy trung bình sai ý nghĩa"],
               ["Include forecast statistics", "Thêm bảng thống kê (Alpha, Beta, MAE…) cạnh dự báo", "Muốn so sánh cấu hình hoặc giải thích với người xem"],
               ["Seasonality: Detect Automatically / Set Manually", "Chọn chu kỳ mùa vụ", "Đặt tay 7 cho dữ liệu hằng ngày có nhịp tuần"]], [5.2, 5.2, 5.6]),
     shot=("Hộp Create Forecast Worksheet với Options mở rộng: Forecast Start, Confidence Interval 95%, Seasonality = Set Manually 7, Fill Missing Points Using, Aggregate Duplicates Using",
           "Demo/B6_Demo_DuBao_HocVien.xlsx · tblNgay · Data → Forecast Sheet → Options", "H07_forecast_options.png"),
     tip="Chuỗi ngày phải đều, không trùng, đúng thứ tự thời gian. Muốn chắc: tạo bảng Ngày–Doanh thu bằng Group By trong Power Query rồi đối chiếu với bảng Calendar.")
rich(d, [("Checklist dự báo: ", {"bold": True, "color": NAVY}),
         ("☐ chuỗi ngày đều, không trùng · ☐ ghi rõ dữ liệu dùng (từ ngày nào đến ngày nào) · ☐ ghi rõ mùa vụ đã chọn · ☐ luôn hiển thị dải tin cậy · "
          "☐ ghi rõ sự kiện không được tính (Tết đổi ngày theo âm lịch: " + ", ".join(f"{v[8:10]}/{v[5:7]}/{k}" for k, v in M['forecast']['tet_days'].items()) + ") · ☐ nói rõ đây là mẫu hình quá khứ cộng giả định, không phải cam kết.", {})])

# ---------------------------------------------------------------- 7 Power BI
d.add_heading("7. Sneak peek Power BI: khi nào nên chuyển công cụ", level=2)
para(d, "Từ Data Model có bốn cách dựng báo cáo và hình ảnh hóa (Power Query / Power Pivot / DAX PDF trang 111–112): (1) PivotChart và Conditional Formatting, (2) dashboard dạng bảng tính dùng hàm CUBE, "
        "(3) Power View, Power Map (nay Microsoft khuyên dùng Power BI thay thế), (4) Power BI là sản phẩm độc lập, gồm bản máy tính (Power BI Desktop) và dịch vụ trên web. Hai cách đầu nằm trong Excel và là phạm vi khóa học này.")
table(d, [["Tiêu chí", "Excel (khóa học này)", "Power BI"],
          ["Cập nhật tự động theo lịch", "Không: chỉ refresh khi có người mở file (Refresh on open) hoặc bấm Refresh", "Có: đặt lịch refresh trên dịch vụ"],
          ["Số người xem cùng lúc, phân quyền", "Chia sẻ file qua OneDrive hoặc SharePoint; ai mở được file thấy toàn bộ dữ liệu trong Data Model", "Chia sẻ báo cáo; hỗ trợ phân quyền theo hàng (RLS) và theo workspace"],
          ["Tương tác giữa các visual", "Slicer, Report Connections; CUBEVALUE lọc bằng đối số", "Mọi visual lọc chéo lẫn nhau theo mặc định"],
          ["Khối lượng dữ liệu", "Data Model nén tốt nhưng giới hạn theo bộ nhớ máy và kích thước file", "Mô hình lớn hơn và nhiều nguồn hơn; DAX và relationship giống Excel"],
          ["Xem trên điện thoại, web", "Hạn chế: Excel for the web hỗ trợ Data Model và Pivot từ Data Model chưa đầy đủ; hỏi bộ phận IT trước khi hứa người xem mở được trên trình duyệt", "Có ứng dụng điện thoại và xem trên web"],
          ["Chi phí", "Có sẵn trong Microsoft 365", "Power BI Desktop miễn phí; chia sẻ trên dịch vụ cần giấy phép (xem powerbi.microsoft.com)"]],
      [3.8, 6.2, 6.0], first_bold=True, size=9)
rich(d, [("Kết luận: ", {"bold": True, "color": NAVY}),
         ("đây không phải \"Excel kém\". Mô hình dữ liệu và DAX bạn học ở Buổi 5 dùng nguyên trong Power BI; thứ cần đổi là lớp trình bày và phát hành. "
          "Chuyển sang Power BI khi có một trong ba nhu cầu: cập nhật tự động hằng ngày, nhiều người xem với quyền khác nhau, hoặc dữ liệu vượt giới hạn của file Excel.", {})])

# ---------------------------------------------------------------- 8 Checklist + HW
d.add_heading("8. Checklist chấm chéo, bài tập về nhà (bắt buộc) và checklist trước Buổi 7", level=2)
d.add_heading("Checklist chấm chéo dashboard", level=3)
para(d, "Dùng để chấm dashboard của bạn bên cạnh (chỉ chấm, không sửa) và để tự kiểm tra trước Buổi 7. Đây cũng là tiêu chí C5 của capstone.")
for s in ["Mỗi visual trả lời một câu hỏi trong brief; không có visual \"cho đẹp\".",
          "KPI quan trọng nhất ở góc trên trái; thứ tự KPI → Trend → Detail.",
          "5–8 visual trên một trang.",
          "Mọi KPI có mốc so sánh (target hoặc kỳ trước) và trạng thái RAG, kèm biểu tượng hoặc nhãn.",
          "Mọi con số lấy từ measure (CUBEVALUE hoặc Pivot); không có công thức KPI ở tầng hiển thị.",
          "Trục bar và column bắt đầu từ 0; không 3D; không pie quá 5 phần.",
          "Slicer hoặc Timeline điều khiển mọi visual (Report Connections; CUBEVALUE lọc bằng đối số).",
          "Có chú thích trên chart chỉ vào điểm bất thường.",
          "Bài kiểm tra 5 giây: người lạ thấy ngay KPI chính và biết đạt hay chưa đạt target.",
          "Nếu có dự báo: hiển thị dải tin cậy và ghi rõ giả định (dữ liệu dùng, mùa vụ, sự kiện không được tính).",
          "Phát hành: Refresh on open, ẩn sheet kỹ thuật, Protect Sheet (Slicer không bị khóa), alt text cho chart, xuất PDF thử."]:
    bullet(d, "☐  " + s)

d.add_heading("Bài tập về nhà B6 (bắt buộc)", level=3)
rich(d, [("Tình huống: ", {"bold": True}), ("Marketing Director xin một dashboard theo dõi khách hàng tháng 4/2025. Bạn nhận một dashboard cũ do đồng nghiệp làm (sheet Dashboard_Cu), "
                                              "phải kiểm tra rồi làm lại.", {})])
rich(d, [("File đề: ", {"bold": True}), ("Homework/B6_BTVN_MCI_Mart_Dashboard_Marketing.xlsx. Làm ngay trên file đề, nộp trước Buổi 7, tên file B6_BTVN_<HoTen>.xlsx. Sheet De_bai ghi đủ yêu cầu.", {})])
rich(d, [("Nộp bài: ", {"bold": True, "color": NAVY}), ("https://hocvien.mcivietnam.com/auth#/ — tài khoản như Buổi 4.", {})])
for s in ["Đối soát Dashboard_Cu với measure; ghi sheet Nhat_ky_doi_soat: lỗi · số cũ · số đúng · cách sửa (tập D2).",
          "Viết brief cho Marketing Director: người xem, 3–5 câu hỏi đo lường được, tần suất refresh, quyết định (tiêu chí C1).",
          "Dựng dashboard một trang, tối đa 8 visual: thẻ KPI bằng CUBEVALUE có target và RAG kèm biểu tượng, Slicer nối đủ, alt text, tiêu đề chart là kết luận (tiêu chí C5).",
          "Viết insight theo SCR, 3 câu, kèm hành động cho Marketing Director (tiêu chí C6)."]:
    bullet(d, s)
rich(d, [("Tự kiểm tra trước khi nộp (tháng 4/2025): ", {"bold": True, "color": NAVY}),
         (f"Active Customers {vn(HW['cust_t4_total'])} (target {vn(HW['cust_target'])} → {pct(HW['cust_t4_total'] / HW['cust_target'])}); "
          f"theo phân khúc: " + " · ".join(f"{k} {vn(v['cust'])}" for k, v in HW['cust_t4_by_segment'].items()) + f" (cộng lại {vn(HW['cust_t4_sum_segments'])}). "
          f"Cảnh giác khi cộng các miền: " + " + ".join(str(v) for v in HW['cust_t4_by_region'].values()) + f" = {vn(HW['cust_t4_sum_regions'])}, khác {vn(HW['cust_t4_total'])}: "
          "số khách duy nhất không cộng dồn được vì một khách có thể mua ở nhiều miền.", {})])

d.add_heading("Checklist trước Buổi 7", level=3).paragraph_format.keep_with_next = True
for s in ["Dashboard một trang trong B6_<HoTen>.xlsx: sáu thẻ KPI (CUBEVALUE, target, RAG có biểu tượng), Slicer nối đủ, 5–8 visual, đã ẩn sheet kỹ thuật.",
          f"Số tháng 4/2025 khớp: doanh thu {vn(c4['rev'])} ({pct(c4['att']['rev'])} target), {vn(c4['orders'])} đơn, AOV {vn(c4['aov'])}, GM {pct(c4['gm'])}, {vn(c4['cust'])} khách.",
          "Đã viết brief (sheet Brief) và insight SCR 3 câu cho tháng bạn chọn.",
          "Đã đọc mục 1–7 của handout; thử ít nhất CUBEVALUE với Slicer và Rules Manager trên bản sao file.",
          "Đã làm và nộp bài tập về nhà B6 (B6_BTVN_<HoTen>.xlsx).",
          "Đã nhận brief, dataset capstone và rubric bản học viên (Rubric_B7_Capstone.md); đọc kỹ tiêu chí C1–C6.",
          "Nếu file chưa đạt: Buổi 7 lớp phát file đáp án B6 (Practice/B6_MCI_Mart_Dashboard_Answer.xlsx) để vẫn làm tiếp được."]:
    bullet(d, "☐  " + s).paragraph_format.keep_with_next = True
d.paragraphs[-1].paragraph_format.keep_with_next = False
d.save(f"{OUT}/B6_Handout.docx")
print("ok;", f"còn {len(MISSING)} khung [CHỤP] chưa có ảnh:" if MISSING else "đủ ảnh.", " ".join(MISSING))
