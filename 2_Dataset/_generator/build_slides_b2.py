"""Dựng deck Buổi 2: Slide_B2.pptx từ Slide_Script_B2.md, theo style của Slide_B1.pptx.

- Master, layout, theme và slide bìa lấy từ Slide_B1.pptx (không sửa file B1; chỉ đọc làm khuôn).
- Chữ trên slide khai báo trong script này, khớp phần "Trên slide" của kịch bản.
- Speaker notes lấy NGUYÊN VĂN từ kịch bản (Lời giảng, Ghi chú giảng viên), bỏ định dạng markdown,
  nên sửa lời giảng thì sửa Slide_Script_B2.md rồi chạy lại script này.
- Quiz: shape đáp án tên "Answer …", lớp mờ "Dim …", hiện khi bấm (build_student_slides.py sẽ xóa ở bản _HocVien).

    python3 _generator/build_slides_b2.py     # cần python-pptx
    python3 _generator/build_student_slides.py B2
"""
import copy, os, re
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

ROOT = os.path.normpath(os.environ.get("MCI_ROOT") or os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
TEMPLATE = f"{ROOT}/1_Sessions/B1_Nen_tang_Excel_Du_lieu/Slide_B1.pptx"
B2 = f"{ROOT}/1_Sessions/B2_Ham_va_KPI"
SCRIPT = f"{B2}/Slide_Script_B2.md"
OUT = f"{B2}/Slide_B2.pptx"
IMG = f"{ROOT}/Materials/B2/Images"        # ảnh chụp thật: S<số slide>_<chủ đề>.png (danh sách: _DANH_SACH_ANH.md)
MISSING = []

NAVY, INK, INK2, MUTED = "203864", "1B2230", "585858", "6B6F76"
RED, GOLD, ALT, HAIR, WHITE = "C00000", "FFC000", "F3F5F9", "D9DEE8", "FFFFFF"
GOOD, GOOD_BG, BAD_TXT, BAD_BG, ACCENT = "2E7D32", "E8F5E9", "9C0006", "FFC7CE", "3398DA"
L, R, W = 1.67, 12.74, 11.07          # lề trái, lề phải, bề rộng vùng nội dung (in)

# ---------------------------------------------------------------------
# Khuôn: mở B1, giữ slide bìa, xóa các slide còn lại
# ---------------------------------------------------------------------
prs = Presentation(TEMPLATE)
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst)[1:]:
    prs.part.drop_rel(sldId.rId); sldIdLst.remove(sldId)
LAYOUT = {l.name: l for l in prs.slide_layouts}

def rgb(h): return RGBColor.from_string(h)

def _runs(p, parts, size, color, bold):
    for part in parts:
        if isinstance(part, str): part = (part, {})
        text, f = part
        r = p.add_run(); r.text = text
        r.font.size = Pt(f.get("size", size)); r.font.bold = f.get("b", bold)
        r.font.color.rgb = rgb(f.get("color", color)); r.font.name = "Arial"
        if f.get("i"): r.font.italic = True

def tb(s, x, y, w, h, paras, size=18, color=INK, bold=False, anchor="t", align=None, name="TextBox", bullet=False, space=6):
    """paras: str | list[para]; para: str | list[run]; run: str | (text, {b, color, size, i})."""
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); sh.name = name
    tf = sh.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "ctr": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str): paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align: p.alignment = {"l": PP_ALIGN.LEFT, "ctr": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        if i: p.space_before = Pt(space)
        if isinstance(para, str): para = [para]
        if bullet: para = [("•  ", {"color": NAVY, "b": True})] + list(para)
        _runs(p, para, size, color, bold)
    return sh

def box(s, x, y, w, h, fill=None, line=None, lw=1, shape=MSO_SHAPE.RECTANGLE, radius=None, name="Box", alpha=None):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h)); sh.name = name
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0] = radius
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    else: sh.fill.background()
    if line: sh.line.color.rgb = rgb(line); sh.line.width = Pt(lw)
    else: sh.line.fill.background()
    sh.shadow.inherit = False
    if alpha is not None:
        clr = sh.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        a = etree.SubElement(clr, qn("a:alpha")); a.set("val", str(alpha))
    return sh

def chip(s, x, y, w, text, fill=None, line=NAVY, color=NAVY, size=14, name="Chip", h=0.4):
    sh = box(s, x, y, w, h, fill=fill, line=line, lw=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2, name=name)
    tf = sh.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; _runs(p, [text], size, color, True)
    return sh

def vline(s, x, y, h):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x), Inches(y + h))
    ln.line.color.rgb = rgb(HAIR); ln.line.width = Pt(1); return ln

def content(title, rail):
    s = prs.slides.add_slide(LAYOUT["Title Only"])
    s.shapes.title.text = title
    tb(s, 0, 1.5, 1.07, 5.25, rail, size=14, color=WHITE, bold=True, anchor="ctr", align="ctr", name="Rail label")
    s.shapes[-1].text_frame._bodyPr.set("vert", "vert270")
    return s

def section(title, body=None):
    s = prs.slides.add_slide(LAYOUT["Section Header"])
    s.shapes.title.text = title
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx == 1:
            if body: ph.text = body
            else: ph._element.getparent().remove(ph._element)
    return s

def tip(s, x, y, w, h, text, pro=False):
    box(s, x, y, w, h, fill=ALT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, name=("PRO TIP" if pro else "TIP") + " callout")
    if pro:
        chip(s, x + 0.12, y + 0.12, 1.1, "PRO TIP", fill=GOLD, line=None, color=INK, size=12, name="PRO TIP chip", h=0.35)
        tb(s, x + 1.32, y + 0.1, w - 1.44, h - 0.2, text, size=16, name="PRO TIP text")
    else:
        tb(s, x + 0.12, y + 0.12, 0.6, 0.35, "TIP", size=14, color=NAVY, bold=True, anchor="ctr")
        tb(s, x + 0.75, y + 0.1, w - 0.87, h - 0.2, text, size=16, name="TIP text")

def tool(s, x, y, w, name, purpose, h_purpose=0.7, size=24):
    tb(s, x, y, w, 0.5, name, size=size, color=RED, bold=True, name="Tool name")
    tb(s, x, y + 0.5, w, h_purpose, purpose, size=16, color=INK2, name="Tool purpose")

def table(s, x, y, colw, rows, size=16, head_size=None, rowh=0.45, bold_first=False, name="Table",
          header=True, fills=None, colors=None, align=None):
    """rows[0] là tiêu đề (nền navy). fills/colors: {(r,c): hex}. Chữ trong ô có thể là list run."""
    nr, nc = len(rows), len(rows[0])
    gf = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(sum(colw)), Inches(rowh * nr)); gf.name = name
    tbl = gf.table
    tbl._tbl.tblPr.set("firstRow", "1"); tbl._tbl.tblPr.set("bandRow", "0")
    style = tbl._tbl.tblPr.find(qn("a:tableStyleId"))
    if style is not None: tbl._tbl.tblPr.remove(style)
    for j, cw in enumerate(colw): tbl.columns[j].width = Inches(cw)
    for i in range(nr):
        tbl.rows[i].height = Inches(rowh)
        for j in range(nc):
            c = tbl.cell(i, j); v = rows[i][j]
            c.margin_left = c.margin_right = Inches(0.08); c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]
            if align and align[j]: p.alignment = {"l": PP_ALIGN.LEFT, "ctr": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align[j]]
            head = header and i == 0
            col = WHITE if head else (colors or {}).get((i, j), INK)
            parts = v if isinstance(v, list) else [str(v)]
            _runs(p, parts, head_size or size, col, head or (bold_first and j == 0))
            tcPr = c._tc.get_or_add_tcPr()
            for side in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                ln = etree.SubElement(tcPr, qn(side))
                if side in ("a:lnL", "a:lnR"):
                    ln.set("w", "0"); etree.SubElement(ln, qn("a:noFill"))
                else:
                    ln.set("w", "12700"); sf = etree.SubElement(ln, qn("a:solidFill"))
                    etree.SubElement(sf, qn("a:srgbClr")).set("val", HAIR)
            f = NAVY if head else (fills or {}).get((i, j))
            if f:
                sf = etree.SubElement(tcPr, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr")).set("val", f)
            else:
                etree.SubElement(tcPr, qn("a:noFill"))
    return gf

def fx(s, x, y, w, formula, size=16, name="Formula bar"):
    """Thanh công thức mô phỏng Excel: ô 'fx' + công thức."""
    box(s, x, y, w, 0.45, fill=WHITE, line=HAIR, name=name)
    tb(s, x + 0.1, y, 0.45, 0.45, [[("fx", {"i": True})]], size=14, color=MUTED, anchor="ctr")
    tb(s, x + 0.6, y, w - 0.7, 0.45, formula, size=size, color=INK, anchor="ctr", name="Formula")

def stat(s, x, y, w, h, label, value, sub=None, name="Stat card", vcolor=NAVY, vsize=32):
    box(s, x, y, w, h, fill=WHITE, line=HAIR, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, name=name)
    paras = [[(label, {"size": 14, "color": INK2})], [(value, {"size": vsize, "b": True, "color": vcolor})]]
    if sub: paras.append([(sub, {"size": 14, "color": INK})])
    tb(s, x + 0.2, y + 0.1, w - 0.4, h - 0.2, paras, anchor="ctr", space=2)

def shot(s, x, y, w, h, what, where, img, name="Image placeholder [CHỤP]"):
    """Ảnh chụp thật (Excel 365 Windows, zoom 100%) — xem CLAUDE.md, mục Ảnh minh họa.
    Có file Materials/B2/Images/<img> thì chèn ảnh vào khung (giữ tỉ lệ, canh giữa, viền hairline);
    chưa có thì để khung chờ [CHỤP] ghi nội dung, file · sheet · vùng và tên file ảnh cần lưu."""
    path = f"{IMG}/{img}"
    if os.path.exists(path):
        pic = s.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))
        if pic.height > Inches(h):
            ratio = Inches(h) / pic.height; pic.height = Inches(h); pic.width = int(pic.width * ratio)
        pic.left = Inches(x) + (Inches(w) - pic.width) // 2; pic.top = Inches(y) + (Inches(h) - pic.height) // 2
        pic.line.color.rgb = rgb(HAIR); pic.line.width = Pt(0.75); pic.name = f"Screenshot {img[:-4]}"
        return pic
    sh = box(s, x, y, w, h, fill=ALT, line="A6A6A6", lw=1.25, name=name)
    sh.line.dash_style = MSO_LINE.DASH
    tb(s, x + 0.2, y + 0.1, w - 0.4, h - 0.2, [[("[CHỤP] ", {"b": True, "color": NAVY}), (what, {"color": INK2})],
                                               [(f"{where} → {img}", {"color": MUTED, "size": 12})]], size=14, anchor="ctr", align="ctr", space=4)
    MISSING.append(img)
    return sh

def footnote(s, text):
    tb(s, L, 6.8, W, 0.35, text, size=12, color=MUTED, anchor="b", name="Footnote")

# ---------------------------------------------------------------------
# Speaker notes từ kịch bản
# ---------------------------------------------------------------------
md = open(SCRIPT, encoding="utf-8").read()
blocks = re.split(r"^### Slide (\d+) · ", md, flags=re.M)
NOTES = {}

def clean(t):
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t); t = re.sub(r"(?<!\w)\*(.+?)\*(?!\w)", r"\1", t)
    return t.replace("`", "").strip()

def fmt_min(m):
    return (f"{m:g}").replace(".", ",")

for k in range(1, len(blocks), 2):
    n = int(blocks[k]); body = blocks[k + 1]
    head, rest = body.split("\n", 1)
    body = re.split(r"^(?:---|## )", rest, flags=re.M)[0]
    m = re.search(r"· ([\d.]+)'\s*$", head) or re.search(r"\((\d+) phút\)", head)
    out = []
    ans = re.search(r"\*\*([A-D])\)", body)
    if head.startswith("QUESTION"):
        out.append(f"Khoảng 40 giây. Đáp án: {ans.group(1)}.")
    elif m:
        out.append(f"Thời lượng: {fmt_min(float(m.group(1)))} phút.")
    for sec in re.finditer(r"^\*\*(Lời giảng[^*]*|Ghi chú giảng viên[^*]*):\*\*\n(.*?)(?=^\*\*(?:Lời giảng|Ghi chú giảng viên)|\Z)", body, flags=re.M | re.S):
        label, txt = sec.group(1), sec.group(2)
        head_, _, tail_ = label.partition(" (")
        out.append(head_.upper() + (f" ({tail_}" if tail_ else ""))
        for line in txt.splitlines():
            line = line.rstrip()
            if line.startswith(">"): line = line[1:].strip()
            elif line.lstrip().startswith("- "): line = line.lstrip()[2:]
            else: line = line.strip()
            if line: out.append(clean(line))
    NOTES[n] = "\n".join(out)
assert len(NOTES) == 41, sorted(NOTES)

def notes(s, n):
    s.notes_slide.notes_text_frame.text = NOTES[n]

# ---------------------------------------------------------------------
# Quiz: 4 phương án, đáp án hiện khi bấm
# ---------------------------------------------------------------------
def quiz(n, topic, question, options, answer, rail="PHẦN 7 · QUIZ & TỔNG KẾT", qsize=22):
    s = content(f"Question {n - 32}: {topic}", rail)
    tb(s, L, 1.65, W, 1.0, question, size=qsize, color=INK)
    pos = {"A": (L, 2.85), "B": (7.33, 2.85), "C": (L, 4.35), "D": (7.33, 4.35)}
    ow, oh = 5.41, 1.3
    for letter, (x, y) in pos.items():
        box(s, x, y, ow, oh, fill=WHITE, line=HAIR, name=f"Option {letter}")
        if letter == answer:
            box(s, x, y, ow, oh, fill=GOOD_BG, line=GOOD, lw=3, name=f"Answer highlight {letter}")
        lt = box(s, x + 0.2, y + (oh - 0.6) / 2, 0.6, 0.6, fill=NAVY, shape=MSO_SHAPE.OVAL, name=f"Letter {letter}")
        _center_text(lt, letter)
        if letter == answer:
            al = box(s, x + 0.2, y + (oh - 0.6) / 2, 0.6, 0.6, fill=GOOD, shape=MSO_SHAPE.OVAL, name=f"Answer letter {letter}")
            _center_text(al, letter)
        tb(s, x + 1.0, y, ow - 1.1, oh, options[letter], size=18, anchor="ctr")
    for letter, (x, y) in pos.items():
        if letter != answer: box(s, x, y, ow, oh, fill=WHITE, name=f"Dim {letter}", alpha=60000)
    ax, ay = pos[answer]
    chip(s, ax + ow - 1.6, ay - 0.2, 1.45, "✓ Đáp án", fill=GOOD, line=None, color=WHITE, size=14, name="Answer chip", h=0.32)
    _reveal(s, [sh for sh in s.shapes if sh.name.startswith(("Answer", "Dim"))])
    notes(s, n)

def _center_text(sh, text):
    tf = sh.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; _runs(p, [text], 20, WHITE, True)

def _reveal(s, shapes):
    """Một lần bấm: mọi shape đáp án/làm mờ hiện dần (fade), giống deck B1."""
    ids = [sh.shape_id for sh in shapes]
    cid = [3]
    def nid(): cid[0] += 1; return cid[0]
    pars = []
    for k, spid in enumerate(ids):
        a, b = nid(), nid(); c = nid()
        node = "clickEffect" if k == 0 else "withEffect"
        pars.append(f'<p:par><p:cTn id="{a}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="{node}">'
                    f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:set><p:cBhvr><p:cTn id="{b}" dur="1" fill="hold">'
                    f'<p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
                    f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
                    f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{c}" dur="400"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
                    f'</p:childTnLst></p:cTn></p:par>')
    bld = "".join(f'<p:bldP spid="{i}" grpId="0" animBg="1"/>' for i in ids)
    xml = ('<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par>'
           '<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek">'
           '<p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst><p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>'
           f'<p:childTnLst><p:par><p:cTn id="{nid() + 900}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{"".join(pars)}</p:childTnLst></p:cTn></p:par>'
           '</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>'
           f'<p:bldLst>{bld}</p:bldLst></p:timing>')
    s._element.append(etree.fromstring(xml))

# =====================================================================
# SLIDES
# =====================================================================
# 1 · Bìa (giữ slide bìa của B1, đổi chữ)
s1 = prs.slides[0]
for sh in s1.shapes:
    if sh.name.startswith("subTitle"):
        sh.text_frame.paragraphs[0].runs[0].text = "Buổi 2: Hàm cho phân tích — từ dữ liệu giao dịch đến KPI"
        for r in sh.text_frame.paragraphs[0].runs[1:]: r.text = ""
    elif sh.has_text_frame and sh.text_frame.text.startswith("Yêu cầu máy"):
        p = sh.text_frame.paragraphs[0]
        p.runs[0].text = "Mở sẵn: file khởi đầu B1_MCI_Mart_Q1_2025_Answer.xlsx · B2_Product_DuplicateKey.xlsx"
        for r in p.runs[1:]: r.text = ""
s1.notes_slide.notes_text_frame.text = NOTES[1]

# 2 · Câu hỏi mở buổi
s = content("Câu hỏi mở buổi", "PHẦN 0 · MỞ ĐẦU")
chip(s, 10.23, 0.6, 2.5, "Tầng 1 — Descriptive", name="Tier chip")
tb(s, L, 1.8, 5.4, 2.2, [[("“Doanh thu, số đơn, biên lợi nhuận của từng vùng tháng này là bao nhiêu?”", {"b": True, "color": NAVY})]], size=28)
tb(s, L, 4.2, 5.4, 0.4, "— Giám đốc bán hàng, MCI Mart", size=16, color=INK2)
for k, t in enumerate(["Doanh thu", "Số đơn", "Biên lợi nhuận"]):
    chip(s, L + k * 1.8, 4.9, 1.65, t, name="Keyword")
tb(s, L, 5.5, 5.4, 0.8, "… của từng vùng … tháng này: Dimension nào? Measure nào? Đếm cái gì?", size=16, color=INK2)
table(s, 7.4, 1.9, [1.5, 1.3, 1.3, 1.25], [["Miền", "Doanh thu", "Số đơn", "Biên LN"]] + [[r, "?", "?", "?"] for r in ("Bắc", "Trung", "Nam")],
      size=20, head_size=16, rowh=0.7, bold_first=True, align=[None, "ctr", "ctr", "ctr"], name="Empty KPI table")
tb(s, 7.4, 5.0, 5.35, 0.8, "Bảng KPI cần điền sau buổi học", size=12, color=MUTED)
notes(s, 2)

# 3 · Recap B1
s = content("Recap Buổi 1 & file khởi đầu", "PHẦN 0 · MỞ ĐẦU")
tb(s, L, 1.75, 6.9, 3.6, [
    [("Viên gạch B1: ", {"b": True, "color": NAVY}), "tblSales · 6.187 dòng · 941.454.000 · không dòng tổng, không ô gộp."],
    [("Hôm nay nhận file khởi đầu ", {"b": True, "color": NAVY}), "= đáp án B1 (B1_MCI_Mart_Q1_2025_Answer.xlsx). Mọi người làm trên cùng một nền; lưu thành B2_<HoTen>.xlsx."],
    [("Hai thứ dùng lại hôm nay: ", {"b": True, "color": NAVY}), "dấu $ (F4) và công thức theo tên cột [@Cột] trong Table."],
], size=20, bullet=True, space=14)
box(s, 9.07, 4.6, 3.67, 1.0, fill=NAVY, name="Brick B1"); _t = s.shapes[-1]
tb(s, 9.17, 4.6, 3.47, 1.0, "✓  B1 · Dữ liệu tidy", size=18, color=WHITE, bold=True, anchor="ctr")
box(s, 9.07, 3.3, 3.67, 1.0, fill=WHITE, line="2A78D6", lw=2.25, name="Brick B2")
tb(s, 9.17, 3.3, 3.47, 1.0, "B2 · Bảng KPI (SUMIFS)", size=18, color=NAVY, bold=True, anchor="ctr")
footnote(s, "Demo: ../B1_Nen_tang_Excel_Du_lieu/Practice/B1_MCI_Mart_Q1_2025_Answer.xlsx, sheet Kiem_tra (8/8 ✔)")
notes(s, 3)

# 4 · Dimension / Measure / Grain
s = content("Dimension · Measure · Grain", "PHẦN 0 · MỞ ĐẦU")
table(s, L, 1.7, [2.0, 2.4, 6.67], [
    ["Vai trò", "Câu hỏi", "Cột trong tblSales"],
    [[("Dimension", {"b": True, "color": "2A78D6"})], "“Theo cái gì?”", "Region, Store_Code, Order_Date (→ Tháng), Product_Code (→ Category)"],
    [[("Measure", {"b": True, "color": "EB6834"})], "“Đo cái gì?”", "Quantity, Revenue (→ COGS, Gross Margin)"],
    [[("Grain", {"b": True, "color": NAVY})], "“Mỗi dòng là gì?”", [("1 dòng = 1 sản phẩm trong 1 đơn", {"b": True}), " (Order_ID × Line_No)"]],
], size=18, rowh=0.62)
shot(s, L, 4.45, 6.2, 1.9, "Đơn HD2501-00002 có 4 dòng (Line_No 1–4), chọn dòng 3–6", "../B1_Nen_tang_Excel_Du_lieu/Practice/B1_MCI_Mart_Q1_2025_Answer.xlsx · sheet Sales · A1:K7 (ẩn tạm cột D–F, I–J)", "S04_grain.png")
stat(s, 8.2, 4.45, 4.54, 1.75, "Đếm dòng ≠ đếm đơn", "6.187 ≠ 3.144", "dòng trong tblSales · đơn hàng Q1/2025", vsize=30)
notes(s, 4)

# 5 · KPI Definition Card
s = content("KPI Definition Card", "PHẦN 0 · MỞ ĐẦU")
table(s, L, 1.65, [2.3, 5.0], [
    ["Trường", "Card 1 · Total Revenue"],
    ["Mục đích", "Tháng này bán được bao nhiêu tiền, ở đâu?"],
    ["Công thức", "SUM(Quantity × Unit_Price)"],
    ["Nguồn dữ liệu", "tblSales"],
    ["Grain", "Dữ liệu: 1 sản phẩm trong 1 đơn · KPI: tháng"],
    ["Dimensions", "Miền, cửa hàng, ngành hàng, tháng"],
    ["Owner", "Giám đốc bán hàng"],
    ["Target / Ngưỡng", "Bảng Target (từ B4) · Warning 90%, Critical 80%"],
    ["Ghi chú giới hạn", "Chưa trừ trả hàng, khuyến mại"],
], size=15, head_size=15, rowh=0.5, bold_first=True)
tip(s, 9.2, 1.65, 3.54, 2.2, "Viết card TRƯỚC khi viết công thức. Card là “hợp đồng” giữa bạn và người đọc báo cáo.")
tb(s, 9.2, 4.1, 3.54, 1.8, [[("Hôm nay viết 3 card", {"b": True, "color": NAVY})], "1 · Total Revenue (cùng lớp)", "2 · Total Orders (thực hành)", "3 · Gross Margin % (thực hành)"], size=16, space=6)
footnote(s, "Mẫu in: Handout/B2_KPI_Definition_Card.docx · 10 trường theo KPI dictionary (Course Plan, Phụ lục C)")
notes(s, 5)

# 6 · Chia phần 1
notes(section("1. Hàm logic: dạy Excel ra quyết định"), 6)

# 7 · IF, IFS
s = content("IF · IFS", "PHẦN 1 · HÀM LOGIC")
tool(s, L, 1.6, 4.2, "IF", "Một điều kiện, hai kết quả.", 0.45)
fx(s, L, 2.6, 4.2, '=IF(F2>=35%,"Biên cao","Biên thấp")', size=14)
vline(s, 6.05, 1.6, 1.5)
tool(s, 6.2, 1.6, 6.54, "IFS", "Nhiều điều kiện, dừng ở điều kiện ĐÚNG ĐẦU TIÊN.", 0.45)
fx(s, 6.2, 2.6, 6.54, '=IFS(D2>=100000,"Cao",D2>=30000,"Trung",TRUE,"Thấp")', size=14)
shot(s, L, 3.3, W, 2.2, "Cột Margin_%, IF (G), IFS (H) sau khi kéo công thức; P041 ra #DIV/0! và “Thấp”", "Demo/B2_Demo_Logic.xlsx · sheet Demo_IF · A1:H6 và dòng 42–43 (ẩn tạm dòng 7–41)", "S07_if_ifs.png")
tip(s, L, 5.65, W, 0.95, "IFS xét từ trái sang phải: đặt ngưỡng cao trước. TRUE ở cuối = “các trường hợp còn lại”. Ô trống được tính là 0 → P041 chưa có giá vẫn ra “Thấp”.")
footnote(s, "Ảnh: Demo/B2_Demo_Logic.xlsx, sheet Demo_IF")
notes(s, 7)

# 8 · AND / OR / NOT
s = content("AND · OR · NOT", "PHẦN 1 · HÀM LOGIC")
cols = [(L, "AND", "TRUE khi MỌI điều kiện đúng.", '=AND(C2="Bánh kẹo",F2>=35%)', "7 sản phẩm “Ưu tiên Tết”"),
        (5.47, "OR", "TRUE khi ÍT NHẤT MỘT điều kiện đúng.", '=OR(C2="Sữa",C2="Thực phẩm khô")', "13 sản phẩm thiết yếu"),
        (9.28, "NOT", "Đảo TRUE ↔ FALSE.", "=NOT(J2)", "29 sản phẩm không thiết yếu")]
for k, (x, n_, p_, f_, res) in enumerate(cols):
    tool(s, x, 1.6, 3.46, n_, p_, 0.7)
    fx(s, x, 2.9, 3.46, f_, size=12)
    box(s, x, 3.5, 3.46, 0.75, fill=WHITE, line=HAIR, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1, name="Stat card")
    tb(s, x + 0.15, 3.5, 3.16, 0.75, [[(res.split(" ", 1)[0] + "  ", {"b": True, "size": 26, "color": NAVY}), (res.split(" ", 1)[1], {"size": 14, "color": INK2})]], anchor="ctr")
    if k < 2: vline(s, x + 3.63, 1.6, 2.65)
shot(s, L, 4.4, W, 1.45, "Cột I, J, K (AND, OR, NOT) ra TRUE/FALSE trên các dòng Sữa và Bánh kẹo", "Demo/B2_Demo_Logic.xlsx · sheet Demo_IF · A1:K1 + dòng 9–13 và 22–28 (ẩn tạm các dòng khác)", "S08_and_or_not.png")
tip(s, L, 5.95, W, 0.75, "Thường lồng vào IF để ra nhãn dễ đọc: =IF(AND(C2=\"Bánh kẹo\",F2>=35%),\"Ưu tiên\",\"\"). Mẹo nhớ: AND khó tính, OR dễ tính.")
footnote(s, "Ảnh: Demo/B2_Demo_Logic.xlsx, sheet Demo_IF (cột I, J, K)")
notes(s, 8)

# 9 · IFERROR
s = content("IFERROR", "PHẦN 1 · HÀM LOGIC")
tool(s, L, 1.6, W, "IFERROR", "Nếu công thức ra lỗi thì trả về giá trị thay thế.", 0.45)
fx(s, L, 2.6, W, '=IFERROR((D2-E2)/D2,"Chưa có giá")')
shot(s, L, 3.3, 8.1, 1.55, "Dòng P041, P042: Margin_% #DIV/0!, cột L “Chưa có giá”; L42 sửa thành IFERROR(…,0) ra 0%", "Demo/B2_Demo_Logic.xlsx · sheet Demo_IF · A41:L43", "S09_iferror.png")
tip(s, L + 8.1, 3.3, 2.97, 1.5, "Có giá (25.000) → mọi cột tự tính lại: 44,0%.")
tip(s, L, 5.0, W, 1.0, "IFERROR CHE lỗi chứ không SỬA lỗi. Thay bằng chữ có nghĩa, đừng thay bằng 0: “0%” là con số sai trông như thật.", pro=True)
footnote(s, "Ảnh: Demo/B2_Demo_Logic.xlsx, sheet Demo_IF, dòng 42–43")
notes(s, 9)

# 10 · Chia phần 2
notes(section("2. Thống kê mô tả: một con số “trung bình” có thể nói dối"), 10)

# 11 · COUNT / AVERAGE vs MEDIAN / MIN MAX
s = content("Thống kê mô tả: trung bình hay trung vị?", "PHẦN 2 · THỐNG KÊ")
for k, (n_, p_) in enumerate([("COUNT / COUNTA", "Đếm ô số / ô không trống."), ("AVERAGE vs MEDIAN", "Trung bình cộng vs giá trị ở giữa."), ("MIN / MAX", "Nhỏ nhất / lớn nhất.")]):
    x = L + k * 3.8
    tool(s, x, 1.6, 3.46, n_, p_, 0.45, size=20)
    if k < 2: vline(s, x + 3.63, 1.6, 1.0)
cards = [("Số đơn Q1/2025", "3.144", "COUNT"), ("Trung bình", "299.445", "AVERAGE"), ("Trung vị", "198.000", "MEDIAN"), ("Nhỏ nhất – lớn nhất", "10.000 – 2.464.000", "MIN – MAX")]
cw = [2.4, 2.6, 2.6, 3.17]; x = L
for (lab, val, fn), w in zip(cards, cw):
    stat(s, x, 2.85, w - 0.15, 1.3, lab, val, fn, vsize=22 if len(val) > 10 else 28, vcolor=NAVY)
    x += w
shot(s, L, 4.3, 6.6, 2.3, "Khối I2:I6 (COUNT, AVERAGE, MEDIAN, MIN, MAX) và Status Bar khi chọn cột E: Average 299.445", "Demo/B2_Demo_ThongKe.xlsx · sheet Demo_ThongKe · E1:I13 + Status Bar", "S11_average_median.png")
tip(s, 8.45, 4.3, 4.29, 2.3, "Trung bình lớn hơn trung vị khoảng 51% → vài đơn rất lớn kéo trung bình lên (lệch phải). Chọn con số theo câu hỏi: “khách thường chi bao nhiêu” → trung vị; AOV nhân ngược ra doanh thu → trung bình.")
footnote(s, "Ảnh: Demo/B2_Demo_ThongKe.xlsx, sheet Demo_ThongKe (I2:I6)")
notes(s, 11)

# 12 · STDEV, skew, outlier
s = content("STDEV · SKEW · Outlier", "PHẦN 2 · THỐNG KÊ")
for k, (n_, p_) in enumerate([("STDEV.S", "Độ phân tán quanh trung bình."), ("SKEW", "> 0: đuôi dài bên phải."), ("QUARTILE.INC", "Q1, Q3 → ngưỡng outlier.")]):
    x = L + k * 3.8
    tool(s, x, 1.6, 3.46, n_, p_, 0.45)
    if k < 2: vline(s, x + 3.63, 1.6, 1.0)
# trục số
ax_y, x0, x1, vmax = 4.1, L + 0.2, R - 0.2, 2464000
X = lambda v: x0 + (x1 - x0) * v / vmax
ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(ax_y), Inches(x1), Inches(ax_y)); ln.line.color.rgb = rgb("C3C7CF"); ln.line.width = Pt(2)
box(s, X(89000), ax_y - 0.25, X(412000) - X(89000), 0.5, fill="CDE2FB", line="2A78D6", name="IQR box")
for v, lab, col in [(89000, "Q1\n89.000", INK2), (198000, "Trung vị\n198.000", NAVY), (412000, "Q3\n412.000", INK2), (896500, "Ngưỡng\n896.500", BAD_TXT), (2464000, "Max\n2.464.000", INK2)]:
    tk = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(X(v)), Inches(ax_y - 0.35), Inches(X(v)), Inches(ax_y + 0.35))
    tk.line.color.rgb = rgb(col); tk.line.width = Pt(2 if col == BAD_TXT else 1.25)
    ly = ax_y - 1.0 if v == 198000 else ax_y + 0.42
    tb(s, X(v) - 0.8, ly, 1.6, 0.6, lab, size=13, color=col, bold=col != INK2, align="ctr", anchor="b" if v == 198000 else "t")
box(s, X(896500), ax_y - 0.3, X(2464000) - X(896500), 0.6, fill=BAD_BG, name="Outlier zone", alpha=55000)
tb(s, X(896500) + 0.2, ax_y - 0.85, 5.0, 0.45, "138 đơn vượt ngưỡng = Q3 + 1,5 × (Q3 − Q1)", size=15, color=BAD_TXT, bold=True)
tb(s, L, 2.55, W, 0.5, [["Độ lệch chuẩn ", ("289.919", {"b": True, "color": NAVY}), " (gần bằng trung bình) · SKEW ", ("1,77", {"b": True, "color": NAVY}), " (lệch phải)"]], size=17)
shot(s, L, 5.2, 6.6, 1.45, "Khối I7:I13: STDEV.S, SKEW, Q1, Q3, IQR, ngưỡng 896.500, 138 đơn vượt ngưỡng", "Demo/B2_Demo_ThongKe.xlsx · sheet Demo_ThongKe · G7:I13", "S12_stdev_outlier.png")
tip(s, 8.45, 5.2, 4.29, 1.45, "Outlier không phải “lỗi để xóa”: có thể là khách mua quà Tết, mua sỉ. Giải thích trước, quyết định sau.", pro=True)
footnote(s, "Ảnh: Demo/B2_Demo_ThongKe.xlsx, sheet Demo_ThongKe (I7:I13)")
notes(s, 12)

# 13 · RANK.EQ
s = content("RANK.EQ", "PHẦN 2 · THỐNG KÊ")
tool(s, L, 1.6, 5.4, "RANK.EQ", "Xếp hạng một giá trị trong danh sách (0 = giảm dần).", 0.7)
fx(s, L, 2.85, 5.4, "=RANK.EQ(D2,$D$2:$D$13,0)")
tip(s, L, 3.6, 5.4, 1.3, "Vùng xếp hạng phải khóa $ — nếu không, kéo xuống vùng sẽ trôi.")
tb(s, L, 5.1, 5.4, 1.2, "Xếp hạng chỉ có nghĩa khi so cùng loại (Mall vs mặt phố, trung tâm vs tỉnh). “Vì sao” → Buổi 3.", size=16, color=INK2)
rank = [("1", "ST08 Quận 1", "130.448.000"), ("2", "ST01 Cầu Giấy", "112.234.000"), ("3", "ST02 Hoàn Kiếm", "94.421.000"), ("4", "ST09 Thủ Đức", "92.643.000"),
        ("…", "…", "…"), ("11", "ST06 Huế", "47.619.000"), ("12", "ST04 Hạ Long", "45.105.000")]
shot(s, 7.4, 1.6, 5.34, 4.7, "Cột E sau khi kéo RANK.EQ: ST08 hạng 1, ST04 hạng 12; F2 hiện vùng $D$2:$D$13", "Demo/B2_Demo_ThongKe.xlsx · sheet Demo_RANK · A1:E13", "S13_rank.png")
footnote(s, "Ảnh: Demo/B2_Demo_ThongKe.xlsx, sheet Demo_RANK")
notes(s, 13)

# 14 · Chia phần 3
notes(section("3. SUMIFS: mỗi tiêu chí là một bộ lọc"), 14)

# 15 · Xương sống Chuỗi C
s = content("Xương sống Chuỗi C: bộ lọc + phép tổng", "PHẦN 3 · SUMIFS")
box(s, L, 2.3, 2.6, 2.6, fill=ALT, line=HAIR, name="Source column")
tb(s, L, 2.3, 2.6, 2.6, [[("tblSales[Revenue]", {"b": True, "color": NAVY})], "6.187 dòng", [("∑ phép tổng", {"b": True})]], size=18, anchor="ctr", align="ctr", space=10)
for k, t in enumerate(["Vùng = Bắc", "Tháng = 3"]):
    y = 2.35 + k * 1.4
    f = box(s, 5.0, y, 3.0, 1.0, fill=WHITE, line="2A78D6", lw=2, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.15, name=f"Filter {k + 1}")
    tb(s, 5.0, y, 3.0, 1.0, [[("Bộ lọc  ", {"color": INK2, "size": 14}), (t, {"b": True, "color": NAVY})]], size=20, anchor="ctr", align="ctr")
for y in (2.85, 4.25):
    a = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(4.35), Inches(3.6), Inches(5.0), Inches(y)); a.line.color.rgb = rgb(ACCENT); a.line.width = Pt(2)
    b = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(8.0), Inches(y), Inches(9.3), Inches(3.6)); b.line.color.rgb = rgb(ACCENT); b.line.width = Pt(2)
box(s, 9.3, 2.75, 3.44, 1.7, fill=NAVY, name="Result cell")
tb(s, 9.3, 2.75, 3.44, 1.7, [[("Một ô kết quả", {"size": 14, "color": "CDE2FB"})], [("90.210.000", {"b": True, "size": 30, "color": WHITE})]], anchor="ctr", align="ctr", space=4)
table(s, L, 5.25, [3.69, 3.69, 3.69], [["B2 · SUMIFS", "B3 · PivotTable", "B5 · DAX"], ["criteria", "tọa độ ô Pivot", "filter context"]],
      size=17, rowh=0.42, align=["ctr", "ctr", "ctr"], name="Chain C")
tb(s, L, 1.65, W, 0.5, "Cùng một ý, ba cái tên. Kiểm chứng: Filter Bắc + tháng 3 → Status Bar Sum = 90.210.000 (635 dòng).", size=16, color=INK2)
footnote(s, "Hình này dùng lại nguyên vẹn ở B3 và B5 · Demo/B2_Demo_SUMIFS.xlsx, sheet Sales và Demo_SUMIFS (D4)")
notes(s, 15)

# 16 · SUMIFS / COUNTIFS / AVERAGEIFS
s = content("SUMIFS · COUNTIFS · AVERAGEIFS", "PHẦN 3 · SUMIFS")
table(s, L, 1.6, [2.3, 2.3, 6.47], [
    ["Hàm", "Mục đích", "Ví dụ: Bắc, tháng 3"],
    [[("SUMIFS", {"b": True, "color": RED})], "Tổng có điều kiện", [("=SUMIFS(tblSales[Revenue],tblSales[Region],\"Bắc\",tblSales[Thang],3)", {"size": 13}), ("  → 90.210.000", {"b": True, "color": NAVY})]],
    [[("COUNTIFS", {"b": True, "color": RED})], "Đếm có điều kiện", [("=COUNTIFS(tblSales[Region],\"Bắc\",tblSales[Thang],3)", {"size": 13}), ("  → 635 dòng", {"b": True, "color": NAVY}),
                                                                         ("   + tblSales[Line_No],1 → 321 đơn", {"b": True, "color": NAVY, "size": 13})]],
    [[("AVERAGEIFS", {"b": True, "color": RED})], "Trung bình có điều kiện", [("=AVERAGEIFS(tblSales[Revenue],…)", {"size": 13}), ("  → 142.063 / dòng", {"b": True, "color": NAVY})]],
], size=15, rowh=0.72, name="Functions")
tb(s, L, 4.6, 5.3, 0.45, [[("① ", {"b": True, "color": ACCENT}), "cột cần cộng (ĐẦU) · ", ("② ", {"b": True, "color": ACCENT}), "cặp cột lọc, giá trị · ", ("③ ", {"b": True, "color": ACCENT}), "thêm cặp"]], size=13)
shot(s, L, 5.05, 5.3, 1.6, "ScreenTip khi gõ =SUMIFS( trên tblSales: sum_range, criteria_range1, criteria1, …", "Demo/B2_Demo_SUMIFS.xlsx · sheet Demo_SUMIFS · ô trống, gõ =SUMIFS(", "S16_sumifs_screentip.png")
tip(s, 7.15, 4.65, 5.59, 0.95, "Dùng cột của Table: các vùng luôn cùng kích thước, không bị #VALUE!.")
tip(s, 7.15, 5.7, 5.59, 0.95, "Grain là dòng sản phẩm: đếm ĐƠN thì COUNTIFS thêm Line_No = 1.", pro=True)
footnote(s, "Ảnh: Demo/B2_Demo_SUMIFS.xlsx, sheet Demo_SUMIFS (D4:D7)")
notes(s, 16)

# 17 · Bảng 2 chiều & trung bình của trung bình
s = content("Bảng KPI 2 chiều · tổng ÷ tổng", "PHẦN 3 · SUMIFS")
fx(s, L, 1.6, W, [[("=SUMIFS(tblSales[Revenue],tblSales[Region],", {}), ("$A12", {"b": True, "color": "1F6FB2"}), (",tblSales[Thang],", {}), ("B$11", {"b": True, "color": "1F6FB2"}), (")", {})]], size=16)
shot(s, L, 2.3, 7.5, 2.45, "Bảng Miền × Tháng sau khi kéo từ B12; F2 ở C13 hiện $A13 và C$11; E15 = 941.454.000", "Demo/B2_Demo_SUMIFS.xlsx · sheet Demo_SUMIFS · A11:E15", "S17_sumifs_grid.png")
tb(s, L, 4.85, 7.5, 0.8, "Một công thức, kéo ra 9 ô: khóa cột $A (Miền), khóa dòng B$11 (Tháng). Đối soát: Tổng = tổng cột Revenue.", size=15, color=INK2)
stat(s, 9.4, 2.3, 3.34, 1.45, "AVERAGE GM% 6 ngành", "31,2%  ✘", None, vcolor=BAD_TXT, vsize=30)
stat(s, 9.4, 3.85, 3.34, 1.45, "Tổng lãi gộp ÷ tổng doanh thu", "34,2%  ✔", None, vcolor=GOOD, vsize=30)
tip(s, L, 5.65, W, 0.95, "KPI tỷ lệ (AOV, Margin %) luôn tính TỔNG ÷ TỔNG. Ngành Sữa (16,6 triệu) không được “một phiếu” ngang Bánh kẹo (338,7 triệu).")
footnote(s, "Ảnh: Demo/B2_Demo_SUMIFS.xlsx, sheet Demo_SUMIFS (B12:E15, dòng 20–28)")
notes(s, 17)

# 18 · Nghỉ
notes(section("Nghỉ 10 phút", "Trước khi quay lại: mở file khởi đầu và B2_Product_DuplicateKey.xlsx"), 18)

# 19 · Chia phần 4
notes(section("4. Nối bảng: lấy thông tin từ bảng khác"), 19)

# 20 · VLOOKUP
s = content("VLOOKUP", "PHẦN 4 · LOOKUP")
tool(s, L, 1.6, W, "VLOOKUP", "Tìm ở CỘT ĐẦU của bảng, trả về cột thứ n bên phải.", 0.45)
fx(s, L, 2.6, W, "=VLOOKUP(A2,Product!$A$2:$E$42,5,FALSE)")
shot(s, L, 3.3, W, 2.3, "Cột B (FALSE) và C (bỏ trống): P041, P21 ra #N/A ở cột B, 71.500 ở cột C", "Demo/B2_Demo_Lookup.xlsx · sheet Demo_VLOOKUP · A1:D7", "S20_vlookup.png")
tip(s, L, 5.75, W, 0.85, "LUÔN ghi đối số cuối là FALSE (hoặc 0). Bỏ trống = khớp gần đúng: không thấy mã thì lấy mã gần nhất nhỏ hơn (P040), không báo lỗi.")
footnote(s, "Ảnh: Demo/B2_Demo_Lookup.xlsx, sheet Demo_VLOOKUP")
notes(s, 20)

# 21 · INDEX / MATCH
s = content("INDEX / MATCH", "PHẦN 4 · LOOKUP")
tool(s, L, 1.6, W, "INDEX / MATCH", "MATCH tìm VỊ TRÍ, INDEX lấy giá trị Ở VỊ TRÍ ĐÓ.", 0.45)
fx(s, L, 2.6, W, "=INDEX(Product!$E$2:$E$42,MATCH(A2,Product!$A$2:$A$42,0))")
shot(s, L, 3.3, 5.1, 2.3, "Cột B = INDEX/MATCH, F2 ở B2: chỉ ô A2 được tô màu (vùng trên sheet Product không tô)", "Demo/B2_Demo_Lookup.xlsx · sheet Demo_XLOOKUP · A1:B7", "S21_index_match.png")

tb(s, 7.0, 3.35, 5.74, 2.2, [[("MATCH(\"P005\", cột mã, 0)", {"b": True}), " → dòng 5"], [("INDEX(cột Unit_Cost, 5)", {"b": True}), " → 83.500"],
                              "Cột tìm và cột trả về tách rời: tìm được ở mọi cột, trả về bên trái, chèn cột không sai."], size=17, space=10)
tip(s, L, 5.75, W, 0.85, "Chạy trên MỌI phiên bản Excel. Dùng khi máy chưa có XLOOKUP (Excel 2016/2019).")
footnote(s, "Ảnh: Demo/B2_Demo_Lookup.xlsx, sheet Demo_XLOOKUP (cột B)")
notes(s, 21)

# 22 · XLOOKUP
s = content("XLOOKUP", "PHẦN 4 · LOOKUP")
tool(s, L, 1.6, W, "XLOOKUP", "Tìm trong một cột, trả về cột khác; mặc định khớp chính xác.", 0.45)
fx(s, L, 2.6, W, '=XLOOKUP(A2,Product!$A$2:$A$42,Product!$E$2:$E$42,"Không có mã")')
table(s, L, 3.25, [0.5, 2.6, 3.0], [
    ["#", "Đối số", "Ý nghĩa"],
    ["①", "lookup_value", "Tìm cái gì"], ["②", "lookup_array", "Tìm trong cột nào"], ["③", "return_array", "Trả về cột nào"],
    ["④", "[if_not_found]", "Không thấy thì hiện gì"], ["⑤ ⑥", "[match_mode] [search_mode]", "Kiểu khớp · chiều tìm"],
], size=14, rowh=0.4, name="Arguments")
shot(s, 8.0, 3.25, 4.74, 1.6, "Cột C (“Không có mã”) và cột F (tra ngược ra P014, P039, P027)", "Demo/B2_Demo_Lookup.xlsx · sheet Demo_XLOOKUP · A1:F7", "S22_xlookup.png")
tip(s, 8.0, 4.95, 4.74, 1.5, "Trong Table: =XLOOKUP([@Product_Code], tblProduct[Product_Code], tblProduct[Unit_Cost])")
tip(s, L, 5.75, 6.1, 0.9, "Chỉ dùng if_not_found khi đã hiểu vì sao không tìm thấy; đừng dùng để che lỗi dữ liệu.", pro=True)
footnote(s, "Ảnh: Demo/B2_Demo_Lookup.xlsx, sheet Demo_XLOOKUP (cột C, F)")
notes(s, 22)

# 23 · Laws of Lookups
s = content("Laws of Lookups: khóa phải duy nhất", "PHẦN 4 · LOOKUP")
tb(s, L, 1.6, 6.0, 2.3, [
    [("Luật 1 · ", {"b": True, "color": NAVY}), "Khóa ở bảng tra phải DUY NHẤT. Trùng khóa → lấy dòng ĐẦU TIÊN, không báo lỗi."],
    [("Luật 2 · ", {"b": True, "color": NAVY}), "Khóa hai bên cùng kiểu, cùng cách viết (slide sau)."],
    [("Luật 3 · ", {"b": True, "color": NAVY}), "Không thấy thì phải thấy #N/A, đừng khớp gần đúng."]], size=16, space=8)
shot(s, L, 3.95, 6.1, 1.6, "Dòng 22–23: P021 hai lần, bản ghi cũ Unit_Cost 62.000 đứng trước", "Demo/B2_Demo_Lookup.xlsx · sheet Product · A1:E23 (ẩn tạm dòng 2–20)", "S23_p021_duplicate.png")
stat(s, 8.1, 1.6, 4.64, 1.55, "COGS của P021 trong Q1 (143 dòng)", "15.872.000", "đúng: 17.664.000 · thiếu 1.792.000", vcolor=BAD_TXT, vsize=28)
stat(s, 8.1, 3.3, 4.64, 1.55, "Gross Margin % toàn công ty", "34,4% vs 34,2%", "đẹp hơn thực tế — không một ô lỗi", vcolor=BAD_TXT, vsize=26)
tip(s, 8.1, 5.0, 4.64, 1.3, "Kiểm tra trùng: =COUNTIF(tblProduct[Product_Code],[@Product_Code]) > 1")
tb(s, L, 5.65, 6.1, 0.8, "Không lỗi ≠ không sai. Doanh thu vẫn đúng (giá bán lấy từ Sales); sai là GIÁ VỐN.", size=16, color=NAVY, bold=True)
footnote(s, "Ảnh: Demo/B2_Demo_Lookup.xlsx, sheet Demo_KhoaTrung và Product (dòng 22–23)")
notes(s, 23)

# 24 · TRIM / VALUE
s = content("Khóa khác kiểu: TRIM · VALUE", "PHẦN 4 · LOOKUP")
tool(s, L, 1.6, 5.3, "TRIM", "Bỏ dấu cách thừa ở đầu, cuối, giữa chuỗi.", 0.45)
vline(s, 7.2, 1.6, 1.0)
tool(s, 7.4, 1.6, 5.34, "VALUE / TEXT", "Đổi số dạng chữ \"1\" ↔ số 1.", 0.45)
shot(s, L, 2.75, W, 2.55, "Cột C ra #N/A với \"1\", \"P005 \", \" P013\"; cột E đúng sau VALUE / TRIM; ô =LEN(A11) = 5", "Demo/B2_Demo_Lookup.xlsx · sheet Demo_KieuKhoa · A3:E14", "S24_trim_value.png")
tip(s, L, 5.5, W, 0.95, "#N/A dù “nhìn thấy có”? Kiểm tra =LEN(ô) (\"P005 \" ra 5, không phải 4) và căn lề: số căn phải, chữ căn trái.")
footnote(s, "Ảnh: Demo/B2_Demo_Lookup.xlsx, sheet Demo_KieuKhoa")
notes(s, 24)

# 25 · Chia phần 5
notes(section("5. Text & Date: chuẩn hóa để phân tích theo thời gian"), 25)

# 26 · Text
s = content("TRIM · LEFT / MID / RIGHT · TEXT · &", "PHẦN 5 · TEXT & DATE")
table(s, L, 1.6, [2.5, 4.6, 3.97], [
    ["Hàm", "Ví dụ", "Kết quả"],
    [[("LEFT / MID / RIGHT", {"b": True, "color": RED})], "=MID(\"HD2501-00001\",3,4)", "2501  (năm 25, tháng 01)"],
    [[("TRIM", {"b": True, "color": RED})], "=TRIM(\"  MCI Mart   Cầu Giấy\")", "MCI Mart Cầu Giấy  (LEN 21 → 17)"],
    [[("&", {"b": True, "color": RED})], "=B2&\" – \"&TRIM(C2)", "ST01 – MCI Mart Cầu Giấy"],
    [[("TEXT", {"b": True, "color": RED})], "=TEXT(D2,\"#,##0\")&\" đ\"", "19,000 đ  (theo cài đặt vùng)"],
], size=15, rowh=0.5, name="Functions")
shot(s, L, 4.25, 6.6, 2.15, "Cột E–J (LEFT, MID, RIGHT, TRIM, &, TEXT) cho 6 cửa hàng", "Demo/B2_Demo_Text_Date.xlsx · sheet Demo_Text · A1:J7", "S26_text.png")
tip(s, 8.45, 4.25, 4.29, 2.15, "Kết quả của TEXT là CHỮ: =J2+1 ra #VALUE!. Tính toán bằng số, chỉ dùng TEXT ở bước trình bày.")
footnote(s, "Ảnh: Demo/B2_Demo_Text_Date.xlsx, sheet Demo_Text")
notes(s, 26)

# 27 · Date
s = content("Hàm ngày lõi", "PHẦN 5 · TEXT & DATE")
table(s, L, 1.6, [2.3, 5.2, 3.57], [
    ["Hàm", "Ví dụ (A2 = 15/01/2025)", "Kết quả"],
    [[("EOMONTH", {"b": True, "color": RED})], "=EOMONTH(A2,0)  ·  =EOMONTH(A2,-1)+1", "31/01/2025 · 01/01/2025"],
    [[("MONTH / Quý", {"b": True, "color": RED})], "=MONTH(A2)  ·  =\"Q\"&ROUNDUP(MONTH(A2)/3,0)", "1 · Q1"],
    [[("TEXT / WEEKDAY", {"b": True, "color": RED})], "=TEXT(A2,\"ddd\")  ·  =WEEKDAY(A2,2)", "Wed (máy tiếng Việt: T4) · 3"],
    [[("NETWORKDAYS", {"b": True, "color": RED})], "=NETWORKDAYS(B10,C10,$H$10:$H$15)", [("Tháng 1/2025: 23 → 17 ngày làm việc", {"b": True, "color": NAVY})]],
], size=14, rowh=0.5, name="Functions")
shot(s, L, 4.25, 6.6, 2.15, "Bảng ngày B2:G5 và khối NETWORKDAYS: tháng 1 ra 23 → 17", "Demo/B2_Demo_Text_Date.xlsx · sheet Demo_Date · A1:H15", "S27_date.png")
tip(s, 8.45, 4.25, 4.29, 2.15, "Ngày trong Excel là SỐ (01/01/2025 = 45658), nên cộng trừ được. Ngày căn trái = ngày lưu dạng chữ, không group theo tháng được.")
footnote(s, "Ảnh: Demo/B2_Demo_Text_Date.xlsx, sheet Demo_Date · ngày lễ: 01/01 và 27–31/01/2025")
notes(s, 27)

# 28 · Chọn hàm theo câu hỏi
s = content("Chọn hàm theo câu hỏi", "PHẦN 5 · TEXT & DATE")
table(s, L, 1.7, [6.2, 4.87], [
    ["Câu hỏi kinh doanh", "Nhóm hàm"],
    ["“Sản phẩm này thuộc nhóm nào?”", "IF · IFS · AND · OR"],
    ["“Khách thường chi bao nhiêu? Có gì bất thường?”", "MEDIAN · STDEV · QUARTILE"],
    ["“Bao nhiêu, theo vùng / tháng?”", "SUMIFS · COUNTIFS · AVERAGEIFS"],
    ["“Lấy thông tin từ bảng khác”", "XLOOKUP (INDEX / MATCH)"],
    ["“Theo tháng / quý / thứ”", "MONTH · EOMONTH · TEXT · WEEKDAY"],
], size=18, rowh=0.66, name="Chooser", colors={(i, 1): NAVY for i in range(1, 6)})
tb(s, L, 5.95, W, 0.5, "Bắt đầu từ câu hỏi, rồi mới chọn hàm.", size=18, color=NAVY, bold=True)
notes(s, 28)

# 29 · Thực hành 1
s = content("Thực hành 1: Làm giàu tblSales", "PHẦN 6 · THỰC HÀNH")
chip(s, L, 1.6, 1.53, "15 phút", name="Time chip")
tb(s, L, 2.2, 6.9, 4.5, [
    [("1. ", {"b": True, "color": NAVY}), "Mở file khởi đầu, lưu thành B2_<HoTen>.xlsx. Chép sheet Product từ B2_Product_DuplicateKey.xlsx: chuột phải tab → Move or Copy → tick Create a copy."],
    [("2. ", {"b": True, "color": NAVY}), "Kiểm tra khóa: =COUNTIF(tblProduct[Product_Code],[@Product_Code]) → xóa bản ghi P021 cũ (dòng 22)."],
    [("3. ", {"b": True, "color": NAVY}), "tblSales thêm: Category, Unit_Cost (XLOOKUP) · COGS = [@Quantity]*[@Unit_Cost] · Gross_Margin = [@Revenue]-[@COGS]."],
    [("4. ", {"b": True, "color": NAVY}), "Thang = MONTH · Quy = \"Q\"&ROUNDUP(MONTH(…)/3,0) · Thu = TEXT(…,\"ddd\") · Phan_khuc = IFS (≥100.000 Cao, ≥30.000 Trung, còn lại Thấp)."],
], size=15, space=10)
box(s, 8.87, 1.6, 3.87, 5.0, fill=ALT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03, name="Check box")
tb(s, 9.07, 1.75, 3.47, 4.7, [[("Kiểm tra", {"b": True, "color": NAVY, "size": 16})], "Không còn #N/A.", [("Tổng COGS", {"b": True, "color": NAVY})],
                              [("619.663.000", {"b": True, "color": NAVY, "size": 30})], "Nếu ra 617.871.000: còn P021 cũ.",
                              [("Phan_khuc (dòng)", {"b": True, "color": NAVY})], "Cao 1.814 · Trung 3.005 · Thấp 1.368"], size=15, space=8)
footnote(s, "File khởi đầu: ../B1_Nen_tang_Excel_Du_lieu/Practice/B1_MCI_Mart_Q1_2025_Answer.xlsx · Practice/B2_Product_DuplicateKey.xlsx")
notes(s, 29)

# 30 · Thực hành 2
s = content("Thực hành 2: Bảng KPI & KPI Card", "PHẦN 6 · THỰC HÀNH")
chip(s, L, 1.6, 1.53, "12 phút", name="Time chip")
tb(s, L, 2.2, 6.9, 4.5, [
    [("1. ", {"b": True, "color": NAVY}), "Sheet KPI_SUMIFS: Doanh thu Miền (A6:A8) × Tháng (B5:D5 = 1, 2, 3). MỘT công thức ở B6, kéo ra 9 ô ($A6, B$5); thêm cột Quý 1 và dòng Tổng."],
    [("2. ", {"b": True, "color": NAVY}), "Số đơn: COUNTIFS thêm tblSales[Line_No],1."],
    [("3. ", {"b": True, "color": NAVY}), "AOV = Doanh thu ÷ Số đơn · Gross Margin % = SUMIFS(Gross_Margin) ÷ Doanh thu. Dòng Tổng: tổng ÷ tổng."],
    [("4. ", {"b": True, "color": NAVY}), "Viết 2 card: Total Orders, Gross Margin %."],
], size=15, space=10)
box(s, 8.87, 1.6, 3.87, 5.0, fill=ALT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03, name="Check box")
tb(s, 9.07, 1.75, 3.47, 4.7, [[("Kiểm tra", {"b": True, "color": NAVY, "size": 16})], [("941.454.000", {"b": True, "color": NAVY, "size": 28})],
                              "Tổng bảng doanh thu", "Tháng 1: 425.850.000", "Số đơn: 3.144 (không phải 6.187)", "AOV: 299.445", "Gross Margin: 34,2%"], size=15, space=8)
footnote(s, "Mẫu card: Handout/B2_KPI_Definition_Card.docx")
notes(s, 30)

# 31 · Chữa bài
s = content("Chữa bài nhanh", "PHẦN 6 · THỰC HÀNH")
items = ["tblSales có Category, Unit_Cost, COGS, Gross_Margin, Thang, Quy, Thu, Phan_khuc; không còn #N/A",
         "tblProduct không còn mã trùng",
         "Bảng KPI Miền × Tháng: tổng = 941.454.000; số đơn = 3.144",
         "AOV, Gross Margin % dòng Tổng tính tổng ÷ tổng",
         "3 KPI Definition Card"]
for k, t in enumerate(items):
    y = 1.8 + k * 0.9
    box(s, L, y + 0.05, 0.35, 0.35, fill=WHITE, line=NAVY, lw=1.5, name=f"Checkbox {k + 1}")
    tb(s, L + 0.6, y, W - 0.6, 0.5, t, size=20, anchor="ctr")
    if k < 4:
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(L), Inches(y + 0.72), Inches(R), Inches(y + 0.72)); ln.line.color.rgb = rgb(HAIR)
footnote(s, "File đáp án Practice/B2_MCI_Mart_Q1_2025_Answer.xlsx (sheet Kiem_tra: 10 dòng đối soát) là file khởi đầu Buổi 3.")
notes(s, 31)

# 32 · Quiz
notes(section("Excel Quiz Challenge", "7 câu hỏi nhanh · câu 6–7 là tình huống"), 32)
quiz(33, "IFS", "Giá bán 135.000 ra gì với =IFS(D2>=30000,\"Trung\",D2>=100000,\"Cao\",TRUE,\"Thấp\")?",
     {"A": "Cao", "B": "Trung", "C": "Thấp", "D": "#N/A"}, "B")
quiz(34, "Statistics", "Giá trị đơn Q1: trung bình 299.445, trung vị 198.000. Điều này cho biết gì?",
     {"A": "Dữ liệu có lỗi", "B": "Có một số đơn rất lớn kéo trung bình lên", "C": "Nửa số đơn trên 299.445", "D": "Trung vị tính sai"}, "B")
quiz(35, "SUMIFS", "Trong SUMIFS, đối số ĐẦU TIÊN là gì?",
     {"A": "Vùng tiêu chí đầu tiên", "B": "Tiêu chí đầu tiên", "C": "Vùng cần cộng", "D": "Số tiêu chí"}, "C")
quiz(36, "VLOOKUP", "=VLOOKUP(A2,Product!$A$2:$E$42,5) bỏ trống đối số cuối thì tra kiểu gì?",
     {"A": "Khớp chính xác", "B": "Khớp gần đúng", "C": "Báo lỗi", "D": "Tìm từ dưới lên"}, "B")
quiz(37, "Date", "=EOMONTH(DATE(2025,2,15),0) trả về?",
     {"A": "01/02/2025", "B": "15/03/2025", "C": "28/02/2025", "D": "31/03/2025"}, "C")
quiz(38, "Tình huống — Lookup", "Sau khi XLOOKUP giá vốn, Gross Margin % cao hơn tháng trước dù không có ô lỗi nào. Việc ĐẦU TIÊN nên làm?",
     {"A": "Đổi sang VLOOKUP", "B": "Kiểm tra khóa trùng trong bảng Product (COUNTIF)", "C": "Bọc IFERROR", "D": "Định dạng lại cột %"}, "B", qsize=20)
quiz(39, "Tình huống — KPI tỷ lệ", "Bạn có AOV của 3 miền; sếp hỏi AOV toàn công ty. Tính thế nào?",
     {"A": "AVERAGE 3 AOV", "B": "Tổng doanh thu ÷ tổng số đơn", "C": "MEDIAN 3 AOV", "D": "Lấy AOV miền lớn nhất"}, "B")

# 40 · Tổng kết + insight
s = content("Tổng kết & insight của bạn", "PHẦN 7 · QUIZ & TỔNG KẾT")
tb(s, L, 1.65, 5.0, 0.45, "Điểm chính cần nhớ", size=18, color=NAVY, bold=True)
tb(s, L, 2.2, 5.0, 4.3, ["(1) Định nghĩa KPI (card) trước, công thức sau.", "(2) Biết grain: 6.187 dòng ≠ 3.144 đơn.", "(3) Mỗi tiêu chí SUMIFS là một bộ lọc.",
                        "(4) Khóa lookup phải duy nhất; không lỗi ≠ không sai.", "(5) KPI tỷ lệ: tổng ÷ tổng."], size=18, space=12)
box(s, 6.97, 1.65, 5.77, 3.8, fill=ALT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03, name="Insight box")
tb(s, 7.17, 1.8, 5.37, 0.45, "Insight của bạn (1 câu)", size=18, color=NAVY, bold=True)
tb(s, 7.17, 2.35, 5.37, 2.2, "“Doanh thu tháng ___ là ___ (con số, có so sánh), vì ___ (nguyên nhân khả dĩ), nên ___ (hàm ý).”", size=22)
for k, t in enumerate(["Con số", "Nguyên nhân", "Hàm ý"]):
    chip(s, 7.17 + k * 1.84, 4.8, 1.69, t, fill=WHITE, line=HAIR, name=f"Tag {t}", h=0.45)
notes(s, 40)

# 41 · Preview B3
s = content("Buổi sau: PivotTable", "PHẦN 7 · QUIZ & TỔNG KẾT")
tb(s, L, 1.65, 7.0, 1.1, "“Doanh thu biến động thế nào theo cửa hàng, ngày trong tuần và ngành hàng — và vì sao tháng 1 cao hơn hẳn tháng 2?”", size=19, color=NAVY, bold=True)
tb(s, L, 2.85, 7.0, 1.0, [[("Giới hạn hôm nay: ", {"b": True}), "mỗi chiều phân tích mới lại phải viết lại SUMIFS → PivotTable (Tầng 2–3)."]], size=16)
tb(s, L, 3.85, 7.0, 2.8, [[("Về nhà", {"b": True, "color": NAVY})], "Đọc handout B2 · hoàn thành 3 KPI card",
                          "Bài tập về nhà (bắt buộc): B2_BTVN_MCI_Mart_T4_2025.xlsx — nộp trước Buổi 3 tại hocvien.mcivietnam.com, tên file B2_BTVN_<HoTen>.xlsx",
                          "Tick “Checklist trước Buổi 3” (mục 10 handout)"], size=15, space=6)
box(s, 9.07, 4.6, 3.67, 1.0, fill=NAVY, name="Brick B2")
tb(s, 9.17, 4.6, 3.47, 1.0, "✓  B2 · Bảng KPI (SUMIFS)", size=18, color=WHITE, bold=True, anchor="ctr")
box(s, 9.07, 3.3, 3.67, 1.0, fill=WHITE, line="2A78D6", lw=2.25, name="Brick B3")
tb(s, 9.17, 3.3, 3.47, 1.0, "B3 · Báo cáo Pivot", size=18, color=NAVY, bold=True, anchor="ctr")
tb(s, 9.07, 5.7, 3.67, 0.5, "Giữ bảng KPI_SUMIFS: đáp án đối soát ở B3, B5.", size=13, color=INK2)
notes(s, 41)

assert len(prs.slides) == 41, len(prs.slides)
prs.save(OUT)
print("ok", OUT, len(prs.slides), "slide;", f"còn {len(MISSING)} khung [CHỤP] chưa có ảnh:" if MISSING else "đủ ảnh.", " ".join(MISSING))
