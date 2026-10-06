#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Chấm bài tập về nhà Buổi 1 (BTVN_MCI_Mart_T4_2025): làm sạch báo cáo, Table, Revenue,
Pct_Share, Data Validation, Circle Invalid Data, Insight.

Chỉ ĐỌC file học viên (openpyxl, không save) và đọc thẳng XML trong file zip cho Data Validation
(kể cả extLst / x14). Đáp án lấy từ file _Answer (sheet Sales, Store, Product, Kiem_tra cột D).

Chấm điểm (thang tạm 10 điểm, trọng số nguyên: việc 1-7 mỗi việc 1 điểm, việc 8 (insight) 3 điểm = 1 điểm mỗi thành phần; Đạt 100%, Một phần 50%, Chưa đạt 0).
Không còn kết luận "nộp lại"/"cần sửa nhỏ": khóa không yêu cầu nộp lại, chỉ chấm điểm.
Tên sheet Kiem_tra / Kiem tra / Kiểm tra coi là giống nhau.

Dùng:
  python grade_b1.py <bai_nop.xlsx> <_Answer.xlsx>            -> in JSON một bài
  python grade_b1.py --folder <thu_muc_bai_nop> <_Answer.xlsx> [--insight-file in.json] -> in JSON cả lớp
  python grade_b1.py --check-seed <de_chua_sach.xlsx> <_Answer.xlsx> -> kiểm danh sách lỗi cài sẵn (SEEDED) với file đề

Phần chữ (câu insight) chỉ được trích nguyên văn, kèm gợi ý máy; người chấm đọc và quyết định.
"""
import sys, os, re, json, glob, zipfile, unicodedata, datetime, warnings, difflib
import xml.etree.ElementTree as ET
import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string

warnings.filterwarnings("ignore")

# --------------------------------------------------------------------------------------
# Lỗi cài sẵn trong file đề (lấy bằng cách so file đề chưa sạch với _Answer theo khóa
# Order_ID + Line_No; kiểm lại bằng --check-seed). Giá trị đúng lấy từ _Answer khi chấm.
# --------------------------------------------------------------------------------------
SEEDED = {
    "Quantity dạng text": [["HD2504-00290", 2], ["HD2504-00388", 1], ["HD2504-00130", 3], ["HD2504-00489", 1],
                           ["HD2504-00653", 1], ["HD2504-00769", 1], ["HD2504-00057", 3], ["HD2504-00406", 3],
                           ["HD2504-00344", 1], ["HD2504-00280", 4]],
    "Order_Date dạng text dd/mm/yyyy": [["HD2504-00117", 1], ["HD2504-00192", 1], ["HD2504-00504", 1],
                                        ["HD2504-00559", 1], ["HD2504-00641", 3], ["HD2504-00261", 2],
                                        ["HD2504-00075", 3], ["HD2504-00655", 1], ["HD2504-00465", 1],
                                        ["HD2504-00053", 4], ["HD2504-00629", 1], ["HD2504-00213", 1]],
    "Store_Code sai (ST5)": [["HD2504-00645", 2], ["HD2504-00684", 1]],
    "Product_Code thiếu số 0": [["HD2504-00601", 1], ["HD2504-00383", 2], ["HD2504-00413", 4]],
    "Quantity âm": [["HD2504-00200", 1], ["HD2504-00632", 2]],
}

SEED_FIELD = {
    "Quantity dạng text": "Quantity",
    "Order_Date dạng text dd/mm/yyyy": "Order_Date",
    "Store_Code sai (ST5)": "Store_Code",
    "Product_Code thiếu số 0": "Product_Code",
    "Quantity âm": "Quantity",
}

VIEC_TEN = {
    1: "Làm sạch cấu trúc sheet",
    2: "Sửa kiểu dữ liệu (Quantity, Order_Date)",
    3: "Excel Table tblSales_T4, sheet Sales",
    4: "Cột Revenue",
    5: "Kiem_tra!C5 và Pct_Share",
    6: "Data Validation",
    7: "Circle Invalid Data (sửa lỗi, không còn ô sai)",
    8: "Insight",
}

CANON = {"orderid": "Order_ID", "lineno": "Line_No", "orderdate": "Order_Date", "storecode": "Store_Code",
         "customerid": "Customer_ID", "productcode": "Product_Code", "productname": "Product_Name",
         "quantity": "Quantity", "unitprice": "Unit_Price", "revenue": "Revenue", "pctshare": "Pct_Share"}

VI_CHARS = re.compile(r"[àáảãạăằắẳẵặâầấẩẫậèéẻẽẹêềếểễệìíỉĩịòóỏõọôồốổỗộơờớởỡợùúủũụưừứửữựỳýỷỹỵđ]", re.I)
NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def strip_acc(s):
    s = unicodedata.normalize("NFD", str(s))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return s.replace("đ", "d").replace("Đ", "D")


def norm(s):
    return re.sub(r"[^a-z0-9]", "", strip_acc(s).lower())


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def tonum(s):
    s = re.sub(r"[.,]", "", s)
    return int(s) if s.isdigit() else None


def localname(tag):
    return tag.split("}")[-1]


# --------------------------------------------------------------------------------------
# Đọc sheet dữ liệu
# --------------------------------------------------------------------------------------
def find_header(ws, max_rows=15):
    for r in range(1, min(ws.max_row, max_rows) + 1):
        cols = {}
        for c in ws[r]:
            if c.value is not None and isinstance(c.value, str):
                cols.setdefault(norm(c.value), c.column)
        if "orderid" in cols and "quantity" in cols:
            return r, {CANON[k]: v for k, v in cols.items() if k in CANON}, \
                   {v: ws.cell(r, v).value for k, v in cols.items() if k in CANON}
    return None


def classify(oid, vals):
    if oid is None:
        return "blank" if all(x is None for x in vals) else "other"
    s = str(oid).strip()
    ns = strip_acc(s).lower()
    if re.match(r"^HD\d{4}-\d{5}$", s):
        return "data"
    if ns.startswith("cua hang"):
        return "group"
    if ns.startswith("cong"):
        return "sub"
    if ns.startswith("tong"):
        return "total"
    if norm(s) == "orderid":
        return "dup_header"
    return "other"


class Book:
    """Gói workbook: 2 lượt (công thức, giá trị đã lưu), bảng dữ liệu, sheet phụ."""

    def __init__(self, path):
        self.path = path
        self.wf = openpyxl.load_workbook(path, data_only=False)
        self.wv = openpyxl.load_workbook(path, data_only=True)
        self.sheets = self.wf.sheetnames
        self.recalc = []  # ghi chú những chỗ phải tự tính lại vì không có giá trị đã lưu
        self.ws_name = None
        for ws in self.wf.worksheets:
            h = find_header(ws)
            if h:
                self.ws_name = ws.title
                self.hdr_row, self.cols, self.hdr_text = h
                break
        if self.ws_name is None:
            raise ValueError("Không tìm thấy bảng dữ liệu có tiêu đề Order_ID, Quantity")
        self.wsf = self.wf[self.ws_name]
        self.wsv = self.wv[self.ws_name]
        self.table = None
        for t in self.wsf.tables.values():
            m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", t.ref)
            if m and int(m.group(2)) == self.hdr_row:
                tot = int(getattr(t, "totalsRowCount", 0) or 0)
                self.table = {"name": t.name, "ref": t.ref, "min_row": int(m.group(2)),
                              "max_row": int(m.group(4)) - tot, "totals_row": tot > 0, "table": t}
                break
        self._read_rows()

    def _read_rows(self):
        maxc = max(self.cols.values())
        rf = list(self.wsf.iter_rows(min_row=1, max_row=self.wsf.max_row, max_col=max(maxc, 12)))
        rv = list(self.wsv.iter_rows(min_row=1, max_row=self.wsv.max_row, max_col=max(maxc, 12)))
        self.rows = []  # mỗi phần tử: dict(row, kind, f{canon}, v{canon})
        for i in range(self.hdr_row + 1, self.wsf.max_row + 1):
            cf, cv = rf[i - 1], rv[i - 1]
            f, v = {}, {}
            for name, ci in self.cols.items():
                f[name] = cf[ci - 1].value
                v[name] = cv[ci - 1].value
            vals = [v[k] for k in self.cols]
            kind = classify(v.get("Order_ID"), vals)
            if self.table and self.table["totals_row"] and i == self.table["max_row"] + 1:
                kind = "bang_tong"  # dòng Total Row của Table: hợp lệ, không phải dòng rác
            self.rows.append({"row": i, "kind": kind, "f": f, "v": v})
        # bỏ dòng trống cuối sheet
        while self.rows and self.rows[-1]["kind"] == "blank":
            self.rows.pop()
        # vùng bảng
        if self.table:
            lo, hi = self.table["min_row"] + 1, self.table["max_row"]
            self.region = [r for r in self.rows if lo <= r["row"] <= hi]
        else:
            self.region = list(self.rows)
        # các dòng phía trên tiêu đề có nội dung
        self.above = []
        for i in range(1, self.hdr_row):
            vals = [c.value for c in rf[i - 1] if c.value is not None]
            if vals:
                self.above.append((i, str(vals[0])[:80]))

    # giá trị Revenue / Pct dùng cho tổng (giá trị đã lưu; thiếu thì tự tính lại)
    def num_rev(self, r):
        v = r["v"].get("Revenue")
        if is_num(v):
            return v
        f = r["f"].get("Revenue")
        q, p = r["v"].get("Quantity"), r["v"].get("Unit_Price")
        if isinstance(f, str) and f.startswith("=") and is_num(q) and is_num(p):
            if "revenue" not in " ".join(self.recalc):
                self.recalc.append("revenue: ô công thức không có giá trị đã lưu, tính lại = Quantity*Unit_Price")
            return q * p
        return None

    def total_rev(self):
        return sum(x for x in (self.num_rev(r) for r in self.region) if x is not None)

    def num_pct(self, r, total):
        v = r["v"].get("Pct_Share")
        if is_num(v):
            return v
        f = r["f"].get("Pct_Share")
        rev = self.num_rev(r)
        if isinstance(f, str) and f.startswith("=") and rev is not None and total:
            if not any(s.startswith("pct_share") for s in self.recalc):
                self.recalc.append("pct_share: ô công thức không có giá trị đã lưu, tính lại = Revenue / tổng")
            return rev / total
        return None


def load_answer(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Sales"]
    A = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        A[(r[0], r[1])] = {"Order_Date": r[2], "Store_Code": r[3], "Product_Code": r[5],
                           "Quantity": r[7], "Unit_Price": r[8]}
    stores = {str(r[0]).lower() for r in wb["Store"].iter_rows(min_row=2, values_only=True) if r[0]}
    prods = {str(r[0]).lower() for r in wb["Product"].iter_rows(min_row=2, values_only=True) if r[0]}
    kt = []
    for r in wb["Kiem_tra"].iter_rows(min_row=5, values_only=True):
        if r[0] and r[3] is not None:
            kt.append((r[0], r[3]))
    total = sum(a["Quantity"] * a["Unit_Price"] for a in A.values())
    return {"rows": A, "stores": stores, "prods": prods, "kiem_tra": kt, "tong_dt": total}


# --------------------------------------------------------------------------------------
# Data Validation: đọc thẳng XML (cả <dataValidations> thường lẫn extLst x14)
# --------------------------------------------------------------------------------------
def read_dv(path, sheet_name):
    z = zipfile.ZipFile(path)
    wb = ET.fromstring(z.read("xl/workbook.xml"))
    rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    rid = None
    for s in wb.iter("{%s}sheet" % NS_MAIN):
        if s.get("name") == sheet_name:
            rid = s.get("{%s}id" % NS_REL)
    target = None
    for r in rels:
        if r.get("Id") == rid:
            target = r.get("Target")
    if not target:
        return []
    target = target.lstrip("/")
    if not target.startswith("xl/"):
        target = "xl/" + target
    root = ET.fromstring(z.read(target))
    out = []
    for el in root.iter():
        if localname(el.tag) != "dataValidation":
            continue
        d = {k: el.get(k) for k in ("type", "operator", "allowBlank", "showErrorMessage", "showDropDown",
                                    "errorTitle", "error", "prompt", "promptTitle", "sqref")}
        d["nguon"] = "xml thường" if el.tag.startswith("{%s}" % NS_MAIN) else "extLst (x14)"
        f1 = f2 = None
        for ch in el.iter():
            ln = localname(ch.tag)
            if ln == "sqref" and ch.text:
                d["sqref"] = ch.text
            if ln == "formula1":
                f1 = "".join(ch.itertext()).strip()
            if ln == "formula2":
                f2 = "".join(ch.itertext()).strip()
        d["formula1"] = f1.lstrip("=") if f1 else f1
        d["formula2"] = f2.lstrip("=") if f2 else f2
        out.append(d)
    return out


def sqref_rows_for_col(sqref, col_idx):
    """Tập số dòng được phủ trong cột col_idx bởi sqref."""
    rows = set()
    for part in (sqref or "").split():
        m = re.match(r"^\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?$", part)
        if not m:
            continue
        c1 = column_index_from_string(m.group(1))
        r1 = int(m.group(2))
        c2 = column_index_from_string(m.group(3)) if m.group(3) else c1
        r2 = int(m.group(4)) if m.group(4) else r1
        if c1 <= col_idx <= c2:
            rows.update(range(r1, r2 + 1))
    return rows


def dv_source_info(book, f1):
    """Cho biết công thức nguồn List trỏ tới đâu."""
    if f1 is None:
        return {"loai": "không có", "sheet": None}
    if f1.startswith('"') or ("," in f1 and "!" not in f1 and not re.search(r"\$?[A-Z]+\$?\d+", f1)):
        return {"loai": "danh sách gõ tay", "sheet": None, "f": f1}
    name = f1
    try:
        dn = book.wf.defined_names
        key = f1.strip()
        if key in dn:
            name = dn[key].attr_text
    except Exception:
        pass
    m = re.match(r"^(?:'([^']+)'|([^'!]+))!", name)
    if m:
        return {"loai": "sheet khác", "sheet": m.group(1) or m.group(2), "f": f1}
    return {"loai": "cùng sheet", "sheet": book.ws_name, "f": f1}


# --------------------------------------------------------------------------------------
# Phân tích công thức Pct_Share
# --------------------------------------------------------------------------------------
STRUCT = re.compile(r"[A-Za-z0-9_À-ỹ]+\[(?:\[[^\]]*\]|[^\]])*\]")
REF = re.compile(r"(?:(?:'([^']+)'|([A-Za-zÀ-ỹ0-9_.]+))!)?(\$?)([A-Z]{1,3})(\$?)(\d+)")


def pct_formula_info(book, f, total):
    if not (isinstance(f, str) and f.startswith("=")):
        return {"la_cong_thuc": False}
    stripped = STRUCT.sub("<T>", f)
    refs = []
    for m in REF.finditer(stripped):
        refs.append({"sheet": m.group(1) or m.group(2) or book.ws_name, "col": m.group(4), "row": int(m.group(6)),
                     "abs": bool(m.group(3) and m.group(5)), "text": m.group(0)})
    return {"la_cong_thuc": True, "refs": refs, "stripped": stripped}


def resolve_cell(book, sheet, col, row):
    if sheet not in book.wf.sheetnames:
        # thử khớp không phân biệt dấu/hoa thường
        cand = [s for s in book.wf.sheetnames if norm(s) == norm(sheet)]
        if not cand:
            return None, None
        sheet = cand[0]
    c = f"{col}{row}"
    return book.wf[sheet][c].value, book.wv[sheet][c].value


DIEM_MOI_VIEC = 1.0  # việc 1-7 mỗi việc 1 điểm, insight 3 điểm (1 điểm/thành phần) = 10 điểm (thang tạm; đề và Kiem_tra chưa quy định điểm)


def diem_viec(kq):
    """Đạt = 100%, Một phần = 50%, Chưa đạt = 0; Không kiểm được = 0 (tạm, cần GV xem)."""
    return DIEM_MOI_VIEC * {"Đạt": 1.0, "Một phần": 0.5}.get(kq, 0.0)


def ap_dung_insight(res, tp):
    """tp: số thành phần insight đạt (0-3), do người chấm đọc. Trả về tổng điểm /10."""
    d8 = float(tp)  # 1 điểm mỗi thành phần
    res["diem_viec"]["8"] = d8
    from decimal import Decimal, ROUND_HALF_UP
    res["tong_diem"] = float(Decimal(repr(round(res["diem_viec_1_7"] + d8, 6))).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
    res["insight_thanh_phan_dat"] = tp
    return res


def name_state(actual, expected):
    """dung | hoa_thuong (chỉ khác chữ hoa/thường, coi như đúng, chỉ ghi chú) | sai"""
    if actual == expected:
        return "dung"
    if actual and actual.lower() == expected.lower():
        return "hoa_thuong"
    return "sai"


# --------------------------------------------------------------------------------------
# Chấm một bài
# --------------------------------------------------------------------------------------
def grade(path, ans_path, ans=None):
    ans = ans or load_answer(ans_path)
    res = {"file": os.path.basename(path), "mo_duoc": True}
    try:
        zipfile.ZipFile(path).testzip()
        bk = Book(path)
    except Exception as e:
        res.update(mo_duoc=False, loi=f"{type(e).__name__}: {e}", tong_diem=0,
                   ly_do="không mở được hoặc không tìm thấy bảng dữ liệu, cần GV xem")
        return res

    z = zipfile.ZipFile(path)
    core = z.read("docProps/core.xml").decode("utf8", "ignore")
    res["nguoi_tao"] = (re.findall(r"<dc:creator>(.*?)</dc:creator>", core) or [""])[0]
    res["nguoi_sua_cuoi"] = (re.findall(r"<cp:lastModifiedBy>(.*?)</cp:lastModifiedBy>", core) or [""])[0]
    res["sheets"] = bk.sheets
    res["sheet_du_lieu"] = bk.ws_name
    res["dong_tieu_de"] = bk.hdr_row
    res["tieu_de_cot"] = bk.hdr_text and list(bk.hdr_text.values())
    res["bang"] = {"ten": bk.table["name"], "ref": bk.table["ref"], "bat_total_row": bk.table["totals_row"]} \
        if bk.table else None

    region, rows = bk.region, bk.rows
    data_rows = [r for r in rows if r["kind"] == "data"]
    kinds = {}
    for r in rows:
        kinds[r["kind"]] = kinds.get(r["kind"], 0) + 1
    junk = {k: v for k, v in kinds.items() if k not in ("data", "bang_tong")}
    junk_in_region = {}
    for r in region:
        if r["kind"] != "data":
            junk_in_region[r["kind"]] = junk_in_region.get(r["kind"], 0) + 1

    # ---- so khóa với _Answer ----
    A = ans["rows"]
    seen, dup = set(), 0
    S = {}
    for r in data_rows:
        key = (r["v"]["Order_ID"], r["v"].get("Line_No"))
        if key in seen:
            dup += 1
        seen.add(key)
        S[key] = r
    matched = [k for k in S if k in A]
    sai_bo = len(matched) < 0.5 * len(A)
    thieu = [k for k in A if k not in S]
    thua = [k for k in S if k not in A]
    fields = ["Order_Date", "Store_Code", "Product_Code", "Quantity", "Unit_Price"]
    lech = {f: [] for f in fields}
    for k in matched:
        v = S[k]["v"]
        a = A[k]
        for f in fields:
            sv, av = v.get(f), a[f]
            if f == "Order_Date":
                ok = isinstance(sv, datetime.datetime) and sv == av
            elif f in ("Quantity", "Unit_Price"):
                ok = is_num(sv) and sv == av
            else:
                ok = isinstance(sv, str) and sv == av
            if not ok:
                lech[f].append([k[0], k[1], str(sv), str(av)])
    res["so_khoa"] = {"dong_du_lieu_hop_le": len(data_rows), "khop_khoa_voi_dap_an": len(matched),
                      "thieu_so_voi_dap_an": len(thieu), "thua_so_voi_dap_an": len(thua), "trung_khoa": dup,
                      "sai_bo_du_lieu": sai_bo,
                      "mau_thieu": [list(k) for k in thieu[:6]], "mau_thua": [list(k) for k in thua[:6]],
                      "so_o_lech": {f: len(lech[f]) for f in fields},
                      "mau_lech": {f: lech[f][:6] for f in fields if lech[f]}}

    # ---- lỗi cài sẵn ----
    seeded = {}
    for name, keys in SEEDED.items():
        f = SEED_FIELD[name]
        st = {"tong": len(keys), "da_sua_dung": 0, "con_loi_nguyen": 0, "sua_sai": 0, "xoa_dong": 0, "chi_tiet_sai": []}
        for k in keys:
            k = tuple(k)
            if k not in A:
                continue
            if k not in S:
                st["xoa_dong"] += 1
                continue
            sv, av = S[k]["v"].get(f), A[k][f]
            ok = (isinstance(sv, datetime.datetime) and sv == av) if f == "Order_Date" else \
                 (is_num(sv) and sv == av) if f == "Quantity" else (isinstance(sv, str) and sv == av)
            if ok:
                st["da_sua_dung"] += 1
            else:
                bad_orig = (isinstance(sv, str) and f in ("Quantity", "Order_Date")) or \
                           (name == "Quantity âm" and is_num(sv) and sv < 0) or \
                           (name == "Store_Code sai (ST5)" and sv == "ST5") or \
                           (name == "Product_Code thiếu số 0" and isinstance(sv, str) and len(sv) == 3)
                if bad_orig:
                    st["con_loi_nguyen"] += 1
                else:
                    st["sua_sai"] += 1
                st["chi_tiet_sai"].append([k[0], k[1], str(sv), str(av)])
        seeded[name] = st
    res["loi_cai_san"] = seeded

    # ---- chỉ tiêu Kiem_tra tính từ dữ liệu của học viên ----
    tong_rev = bk.total_rev()
    tong_qty = sum(r["v"]["Quantity"] for r in region if is_num(r["v"].get("Quantity")))
    tong_pct = 0.0
    for r in region:
        p = bk.num_pct(r, tong_rev)
        if p is not None:
            tong_pct += p
    dreg = [r for r in region if r["kind"] == "data"]  # các chỉ tiêu kiểu / hợp lệ chỉ xét dòng giao dịch
    st_bad = sum(1 for r in dreg if not isinstance(r["v"].get("Store_Code"), str) or
                 r["v"]["Store_Code"].lower() not in ans["stores"])
    pr_bad = sum(1 for r in dreg if not isinstance(r["v"].get("Product_Code"), str) or
                 r["v"]["Product_Code"].lower() not in ans["prods"])
    q_le0 = sum(1 for r in dreg if is_num(r["v"].get("Quantity")) and r["v"]["Quantity"] <= 0)
    q_text = sum(1 for r in dreg if isinstance(r["v"].get("Quantity"), str))
    d_text = sum(1 for r in dreg if isinstance(r["v"].get("Order_Date"), str))
    hv = [tong_rev, len(region), tong_qty, tong_pct, st_bad, pr_bad, q_le0, q_text, d_text]
    nguon = ["Σ Revenue (cột Revenue của học viên)", "số dòng trong Table (hoặc dưới tiêu đề)",
             "Σ Quantity", "Σ Pct_Share", "COUNTIF Store_Code ∉ Store", "COUNTIF Product_Code ∉ Product",
             "Quantity ≤ 0", "Quantity là text", "Order_Date là text"]
    kt = []
    for i, (ten, dap_an) in enumerate(ans["kiem_tra"]):
        val = hv[i] if i < len(hv) else None
        khop = val is not None and round(float(val), 6) == round(float(dap_an), 6)
        kt.append({"chi_tieu": ten, "hoc_vien": val, "dap_an": dap_an, "khop": bool(khop), "nguon": nguon[i]})
    res["kiem_tra"] = kt

    # ---- sheet Kiem_tra của học viên ----
    kt_sheet = next((s for s in bk.sheets if norm(s) == "kiemtra"), None)
    # Quyết định GV: tên sheet Kiem_tra / Kiem tra / Kiểm tra coi là giống nhau, không trừ điểm
    kt_name_state = "dung" if kt_sheet else None
    kt_info = {"sheet": kt_sheet}
    if kt_sheet:
        f5, v5 = bk.wf[kt_sheet]["C5"].value, bk.wv[kt_sheet]["C5"].value
        kt_info.update({"C5_cong_thuc": f5, "C5_gia_tri_luu": v5})
        if v5 is None and isinstance(f5, str) and f5.upper().startswith("=SUM("):
            kt_info["C5_gia_tri_tinh_lai"] = tong_rev
            bk.recalc.append("Kiem_tra!C5: không có giá trị đã lưu, lấy tổng Revenue tính lại")
        # ô khác có hằng số gõ tay trong khối chỉ tiêu
        tay = []
        for row in bk.wf[kt_sheet].iter_rows():
            for c in row:
                if is_num(c.value) and c.column_letter in ("C", "D") and c.coordinate != "C5":
                    tay.append(f"{c.coordinate}={c.value}")
        kt_info["o_so_go_tay_khac"] = tay[:10]
        # tìm ô nào chứa tổng doanh thu trong sheet
        found = []
        for row in bk.wv[kt_sheet].iter_rows():
            for c in row:
                if is_num(c.value) and abs(c.value - tong_rev) < 1 and tong_rev:
                    found.append(c.coordinate)
        kt_info["o_chua_tong_dt"] = found
    res["sheet_kiem_tra"] = kt_info

    # ---- Việc 1: cấu trúc ----
    ev1 = []
    ok_hdr = bk.hdr_row == 1
    if not ok_hdr:
        ev1.append(f"Dòng tiêu đề ở dòng {bk.hdr_row}, phía trên còn {len(bk.above)} dòng: "
                   + "; ".join(f"dòng {i}: {t}" for i, t in bk.above[:3]))
    ten_junk = {"group": "dòng 'Cửa hàng …'", "sub": "dòng 'Cộng …'", "total": "dòng TỔNG CỘNG",
                "dup_header": "dòng tiêu đề lặp", "blank": "dòng trống giữa dữ liệu", "other": "dòng lạ"}
    for k, n in junk.items():
        ev1.append(f"Còn {n} {ten_junk[k]} (dòng " + ", ".join(str(r['row']) for r in rows if r['kind'] == k)[:60] + ")")
    if not junk and ok_hdr:
        v1 = "Đạt"
        ev1.append(f"Tiêu đề ở dòng 1, {len(data_rows)} dòng dữ liệu liền nhau, không còn dòng rác")
    elif not junk and not ok_hdr:
        v1 = "Một phần"
    else:
        tong_junk = sum(junk.values())
        v1 = "Chưa đạt" if (len(junk) >= 2 or tong_junk >= 6) else "Một phần"
    # dòng dữ liệu bị mất
    if thieu and not sai_bo:
        ev1.append(f"Thiếu {len(thieu)} dòng dữ liệu so với đáp án (đã xóa nhầm?)")
    if v1 == "Đạt" and thieu and not sai_bo:
        v1 = "Một phần"

    # ---- Việc 2: kiểu dữ liệu ----
    ev2 = [f"Quantity còn dạng text: {q_text} ô; Order_Date còn dạng text: {d_text} ô (đề cài sẵn 10 + 12)"]
    sw = lech["Order_Date"]
    if sw and not sai_bo:
        ev2.append(f"Order_Date lệch đáp án {len(sw)} ô, ví dụ {sw[0][0]} dòng {sw[0][1]}: {sw[0][2]} vs {sw[0][3]}")
    non_dt = sum(1 for r in region if r["kind"] == "data" and not isinstance(r["v"].get("Order_Date"), datetime.datetime))
    if q_text == 0 and d_text == 0 and not (sw and not sai_bo):
        v2 = "Đạt"
    elif (q_text + d_text) < 22:
        v2 = "Một phần"
    else:
        v2 = "Chưa đạt"
    if q_text == 0 and d_text == 0 and sw and not sai_bo:
        v2 = "Một phần"

    # ---- Việc 3: Table ----
    ev3 = []
    if bk.table:
        tn = bk.table["name"]
        ev3.append(f"Table '{tn}' phạm vi {bk.table['ref']} trên sheet '{bk.ws_name}'"
                   + (" (bật Total Row)" if bk.table["totals_row"] else ""))
        extra = []
        for label, act, exp in (("bảng", tn, "tblSales_T4"), ("sheet", bk.ws_name, "Sales")):
            stt = name_state(act, exp)
            if stt == "sai":
                extra.append(f"tên {label} '{act}' (đề: {exp})")
            elif stt == "hoa_thuong":
                ev3.append(f"Tên {label} '{act}' chỉ khác chữ hoa/thường so với đề ({exp}), không trừ")
        n_tbl_rows = bk.table["max_row"] - bk.table["min_row"]
        if n_tbl_rows != len(A) and not sai_bo:
            extra.append(f"Table có {n_tbl_rows} dòng dữ liệu, đúng là {len(A)}")
        if bk.table["min_row"] != 1:
            ev3.append(f"Table bắt đầu từ dòng {bk.table['min_row']} (hệ quả của việc 1, không tính lại)")
        ev3 += extra
        v3 = "Đạt" if not extra else "Một phần"
    else:
        v3 = "Chưa đạt"
        ev3.append("Không thấy Excel Table trên sheet dữ liệu")

    # ---- Việc 4: Revenue ----
    ev4 = []
    rev_hdr = bk.hdr_text.get(bk.cols.get("Revenue")) if "Revenue" in bk.cols else None
    if rev_hdr is None:
        v4 = "Chưa đạt"
        ev4.append("Không có cột Revenue")
    else:
        nform = nconst = nwrong = nformula_ok = 0
        for r in region:
            if r["kind"] != "data":
                continue
            f = r["f"].get("Revenue")
            q, p = r["v"].get("Quantity"), r["v"].get("Unit_Price")
            if isinstance(f, str) and f.startswith("="):
                nform += 1
                g = f.replace(" ", "")
                uses = ("[@Quantity]" in g and "[@Unit_Price]" in g) or \
                       ("[#ThisRow],[Quantity]" in g and "[#ThisRow],[Unit_Price]" in g) or \
                       bool(re.search(r"[A-Z]+\d+\*[A-Z]+\d+", g))
                if uses:
                    nformula_ok += 1
                rv = r["v"].get("Revenue")
                if is_num(rv) and is_num(q) and is_num(p) and abs(rv - q * p) > 0.5:
                    nwrong += 1
            else:
                nconst += 1
        ev4.append(f"Cột '{rev_hdr}': {nform} ô công thức ({nformula_ok} đúng dạng Quantity*Unit_Price), "
                   f"{nconst} ô gõ số/để trống, {nwrong} ô giá trị lệch Quantity*Unit_Price")
        if rev_hdr != "Revenue":
            ev4.append(f"Tên cột '{rev_hdr}' (đề: Revenue) – chỉ khác chữ hoa/thường")
        calc = None
        if bk.table:
            for c in bk.table["table"].tableColumns:
                if norm(c.name) == "revenue" and c.calculatedColumnFormula is not None:
                    calc = c.calculatedColumnFormula.attr_text
        ev4.append(f"Calculated column của Table: {calc}")
        if nform == 0:
            v4 = "Chưa đạt"
        elif nconst == 0 and nwrong == 0 and nformula_ok == nform:
            v4 = "Đạt"
        else:
            v4 = "Một phần"

    # ---- Việc 5: Kiem_tra!C5 và Pct_Share ----
    ev5 = []
    c5_ok = False
    if kt_sheet:
        f5 = kt_info.get("C5_cong_thuc")
        v5 = kt_info.get("C5_gia_tri_luu")
        if v5 is None:
            v5 = kt_info.get("C5_gia_tri_tinh_lai")
        is_f = isinstance(f5, str) and f5.startswith("=")
        okv = is_num(v5) and round(v5, 6) == round(ans["kiem_tra"][0][1], 6)
        ev5.append(f"Sheet '{kt_sheet}'!C5 = {f5!r} → {v5}")
        if kt_name_state == "sai":
            ev5.append(f"Tên sheet '{kt_sheet}' khác đề (Kiem_tra)")
        if is_f and okv:
            c5_ok = True
        elif not is_f:
            ev5.append("C5 là số gõ tay, không phải công thức tính tổng doanh thu")
        elif not okv:
            ev5.append(f"C5 không ra tổng doanh thu đúng ({ans['kiem_tra'][0][1]})")
    else:
        ev5.append("Không có sheet Kiem_tra")
        # tổng doanh thu đặt ở đâu?
    pct_hdr = bk.hdr_text.get(bk.cols.get("Pct_Share")) if "Pct_Share" in bk.cols else None
    pct_ok = False
    pinfo = None
    if pct_hdr is None:
        ev5.append("Không có cột Pct_Share")
    else:
        sample = None
        n_formula = n_abs = 0
        forms = {}
        for r in region:
            if r["kind"] != "data":
                continue
            f = r["f"].get("Pct_Share")
            info = pct_formula_info(bk, f, tong_rev)
            if info["la_cong_thuc"]:
                n_formula += 1
                forms[f] = forms.get(f, 0) + 1
                if info["refs"] and all(x["abs"] for x in info["refs"]):
                    n_abs += 1
                if sample is None:
                    sample = (f, info)
        nd = len([r for r in region if r["kind"] == "data"])
        ev5.append(f"Cột '{pct_hdr}': {n_formula}/{nd} ô là công thức, {n_abs} ô có tham chiếu tuyệt đối ($) tới ô tổng; "
                   f"{len(forms)} dạng công thức khác nhau")
        target = None
        if sample:
            f, info = sample
            ev5.append(f"Công thức mẫu: {f}")
            if info["refs"]:
                t = info["refs"][0]
                tf, tv = resolve_cell(bk, t["sheet"], t["col"], t["row"])
                if tv is None and isinstance(tf, str) and re.match(r"^=SUM\(.*(Revenue|J|J:J)", tf, re.I):
                    tv = tong_rev
                    bk.recalc.append(f"{t['sheet']}!{t['col']}{t['row']}: không có giá trị đã lưu, lấy tổng Revenue tính lại")
                target = {"o": f"{t['sheet']}!{t['col']}{t['row']}", "cong_thuc": tf, "gia_tri": tv}
                ev5.append(f"Ô tổng được tham chiếu: {target['o']} = {tf!r} → {tv}")
        pinfo = {"n_formula": n_formula, "n_abs": n_abs, "target": target, "dang_cong_thuc": len(forms)}
        pct_data = sum(x for x in (bk.num_pct(r, tong_rev) for r in dreg) if x is not None)
        tp_ok = abs(pct_data - 1) < 1e-6
        ev5.append(f"Tổng cột Pct_Share (các dòng giao dịch) = {pct_data:.6f}")
        rev_data = sum(x for x in (bk.num_rev(r) for r in dreg) if x is not None)
        tgt_ok = bool(target and is_num(target["gia_tri"]) and abs(target["gia_tri"] - rev_data) < 0.5)
        if target and not tgt_ok:
            ev5.append(f"Ô tổng ({target['gia_tri']}) không bằng tổng Revenue các dòng giao dịch ({rev_data})")
        pct_ok = n_formula == nd and n_abs == nd and tp_ok and tgt_ok
    res["pct_info"] = pinfo
    if c5_ok and pct_ok and kt_name_state == "sai":
        v5r = "Một phần"
    elif c5_ok and pct_ok:
        v5r = "Đạt"
    elif pct_hdr is None and not c5_ok:
        v5r = "Chưa đạt"
    elif not kt_sheet and pct_ok:
        v5r = "Một phần"
    elif (c5_ok or pct_ok):
        v5r = "Một phần"
    else:
        v5r = "Chưa đạt" if (pct_hdr is None or not (pinfo and pinfo["n_formula"])) else "Một phần"

    # ---- Việc 6: Data Validation ----
    dvs = read_dv(path, bk.ws_name)
    ev6 = []
    dv_res = {}
    want = {"Store_Code": ("list", "Store"), "Product_Code": ("list", "Product"), "Quantity": ("whole", None)}
    data_rows_set = {r["row"] for r in rows if r["kind"] == "data"}
    for col, (typ, src) in want.items():
        ci = bk.cols.get(col)
        found = None
        for d in dvs:
            rws = sqref_rows_for_col(d["sqref"], ci) if ci else set()
            if rws:
                found = (d, rws)
                break
        if not found:
            dv_res[col] = {"co": False}
            ev6.append(f"{col}: chưa có Data Validation")
            continue
        d, rws = found
        cover = len(rws & data_rows_set)
        item = {"co": True, "type": d["type"], "operator": d["operator"], "formula1": d["formula1"],
                "formula2": d["formula2"], "sqref": d["sqref"], "phu_dong_du_lieu": f"{cover}/{len(data_rows_set)}",
                "thong_bao": d["error"], "tieu_de_thong_bao": d["errorTitle"],
                "show_error": d["showErrorMessage"], "nguon_xml": d["nguon"], "an_mui_ten": d["showDropDown"],
                "input_message": d.get("prompt"), "van_de": []}
        if typ == "list":
            if d["type"] != "list":
                item["van_de"].append(f"loại '{d['type']}', cần List")
            else:
                si = dv_source_info(bk, d["formula1"])
                item["nguon_list"] = si
                if si["loai"] == "danh sách gõ tay":
                    item["van_de"].append("List gõ tay, không lấy từ sheet " + src)
                elif si["sheet"] is None or norm(si["sheet"]) != norm(src):
                    item["van_de"].append(f"List lấy từ '{si['sheet']}' ({d['formula1']}), đề yêu cầu sheet {src}")
        else:
            op = d["operator"] or "between"
            f1 = (d["formula1"] or "").strip()
            okq = d["type"] == "whole" and ((op == "greaterThan" and f1 == "0") or
                                           (op == "greaterThanOrEqual" and f1 == "1") or
                                           (op == "between" and f1 == "1"))
            if not okq:
                item["van_de"].append(f"điều kiện {d['type']}/{op}/{f1}, cần Whole number > 0")
        if cover < len(data_rows_set):
            item["van_de"].append(f"mới phủ {cover}/{len(data_rows_set)} dòng dữ liệu")
        msg = d["error"] or ""
        if typ == "list" and d["showDropDown"] in ("1", "true"):
            item["van_de"].append("đã ẩn mũi tên thả xuống (bỏ tick In-cell dropdown)")
        if not msg.strip() and (d.get("prompt") or "").strip():
            item["van_de"].append(f"nội dung nằm ở Input Message ({d['prompt']!r}), chưa có ở Error Alert")
        if d["showErrorMessage"] not in ("1", "true"):
            item["van_de"].append("chưa bật Show error alert (showErrorMessage)")
        if not msg.strip():
            item["van_de"].append("chưa có nội dung thông báo lỗi")
        elif not VI_CHARS.search(msg):
            item["van_de"].append("thông báo lỗi không có dấu tiếng Việt (cần GV xem)")
        dv_res[col] = item
        ev6.append(f"{col}: {d['type']} {d['operator'] or ''} {d['formula1']} | sqref {d['sqref']} | thông báo: {msg!r}"
                   + (" | VẤN ĐỀ: " + "; ".join(item["van_de"]) if item["van_de"] else ""))
    n_have = sum(1 for v in dv_res.values() if v["co"])
    n_clean = sum(1 for v in dv_res.values() if v["co"] and not v["van_de"])
    v6 = "Đạt" if n_clean == 3 else ("Chưa đạt" if n_have == 0 else "Một phần")
    res["data_validation"] = dv_res
    res["so_dv_trong_sheet"] = len(dvs)

    # ---- Việc 7: Circle Invalid Data ----
    con_sai = st_bad + pr_bad + q_le0  # Quantity text đã tính ở việc 2 (một lỗi gốc chỉ tính một lần)
    ev7 = [f"Còn {st_bad} Store_Code ngoài danh mục, {pr_bad} Product_Code ngoài danh mục, "
           f"{q_le0} Quantity ≤ 0" + (f"; còn {q_text} ô Quantity dạng text vẫn bị khoanh (đã tính ở việc 2)" if q_text else "")]
    sa = {k: v for k, v in seeded.items() if k in ("Store_Code sai (ST5)", "Product_Code thiếu số 0", "Quantity âm")}
    tong7 = sum(v["tong"] for v in sa.values())
    ok7 = sum(v["da_sua_dung"] for v in sa.values())
    ev7.append(f"Đối chiếu 7 ô cài sẵn: {ok7}/{tong7} đã sửa đúng giá trị; "
               + "; ".join(f"{k}: xóa dòng {v['xoa_dong']}, còn nguyên {v['con_loi_nguyen']}, sửa sai {v['sua_sai']}"
                           for k, v in sa.items() if v["xoa_dong"] or v["con_loi_nguyen"] or v["sua_sai"]) or "")
    if con_sai == 0 and ok7 == tong7:
        v7 = "Đạt"
    elif ok7 == tong7:
        v7 = "Một phần"  # đúng giá trị nhưng còn ô sai khác
    elif con_sai == 0:
        v7 = "Một phần"  # hết ô bị khoanh nhưng giá trị không giống đáp án (xóa dòng / đổi sai)
    else:
        v7 = "Chưa đạt" if ok7 == 0 else "Một phần"
    if sai_bo:
        v7 = "Không kiểm được"
        ev7.append("Bài làm trên bộ dữ liệu khác với đề (khóa dòng không khớp), không đối chiếu được từng ô cài sẵn")

    # ---- Việc 8: Insight ----
    ins_sheet = next((s for s in bk.sheets if norm(s).startswith("ins")), None)
    texts = []
    if ins_sheet:
        for row in bk.wv[ins_sheet].iter_rows():
            for c in row:
                if c.value is not None and str(c.value).strip():
                    texts.append(str(c.value).strip())
    ins_text = " ".join(texts)
    nums = [n for n in (tonum(x) for x in re.findall(r"\d[\d.,]*", ins_text)) if n is not None]
    sure = {1919: "số dòng đúng", 3448: "tổng Quantity đúng", 296965000: "tổng doanh thu đúng",
            10309: "SUM Quantity khi chưa sạch", 13: "13 dòng cộng", 10: "10 loại lỗi (Dap_an_loi)"}
    hints = {"con_so": bool(nums), "nguyen_nhan": bool(re.search(r"nguy hiểm|nguyên nhân|vì |do |lỗi|bởi", ins_text, re.I)),
             "ham_y": bool(re.search(r"nếu không|nên |sẽ |dẫn đến|ảnh hưởng|sai theo", ins_text, re.I))}
    res["insight"] = {"sheet": ins_sheet, "ten_sheet": name_state(ins_sheet, "Insight") if ins_sheet else None, "van_ban": ins_text, "cac_so_trong_cau": nums,
                      "so_khop_dap_an": {str(n): sure[n] for n in nums if n in sure},
                      "so_lech": [n for n in nums if n not in sure],
                      "goi_y_may_thanh_phan": hints,
                      "luu_y": "Máy chỉ gợi ý; người chấm đọc nội dung, tra lại con số và quyết định mức D4"}
    if not ins_sheet or not ins_text:
        v8 = "Chưa đạt"
        ev8 = ["Không có sheet Insight hoặc sheet trống"]
    else:
        v8 = "Cần chấm tay"
        ev8 = [f"Sheet '{ins_sheet}': {ins_text[:300]}"]
        if name_state(ins_sheet, "Insight") == "sai":
            ev8.append(f"Tên sheet '{ins_sheet}' khác đề (Insight)")

    viec = [(1, v1, ev1), (2, v2, ev2), (3, v3, ev3), (4, v4, ev4), (5, v5r, ev5), (6, v6, ev6), (7, v7, ev7),
            (8, v8, ev8)]
    if sai_bo:
        ghi = ("Bài làm trên bộ dữ liệu khác với đề (không khớp khóa Order_ID + Line_No nào với đáp án); "
               "các kết quả dưới đây chấm trên dữ liệu học viên dùng")
        for _, _, e in viec:
            e.append(ghi)
    res["viec"] = [{"id": i, "ten": VIEC_TEN[i], "ket_qua": k, "bang_chung": e} for i, k, e in viec]
    res["tong_so"] = {"tong_dt": tong_rev, "so_dong_table": len(region), "tong_qty": tong_qty, "tong_pct": tong_pct}
    res["junk"] = junk
    res["ghi_chu_tinh_lai"] = bk.recalc

    # ---- điểm (thang tạm, chờ GV xác nhận: việc 1-7 x 1 + insight 3 = 10) ----
    kq = [v for _, v, _ in viec]
    res["chi_tieu_lech"] = [x["chi_tieu"] for x in kt if not x["khop"]]
    res["diem_viec"] = {str(i): diem_viec(k) for i, k, _ in viec if i != 8}
    res["diem_viec_1_7"] = round(sum(v for v in res["diem_viec"].values() if v is not None), 4)
    res["viec7_khong_kiem_duoc"] = any(k == "Không kiểm được" for k in kq)

    # ---- dấu hiệu bất thường ----
    bt = []
    for s in bk.sheets:
        if norm(s) in ("dapanloi", "insightmau", "dapan") or "dapan" in norm(s):
            bt.append(f"Có sheet '{s}' của file đáp án")
    for s in bk.wv.worksheets:
        a1 = s["A1"].value
        if isinstance(a1, str) and "ĐÁP ÁN" in a1.upper() and s.title != "De_bai":
            bt.append(f"Sheet '{s.title}' có tiêu đề 'ĐÁP ÁN'")
    res["bat_thuong"] = bt
    return res


# --------------------------------------------------------------------------------------
def check_seed(de_path, ans_path):
    """So file đề (chưa sạch) với _Answer, in số lỗi cài sẵn tìm được và so với SEEDED."""
    ans = load_answer(ans_path)
    wb = openpyxl.load_workbook(de_path)
    ws = next(w for w in wb.worksheets if find_header(w))
    hdr, cols, _ = find_header(ws)
    found = {k: [] for k in SEEDED}
    junk = {}
    qsum = 0
    for row in ws.iter_rows(min_row=hdr + 1):
        g = lambda n: row[cols[n] - 1].value
        oid = g("Order_ID")
        kind = classify(oid, [g(n) for n in cols])
        if is_num(g("Quantity")):
            qsum += g("Quantity")
        if kind != "data":
            junk[kind] = junk.get(kind, 0) + 1
            continue
        key = (oid, g("Line_No"))
        a = ans["rows"][key]
        q = g("Quantity")
        if isinstance(q, str):
            found["Quantity dạng text"].append([key[0], key[1]])
        elif is_num(q) and q < 0:
            found["Quantity âm"].append([key[0], key[1]])
        if isinstance(g("Order_Date"), str):
            found["Order_Date dạng text dd/mm/yyyy"].append([key[0], key[1]])
        if g("Store_Code") != a["Store_Code"]:
            found["Store_Code sai (ST5)"].append([key[0], key[1]])
        if g("Product_Code") != a["Product_Code"]:
            found["Product_Code thiếu số 0"].append([key[0], key[1]])
    out = {"junk": junk, "sum_Quantity_cot_H_chua_sach": qsum, "khop_SEEDED": {}}
    for k in SEEDED:
        out["khop_SEEDED"][k] = {"tim_thay": len(found[k]), "dung_voi_SEEDED": sorted(map(tuple, found[k])) ==
                                 sorted(map(tuple, SEEDED[k]))}
        if not out["khop_SEEDED"][k]["dung_voi_SEEDED"]:
            out["khop_SEEDED"][k]["thuc_te"] = found[k]
    return out


def similarity(results):
    """Đối chiếu độ giống của câu insight giữa các bài (chỉ để báo GV, không kết luận)."""
    items = [(r["file"], r.get("insight", {}).get("van_ban", "")) for r in results if r.get("mo_duoc")]
    out = []
    for i in range(len(items)):
        for j in range(i + 1, len(items)):
            a, b = items[i][1], items[j][1]
            if a and b:
                ratio = difflib.SequenceMatcher(None, a.lower(), b.lower()).ratio()
                if ratio >= 0.6:
                    out.append({"a": items[i][0], "b": items[j][0], "ti_le_giong": round(ratio, 2)})
    return out


def main(argv):
    if len(argv) >= 4 and argv[1] == "--folder":
        ans = load_answer(argv[3])
        files = sorted(glob.glob(os.path.join(argv[2], "**", "*.xlsx"), recursive=True))
        files = [f for f in files if not os.path.basename(f).startswith("~$")]
        results = [grade(f, argv[3], ans) for f in files]
        if len(argv) >= 6 and argv[4] == "--insight-file":  # JSON {ten_file: so_thanh_phan_0_3}
            tay = json.load(open(argv[5], encoding="utf8"))
            for r in results:
                if r.get("mo_duoc") and r["file"] in tay:
                    ap_dung_insight(r, tay[r["file"]])
        names = {}
        for r in results:
            nm = r.get("nguoi_sua_cuoi")
            names.setdefault(nm, []).append(r["file"])
        out = {"bai": results, "giong_nhau_insight": similarity(results), "nguoi_sua_cuoi": names}
    elif False:
        pass
    elif len(argv) >= 4 and argv[1] == "--check-seed":
        out = check_seed(argv[2], argv[3])
    elif len(argv) >= 3:
        out = grade(argv[1], argv[2])
    else:
        print(__doc__)
        return 2
    sys.stdout.reconfigure(encoding="utf-8")
    print(json.dumps(out, ensure_ascii=False, indent=1, default=str))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
