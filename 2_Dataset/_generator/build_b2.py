"""Dựng file Excel của Buổi 2 (Hàm và KPI): đáp án thực hành, file demo (bản giảng viên + _HocVien), bài tập về nhà.

Chỉ đọc file đã có: MCI_Mart_Master.xlsx và B1_MCI_Mart_Q1_2025_Answer.xlsx (file khởi đầu B2). Không cần thư mục build/.
Mọi số liệu tính từ dataset; con số dùng trong kịch bản in ra cuối script (và lưu vào build_b2_meta.json cạnh script).

    python _generator/build_b2.py
"""
import json, os, random
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = os.path.normpath(os.environ.get("MCI_ROOT") or os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
MASTER = f"{ROOT}/2_Dataset/MCI_Mart_Master.xlsx"
B1_ANS = f"{ROOT}/1_Sessions/B1_Nen_tang_Excel_Du_lieu/Practice/B1_MCI_Mart_Q1_2025_Answer.xlsx"
B2 = f"{ROOT}/1_Sessions/B2_Ham_va_KPI"
PRACTICE, DEMO, HW = f"{B2}/Practice", f"{B2}/Demo", f"{B2}/Homework"
for d in (PRACTICE, DEMO, HW): os.makedirs(d, exist_ok=True)

F = "Arial"
HDR_FONT = Font(name=F, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="203864")
TITLE = Font(name=F, size=14, bold=True, color="203864")
BODY = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
INPUT = Font(name=F, size=10, color="0000FF")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
NOTE = Font(name=F, size=10, italic=True, color="585858")
DATE_FMT, NUM_FMT, PCT_FMT = "dd/mm/yyyy", "#,##0", "0.0%"
MONTH_FMT = '"Tháng "0'
vn = lambda x: f"{x:,.0f}".replace(",", ".")
pct = lambda x, d=1: f"{x * 100:.{d}f}%".replace(".", ",")

def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

def header(ws, row, labels, col=1):
    for j, h in enumerate(labels, start=col):
        c = ws.cell(row=row, column=j, value=h); c.font = HDR_FONT; c.fill = HDR_FILL

def py(v):
    if isinstance(v, pd.Timestamp): return v.to_pydatetime()
    if isinstance(v, float) and np.isnan(v): return None
    return v.item() if hasattr(v, "item") else v

def write_rows(ws, df, start_row=2, formats=None, font=BODY):
    for i, row in enumerate(df.itertuples(index=False), start=start_row):
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=py(v)); c.font = font
            if formats and df.columns[j - 1] in formats: c.number_format = formats[df.columns[j - 1]]

def add_table(ws, name, ncols, nrows, start_row=1):
    t = Table(displayName=name, ref=f"A{start_row}:{get_column_letter(ncols)}{start_row + nrows}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)

def table_df(ws, df, name, formats=None, widths=None):
    header(ws, 1, list(df.columns)); write_rows(ws, df, 2, formats)
    add_table(ws, name, len(df.columns), len(df)); ws.freeze_panes = "A2"
    if widths: set_widths(ws, widths)

def guide(ws, title, steps):
    ws["A1"] = title; ws["A1"].font = TITLE
    header(ws, 3, ["Slide", "Thao tác", "Kết quả mong đợi"])
    for i, (slide, action, result) in enumerate(steps, start=4):
        for j, v in enumerate((slide, action, result), start=1):
            c = ws.cell(row=i, column=j, value=v); c.font = BODY
            c.alignment = Alignment(wrap_text=True, vertical="top")
    set_widths(ws, [8, 72, 50])

FORBIDDEN_HV = ("trước lớp", "cho lớp", "giảng viên", "demo", "Demo_")

def save_both(wb, name, work_sheet):
    """Bản giảng viên: mở ra ở sheet làm việc, Huong_dan để cuối. Bản _HocVien: Huong_dan đứng đầu, tiêu đề VÍ DỤ,
    sheet Demo_* đổi thành Vi_du_*, không còn chữ dành cho giảng viên (kiểm tra bằng assert)."""
    g = wb["Huong_dan"]
    wb.move_sheet(g, offset=len(wb.sheetnames) - 1 - wb.sheetnames.index("Huong_dan"))
    wb.active = wb.sheetnames.index(work_sheet)
    for ws in wb.worksheets: ws.sheet_view.tabSelected = ws.title == work_sheet
    wb.save(f"{DEMO}/{name}.xlsx")
    g["A1"] = g["A1"].value.replace("DEMO B2", "VÍ DỤ B2")
    renames = {ws.title: "Vi_du_" + ws.title[5:] for ws in wb.worksheets if ws.title.startswith("Demo_")}
    for ws in wb.worksheets:
        if ws.title in renames: ws.title = renames[ws.title]
    for row in g.iter_rows(min_row=4):
        for c in row:
            if isinstance(c.value, str):
                for a, b in renames.items(): c.value = c.value.replace(a, b)
                low = c.value.lower()
                assert not any(w.lower() in low for w in FORBIDDEN_HV), (name, c.value)
    wb.move_sheet(g, offset=-wb.sheetnames.index("Huong_dan"))
    wb.active = 0
    for ws in wb.worksheets: ws.sheet_view.tabSelected = ws.title == "Huong_dan"
    wb.save(f"{DEMO}/{name}_HocVien.xlsx")

# ---------------------------------------------------------------------
# Dữ liệu
# ---------------------------------------------------------------------
products = pd.read_excel(MASTER, "Product")
stores = pd.read_excel(MASTER, "Store")
m_sales = pd.read_excel(MASTER, "Sales")
b1 = pd.read_excel(B1_ANS, "Sales", usecols="A:J")          # Region … Unit_Price (giá trị), đúng thứ tự dòng của file khởi đầu
N = len(b1)
cost = products.set_index("Product_Code")["Unit_Cost"]
cat = products.set_index("Product_Code")["Category"]
q1 = b1.copy()
q1["Revenue"] = q1.Quantity * q1.Unit_Price
q1["Category"] = q1.Product_Code.map(cat)
q1["Unit_Cost"] = q1.Product_Code.map(cost)
q1["COGS"] = q1.Quantity * q1.Unit_Cost
q1["Gross_Margin"] = q1.Revenue - q1.COGS
q1["Thang"] = q1.Order_Date.dt.month
q1["Phan_khuc"] = np.select([q1.Unit_Price >= 100000, q1.Unit_Price >= 30000], ["Cao", "Trung"], "Thấp")
REV, COGS = int(q1.Revenue.sum()), int(q1.COGS.sum())
GM = REV - COGS
ORDERS = int((q1.Line_No == 1).sum())
assert ORDERS == q1.Order_ID.nunique()
REGIONS, MONTHS = ["Bắc", "Trung", "Nam"], [1, 2, 3]
meta = {"q1": {"rows": N, "revenue": REV, "orders": ORDERS, "cogs": COGS, "gm": GM, "gm_pct": GM / REV,
               "aov": REV / ORDERS}}
grid = q1.groupby(["Region", "Thang"]).agg(rev=("Revenue", "sum"), cogs=("COGS", "sum"),
                                           orders=("Line_No", lambda s: int((s == 1).sum())))
meta["q1"]["grid"] = {f"{r}-{m}": {k: float(grid.loc[(r, m), k]) for k in grid.columns} for r in REGIONS for m in MONTHS}
bym = q1.groupby("Thang").agg(rev=("Revenue", "sum"), cogs=("COGS", "sum"), orders=("Line_No", lambda s: int((s == 1).sum())))
meta["q1"]["by_month"] = {int(m): {"rev": int(r.rev), "orders": int(r.orders), "aov": r.rev / r.orders, "gm_pct": 1 - r.cogs / r.rev}
                          for m, r in bym.iterrows()}
byr = q1.groupby("Region").agg(rev=("Revenue", "sum"), cogs=("COGS", "sum"), orders=("Line_No", lambda s: int((s == 1).sum())))
meta["q1"]["by_region"] = {r: {"rev": int(x.rev), "orders": int(x.orders), "aov": x.rev / x.orders, "gm_pct": 1 - x.cogs / x.rev}
                           for r, x in byr.iterrows()}
meta["q1"]["phan_khuc"] = q1.Phan_khuc.value_counts().to_dict()
byc = q1.groupby("Category").agg(rev=("Revenue", "sum"), cogs=("COGS", "sum"))
byc["gm_pct"] = 1 - byc.cogs / byc.rev
meta["q1"]["by_category_gm"] = byc.gm_pct.to_dict()
meta["q1"]["avg_of_category_gm"] = float(byc.gm_pct.mean())
p21 = q1[q1.Product_Code == "P021"]
OLD_COST = 62000
meta["p021"] = {"lines": len(p21), "qty": int(p21.Quantity.sum()), "cost_right": int(cost["P021"]), "cost_old": OLD_COST,
                "cogs_right": int(p21.COGS.sum()), "cogs_wrong": int(p21.Quantity.sum() * OLD_COST),
                "total_cogs_wrong": COGS - int(p21.Quantity.sum() * (cost["P021"] - OLD_COST))}
meta["p021"]["gm_pct_wrong"] = 1 - meta["p021"]["total_cogs_wrong"] / REV

# =====================================================================
# 1) PRACTICE — B2_MCI_Mart_Q1_2025_Answer.xlsx (brick output B2; file khởi đầu B3)
# =====================================================================
wb = Workbook(); ws = wb.active; ws.title = "Sales"
base_cols = ["Region", "Order_ID", "Line_No", "Order_Date", "Store_Code", "Customer_ID", "Product_Code", "Product_Name", "Quantity", "Unit_Price"]
new_cols = ["Revenue", "Pct_Share", "Category", "Unit_Cost", "COGS", "Gross_Margin", "Thang", "Quy", "Thu", "Phan_khuc"]
header(ws, 1, base_cols + new_cols)
TR = lambda c: f"tblSales[[#This Row],[{c}]]"
FORMULAS = {
    "Revenue": (f"={TR('Quantity')}*{TR('Unit_Price')}", NUM_FMT),
    "Pct_Share": (f"={TR('Revenue')}/Kiem_tra!$C$5", "0.000%"),
    "Category": (f"=_xlfn.XLOOKUP({TR('Product_Code')},tblProduct[Product_Code],tblProduct[Category])", None),
    "Unit_Cost": (f"=_xlfn.XLOOKUP({TR('Product_Code')},tblProduct[Product_Code],tblProduct[Unit_Cost])", NUM_FMT),
    "COGS": (f"={TR('Quantity')}*{TR('Unit_Cost')}", NUM_FMT),
    "Gross_Margin": (f"={TR('Revenue')}-{TR('COGS')}", NUM_FMT),
    "Thang": (f"=MONTH({TR('Order_Date')})", "0"),
    "Quy": (f'="Q"&ROUNDUP(MONTH({TR("Order_Date")})/3,0)', None),
    "Thu": (f'=TEXT({TR("Order_Date")},"ddd")', None),
    "Phan_khuc": (f'=_xlfn.IFS({TR("Unit_Price")}>=100000,"Cao",{TR("Unit_Price")}>=30000,"Trung",TRUE,"Thấp")', None),
}
for i, row in enumerate(b1.itertuples(index=False), start=2):
    for j, v in enumerate(row, start=1):
        c = ws.cell(row=i, column=j, value=py(v)); c.font = BODY
        if j == 4: c.number_format = DATE_FMT
        if j == 10: c.number_format = NUM_FMT
    for j, col in enumerate(new_cols, start=len(base_cols) + 1):
        f, fmt = FORMULAS[col]
        c = ws.cell(row=i, column=j, value=f); c.font = BODY
        if fmt: c.number_format = fmt
add_table(ws, "tblSales", len(base_cols) + len(new_cols), N)
ws.freeze_panes = "A2"
set_widths(ws, [9, 15, 7, 12, 10, 11, 11, 36, 9, 11, 12, 9, 16, 10, 12, 13, 7, 6, 6, 10])
dv = DataValidation(type="list", formula1="=Store!$A$2:$A$13", showErrorMessage=True,
                    errorTitle="Mã cửa hàng không hợp lệ", error="Chỉ nhập mã có trong sheet Store (ST01–ST12).")
dv.add(f"E2:E{N + 1}"); ws.add_data_validation(dv)
dv = DataValidation(type="whole", operator="greaterThan", formula1="0", showErrorMessage=True,
                    errorTitle="Số lượng không hợp lệ", error="Quantity phải là số nguyên lớn hơn 0.")
dv.add(f"I2:I{N + 1}"); ws.add_data_validation(dv)

ws = wb.create_sheet("Product")
table_df(ws, products, "tblProduct", {"Unit_Price": NUM_FMT, "Unit_Cost": NUM_FMT}, [13, 40, 18, 12, 12])
ws = wb.create_sheet("Store")
table_df(ws, stores, "tblStore", {"Open_Date": DATE_FMT}, [11, 24, 13, 9, 11, 12])

# --- KPI_SUMIFS: 4 khối Miền × Tháng, tham chiếu hỗn hợp $A / B$ ---
k = wb.create_sheet("KPI_SUMIFS")
k["A1"] = "BẢNG KPI THEO MIỀN × THÁNG — QUÝ 1/2025 (SUMIFS)"; k["A1"].font = TITLE
k["A2"] = "Mỗi ô = một bộ lọc: Miền (cột A, khóa cột $A) × Tháng (dòng tiêu đề, khóa dòng). Bảng này là “đáp án” để đối soát ở Buổi 3 và Buổi 5."
k["A2"].font = NOTE
blocks = [("Doanh thu (VNĐ)", "rev"), ("Số đơn", "orders"), ("Giá trị đơn trung bình – AOV (VNĐ)", "aov"), ("Gross Margin %", "gm")]
anchor = {}
r0 = 4
for title, key in blocks:
    k.cell(row=r0, column=1, value=title).font = BOLD
    h = r0 + 1
    header(k, h, ["Miền", 1, 2, 3, "Quý 1"])
    for j in range(2, 5): k.cell(row=h, column=j).number_format = MONTH_FMT
    anchor[key] = h
    for i, reg in enumerate(REGIONS + ["Tổng"], start=h + 1):
        k.cell(row=i, column=1, value=reg).font = BOLD if reg == "Tổng" else BODY
        for j in range(2, 6):
            col = get_column_letter(j)
            if key == "rev":
                f = (f"=SUMIFS(tblSales[Revenue],tblSales[Region],$A{i},tblSales[Thang],{col}${h})" if reg != "Tổng" and j < 5
                     else (f"=SUM(B{i}:D{i})" if j == 5 else f"=SUM({col}{h + 1}:{col}{h + 3})"))
                fmt = NUM_FMT
            elif key == "orders":
                f = (f"=COUNTIFS(tblSales[Region],$A{i},tblSales[Thang],{col}${h},tblSales[Line_No],1)" if reg != "Tổng" and j < 5
                     else (f"=SUM(B{i}:D{i})" if j == 5 else f"=SUM({col}{h + 1}:{col}{h + 3})"))
                fmt = NUM_FMT
            elif key == "aov":
                ri, ro = anchor["rev"] + (i - h), anchor["orders"] + (i - h)
                f = f"={col}{ri}/{col}{ro}"; fmt = NUM_FMT
            else:
                ri = anchor["rev"] + (i - h)
                if reg != "Tổng" and j < 5:
                    f = f"=SUMIFS(tblSales[Gross_Margin],tblSales[Region],$A{i},tblSales[Thang],{col}${h})/{col}{ri}"
                elif reg != "Tổng":
                    f = f"=SUMIFS(tblSales[Gross_Margin],tblSales[Region],$A{i})/{col}{ri}"
                elif j < 5:
                    f = f"=SUMIFS(tblSales[Gross_Margin],tblSales[Thang],{col}${h})/{col}{ri}"
                else:
                    f = f"=SUM(tblSales[Gross_Margin])/{col}{ri}"
                fmt = PCT_FMT
            c = k.cell(row=i, column=j, value=f); c.font = BOLD if reg == "Tổng" else BODY; c.number_format = fmt
    r0 = h + 6
k.cell(row=r0, column=1, value="AOV và Gross Margin % của dòng Tổng tính bằng tổng ÷ tổng, KHÔNG lấy trung bình của các miền.").font = NOTE
set_widths(k, [34, 15, 15, 15, 16])
REV_TOTAL_CELL = f"KPI_SUMIFS!E{anchor['rev'] + 4}"
ORD_TOTAL_CELL = f"KPI_SUMIFS!E{anchor['orders'] + 4}"

# --- KPI_Card ---
CARD_FIELDS = ["Tên KPI", "Mục đích (trả lời câu hỏi gì)", "Công thức", "Nguồn dữ liệu", "Grain",
               "Dimensions", "Owner", "Target", "Ngưỡng Warning / Critical", "Ghi chú giới hạn"]
CARDS = [
    ["Total Revenue (Tổng doanh thu)", "Tháng này bán được bao nhiêu tiền, ở miền / cửa hàng nào?",
     "SUM(Quantity × Unit_Price) = SUM(tblSales[Revenue])", "tblSales (giá bán thực tế trong giao dịch)",
     "Dữ liệu: 1 dòng = 1 sản phẩm trong 1 đơn. KPI báo cáo theo tháng", "Miền, cửa hàng, ngành hàng, tháng",
     "Giám đốc bán hàng", "Theo bảng Target (Tháng × Cửa hàng) — dùng từ Buổi 4", "90% / 80% target",
     "Cộng dồn được theo mọi chiều. Chưa trừ trả hàng, khuyến mại."],
    ["Total Orders (Số đơn)", "Có bao nhiêu lượt mua hàng?",
     "Số Order_ID khác nhau = COUNTIFS(…, tblSales[Line_No], 1)", "tblSales (cột Order_ID, Line_No)",
     "Dữ liệu: 1 dòng = 1 sản phẩm trong 1 đơn, nên KHÔNG đếm số dòng. KPI báo cáo theo tháng", "Miền, cửa hàng, tháng",
     "Giám đốc bán hàng", "1.000 đơn/tháng", "90% / 80%",
     "Một đơn có nhiều dòng: đếm dòng Line_No = 1. Không lọc theo ngành hàng được (một đơn có nhiều ngành)."],
    ["Gross Margin % (Biên lợi nhuận gộp)", "Mỗi đồng doanh thu còn lại bao nhiêu sau giá vốn?",
     "(SUM(Revenue) − SUM(COGS)) ÷ SUM(Revenue); COGS = Quantity × Unit_Cost", "tblSales + tblProduct (Unit_Cost qua XLOOKUP)",
     "KPI tỷ lệ, báo cáo theo tháng", "Ngành hàng, miền, tháng", "Finance (Tài chính)", "33%", "90% / 80% (29,7% / 26,4%)",
     "Là tỷ lệ: tổng ÷ tổng, không lấy trung bình các tỷ lệ. Giá vốn lấy một bảng Product duy nhất (giả định)."],
]
kc = wb.create_sheet("KPI_Card")
kc["A1"] = "KPI DEFINITION CARD — 3 KPI của Buổi 2"; kc["A1"].font = TITLE
kc["A2"] = "Viết card TRƯỚC khi viết công thức. Mẫu đầy đủ: Handout/B2_KPI_Definition_Card.docx"; kc["A2"].font = NOTE
header(kc, 4, ["Trường"] + [f"Card {i}" for i in (1, 2, 3)])
for i, fld in enumerate(CARD_FIELDS, start=5):
    kc.cell(row=i, column=1, value=fld).font = BOLD
    for j, card in enumerate(CARDS, start=2):
        c = kc.cell(row=i, column=j, value=card[i - 5]); c.font = BODY; c.alignment = Alignment(wrap_text=True, vertical="top")
kc.cell(row=16, column=1, value="Giá trị Q1/2025").font = BOLD
for j, f in enumerate(["=SUM(tblSales[Revenue])", f"={ORD_TOTAL_CELL}", "=SUM(tblSales[Gross_Margin])/SUM(tblSales[Revenue])"], start=2):
    c = kc.cell(row=16, column=j, value=f); c.font = BOLD; c.number_format = PCT_FMT if j == 4 else NUM_FMT
set_widths(kc, [26, 42, 42, 42])

# --- Kiem_tra ---
kt = wb.create_sheet("Kiem_tra")
kt["A1"] = "KIỂM TRA / ĐỐI SOÁT — Brick output B2"; kt["A1"].font = TITLE
kt["A2"] = "Cột C = công thức tính trên tblSales; cột D = đáp án (tính từ dataset); cột E = khớp? C5 là ô tổng doanh thu mà cột Pct_Share tham chiếu ($C$5)."
kt["A2"].font = NOTE
header(kt, 4, ["Chỉ tiêu", "", "Kết quả (công thức)", "Đáp án", "Khớp?"])
checks = [
    ("Tổng doanh thu Q1/2025 (VNĐ)", "=SUM(tblSales[Revenue])", REV, NUM_FMT),
    ("Số dòng dữ liệu", "=ROWS(tblSales[Order_ID])", N, NUM_FMT),
    ("Số đơn (dòng có Line_No = 1)", "=COUNTIFS(tblSales[Line_No],1)", ORDERS, NUM_FMT),
    ("Tổng COGS (VNĐ)", "=SUM(tblSales[COGS])", COGS, NUM_FMT),
    ("Gross Margin % Q1/2025", "=SUM(tblSales[Gross_Margin])/SUM(tblSales[Revenue])", GM / REV, "0.00%"),
    ("Số mã bị trùng trong tblProduct", "=SUMPRODUCT(--(COUNTIF(tblProduct[Product_Code],tblProduct[Product_Code])>1))", 0, "0"),
    ("Số ô #N/A ở Category và Unit_Cost", "=SUMPRODUCT(--ISNA(tblSales[Category]))+SUMPRODUCT(--ISNA(tblSales[Unit_Cost]))", 0, "0"),
    ("Tổng bảng KPI doanh thu = tổng cột Revenue", f"={REV_TOTAL_CELL}", REV, NUM_FMT),
    ("Tổng bảng KPI số đơn", f"={ORD_TOTAL_CELL}", ORDERS, NUM_FMT),
    ("Tổng cột Pct_Share", "=SUM(tblSales[Pct_Share])", 1, "0.00%"),
]
for r, (label, f, exp, fmt) in enumerate(checks, start=5):
    kt.cell(row=r, column=1, value=label).font = BODY
    c = kt.cell(row=r, column=3, value=f); c.number_format = fmt; c.font = BODY
    c = kt.cell(row=r, column=4, value=exp); c.number_format = fmt; c.font = INPUT
    kt.cell(row=r, column=5, value=f'=IF(ROUND(C{r},6)=ROUND(D{r},6),"✔","✘")').font = BODY
kt["A16"] = (f"Tổng COGS ra {vn(meta['p021']['total_cogs_wrong'])} (thiếu {vn(COGS - meta['p021']['total_cogs_wrong'])}) = còn bản ghi P021 cũ "
             f"(Unit_Cost {vn(OLD_COST)}) đứng trước trong Product. Doanh thu vẫn đúng vì Unit_Price lấy từ Sales.")
kt["A16"].font = NOTE
set_widths(kt, [46, 2, 20, 16, 8])
wb.save(f"{PRACTICE}/B2_MCI_Mart_Q1_2025_Answer.xlsx")

# Product có khóa trùng (đề cho bài Laws of Lookups): dựng lại từ Master, bỏ ghi chú giảng viên ở G1
p_dup = products.copy()
dup = p_dup[p_dup.Product_Code == "P021"].copy()
dup["Unit_Price"] = 99000; dup["Unit_Cost"] = OLD_COST
dup["Product_Name"] = "Bánh quy bơ hộp thiếc 454g (mã cũ – ngừng KD 2024)"
p_dup = pd.concat([p_dup.iloc[:20], dup, p_dup.iloc[20:]], ignore_index=True)   # bản ghi cũ đứng TRƯỚC
wb = Workbook(); ws = wb.active; ws.title = "Product"
table_df(ws, p_dup, "tblProduct", {"Unit_Price": NUM_FMT, "Unit_Cost": NUM_FMT}, [13, 48, 18, 12, 12])
wb.save(f"{PRACTICE}/B2_Product_DuplicateKey.xlsx")

# =====================================================================
# 2) DEMO
# =====================================================================
# ---- B2_Demo_Logic.xlsx — slide 7 (IF, IFS), 8 (AND/OR/NOT), 9 (IFERROR)
wb = Workbook(); ws = wb.active; ws.title = "Demo_IF"
lg = products.copy()
new = pd.DataFrame([["P041", "Trà sữa đóng chai 350ml (hàng mới, chưa có giá)", "Đồ uống", None, 14000],
                    ["P042", "Nồi cơm điện 1,8L (hàng mới, chưa có giá)", "Gia dụng", None, 520000]], columns=lg.columns)
lg = pd.concat([lg, new], ignore_index=True)
cols = ["Product_Code", "Product_Name", "Category", "Unit_Price", "Unit_Cost", "Margin_%", "IF: Biên", "IFS: Phân khúc giá",
        "AND: Ưu tiên Tết", "OR: Thiết yếu", "NOT: Không thiết yếu", "IFERROR: Margin_%"]
header(ws, 1, cols)
L = len(lg)
for i, row in enumerate(lg.itertuples(index=False), start=2):
    for j, v in enumerate(row, start=1):
        c = ws.cell(row=i, column=j, value=py(v)); c.font = INPUT if j == 4 else BODY
        if j in (4, 5): c.number_format = NUM_FMT
    fs = [f"=(D{i}-E{i})/D{i}", f'=IF(F{i}>=35%,"Biên cao","Biên thấp")',
          f'=_xlfn.IFS(D{i}>=100000,"Cao",D{i}>=30000,"Trung",TRUE,"Thấp")',
          f'=AND(C{i}="Bánh kẹo",F{i}>=35%)', f'=OR(C{i}="Sữa",C{i}="Thực phẩm khô")', f"=NOT(J{i})",
          f'=IFERROR((D{i}-E{i})/D{i},"Chưa có giá")']
    for j, f in enumerate(fs, start=6):
        c = ws.cell(row=i, column=j, value=f); c.font = BODY
        if j in (6, 12): c.number_format = PCT_FMT
last = L + 1
ws.cell(row=last + 2, column=1, value="Đếm").font = BOLD
summary = [("Số SP “Biên cao”", f'=COUNTIF(G2:G{last},"Biên cao")'), ("Số ô lỗi ở cột Margin_%", f"=SUMPRODUCT(--ISERROR(F2:F{last}))"),
           ("Số SP phân khúc “Cao”", f'=COUNTIF(H2:H{last},"Cao")'), ("Số SP “Ưu tiên Tết” (TRUE)", f"=COUNTIF(I2:I{last},TRUE)"),
           ("Số SP thiết yếu (TRUE)", f"=COUNTIF(J2:J{last},TRUE)")]
for r, (lab, f) in enumerate(summary, start=last + 3):
    ws.cell(row=r, column=1, value=lab).font = BODY
    ws.cell(row=r, column=4, value=f).font = BOLD
ws.freeze_panes = "B2"
set_widths(ws, [12, 44, 16, 11, 11, 10, 11, 12, 12, 11, 13, 14])
lgm = (products.Unit_Price - products.Unit_Cost) / products.Unit_Price
cnt_high = int((lgm >= 0.35).sum())
cnt_cao = int((products.Unit_Price >= 100000).sum())
cnt_tet = int(((products.Category == "Bánh kẹo") & (lgm >= 0.35)).sum())
cnt_ty = int(products.Category.isin(["Sữa", "Thực phẩm khô"]).sum())
meta["logic"] = {"bien_cao": cnt_high, "bien_thap": 40 - cnt_high, "loi": 2, "phan_khuc_cao_40": cnt_cao,
                 "phan_khuc_cao_42": cnt_cao, "uu_tien_tet": cnt_tet, "thiet_yeu": cnt_ty,
                 "p001_margin": float(lgm.iloc[0]), "p041_row": last - 1, "p042_row": last}
g = wb.create_sheet("Huong_dan")
guide(g, "DEMO B2 — HÀM LOGIC IF / IFS / AND / OR / NOT / IFERROR (slide 7, 8, 9)", [
    (7, "Sheet Demo_IF. Cột F là Margin_% = (Giá − Vốn) ÷ Giá. Chọn G2:G43, nhấn Delete, gõ lại G2: =IF(F2>=35%,\"Biên cao\",\"Biên thấp\") rồi kéo xuống G43.",
     f"{cnt_high} sản phẩm “Biên cao”, {40 - cnt_high} “Biên thấp”; P041, P042 (dòng {last - 1}, {last}) ra #DIV/0!."),
    (7, "H2: =IFS(D2>=100000,\"Cao\",D2>=30000,\"Trung\",TRUE,\"Thấp\"), kéo xuống. Nhìn dòng P041, P042.",
     f"{cnt_cao} sản phẩm phân khúc “Cao”. P041, P042 chưa có giá mà vẫn ra “Thấp”: ô trống được tính là 0."),
    (8, "I2: =AND(C2=\"Bánh kẹo\",F2>=35%); J2: =OR(C2=\"Sữa\",C2=\"Thực phẩm khô\"); K2: =NOT(J2). Kéo xuống.",
     f"AND: {cnt_tet} sản phẩm TRUE (Bánh kẹo và biên ≥ 35%). OR: {cnt_ty} sản phẩm TRUE. NOT đảo ngược cột J."),
    (9, "L2: =IFERROR((D2-E2)/D2,\"Chưa có giá\"), kéo xuống. Thử đổi giá trị thay thế thành 0: =IFERROR((D2-E2)/D2,0).",
     "P041, P042 hiện “Chưa có giá”. Nếu thay bằng 0, hai sản phẩm hiện Margin 0% — con số sai mà trông như thật."),
    (9, "Gõ giá cho P041 ở D" + str(last - 1) + ": 25000.", "Mọi cột của dòng P041 tự tính lại; Margin_% = 44,0%."),
])
save_both(wb, "B2_Demo_Logic", "Demo_IF")

# ---- B2_Demo_ThongKe.xlsx — slide 11 (COUNT/AVERAGE/MEDIAN/MIN/MAX), 12 (STDEV, outlier), 13 (RANK)
orders = q1.groupby("Order_ID").agg(Region=("Region", "first"), Store_Code=("Store_Code", "first"),
                                    Order_Date=("Order_Date", "first"), Order_Value=("Revenue", "sum")).reset_index()
orders = orders.sort_values(["Order_Date", "Order_ID"]).reset_index(drop=True)
wb = Workbook(); ws = wb.active; ws.title = "Demo_ThongKe"
oc = ["Order_ID", "Region", "Store_Code", "Order_Date", "Order_Value"]
header(ws, 1, oc); write_rows(ws, orders[oc], 2, {"Order_Date": DATE_FMT, "Order_Value": NUM_FMT})
no = len(orders); rng = f"$E$2:$E${no + 1}"
header(ws, 1, ["Thống kê giá trị đơn (VNĐ)", "Công thức", "Kết quả"], col=7)
stats = [("Số đơn", f"=COUNT({rng})", "COUNT"), ("Trung bình (mean)", f"=AVERAGE({rng})", "AVERAGE"),
         ("Trung vị (median)", f"=MEDIAN({rng})", "MEDIAN"), ("Nhỏ nhất", f"=MIN({rng})", "MIN"), ("Lớn nhất", f"=MAX({rng})", "MAX"),
         ("Độ lệch chuẩn", f"=_xlfn.STDEV.S({rng})", "STDEV.S"), ("Độ lệch (skew)", f"=SKEW({rng})", "SKEW"),
         ("Tứ phân vị Q1", f"=_xlfn.QUARTILE.INC({rng},1)", "QUARTILE.INC"), ("Tứ phân vị Q3", f"=_xlfn.QUARTILE.INC({rng},3)", "QUARTILE.INC"),
         ("IQR = Q3 − Q1", "=I10-I9", "IQR"), ("Ngưỡng outlier trên = Q3 + 1,5 × IQR", "=I10+1.5*I11", ""),
         ("Số đơn vượt ngưỡng (outlier)", f'=COUNTIF({rng},">"&I12)', "COUNTIF")]
for r, (lab, f, fn) in enumerate(stats, start=2):
    ws.cell(row=r, column=7, value=lab).font = BODY
    ws.cell(row=r, column=8, value=fn).font = NOTE
    c = ws.cell(row=r, column=9, value=f); c.font = BOLD; c.number_format = "0.00" if fn == "SKEW" else NUM_FMT
ws.freeze_panes = "A2"
set_widths(ws, [15, 8, 10, 12, 13, 3, 36, 14, 16])
ov = orders.Order_Value
q1_, q3_ = ov.quantile(.25), ov.quantile(.75); fence = q3_ + 1.5 * (q3_ - q1_)
meta["thongke"] = {"n": no, "mean": float(ov.mean()), "median": float(ov.median()), "min": int(ov.min()), "max": int(ov.max()),
                   "stdev": float(ov.std()), "skew": float(ov.skew()), "q1": float(q1_), "q3": float(q3_), "fence": float(fence),
                   "n_outlier": int((ov > fence).sum()), "max_order": orders.loc[ov.idxmax(), "Order_ID"],
                   "trimmed_mean_no_outlier": float(ov[ov <= fence].mean())}
# RANK cửa hàng
rs = q1.groupby("Store_Code").Revenue.sum().reset_index().merge(stores[["Store_Code", "Store_Name", "Region"]], on="Store_Code")
rs = rs[["Store_Code", "Store_Name", "Region", "Revenue"]]
wr = wb.create_sheet("Demo_RANK")
header(wr, 1, ["Store_Code", "Store_Name", "Region", "Doanh thu Q1/2025", "Hạng (RANK.EQ)"])
write_rows(wr, rs, 2, {"Revenue": NUM_FMT})
for i in range(2, len(rs) + 2):
    c = wr.cell(row=i, column=5, value=f"=_xlfn.RANK.EQ(D{i},$D$2:$D$13,0)"); c.font = BODY
set_widths(wr, [11, 24, 9, 18, 15])
rs["rank"] = rs.Revenue.rank(ascending=False, method="min").astype(int)
meta["rank"] = rs.sort_values("rank")[["Store_Code", "Store_Name", "Revenue", "rank"]].values.tolist()
t = meta["thongke"]
g = wb.create_sheet("Huong_dan")
guide(g, "DEMO B2 — THỐNG KÊ MÔ TẢ (slide 11, 12, 13)", [
    (11, f"Sheet Demo_ThongKe: mỗi dòng là 1 đơn hàng Q1/2025 ({vn(no)} đơn). Xem khối I2:I6 (COUNT, AVERAGE, MEDIAN, MIN, MAX); F2 vào từng ô để đọc công thức.",
     f"Số đơn {vn(no)}; trung bình {vn(t['mean'])}; trung vị {vn(t['median'])}; nhỏ nhất {vn(t['min'])}; lớn nhất {vn(t['max'])}."),
    (11, "So sánh AVERAGE và MEDIAN: chọn cột E, nhìn Status Bar.",
     f"Trung bình cao hơn trung vị khoảng {pct(t['mean'] / t['median'] - 1, 0)}: vài đơn rất lớn kéo trung bình lên."),
    (12, "Xem I7:I13: độ lệch chuẩn, skew, Q1, Q3, IQR, ngưỡng outlier, số đơn vượt ngưỡng. Sort cột E giảm dần (Data → Sort Z→A) để thấy các đơn lớn nhất, xong Ctrl+Z.",
     f"STDEV {vn(t['stdev'])}; skew {t['skew']:.2f}".replace(".", ",") + f"; ngưỡng {vn(t['fence'])}; {t['n_outlier']} đơn outlier; đơn lớn nhất {t['max_order']} = {vn(t['max'])}."),
    (13, "Sheet Demo_RANK: E2 = RANK.EQ(D2,$D$2:$D$13,0), kéo xuống E13 (nhớ khóa vùng bằng F4).",
     "Hạng 1 = " + meta["rank"][0][0] + f" ({vn(meta['rank'][0][2])}); hạng 12 = " + meta["rank"][-1][0] + f" ({vn(meta['rank'][-1][2])})."),
])
save_both(wb, "B2_Demo_ThongKe", "Demo_ThongKe")

# ---- B2_Demo_SUMIFS.xlsx — slide 15 (xương sống Chuỗi C), 16 (SUMIFS/COUNTIFS/AVERAGEIFS), 17 (bảng 2 chiều, trung bình của trung bình)
wb = Workbook(); ws = wb.active; ws.title = "Sales"
sc = ["Region", "Order_ID", "Line_No", "Order_Date", "Store_Code", "Product_Code", "Category", "Quantity", "Unit_Price", "Revenue", "COGS", "Thang"]
table_df(ws, q1[sc], "tblSales", {"Order_Date": DATE_FMT, "Unit_Price": NUM_FMT, "Revenue": NUM_FMT, "COGS": NUM_FMT},
         [9, 15, 7, 12, 10, 11, 16, 9, 11, 12, 12, 7])
d = wb.create_sheet("Demo_SUMIFS")
d["A1"] = "CHUỖI C — MỖI TIÊU CHÍ LÀ MỘT BỘ LỌC"; d["A1"].font = TITLE
header(d, 3, ["Câu hỏi", "Miền", "Tháng", "Kết quả", "Hàm"])
single = [("Doanh thu Miền Bắc, tháng 3", "=SUMIFS(tblSales[Revenue],tblSales[Region],B4,tblSales[Thang],C4)", "SUMIFS"),
          ("Số dòng giao dịch Miền Bắc, tháng 3", "=COUNTIFS(tblSales[Region],B5,tblSales[Thang],C5)", "COUNTIFS"),
          ("Số đơn Miền Bắc, tháng 3 (Line_No = 1)", "=COUNTIFS(tblSales[Region],B6,tblSales[Thang],C6,tblSales[Line_No],1)", "COUNTIFS"),
          ("Doanh thu trung bình 1 dòng, Miền Bắc, tháng 3", "=AVERAGEIFS(tblSales[Revenue],tblSales[Region],B7,tblSales[Thang],C7)", "AVERAGEIFS")]
for r, (qq, f, fn) in enumerate(single, start=4):
    d.cell(row=r, column=1, value=qq).font = BODY
    for col, v in ((2, "Bắc"), (3, 3)):
        c = d.cell(row=r, column=col, value=v); c.font = INPUT; c.fill = INPUT_FILL
    c = d.cell(row=r, column=4, value=f); c.font = BOLD; c.number_format = NUM_FMT
    d.cell(row=r, column=5, value=fn).font = NOTE
d["A9"] = "BẢNG 2 CHIỀU: Doanh thu Miền × Tháng (một công thức ở B12, kéo sang phải và xuống)"; d["A9"].font = BOLD
header(d, 11, ["Miền", 1, 2, 3, "Tổng"])
for j in range(2, 5): d.cell(row=11, column=j).number_format = MONTH_FMT
for i, reg in enumerate(REGIONS + ["Tổng"], start=12):
    d.cell(row=i, column=1, value=reg).font = BOLD if reg == "Tổng" else BODY
    for j in range(2, 6):
        col = get_column_letter(j)
        f = (f"=SUMIFS(tblSales[Revenue],tblSales[Region],$A{i},tblSales[Thang],{col}$11)" if reg != "Tổng" and j < 5
             else (f"=SUM(B{i}:D{i})" if j == 5 else f"=SUM({col}12:{col}14)"))
        c = d.cell(row=i, column=j, value=f); c.font = BOLD if reg == "Tổng" else BODY; c.number_format = NUM_FMT
d["A17"] = "Đối soát: tổng bảng − tổng cột Revenue (phải = 0)"; d["A17"].font = BODY
d["E17"] = "=E15-SUM(tblSales[Revenue])"; d["E17"].font = BOLD; d["E17"].number_format = NUM_FMT
d["A19"] = "TRUNG BÌNH CỦA TRUNG BÌNH: Gross Margin % theo ngành hàng"; d["A19"].font = BOLD
header(d, 20, ["Ngành hàng", "Doanh thu", "COGS", "Gross Margin %"])
cats = list(byc.sort_values("rev", ascending=False).index)
for i, ct in enumerate(cats, start=21):
    d.cell(row=i, column=1, value=ct).font = BODY
    for j, f in ((2, f"=SUMIFS(tblSales[Revenue],tblSales[Category],$A{i})"), (3, f"=SUMIFS(tblSales[COGS],tblSales[Category],$A{i})"),
                 (4, f"=(B{i}-C{i})/B{i}")):
        c = d.cell(row=i, column=j, value=f); c.font = BODY; c.number_format = PCT_FMT if j == 4 else NUM_FMT
e = 21 + len(cats)
d.cell(row=e, column=1, value="SAI: AVERAGE các ngành").font = BOLD
c = d.cell(row=e, column=4, value=f"=AVERAGE(D21:D{e - 1})"); c.font = BOLD; c.number_format = PCT_FMT
d.cell(row=e + 1, column=1, value="ĐÚNG: tổng ÷ tổng").font = BOLD
for j, f in ((2, f"=SUM(B21:B{e - 1})"), (3, f"=SUM(C21:C{e - 1})"), (4, f"=(B{e + 1}-C{e + 1})/B{e + 1}")):
    c = d.cell(row=e + 1, column=j, value=f); c.font = BOLD; c.number_format = PCT_FMT if j == 4 else NUM_FMT
set_widths(d, [46, 15, 15, 16, 12])
da = wb.create_sheet("Demo_DynamicArray")
da["A1"] = "PRO TIP — DYNAMIC ARRAYS (Excel 365). Gõ công thức vào ô vàng, kết quả tự “tràn” (spill) xuống dưới."; da["A1"].font = TITLE
dyn = [("A3", "Danh sách cửa hàng không trùng", "=SORT(UNIQUE(tblSales[Store_Code]))"),
       ("C3", "Các dòng Miền Trung có doanh thu ≥ 1.000.000", '=FILTER(tblSales[[Order_ID]:[Revenue]],(tblSales[Region]="Trung")*(tblSales[Revenue]>=1000000))'),
       ("A20", "Doanh thu từng cửa hàng (cạnh danh sách ở A4)", "=SUMIFS(tblSales[Revenue],tblSales[Store_Code],A4#)")]
for addr, lab, f in dyn:
    c = da[addr]; c.value = lab; c.font = BOLD
    r = da[addr].row + 1; col = da[addr].column
    c = da.cell(row=r, column=col, value="(gõ: " + f + ")"); c.font = NOTE; c.fill = INPUT_FILL
set_widths(da, [30, 18, 16, 10, 12, 12, 10, 16, 16, 12, 12])
big_trung = q1[(q1.Region == "Trung") & (q1.Revenue >= 1000000)]
meta["sumifs"] = {"bac_t3_rev": int(grid.loc[("Bắc", 3), "rev"]), "bac_t3_lines": int(((q1.Region == "Bắc") & (q1.Thang == 3)).sum()),
                  "bac_t3_orders": int(grid.loc[("Bắc", 3), "orders"]),
                  "bac_t3_avg_line": float(q1[(q1.Region == "Bắc") & (q1.Thang == 3)].Revenue.mean()),
                  "big_trung_rows": len(big_trung)}
s = meta["sumifs"]; mq = meta["q1"]
g = wb.create_sheet("Huong_dan")
guide(g, "DEMO B2 — COUNTIFS / SUMIFS / AVERAGEIFS (slide 15, 16, 17)", [
    (15, "Sheet Demo_SUMIFS, ô D4. Nhấn F2 để thấy 2 bộ lọc: Region = B4 (Bắc), Thang = C4 (3).",
     f"Doanh thu Miền Bắc tháng 3 = {vn(s['bac_t3_rev'])}."),
    (15, "Kiểm chứng bằng Filter: sang sheet Sales, lọc Region = Bắc, Thang = 3, chọn cột Revenue, nhìn Status Bar (Sum). Xong thì Clear Filter.",
     f"Status Bar: Sum = {vn(s['bac_t3_rev'])}, Count = {vn(s['bac_t3_lines'])} — đúng bằng D4, D5."),
    (16, "Xem D5 (COUNTIFS), D6 (COUNTIFS thêm Line_No = 1), D7 (AVERAGEIFS). Đổi B4:B7 thành Nam, C4:C7 thành 1.",
     f"Bắc–3: {vn(s['bac_t3_lines'])} dòng, {vn(s['bac_t3_orders'])} đơn, {vn(s['bac_t3_avg_line'])} / dòng. Nam–1: doanh thu {vn(mq['grid']['Nam-1']['rev'])}, {int(mq['grid']['Nam-1']['orders'])} đơn."),
    (17, "Xóa B12:D14 rồi viết lại B12: =SUMIFS(tblSales[Revenue],tblSales[Region],$A12,tblSales[Thang],B$11). Kéo sang D12, rồi kéo xuống dòng 14.",
     f"Tháng 1 = {vn(mq['by_month'][1]['rev'])}; tháng 2 = {vn(mq['by_month'][2]['rev'])}; tháng 3 = {vn(mq['by_month'][3]['rev'])}; tổng E15 = {vn(REV)}; E17 = 0."),
    (17, "Thử bỏ dấu $ (viết $A12 thành A12) rồi kéo sang phải để thấy công thức hỏng; Ctrl+Z.",
     "Kéo sang phải, tiêu chí Miền trôi sang cột B, C (là số) → kết quả ra 0."),
    (17, f"Khối Gross Margin % theo ngành (dòng 20–{e + 1}): so sánh D{e} (AVERAGE các ngành) với D{e + 1} (tổng ÷ tổng).",
     f"AVERAGE = {pct(mq['avg_of_category_gm'])} (sai); tổng ÷ tổng = {pct(GM / REV)} (đúng)."),
    ("PRO TIP", "Sheet Demo_DynamicArray (dùng khi lớp đã biết Excel): gõ các công thức trong ô vàng.",
     f"A4 tràn 12 mã ST01…ST12; C4 tràn {s['big_trung_rows']} dòng Miền Trung có doanh thu ≥ 1.000.000; A21 tràn 12 doanh thu cửa hàng."),
])
save_both(wb, "B2_Demo_SUMIFS", "Demo_SUMIFS")

# ---- B2_Demo_Lookup.xlsx — slide 20 (VLOOKUP), 21 (INDEX/MATCH), 22 (XLOOKUP), 23 (khóa trùng), 24 (kiểu khóa)
wb = Workbook(); ws = wb.active; ws.title = "Product"
header(ws, 1, list(p_dup.columns)); write_rows(ws, p_dup, 2, {"Unit_Price": NUM_FMT, "Unit_Cost": NUM_FMT})
set_widths(ws, [13, 48, 18, 12, 12]); ws.freeze_panes = "A2"
PR = len(p_dup) + 1                                   # dòng cuối của Product (42)
PA, PE, PB = f"Product!$A$2:$A${PR}", f"Product!$E$2:$E${PR}", f"Product!$B$2:$B${PR}"
codes = ["P005", "P013", "P027", "P041", "P21", "P039"]
v = wb.create_sheet("Demo_VLOOKUP")
header(v, 1, ["Product_Code", "VLOOKUP khớp chính xác (FALSE)", "VLOOKUP bỏ trống range_lookup", "Ghi chú"])
notes = {"P041": "Mã mới, chưa có trong danh mục", "P21": "Gõ thiếu số 0 (đúng là P021)"}
for i, cd in enumerate(codes, start=2):
    c = v.cell(row=i, column=1, value=cd); c.font = INPUT; c.fill = INPUT_FILL
    for j, f in ((2, f"=VLOOKUP(A{i},Product!$A$2:$E${PR},5,FALSE)"), (3, f"=VLOOKUP(A{i},Product!$A$2:$E${PR},5)")):
        c = v.cell(row=i, column=j, value=f); c.font = BODY; c.number_format = NUM_FMT
    v.cell(row=i, column=4, value=notes.get(cd, "")).font = NOTE
set_widths(v, [14, 30, 30, 34])
x = wb.create_sheet("Demo_XLOOKUP")
header(x, 1, ["Product_Code", "INDEX / MATCH", "XLOOKUP (có if_not_found)", "", "Product_Name cần tra ngược", "Mã (XLOOKUP sang trái)"])
for i, cd in enumerate(codes, start=2):
    c = x.cell(row=i, column=1, value=cd); c.font = INPUT; c.fill = INPUT_FILL
    for j, f in ((2, f"=INDEX({PE},MATCH(A{i},{PA},0))"), (3, f'=_xlfn.XLOOKUP(A{i},{PA},{PE},"Không có mã")')):
        c = x.cell(row=i, column=j, value=f); c.font = BODY; c.number_format = NUM_FMT
for i, nm in enumerate(["Gạo thơm 5kg", "Chảo chống dính 26cm", "Hộp quà Tết cao cấp"], start=2):
    c = x.cell(row=i, column=5, value=nm); c.font = INPUT; c.fill = INPUT_FILL
    x.cell(row=i, column=6, value=f"=_xlfn.XLOOKUP(E{i},{PB},{PA})").font = BODY
set_widths(x, [14, 18, 26, 3, 26, 22])
kd = wb.create_sheet("Demo_KhoaTrung")
kd["A1"] = "LAWS OF LOOKUPS — KHÓA PHẢI DUY NHẤT"; kd["A1"].font = TITLE
kd["A3"] = "Số lần P021 xuất hiện trong Product"; kd["D3"] = f'=COUNTIF({PA},"P021")'
kd["A4"] = "Số dòng có mã bị trùng trong Product"; kd["D4"] = f"=SUMPRODUCT(--(COUNTIF({PA},{PA})>1))"
kd["A5"] = "Tổng COGS của P021 (theo cột F bên dưới)"; kd["D5"] = "=SUM(F9:F" + str(8 + len(p21)) + ")"
kd["A6"] = "Tổng COGS đúng (Unit_Cost 69.000)"; kd["D6"] = f"=SUM(D9:D{8 + len(p21)})*69000"
for a in ("A3", "A4", "A5", "A6"): kd[a].font = BODY
for a in ("D3", "D4", "D5", "D6"): kd[a].font = BOLD; kd[a].number_format = NUM_FMT
header(kd, 8, ["Order_ID", "Line_No", "Product_Code", "Quantity", "Unit_Cost (XLOOKUP)", "COGS"])
for i, row in enumerate(p21[["Order_ID", "Line_No", "Product_Code", "Quantity"]].itertuples(index=False), start=9):
    for j, val in enumerate(row, start=1): kd.cell(row=i, column=j, value=py(val)).font = BODY
    c = kd.cell(row=i, column=5, value=f"=_xlfn.XLOOKUP(C{i},{PA},{PE})"); c.font = BODY; c.number_format = NUM_FMT
    c = kd.cell(row=i, column=6, value=f"=D{i}*E{i}"); c.font = BODY; c.number_format = NUM_FMT
kd.freeze_panes = "A9"
set_widths(kd, [15, 8, 13, 10, 20, 14])
kk = wb.create_sheet("Demo_KieuKhoa")
kk["A1"] = "KHÓA NHÌN GIỐNG NHAU NHƯNG KHÁC KIỂU"; kk["A1"].font = TITLE
header(kk, 3, ["Mã vùng (số)", "Tên vùng"])
for i, (a_, b_) in enumerate([(1, "Bắc"), (2, "Trung"), (3, "Nam")], start=4):
    kk.cell(row=i, column=1, value=a_).font = BODY; kk.cell(row=i, column=2, value=b_).font = BODY
header(kk, 8, ["Giá trị tra", "Kiểu", "XLOOKUP thẳng", "Cách sửa", "XLOOKUP sau khi sửa"])
cases = [("1", "Text (số lưu dạng chữ, từ file CSV)", "=_xlfn.XLOOKUP(A9,$A$4:$A$6,$B$4:$B$6)", "VALUE(A9)", "=_xlfn.XLOOKUP(VALUE(A9),$A$4:$A$6,$B$4:$B$6)"),
         (2, "Số", "=_xlfn.XLOOKUP(A10,$A$4:$A$6,$B$4:$B$6)", "—", "=_xlfn.XLOOKUP(A10,$A$4:$A$6,$B$4:$B$6)"),
         ("P005 ", "Text có dấu cách thừa ở cuối", f"=_xlfn.XLOOKUP(A11,{PA},{PB})", "TRIM(A11)", f"=_xlfn.XLOOKUP(TRIM(A11),{PA},{PB})"),
         (" P013", "Text có dấu cách thừa ở đầu", f"=_xlfn.XLOOKUP(A12,{PA},{PB})", "TRIM(A12)", f"=_xlfn.XLOOKUP(TRIM(A12),{PA},{PB})")]
for i, cs in enumerate(cases, start=9):
    for j, val in enumerate(cs, start=1):
        c = kk.cell(row=i, column=j, value=val); c.font = INPUT if j == 1 else BODY
        if j == 1: c.fill = INPUT_FILL
kk["A14"] = "Kiểm tra độ dài: =LEN(A11) ra 5 thay vì 4 → có ký tự thừa."; kk["A14"].font = NOTE
set_widths(kk, [14, 36, 22, 12, 26])
mp = meta["p021"]
g = wb.create_sheet("Huong_dan")
guide(g, "DEMO B2 — VLOOKUP → INDEX/MATCH → XLOOKUP, LAWS OF LOOKUPS (slide 20–24)", [
    (20, "Sheet Demo_VLOOKUP: so cột B (FALSE) và cột C (bỏ trống range_lookup). Chú ý dòng P041 và P21.",
     "Cột B: P041, P21 ra #N/A (đúng: không có mã). Cột C: cả hai ra 71.500 — giá vốn của P040, sai mà không báo lỗi."),
    (21, "Sheet Demo_XLOOKUP, cột B: =INDEX(cột Unit_Cost, MATCH(A2, cột Product_Code, 0)).",
     "Kết quả giống cột B của Demo_VLOOKUP; P041, P21 ra #N/A."),
    (22, "Cột C: =XLOOKUP(A2, cột Product_Code, cột Unit_Cost, \"Không có mã\"). Cột F: tra ngược từ Product_Name ra mã (cột trả về nằm bên trái).",
     "P005 = 83.500; P041, P21 hiện “Không có mã”. F2:F4 = P014, P039, P027."),
    (23, "Sheet Demo_KhoaTrung. Xem cột E (XLOOKUP Unit_Cost của P021) và ô D3:D6. Sang sheet Product, xem dòng 22–23.",
     f"{mp['lines']} dòng P021 đều lấy Unit_Cost {vn(OLD_COST)} (bản ghi cũ đứng trước). D3 = 2, D4 = 2. "
     f"COGS P021 = {vn(mp['cogs_wrong'])} thay vì {vn(mp['cogs_right'])} (thiếu {vn(mp['cogs_right'] - mp['cogs_wrong'])})."),
    (23, "Sửa: ở sheet Product xóa dòng 22 (bản ghi “mã cũ – ngừng KD 2024”). Quay lại Demo_KhoaTrung.",
     f"Cột E thành 69.000; D5 = D6 = {vn(mp['cogs_right'])}; D3 = 1, D4 = 0."),
    (24, "Sheet Demo_KieuKhoa: xem cột C (tra thẳng) và cột E (sau khi sửa bằng VALUE / TRIM). Gõ =LEN(A11) ở một ô trống.",
     "C9, C11, C12 ra #N/A; E9:E12 ra Bắc, Trung, Cà phê rang xay 500g, Mì ăn liền tôm chua cay (thùng 30 gói). LEN(A11) = 5."),
])
save_both(wb, "B2_Demo_Lookup", "Demo_VLOOKUP")

# ---- B2_Demo_Text_Date.xlsx — slide 26 (Text), 27 (Date)
wb = Workbook(); ws = wb.active; ws.title = "Demo_Text"
smp = q1.drop_duplicates("Store_Code").sort_values("Store_Code").head(6)
sname = stores.set_index("Store_Code")["Store_Name"]
header(ws, 1, ["Order_ID", "Store_Code", "Store_Name (thô)", "Unit_Price", "Tiền tố LEFT", "Kỳ yymm MID", "Số thứ tự RIGHT",
               "TRIM tên", "Nhãn ghép &", "TEXT giá"])
# Tên cửa hàng "thô" như gõ tay: thêm dấu cách thừa ở đầu, giữa, cuối (TRIM đưa về đúng tên trong Store)
_pad = [("  ", "   ", ""), ("", " ", "  "), ("", "  ", ""), (" ", " ", ""), ("", "   ", " "), ("", " ", "")]
raws = [a + nm.replace("Mart ", "Mart" + b, 1) + c for nm, (a, b, c) in zip([sname[c_] for c_ in smp.Store_Code], _pad)]
for i, (row, raw) in enumerate(zip(smp.itertuples(index=False), raws), start=2):
    for j, val in enumerate([row.Order_ID, row.Store_Code, raw, int(row.Unit_Price)], start=1):
        c = ws.cell(row=i, column=j, value=val); c.font = INPUT; c.fill = INPUT_FILL
        if j == 4: c.number_format = NUM_FMT
    for j, f in enumerate([f"=LEFT(A{i},2)", f"=MID(A{i},3,4)", f"=RIGHT(A{i},5)", f"=TRIM(C{i})",
                           f'=B{i}&" – "&TRIM(C{i})', f'=TEXT(D{i},"#,##0")&" đ"'], start=5):
        ws.cell(row=i, column=j, value=f).font = BODY
set_widths(ws, [15, 11, 24, 11, 11, 12, 14, 22, 30, 14])
assert [sname[c] for c in smp.Store_Code] == [" ".join(r.split()) for r in raws], "tên cửa hàng mẫu lệch Store"
dt = wb.create_sheet("Demo_Date")
header(dt, 1, ["Order_Date", "Cuối tháng EOMONTH(,0)", "Đầu tháng EOMONTH(,-1)+1", "Tháng TEXT mm/yyyy", "Quý", "Thứ TEXT ddd", "WEEKDAY(,2)"])
dates = [pd.Timestamp("2025-01-15"), pd.Timestamp("2025-01-29"), pd.Timestamp("2025-02-28"), pd.Timestamp("2025-03-08")]
for i, dd in enumerate(dates, start=2):
    c = dt.cell(row=i, column=1, value=dd.to_pydatetime()); c.font = INPUT; c.fill = INPUT_FILL; c.number_format = DATE_FMT
    for j, (f, fmt) in enumerate([(f"=EOMONTH(A{i},0)", DATE_FMT), (f"=EOMONTH(A{i},-1)+1", DATE_FMT), (f'=TEXT(A{i},"mm/yyyy")', None),
                                  (f'="Q"&ROUNDUP(MONTH(A{i})/3,0)', None), (f'=TEXT(A{i},"ddd")', None), (f"=WEEKDAY(A{i},2)", "0")], start=2):
        c = dt.cell(row=i, column=j, value=f); c.font = BODY
        if fmt: c.number_format = fmt
dt["A8"] = "NGÀY LÀM VIỆC TRONG THÁNG (NETWORKDAYS)"; dt["A8"].font = BOLD
header(dt, 9, ["Tháng", "Từ ngày", "Đến ngày", "Ngày làm việc", "Trừ ngày lễ"])
hol = [pd.Timestamp("2025-01-01")] + list(pd.date_range("2025-01-27", "2025-01-31"))
for i, m in enumerate(MONTHS, start=10):
    dt.cell(row=i, column=1, value=m).number_format = MONTH_FMT
    c = dt.cell(row=i, column=2, value=f"=DATE(2025,A{i},1)"); c.number_format = DATE_FMT
    c = dt.cell(row=i, column=3, value=f"=EOMONTH(B{i},0)"); c.number_format = DATE_FMT
    dt.cell(row=i, column=4, value=f"=NETWORKDAYS(B{i},C{i})")
    dt.cell(row=i, column=5, value=f"=NETWORKDAYS(B{i},C{i},$H$10:$H${9 + len(hol)})")
    for j in range(1, 6): dt.cell(row=i, column=j).font = BODY
header(dt, 9, ["Ngày lễ 2025 (Tết Dương lịch, nghỉ Tết Âm lịch)"], col=8)
for i, hd in enumerate(hol, start=10):
    c = dt.cell(row=i, column=8, value=hd.to_pydatetime()); c.font = INPUT; c.fill = INPUT_FILL; c.number_format = DATE_FMT
set_widths(dt, [13, 20, 22, 18, 8, 12, 12, 40])
wd_plain = [int(np.busday_count(f"2025-{m:02d}-01", (pd.Timestamp(f"2025-{m:02d}-01") + pd.offsets.MonthEnd(0) + pd.Timedelta(days=1)).strftime("%Y-%m-%d"))) for m in MONTHS]
wd_hol = [int(np.busday_count(f"2025-{m:02d}-01", (pd.Timestamp(f"2025-{m:02d}-01") + pd.offsets.MonthEnd(0) + pd.Timedelta(days=1)).strftime("%Y-%m-%d"),
                              holidays=[h.strftime("%Y-%m-%d") for h in hol])) for m in MONTHS]
meta["date"] = {"networkdays": wd_plain, "networkdays_hol": wd_hol,
                "weekday": [dd.day_name() for dd in dates], "weekday_num": [dd.dayofweek + 1 for dd in dates]}
first = smp.iloc[0]
g = wb.create_sheet("Huong_dan")
guide(g, "DEMO B2 — HÀM TEXT VÀ DATE (slide 26, 27)", [
    (26, "Sheet Demo_Text: F2 vào E2, F2, G2 (LEFT, MID, RIGHT tách Order_ID).",
     f"{first.Order_ID} → “HD”, “{first.Order_ID[2:6]}” (năm 25, tháng 01), “{first.Order_ID[-5:]}”."),
    (26, "Cột H (TRIM) và I (ghép bằng &). So độ dài: =LEN(C2) và =LEN(H2).",
     f"“{raws[0]}” → “{sname[smp.Store_Code.iloc[0]]}”; I2 = “{smp.Store_Code.iloc[0]} – {sname[smp.Store_Code.iloc[0]]}”; LEN {len(raws[0])} → {len(sname[smp.Store_Code.iloc[0]])}."),
    (26, "Cột J: =TEXT(D2,\"#,##0\")&\" đ\". Thử cộng J2 với 1.",
     "J2 hiện dạng “19,000 đ” hoặc “19.000 đ” tùy cài đặt vùng; J2+1 ra #VALUE! vì TEXT trả về chữ, không còn là số."),
    (27, "Sheet Demo_Date: xem B2:G5 (EOMONTH, TEXT, Quý, Thứ, WEEKDAY).",
     "15/01/2025 → cuối tháng 31/01/2025, đầu tháng 01/01/2025, “01/2025”, Q1, thứ Tư (WEEKDAY = 3). 29/01/2025 (mùng 1 Tết) là thứ Tư. "
     "28/02/2025 → cuối tháng 28/02/2025. 08/03/2025 là thứ Bảy (6)."),
    (27, "Khối NETWORKDAYS (A10:E12): cột D không tính ngày lễ, cột E trừ danh sách ngày lễ ở H10:H15.",
     f"Tháng 1: {wd_plain[0]} → {wd_hol[0]} ngày; tháng 2: {wd_plain[1]}; tháng 3: {wd_plain[2]}."),
])
save_both(wb, "B2_Demo_Text_Date", "Demo_Text")

# =====================================================================
# 3) HOMEWORK — B2_BTVN_MCI_Mart_T4_2025.xlsx (+ _Answer)
# =====================================================================
rnd = random.Random(42)
t4 = m_sales[m_sales.Order_ID.str.startswith("HD2504")].copy()
t4["Product_Name"] = t4.Product_Code.map(products.set_index("Product_Code")["Product_Name"])
t4 = t4.sort_values(["Store_Code", "Order_Date", "Order_ID", "Line_No"]).reset_index(drop=True)
hc = ["Order_ID", "Line_No", "Order_Date", "Store_Code", "Customer_ID", "Product_Code", "Product_Name", "Quantity", "Unit_Price"]
t4 = t4[hc]
NT = len(t4)
space_idx = sorted(rnd.sample(list(t4.index), 4))
t4d = t4.copy(); t4d.loc[space_idx, "Product_Code"] = t4d.loc[space_idx, "Product_Code"] + " "
t4["Region"] = t4.Store_Code.map(stores.set_index("Store_Code")["Region"])
t4["Category"] = t4.Product_Code.map(cat); t4["Unit_Cost"] = t4.Product_Code.map(cost)
t4["Revenue"] = t4.Quantity * t4.Unit_Price; t4["COGS"] = t4.Quantity * t4.Unit_Cost
T4REV, T4COGS = int(t4.Revenue.sum()), int(t4.COGS.sum()); T4ORD = int((t4.Line_No == 1).sum())
CATS = ["Đồ uống", "Sữa", "Thực phẩm khô", "Bánh kẹo", "Chăm sóc cá nhân", "Gia dụng"]
t4r = t4.groupby("Region").agg(rev=("Revenue", "sum"), cogs=("COGS", "sum"), orders=("Line_No", lambda s: int((s == 1).sum())))
t4c = t4.groupby("Category").agg(rev=("Revenue", "sum"), cogs=("COGS", "sum")); t4c["gm"] = 1 - t4c.cogs / t4c.rev
meta["hw"] = {"rows": NT, "revenue": T4REV, "orders": T4ORD, "cogs": T4COGS, "gm_pct": 1 - T4COGS / T4REV, "aov": T4REV / T4ORD,
              "space_rows_excel": [i + 2 for i in space_idx], "space_codes": [t4.Product_Code[i] for i in space_idx],
              "by_region": {r: {"rev": int(x.rev), "orders": int(x.orders), "aov": x.rev / x.orders, "gm": 1 - x.cogs / x.rev} for r, x in t4r.iterrows()},
              "by_cat_gm": t4c.gm.to_dict(), "avg_cat_gm": float(t4c.gm.mean()),
              "region_cat": t4.pivot_table(index="Region", columns="Category", values="Revenue", aggfunc="sum").fillna(0).astype(int).to_dict()}
mh = meta["hw"]

def hw_tables(wb, dirty):
    ws = wb.create_sheet("Sales")
    src = t4d if dirty else t4[hc]
    header(ws, 1, hc + ["Revenue"]); write_rows(ws, src, 2, {"Order_Date": DATE_FMT, "Unit_Price": NUM_FMT})
    for i in range(2, NT + 2):
        c = ws.cell(row=i, column=10, value="=tblSales_T4[[#This Row],[Quantity]]*tblSales_T4[[#This Row],[Unit_Price]]")
        c.font = BODY; c.number_format = NUM_FMT
    return ws

wb = Workbook(); de = wb.active; de.title = "De_bai"
ws = hw_tables(wb, True); add_table(ws, "tblSales_T4", 10, NT); ws.freeze_panes = "A2"
set_widths(ws, [15, 7, 12, 10, 11, 12, 40, 9, 11, 12])
s_ = wb.create_sheet("Store"); table_df(s_, stores, "tblStore", {"Open_Date": DATE_FMT}, [11, 24, 13, 9, 11, 12])
p_ = wb.create_sheet("Product"); table_df(p_, products, "tblProduct", {"Unit_Price": NUM_FMT, "Unit_Cost": NUM_FMT}, [13, 40, 18, 12, 12])
tasks = [
    ("1", "Tra cứu (XLOOKUP)", "Trong tblSales_T4 thêm 3 cột: Region (từ tblStore), Category và Unit_Cost (từ tblProduct). Nếu có ô #N/A: tìm nguyên nhân (gợi ý: LEN của mã) và sửa, không dùng IFERROR để che."),
    ("2", "Cột tính toán", "Thêm COGS = Quantity × Unit_Cost, Gross_Margin = Revenue − COGS, Phan_khuc = IFS theo Unit_Price (≥ 100.000 “Cao”, ≥ 30.000 “Trung”, còn lại “Thấp”)."),
    ("3", "Bảng KPI 1 — Doanh thu Miền × Ngành hàng", "Tạo sheet KPI. Dựng bảng 3 dòng (Bắc, Trung, Nam) × 6 cột ngành hàng, CHỈ viết một công thức SUMIFS ở ô đầu rồi kéo (dùng $ hỗn hợp). Thêm dòng và cột Tổng."),
    ("4", "Bảng KPI 2 — theo Miền", "Cùng sheet KPI: Doanh thu, Số đơn (COUNTIFS với Line_No = 1), AOV = Doanh thu ÷ Số đơn, Gross Margin % = Gross_Margin ÷ Doanh thu. Dòng Tổng tính tổng ÷ tổng."),
    ("5", "KPI Definition Card", "Tạo sheet KPI_Card, viết card cho KPI Avg Order Value (AOV) đủ 10 trường theo mẫu B2_KPI_Definition_Card."),
    ("6", "Viết insight", "Tạo sheet Insight, viết 1 câu có đủ: con số (có so sánh), nguyên nhân khả dĩ, hàm ý."),
]
de["A1"] = "BÀI TẬP VỀ NHÀ BUỔI 2 (BẮT BUỘC) — LÀM GIÀU DỮ LIỆU VÀ BẢNG KPI THÁNG 4/2025"; de["A1"].font = TITLE
de["A2"] = "Nộp trước Buổi 3. Đặt tên file: B2_BTVN_<HoTen>.xlsx. Dữ liệu tháng 4/2025 đã làm sạch (kết quả bài về nhà Buổi 1), đã là Table tblSales_T4."; de["A2"].font = NOTE
de["A3"] = "Nộp bài tại https://hocvien.mcivietnam.com/auth#/ — Tài khoản: email bạn đăng ký tại MCI · Mật khẩu: hv12345678"; de["A3"].font = BOLD
header(de, 5, ["#", "Việc cần làm", "Yêu cầu"])
for i, row in enumerate(tasks, start=6):
    for j, val in enumerate(row, start=1):
        c = de.cell(row=i, column=j, value=val); c.font = BODY; c.alignment = Alignment(wrap_text=True, vertical="top")
kr = 6 + len(tasks) + 1
de.cell(row=kr, column=1, value="Tự kiểm tra trước khi nộp").font = BOLD
for i, tx in enumerate([f"Không còn ô #N/A; {vn(NT)} dòng", f"Tổng doanh thu = {vn(T4REV)} VNĐ (bảng KPI 1 và KPI 2 đều ra đúng số này)",
                        f"Tổng số đơn = {vn(T4ORD)}", f"Tổng COGS = {vn(T4COGS)} VNĐ", f"Gross Margin % tháng 4 = {pct(1 - T4COGS / T4REV)}"], start=kr + 1):
    de.cell(row=i, column=2, value="☐ " + tx).font = BODY
set_widths(de, [5, 34, 100])
wb.active = 0
wb.save(f"{HW}/B2_BTVN_MCI_Mart_T4_2025.xlsx")

# --- Đáp án
wb = Workbook(); wb.remove(wb.active)
ws = wb.create_sheet("Sales")
extra = ["Region", "Category", "Unit_Cost", "COGS", "Gross_Margin", "Phan_khuc"]
header(ws, 1, hc + ["Revenue"] + extra); write_rows(ws, t4[hc], 2, {"Order_Date": DATE_FMT, "Unit_Price": NUM_FMT})
TT = lambda c_: f"tblSales_T4[[#This Row],[{c_}]]"
hf = [f"={TT('Quantity')}*{TT('Unit_Price')}",
      f"=_xlfn.XLOOKUP({TT('Store_Code')},tblStore[Store_Code],tblStore[Region])",
      f"=_xlfn.XLOOKUP({TT('Product_Code')},tblProduct[Product_Code],tblProduct[Category])",
      f"=_xlfn.XLOOKUP({TT('Product_Code')},tblProduct[Product_Code],tblProduct[Unit_Cost])",
      f"={TT('Quantity')}*{TT('Unit_Cost')}", f"={TT('Revenue')}-{TT('COGS')}",
      f'=_xlfn.IFS({TT("Unit_Price")}>=100000,"Cao",{TT("Unit_Price")}>=30000,"Trung",TRUE,"Thấp")']
for i in range(2, NT + 2):
    for j, f in enumerate(hf, start=10):
        c = ws.cell(row=i, column=j, value=f); c.font = BODY
        if j in (10, 13, 14, 15): c.number_format = NUM_FMT
add_table(ws, "tblSales_T4", len(hc) + 1 + len(extra), NT); ws.freeze_panes = "A2"
set_widths(ws, [15, 7, 12, 10, 11, 12, 40, 9, 11, 12, 9, 16, 10, 12, 13, 10])
s_ = wb.create_sheet("Store"); table_df(s_, stores, "tblStore", {"Open_Date": DATE_FMT}, [11, 24, 13, 9, 11, 12])
p_ = wb.create_sheet("Product"); table_df(p_, products, "tblProduct", {"Unit_Price": NUM_FMT, "Unit_Cost": NUM_FMT}, [13, 40, 18, 12, 12])
kp = wb.create_sheet("KPI")
kp["A1"] = "KPI THÁNG 4/2025"; kp["A1"].font = TITLE
kp["A3"] = "Bảng 1 — Doanh thu (VNĐ) Miền × Ngành hàng"; kp["A3"].font = BOLD
header(kp, 4, ["Miền"] + CATS + ["Tổng"])
for i, reg in enumerate(REGIONS + ["Tổng"], start=5):
    kp.cell(row=i, column=1, value=reg).font = BOLD if reg == "Tổng" else BODY
    for j in range(2, 9):
        col = get_column_letter(j)
        f = (f"=SUMIFS(tblSales_T4[Revenue],tblSales_T4[Region],$A{i},tblSales_T4[Category],{col}$4)" if reg != "Tổng" and j < 8
             else (f"=SUM(B{i}:G{i})" if j == 8 else f"=SUM({col}5:{col}7)"))
        c = kp.cell(row=i, column=j, value=f); c.font = BOLD if reg == "Tổng" else BODY; c.number_format = NUM_FMT
kp["A11"] = "Bảng 2 — Theo Miền"; kp["A11"].font = BOLD
header(kp, 12, ["Miền", "Doanh thu", "Số đơn", "AOV", "Gross Margin %"])
for i, reg in enumerate(REGIONS + ["Tổng"], start=13):
    kp.cell(row=i, column=1, value=reg).font = BOLD if reg == "Tổng" else BODY
    if reg != "Tổng":
        fs = [f"=SUMIFS(tblSales_T4[Revenue],tblSales_T4[Region],$A{i})", f"=COUNTIFS(tblSales_T4[Region],$A{i},tblSales_T4[Line_No],1)",
              f"=B{i}/C{i}", f"=SUMIFS(tblSales_T4[Gross_Margin],tblSales_T4[Region],$A{i})/B{i}"]
    else:
        fs = ["=SUM(B13:B15)", "=SUM(C13:C15)", f"=B{i}/C{i}", f"=SUM(tblSales_T4[Gross_Margin])/B{i}"]
    for j, f in enumerate(fs, start=2):
        c = kp.cell(row=i, column=j, value=f); c.font = BOLD if reg == "Tổng" else BODY; c.number_format = PCT_FMT if j == 5 else NUM_FMT
set_widths(kp, [10, 14, 14, 16, 16, 18, 14, 16])
kc = wb.create_sheet("KPI_Card")
kc["A1"] = "KPI DEFINITION CARD — Avg Order Value (đáp án tham khảo)"; kc["A1"].font = TITLE
aov_card = ["Avg Order Value – AOV (Giá trị đơn trung bình)", "Mỗi lượt mua, khách chi bao nhiêu tiền?",
            "Tổng doanh thu ÷ Số đơn = SUM(Revenue) ÷ COUNTIFS(Line_No, 1)", "tblSales_T4 (Revenue, Order_ID, Line_No)",
            "Dữ liệu: 1 dòng = 1 sản phẩm trong 1 đơn. KPI báo cáo theo tháng", "Miền, cửa hàng, tháng (không theo ngành hàng: một đơn có nhiều ngành)",
            "Giám đốc bán hàng", "300.000 VNĐ", "90% / 80% (270.000 / 240.000)",
            "Là tỷ lệ: tổng ÷ tổng, không lấy AVERAGE của AOV các miền. Bị kéo lệch bởi vài đơn rất lớn → xem thêm MEDIAN."]
header(kc, 3, ["Trường", "Nội dung"])
for i, (fld, val) in enumerate(zip(CARD_FIELDS, aov_card), start=4):
    kc.cell(row=i, column=1, value=fld).font = BOLD
    c = kc.cell(row=i, column=2, value=val); c.font = BODY; c.alignment = Alignment(wrap_text=True, vertical="top")
set_widths(kc, [28, 90])
kt = wb.create_sheet("Kiem_tra")
kt["A1"] = "KIỂM TRA / ĐỐI SOÁT — Bài tập về nhà B2"; kt["A1"].font = TITLE
kt["A2"] = "Chấm bài: so từng chỉ tiêu của học viên với cột Đáp án."; kt["A2"].font = NOTE
header(kt, 4, ["Chỉ tiêu", "", "Kết quả (công thức)", "Đáp án", "Khớp?"])
hchecks = [("Số dòng dữ liệu", "=ROWS(tblSales_T4[Order_ID])", NT, NUM_FMT),
           ("Số ô #N/A (Region, Category, Unit_Cost)", "=SUMPRODUCT(--ISNA(tblSales_T4[Region]))+SUMPRODUCT(--ISNA(tblSales_T4[Category]))+SUMPRODUCT(--ISNA(tblSales_T4[Unit_Cost]))", 0, "0"),
           ("Tổng doanh thu (VNĐ)", "=SUM(tblSales_T4[Revenue])", T4REV, NUM_FMT),
           ("Tổng bảng KPI 1 (Miền × Ngành)", "=KPI!H8", T4REV, NUM_FMT),
           ("Tổng số đơn", "=KPI!C16", T4ORD, NUM_FMT),
           ("Tổng COGS (VNĐ)", "=SUM(tblSales_T4[COGS])", T4COGS, NUM_FMT),
           ("Gross Margin % tháng 4", "=KPI!E16", 1 - T4COGS / T4REV, "0.00%"),
           ("AOV tháng 4 (VNĐ)", "=KPI!D16", T4REV / T4ORD, NUM_FMT)]
for r, (label, f, exp, fmt) in enumerate(hchecks, start=5):
    kt.cell(row=r, column=1, value=label).font = BODY
    c = kt.cell(row=r, column=3, value=f); c.number_format = fmt; c.font = BODY
    c = kt.cell(row=r, column=4, value=exp); c.number_format = fmt; c.font = INPUT
    kt.cell(row=r, column=5, value=f'=IF(ROUND(C{r},6)=ROUND(D{r},6),"✔","✘")').font = BODY
set_widths(kt, [44, 2, 20, 16, 8])
ea = wb.create_sheet("Dap_an_loi")
ea["A1"] = "ĐÁP ÁN — LỖI CÀI SẴN VÀ LỖI HAY GẶP"; ea["A1"].font = TITLE
rows = [("#", "Lỗi", "Vị trí / dấu hiệu", "Cách sửa"),
        (1, f"{len(space_idx)} mã Product_Code có dấu cách thừa ở cuối", "Dòng " + ", ".join(str(i + 2) for i in space_idx) + " của sheet Sales (mã " + ", ".join(t4.Product_Code[i] for i in space_idx) + "); XLOOKUP ra #N/A; LEN = 5",
         "TRIM (cột phụ hoặc Find & Replace dấu cách); không dùng IFERROR để che"),
        (2, "Tổng Miền × Ngành không khớp tổng doanh thu", "Quên $ hỗn hợp: kéo sang phải, tiêu chí Miền trôi", "$A5 (khóa cột) và B$4 (khóa dòng)"),
        (3, "Số đơn = số dòng", f"Ra {vn(NT)} thay vì {vn(T4ORD)}", "COUNTIFS thêm điều kiện Line_No = 1"),
        (4, "AOV / GM% dòng Tổng = AVERAGE các miền", f"GM% trung bình các ngành = {pct(mh['avg_cat_gm'])} ≠ {pct(1 - T4COGS / T4REV)}", "Tổng ÷ tổng")]
for i, rr in enumerate(rows, start=3):
    for j, val in enumerate(rr, start=1):
        c = ea.cell(row=i, column=j, value=val); c.font = HDR_FONT if i == 3 else BODY
        if i == 3: c.fill = HDR_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
set_widths(ea, [4, 40, 56, 44])
br = mh["by_region"]
top_r = max(br, key=lambda r: br[r]["rev"])
ins = wb.create_sheet("Insight_mau")
ins["A1"] = "INSIGHT MẪU (tham khảo khi chấm; học viên viết theo cách của mình)"; ins["A1"].font = TITLE
best_cat, worst_cat = max(mh["by_cat_gm"], key=mh["by_cat_gm"].get), min(mh["by_cat_gm"], key=mh["by_cat_gm"].get)
ins["A3"] = (f"\"Tháng 4/2025 MCI Mart bán {vn(T4REV)} VNĐ qua {vn(T4ORD)} đơn (AOV {vn(T4REV / T4ORD)}), Gross Margin {pct(1 - T4COGS / T4REV)}; "
             f"Miền {top_r} đóng góp nhiều nhất ({pct(br[top_r]['rev'] / T4REV)} doanh thu). Biên chênh lệch lớn giữa các ngành: {best_cat} {pct(mh['by_cat_gm'][best_cat])} "
             f"so với {worst_cat} {pct(mh['by_cat_gm'][worst_cat])}, nên cơ cấu ngành hàng của từng miền sẽ quyết định biên lợi nhuận — cần xem theo cửa hàng và theo ngày ở Buổi 3.\"")
ins["A3"].font = BODY; ins["A3"].alignment = Alignment(wrap_text=True, vertical="top")
ins["A5"] = "Chấm: đủ 3 thành phần (con số có so sánh · nguyên nhân khả dĩ · hàm ý) và con số khớp dữ liệu."; ins["A5"].font = NOTE
set_widths(ins, [120])
wb.save(f"{HW}/B2_BTVN_MCI_Mart_T4_2025_Answer.xlsx")

json.dump(meta, open(os.path.join(os.path.dirname(__file__), "build_b2_meta.json"), "w"), ensure_ascii=False, indent=1, default=str)
print(json.dumps(meta, ensure_ascii=False, indent=1, default=str))
