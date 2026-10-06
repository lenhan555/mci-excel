"""Nội dung sheet De_bai của bài tập về nhà Buổi 3 (bản học viên).

Dùng chung cho hai nơi, để đề luôn một nguồn:
- build_b3.py: dựng file đề từ đầu (gọi fill_de_bai).
- rewrite_b3_debai.py: viết lại riêng sheet De_bai trong file đề đã có, không dựng lại cả Buổi 3.

Mỗi câu gồm một dòng tiêu đề và các dòng "Bước n". Cột A: bước. Cột B: thao tác. Cột C: kết quả bạn sẽ thấy.
Văn phong theo .claude/skills/viet-van-phong-mci (sắc thái đào tạo, xưng "bạn"). Chưa kèm đáp án.
"""
import math
from openpyxl.styles import Font, PatternFill, Alignment

F = "Arial"
TITLE = Font(name=F, size=14, bold=True, color="203864")
BODY = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
NOTE = Font(name=F, size=10, italic=True, color="585858")
HDR_FONT = Font(name=F, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="203864")
Q_FILL = PatternFill("solid", fgColor="D9E2F3")
Q_FONT = Font(name=F, size=11, bold=True, color="203864")
WIDTHS = {"A": 10, "B": 100, "C": 52}
TOP_WRAP = Alignment(wrap_text=True, vertical="top")


def build_steps(f):
    """f: dict số liệu đã định dạng kiểu Việt: rows, qty4, rev, rev4, mom4."""
    return [
        ("CÂU 0. Chuẩn bị", None, [
            ("Mở file này, bấm File > Save As và lưu thành B3_BTVN_<HoTen>.xlsx (thay <HoTen> bằng tên bạn, không dấu). Làm tiếp trên chính file đã lưu.",
             "Tên file đúng mẫu, ví dụ B3_BTVN_NguyenVanA.xlsx."),
            (f"Xem qua ba sheet có sẵn. Sheet Sales là bảng tblSales4T, {f['rows']} dòng, tháng 1 đến tháng 4 năm 2025. "
             "Sheet So_lieu_he_thong là số tổng theo tháng do hệ thống bán hàng xuất ra. Sheet Store là danh sách cửa hàng.",
             "Chỉ sửa dữ liệu ở sheet Sales, và chỉ khi câu 1 yêu cầu. Không sửa sheet So_lieu_he_thong."),
            ("Mỗi câu dưới đây dùng một sheet riêng. Đặt đúng tên sheet ghi trong đề: bấm đúp vào tên sheet, gõ tên mới, Enter.",
             "Cuối bài có 6 sheet chính: PV_DoiSoat, PV_Mien, PV_Nganh, PV_CuaHang, PV_AOV, Insight (cộng hai sheet chi tiết ST07 ở câu 4)."),
        ]),
        ("CÂU 1. Đối soát trước khi phân tích (sheet PV_DoiSoat)", None, [
            ("Bấm vào một ô bất kỳ trong bảng ở sheet Sales. Vào Insert > PivotTable. Trong hộp thoại, mục Table/Range phải ghi tblSales4T. "
             "Chọn New Worksheet, bấm OK.",
             "Excel mở một sheet mới có khung PivotTable Fields bên phải."),
            ("Đổi tên sheet mới thành PV_DoiSoat.", None),
            ("Trong PivotTable Fields, kéo Thang xuống ô Rows. Kéo Revenue, So_don và Quantity xuống ô Values, mỗi trường một lần.",
             "Pivot có 4 dòng tháng và dòng Grand Total, cộng 3 cột giá trị."),
            ("Nhìn tiêu đề ba cột giá trị: cả ba phải ghi Sum of ... Cột nào Excel tự ghi Count of ... thì bấm chuột phải vào một ô của cột đó, "
             "chọn Summarize Values By > Sum.",
             "Ba cột đều là Sum of Revenue, Sum of So_don, Sum of Quantity."),
            ("Đặt Pivot cạnh bảng ở sheet So_lieu_he_thong rồi so từng tháng. Ba cột nhân bốn tháng là 12 ô cần so: doanh thu, số đơn, Quantity. "
             "Ghi lại ô nào lệch và lệch bao nhiêu.",
             "Có ít nhất một cột lệch. Các cột còn lại khớp đúng."),
            ("Tìm nguyên nhân ở dữ liệu gốc. Quay về sheet Sales, nhìn cột của chỉ tiêu bị lệch và tìm các dòng của tháng lệch. "
             "Hai dấu hiệu dễ thấy: số trong ô thường căn lề phải, ô nào căn lề trái là chữ; góc trên bên trái ô có thể có tam giác xanh. "
             "Đếm xem có bao nhiêu ô như vậy và nằm ở dòng nào.",
             "Biết lỗi nằm ở cột nào, bao nhiêu ô, thuộc tháng nào."),
            ("Sửa ngay trên bảng tblSales4T ở sheet Sales: chọn các ô lỗi, bấm biểu tượng cảnh báo hiện cạnh ô, chọn Convert to Number. "
             "Không sửa số trên Pivot.",
             "Các ô vừa sửa căn sang lề phải, tam giác xanh biến mất."),
            ("Quay về PV_DoiSoat, bấm chuột phải vào Pivot, chọn Refresh (hoặc Ctrl+Alt+F5). Kiểm tra lại ba cột vẫn là Sum of ...",
             "Số của Pivot đổi sau khi Refresh."),
            ("So lại 12 ô với So_lieu_he_thong. Còn ô lệch thì quay lại bước tìm nguyên nhân.",
             f"Cả 12 ô khớp. Tổng Quantity tháng 4 là {f['qty4']}."),
            ("Bấm Insert > New Worksheet (hoặc dấu + cạnh tên sheet), đặt tên Insight. Ở ô A1 gõ: Lỗi dữ liệu đã sửa. Ở ô A2 viết 1 đến 2 câu: "
             "lỗi nằm ở cột nào, bao nhiêu ô, thuộc tháng nào, làm Pivot lệch ra sao, bạn đã sửa thế nào.",
             "Sheet Insight có tiêu đề ở A1 và đoạn viết ở A2."),
        ]),
        ("CÂU 2. Miền nào tăng mạnh nhất tháng 4 so với tháng 3? (sheet PV_Mien)", None, [
            ("Bấm dấu + tạo sheet mới, đổi tên thành PV_Mien. Tại Sales, bấm vào tblSales4T > Insert > PivotTable. Chọn Existing Worksheet, Location: PV_Mien!$A$12, OK. "
             "Từ câu này trở đi, mọi Pivot đều tạo từ tblSales4T để câu 6 nối được Slicer (nút lọc bấm được).",
             None),
            ("Kéo Region xuống Rows, Thang xuống Columns, Revenue xuống Values.",
             "Bảng doanh thu theo miền, mỗi tháng một cột."),
            ("Tạo Pivot thứ hai ngay trên sheet này: Insert > PivotTable, Table/Range ghi tblSales4T, chọn Existing Worksheet, ở ô Location gõ PV_Mien!$A$4, OK.",
             "Pivot thứ hai xuất hiện từ A4, nằm trên Pivot doanh thu ở A12."),
            ("Ở Pivot thứ hai, kéo Region xuống Rows, Thang xuống Columns, Revenue xuống Values.",
             "Hai Pivot giống nhau."),
            ("Bấm chuột phải vào một ô số của Pivot thứ hai, chọn Show Values As > % Difference From. Trong hộp thoại: Base field chọn Thang, "
             "Base item chọn (previous), bấm OK.",
             "Cột Tháng 1 để trống vì chưa có tháng trước. Cột Tháng 2, 3, 4 là phần trăm tăng giảm so với tháng liền trước."),
            ("Đọc cột Tháng 4 của Pivot thứ hai. Ghi lại tên miền có phần trăm cao nhất và con số đó.",
             "Một miền, một phần trăm. Dùng lại ở câu 6 và câu 7."),
        ]),
        ("CÂU 3. Cơ cấu ngành hàng (sheet PV_Nganh)", None, [
            ("Tạo Pivot mới từ tblSales4T vào New Worksheet, đổi tên sheet thành PV_Nganh.", None),
            ("Kéo Category xuống Rows, Thang xuống Columns, Revenue xuống Values.", "Bảng doanh thu theo ngành hàng và tháng."),
            ("Bấm chuột phải vào một ô số, chọn Show Values As > % of Column Total.",
             "Mỗi cột tháng cộng lại đúng 100% ở dòng Grand Total."),
            ("Tìm dòng Bánh kẹo. Ghi lại tỷ trọng tháng 1 và tháng 4, rồi tính chênh lệch giữa hai tháng, đơn vị là điểm phần trăm.",
             "Hai con số và một chênh lệch. Dùng lại ở câu 6 và câu 7."),
        ]),
        ("CÂU 4. ST07 đã hồi phục chưa? (sheet PV_CuaHang)", None, [
            ("Tạo Pivot mới từ tblSales4T vào New Worksheet, đổi tên sheet thành PV_CuaHang. Kéo Store_Code xuống Rows, Thang xuống Columns, Revenue xuống Values.",
             "Doanh thu từng cửa hàng theo tháng."),
            ("Tạo Pivot thứ hai trên cùng sheet: Insert > PivotTable, Table/Range ghi tblSales4T, chọn Existing Worksheet, Location gõ PV_CuaHang!$H$4, OK. "
             "Kéo Store_Code xuống Rows, Thang xuống Columns, Revenue xuống Values.",
             "Pivot thứ hai nằm bên phải Pivot thứ nhất, không chồng lên nhau."),
            ("Ở Pivot thứ hai, bấm chuột phải vào một ô số, chọn Show Values As > % Difference From, Base field là Thang, Base item là (previous), OK.",
             "Cột tháng 2, 3, 4 hiện phần trăm so với tháng trước."),
            ("Ở Pivot thứ nhất, bấm đúp vào ô giao của dòng ST07 và cột Tháng 4.",
             "Excel mở sheet mới, liệt kê các dòng bán hàng của ST07 trong tháng 4."),
            ("Trong sheet vừa mở: bấm vào chữ cột So_don để chọn cả cột, nhìn góc phải thanh trạng thái phía dưới, đọc Sum. Đó là số đơn của ST07 tháng 4. "
             "Rồi bấm một ô trống bên cạnh, gõ =COUNTA(UNIQUE( , quét đúng các ô ngày trong cột Order_Date (bỏ tiêu đề và ô trống), gõ )) và Enter. "
             "Cần Excel 365; nếu không có UNIQUE, sao chép cột ngày sang vùng trống, Data > Remove Duplicates rồi đếm các ngày còn lại.",
             "Hai con số: số đơn, số ngày có bán của ST07 tháng 4."),
            ("Làm lại hai bước trên cho ô giao của dòng ST07 và cột Tháng 3 (sheet chi tiết mới). Đổi tên hai sheet chi tiết thành CT_ST07_T3 và CT_ST07_T4 cho dễ tìm.",
             "Hai con số của tháng 3. Giờ có bốn số để so."),
            ("Đọc phần trăm của ST07 ở cột Tháng 4 trong Pivot thứ hai. Đặt cạnh số đơn và số ngày có bán của hai tháng, rồi nghĩ: doanh thu tháng 4 tăng nhờ đâu?",
             "Biết ST07 tháng 4 so với tháng 3 tăng hay giảm bao nhiêu phần trăm, và vì sao."),
        ]),
        ("CÂU 5. AOV bằng Calculated Field (sheet PV_AOV)", None, [
            ("Tạo Pivot mới từ tblSales4T vào New Worksheet, đổi tên sheet thành PV_AOV. Kéo Thang xuống Rows, Revenue và So_don xuống Values.",
             "Hai cột Sum of Revenue và Sum of So_don theo tháng."),
            ("Bấm vào một ô trong Pivot. Vào tab PivotTable Analyze > Fields, Items & Sets > Calculated Field.",
             "Hộp thoại Insert Calculated Field mở ra."),
            ("Ô Name gõ AOV. Ô Formula xóa số 0 mặc định, gõ =Revenue/So_don (có thể chọn Revenue trong danh sách Fields rồi bấm Insert Field). "
             "Bấm Add, rồi OK.",
             "Cột Sum of AOV xuất hiện cạnh hai cột cũ."),
            ("Bấm chuột phải vào một ô của cột AOV, chọn Value Field Settings. Custom Name gõ AOV (VNĐ/đơn). Bấm Number Format, chọn Number, bỏ số thập phân, tick Use 1000 Separator, OK.",
             "Cột AOV hiện số nguyên có dấu phân cách nghìn."),
            ("Đọc AOV của tháng 4. Sau đó bấm mũi tên cạnh Row Labels, bỏ tick tháng 4, bấm OK. Lúc này dòng Grand Total là AOV của quý 1 (tháng 1 đến 3). "
             "Ghi lại số đó, rồi mở lại bộ lọc và tick Select All để trả đủ bốn tháng.",
             "Hai con số: AOV tháng 4 và AOV quý 1. Pivot trở lại đủ bốn tháng."),
            ("So hai số: AOV tháng 4 cao hơn hay thấp hơn quý 1, chênh bao nhiêu VNĐ. AOV là tổng doanh thu chia tổng số đơn, nên không lấy trung bình của các AOV từng tháng.",
             "Một câu so sánh có số."),
        ]),
        ("CÂU 6. PivotChart và Slicer", None, [
            ("Đặt tên cho từng Pivot để chọn cho dễ ở bước nối. Bấm vào một ô trong Pivot, vào tab PivotTable Analyze, gõ tên vào ô PivotTable Name ở góc trái trên rồi Enter. "
             "PV_Mien: Pivot thứ nhất đặt pvtMienDT, Pivot thứ hai đặt pvtMienPct. PV_Nganh: pvtNganh. PV_CuaHang: Pivot thứ nhất pvtCuaHangDT, Pivot thứ hai pvtCuaHangPct.",
             "Năm Pivot của câu 2, 3, 4 có tên riêng."),
            ("Câu 2: bấm vào một ô trong pvtMienDT (Pivot doanh thu), vào PivotTable Analyze > PivotChart, chọn Clustered Column, OK. "
             "Bấm vào tiêu đề biểu đồ, xóa chữ cũ và gõ câu kết luận có số. Khuôn mẫu: Tháng 4: Miền ... tăng mạnh nhất (+...% so với tháng 3).",
             "Biểu đồ cột theo miền, tiêu đề là câu kết luận."),
            ("Câu 3: làm tương tự với pvtNganh. Khuôn mẫu tiêu đề: Bánh kẹo ...% ở tháng 1, ...% ở tháng 4.",
             "Biểu đồ cột tỷ trọng ngành hàng, tiêu đề là câu kết luận."),
            ("Câu 4: làm tương tự với pvtCuaHangDT. Khuôn mẫu tiêu đề: ST07: tháng 4 đạt ... triệu, ...% so với tháng 3.",
             "Biểu đồ cột theo cửa hàng, tiêu đề là câu kết luận."),
            ("Bấm vào một ô trong pvtMienDT, vào PivotTable Analyze > Insert Slicer, tick Region, OK.",
             "Một Slicer có ba nút Bắc, Trung, Nam hiện trên sheet PV_Mien."),
            ("Bấm chuột phải vào Slicer, chọn Report Connections. Tick cả năm Pivot: pvtMienDT, pvtMienPct, pvtNganh, pvtCuaHangDT, pvtCuaHangPct. Bấm OK. "
             "Nếu thiếu Pivot, kiểm lại nguồn tblSales4T và Pivot cache; có thể sao chép một Pivot đã nối được rồi đổi trường để dùng chung cache.",
             "Danh sách có đủ năm Pivot, cả năm đều được tick."),
            ("Thử bấm vào một miền trên Slicer rồi mở các sheet PV_Mien, PV_Nganh, PV_CuaHang. Pivot và biểu đồ đổi theo. Bấm biểu tượng phễu gạch chéo ở góc phải trên của Slicer để bỏ lọc.",
             "Trước khi nộp, Slicer ở trạng thái không lọc (cả ba miền sáng)."),
        ]),
        ("CÂU 7. Viết insight (sheet Insight)", None, [
            ("Mở sheet Insight đã tạo ở câu 1. Ở ô A4 gõ: Insight. Ở ô A5 viết insight: 1 câu dài hoặc 2 câu liền ý.",
             "A1 và A2 là lỗi đã sửa, A4 và A5 là insight."),
            ("Insight có ba phần, theo thứ tự này. Một: con số có so sánh, ví dụ doanh thu tháng 4 so với tháng 3, miền tăng mạnh nhất, hoặc ST07. "
             "Hai: nguyên nhân khả dĩ, rút ra từ những gì bạn đã thấy trong Pivot và sheet chi tiết. "
             "Ba: hàm ý, tức người đọc cần làm gì hoặc cần cẩn thận gì khi dùng các con số này.",
             "Đủ ba phần, không bỏ phần nào."),
            ("Đọc lại. Mỗi con số trong insight phải khớp một ô trong Pivot của bạn, và có ít nhất một phép so sánh (so với tháng trước, hoặc giữa hai miền, hai cửa hàng).",
             None),
            ("Bấm Ctrl+S để lưu. Nộp file tại địa chỉ ghi ở dòng 3 phía trên, trước Buổi 4.",
             "File đã nộp tên B3_BTVN_<HoTen>.xlsx."),
        ]),
    ]


def fill_de_bai(de, f):
    """Ghi toàn bộ sheet De_bai. de: worksheet rỗng. f: số liệu đã định dạng (xem build_steps)."""
    de["A1"] = "BÀI TẬP VỀ NHÀ BUỔI 3 (BẮT BUỘC) - BÁO CÁO PIVOT 4 THÁNG ĐẦU NĂM 2025"; de["A1"].font = TITLE
    de["A2"] = (f"Nộp trước Buổi 4. Dữ liệu tháng 1 đến 4/2025 đã làm giàu (tblSales4T, {f['rows']} dòng). "
                "Phòng IT vừa ghép thêm tháng 4 vào, hãy kiểm tra trước khi tin các con số.")
    de["A2"].font = NOTE
    de["A3"] = "Nộp bài tại https://hocvien.mcivietnam.com/auth#/ - Tài khoản: email bạn đăng ký tại MCI · Mật khẩu: hv12345678"; de["A3"].font = BOLD
    de["A4"] = "Cách làm: đi lần lượt từ Câu 0 đến Câu 7, mỗi dòng Bước là một thao tác. Làm xong bước nào thì đối chiếu cột \"Bạn sẽ thấy\" rồi mới sang bước tiếp."
    de["A4"].font = BODY

    r = 6
    for col, lab in zip("ABC", ("Bước", "Thao tác", "Bạn sẽ thấy")):
        c = de[f"{col}{r}"]; c.value = lab; c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = TOP_WRAP
    r += 1
    for title, _, steps in build_steps(f):
        for col in "ABC":
            de[f"{col}{r}"].fill = Q_FILL
        de[f"A{r}"] = title; de[f"A{r}"].font = Q_FONT
        de.row_dimensions[r].height = 20
        r += 1
        for n, (act, res) in enumerate(steps, start=1):
            de[f"A{r}"] = f"Bước {n}"; de[f"B{r}"] = act; de[f"C{r}"] = res
            for col in "ABC":
                de[f"{col}{r}"].font = BODY; de[f"{col}{r}"].alignment = TOP_WRAP
            lines = max(math.ceil(len(act) / 105), math.ceil(len(res or "") / 52), 1)
            de.row_dimensions[r].height = 13.5 * lines + 4
            r += 1

    r += 1
    de[f"A{r}"] = "Tự kiểm tra trước khi nộp"; de[f"A{r}"].font = BOLD
    checks = [
        f"Pivot theo tháng khớp So_lieu_he_thong ở cả 3 cột Revenue, So_don, Quantity (tổng Quantity tháng 4 = {f['qty4']})",
        f"Tổng doanh thu 4 tháng = {f['rev']} VNĐ; tháng 4 = {f['rev4']} VNĐ",
        f"Tháng 4 so với tháng 3: {f['mom4']}",
        "File có đủ 6 sheet mới: PV_DoiSoat, PV_Mien, PV_Nganh, PV_CuaHang, PV_AOV, Insight",
        "Mỗi PivotChart có tiêu đề là câu kết luận; Slicer Region lọc được các Pivot của câu 2, 3, 4",
        "Sheet Insight có đủ hai phần: lỗi dữ liệu đã sửa (A1, A2) và insight (A4, A5)",
    ]
    for tx in checks:
        r += 1
        de[f"B{r}"] = "[ ] " + tx; de[f"B{r}"].font = BODY

    for col, w in WIDTHS.items():
        de.column_dimensions[col].width = w
    de.freeze_panes = "A7"
    return de
