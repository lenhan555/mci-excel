"""Build instructor demo files (thư mục Demo/ của từng buổi) from the course files already built.

Mỗi demo "(→ demo …)" trong Slide_Script_B<n>.md phải có file demo tương ứng (xem CLAUDE.md ở gốc dự án).
Script chỉ đọc file đã có (Answer, Dirty, MCI_Mart_Master), không cần thư mục build/.
"""
import os
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.normpath(os.environ.get("MCI_ROOT") or os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
B1 = f"{ROOT}/1_Sessions/B1_Nen_tang_Excel_Du_lieu"
B1_DEMO = f"{B1}/Demo"
MASTER = f"{ROOT}/2_Dataset/MCI_Mart_Master.xlsx"
os.makedirs(B1_DEMO, exist_ok=True)

F = "Arial"
HDR_FONT = Font(name=F, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="203864")
TITLE = Font(name=F, size=14, bold=True, color="203864")
BODY = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
INPUT = Font(name=F, size=10, color="0000FF")
NOTE = Font(name=F, size=10, italic=True, color="585858")
DATE_FMT = "dd/mm/yyyy"
NUM_FMT = "#,##0"
PCT_FMT = "0.0%"

def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

def header(ws, row, labels):
    for j, h in enumerate(labels, start=1):
        c = ws.cell(row=row, column=j, value=h); c.font = HDR_FONT; c.fill = HDR_FILL

def guide(ws, title, steps):
    """Sheet hướng dẫn cho giảng viên: tiêu đề + các bước theo slide."""
    ws["A1"] = title; ws["A1"].font = TITLE
    header(ws, 3, ["Slide", "Thao tác", "Kết quả mong đợi"])
    for i, (slide, action, result) in enumerate(steps, start=4):
        for j, v in enumerate((slide, action, result), start=1):
            c = ws.cell(row=i, column=j, value=v); c.font = BODY
            c.alignment = Alignment(wrap_text=True, vertical="top")
    set_widths(ws, [8, 70, 45])

def save_both(wb, name, work_sheet, hv_replace=()):
    """Lưu 2 bản: bản giảng viên (mở ra là sheet làm việc, Huong_dan để cuối, tránh chiếu lên màn hình)
    và bản học viên `_HocVien` (Huong_dan đứng đầu; bỏ các câu chỉ dành cho giảng viên)."""
    g = wb["Huong_dan"]
    wb.move_sheet(g, offset=len(wb.sheetnames) - 1 - wb.sheetnames.index("Huong_dan"))
    wb.active = wb.sheetnames.index(work_sheet)
    for ws in wb.worksheets: ws.sheet_view.tabSelected = ws.title == work_sheet
    wb.save(f"{B1_DEMO}/{name}.xlsx")
    g["A1"] = g["A1"].value.replace("DEMO B1", "VÍ DỤ B1")
    for row in g.iter_rows(min_row=4):
        for c in row:
            if isinstance(c.value, str):
                for a, b in hv_replace: c.value = c.value.replace(a, b)
    for ws in wb.worksheets:
        if ws.title.startswith("Demo_"): ws.title = "Vi_du_" + ws.title[len("Demo_"):]
    wb.move_sheet(g, offset=-wb.sheetnames.index("Huong_dan"))
    wb.active = 0
    for ws in wb.worksheets: ws.sheet_view.tabSelected = ws.title == "Huong_dan"
    wb.save(f"{B1_DEMO}/{name}_HocVien.xlsx")

# Dữ liệu gốc: sheet Sales của file Answer (cột A:J là giá trị, K:L là công thức)
ans = pd.read_excel(f"{B1}/Practice/B1_MCI_Mart_Q1_2025_Answer.xlsx", "Sales", usecols="A:J")
stores = pd.read_excel(f"{B1}/Practice/B1_MCI_Mart_Q1_2025_Answer.xlsx", "Store")
revenue = ans.Quantity * ans.Unit_Price

# =====================================================================
# B1_Demo_PctShare.xlsx — slide 19 (tham chiếu $), 21 (Trace), 22 (Evaluate Formula, Watch Window)
# Bố cục: tiêu đề cột ở dòng 1, ô tổng ở B5, để công thức =IFERROR(B2/$B$5,0) trong kịch bản dùng được nguyên văn.
# B6, B7 phải để trống: bản sai kéo xuống sẽ chia cho hai ô này và ra #DIV/0!.
# =====================================================================
wb = Workbook(); ws = wb.active; ws.title = "Demo_PctShare"
header(ws, 1, ["Miền", "Doanh thu (VNĐ)", "% sai (không $)", "% đúng (có $)", "IFERROR (slide 22)"])
by_region = revenue.groupby(ans.Region).sum()
for k, reg in enumerate(["Bắc", "Trung", "Nam"], start=2):
    ws.cell(row=k, column=1, value=reg).font = BODY
    c = ws.cell(row=k, column=2, value=int(by_region[reg])); c.font = INPUT; c.number_format = NUM_FMT
    c = ws.cell(row=k, column=3, value=f"=B{k}/B{k + 3}"); c.font = BODY; c.number_format = PCT_FMT
    c = ws.cell(row=k, column=4, value=f"=B{k}/$B$5"); c.font = BODY; c.number_format = PCT_FMT
ws["E2"] = "=IFERROR(B2/$B$5,0)"; ws["E2"].font = BODY; ws["E2"].number_format = PCT_FMT
ws["A5"] = "Tổng"; ws["B5"] = "=SUM(B2:B4)"; ws["D5"] = "=SUM(D2:D4)"
for a in ("A5", "B5", "D5"): ws[a].font = BOLD
ws["B5"].number_format = NUM_FMT; ws["D5"].number_format = PCT_FMT
set_widths(ws, [10, 18, 16, 16, 18])
ws.sheet_view.zoomScale = 160

g = wb.create_sheet("Huong_dan")
guide(g, "DEMO B1 — % ĐÓNG GÓP THEO MIỀN (slide 19, 21, 22)", [
    (19, "Xóa nội dung C2:C4 (chọn C2:C4, phím Delete; không xóa cả cột) rồi làm lại trước lớp: gõ =B2/B5 ở C2, kéo xuống C4. Đứng ở C3, nhấn F2.",
     "Trung, Nam ra #DIV/0!; F2 cho thấy mẫu số trôi xuống B6 (ô trống)."),
    (19, "Sửa C2 thành =B2/$B$5 (con trỏ ở B5 trong công thức, nhấn F4), kéo lại xuống C4.",
     f"34,4% / 20,2% / 45,4%; tổng 100% (giống cột D)."),
    (21, "Chọn D2 → Formulas → Trace Precedents. Chọn B5 → Trace Dependents. Remove Arrows khi xong.",
     "Mũi tên từ B2, B5 tới D2; từ B5 tới mọi ô % dùng ô tổng: cột C (C2 nếu chưa sửa, C2:C4 sau khi sửa), D2:D4 và E2."),
    (22, "Chọn E2 → Formulas → Evaluate Formula → bấm Evaluate từng bước.",
     "B2/$B$5 → 0,344… → IFERROR trả về 0,344…"),
    (22, "Formulas → Watch Window → Add Watch → chọn B5. Sửa B3 thành 0 để thấy tổng đổi, rồi Ctrl+Z.",
     "Watch Window hiện B5 = 941.454.000 và đổi ngay khi sửa B3."),
])
save_both(wb, "B1_Demo_PctShare", "Demo_PctShare",
          hv_replace=[("Xóa nội dung C2:C4 (chọn C2:C4, phím Delete; không xóa cả cột) rồi làm lại trước lớp: gõ =B2/B5 ở C2", "Xóa nội dung C2:C4 (chọn C2:C4, phím Delete; không xóa cả cột), rồi gõ =B2/B5 vào C2")])

# =====================================================================
# B1_Demo_Table_Validation.xlsx — slide 26 (Ctrl+T, tự mở rộng), 27 (Data Validation), 28 (Circle Invalid Data)
# Sheet Sales: dữ liệu đã tidy nhưng CHƯA là Table, chưa có Validation, còn giữ 5 ô sai của file Dirty
# (3 ô ST3, 2 ô Quantity âm) để Circle Invalid Data khoanh đúng 5 ô.
# =====================================================================
dirty_ws = load_workbook(f"{B1}/Practice/B1_MCI_Mart_Q1_2025_Dirty.xlsx", read_only=True).active
bad = {}
for r in dirty_ws.iter_rows(min_row=6, values_only=True):
    oid, line, store, qty = r[1], r[2], r[4], r[8]
    if store == "ST3": bad[(oid, line)] = ("Store_Code", "ST3")
    elif isinstance(qty, (int, float)) and qty < 0: bad[(oid, line)] = ("Quantity", qty)
demo = ans.copy()
for (oid, line), (col, val) in bad.items():
    demo.loc[(demo.Order_ID == oid) & (demo.Line_No == line), col] = val
assert len(bad) == 5, bad

wb = Workbook(); g = wb.active; g.title = "Huong_dan"
ws = wb.create_sheet("Sales")
cols = list(demo.columns)
header(ws, 1, cols)
for i, row in enumerate(demo.itertuples(index=False), start=2):
    for j, v in enumerate(row, start=1):
        if isinstance(v, pd.Timestamp): v = v.to_pydatetime()
        elif hasattr(v, "item"): v = v.item()
        c = ws.cell(row=i, column=j, value=v); c.font = BODY
        if cols[j - 1] == "Order_Date": c.number_format = DATE_FMT
        if cols[j - 1] == "Unit_Price": c.number_format = NUM_FMT
ws.freeze_panes = "A2"
set_widths(ws, [9, 15, 7, 12, 10, 11, 11, 40, 9, 11])

s = wb.create_sheet("Store")
header(s, 1, list(stores.columns))
for i, row in enumerate(stores.itertuples(index=False), start=2):
    for j, v in enumerate(row, start=1):
        if isinstance(v, pd.Timestamp): v = v.to_pydatetime()
        c = s.cell(row=i, column=j, value=v); c.font = BODY
        if j == 6: c.number_format = DATE_FMT
set_widths(s, [11, 24, 13, 9, 11, 12])

# Vài dòng tháng 4/2025 (thật, từ MCI_Mart_Master) để dán ngay dưới bảng khi demo Table tự mở rộng
m_sales = pd.read_excel(MASTER, "Sales")
m_prod = pd.read_excel(MASTER, "Product").set_index("Product_Code")["Product_Name"]
new = m_sales[m_sales.Order_ID.str.startswith("HD2504")].head(5).copy()
new["Region"] = new.Store_Code.map(stores.set_index("Store_Code")["Region"])
new["Product_Name"] = new.Product_Code.map(m_prod)
new = new[cols]
t = wb.create_sheet("Them_dong")
header(t, 1, cols)
for i, row in enumerate(new.itertuples(index=False), start=2):
    for j, v in enumerate(row, start=1):
        if isinstance(v, pd.Timestamp): v = v.to_pydatetime()
        elif hasattr(v, "item"): v = v.item()
        c = t.cell(row=i, column=j, value=v); c.font = BODY
        if cols[j - 1] == "Order_Date": c.number_format = DATE_FMT
        if cols[j - 1] == "Unit_Price": c.number_format = NUM_FMT
set_widths(t, [9, 15, 7, 12, 10, 11, 11, 40, 9, 11])

n = len(demo)
guide(g, "DEMO B1 — EXCEL TABLE & DATA VALIDATION (slide 26, 27, 28)", [
    (26, "Sheet Sales: đứng trong vùng dữ liệu → Ctrl + T → tick My table has headers → OK. Table Design → Table Name = tblSales.",
     f"Table {n:,} dòng".replace(",", ".") + ", tên tblSales."),
    (26, "Gõ Revenue ở K1; K2: =[@Quantity]*[@Unit_Price] → Enter.",
     "Cả cột Revenue tự điền."),
    (26, f"Sheet Them_dong: copy A2:J6 → dán vào A{n + 2} của sheet Sales (ngay dưới bảng).",
     "5 dòng tháng 4 tự thuộc về tblSales; Revenue tự tính."),
    (27, "Chọn cột Store_Code (E2:E…) → Data → Data Validation → List, Source = =Store!$A$2:$A$13; tab Error Alert: viết thông báo tiếng Việt.",
     "Gõ thử ST99 → bị chặn, hiện thông báo tiếng Việt."),
    (27, "Chọn cột Quantity (I2:I…) → Data Validation → Whole number, greater than 0.",
     "Gõ thử -1 hoặc 1,5 → bị chặn."),
    (28, "Data → Data Validation ▾ → Circle Invalid Data.",
     "Đúng 5 ô bị khoanh: 3 ô ST3 (Miền Bắc) và 2 ô Quantity âm."),
])
save_both(wb, "B1_Demo_Table_Validation", "Sales")
print("ok", sorted(bad.items()))
