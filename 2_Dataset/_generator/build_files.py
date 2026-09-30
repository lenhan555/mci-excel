"""Build course files from the generated MCI Mart dataset."""
import json, os, csv, random
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo, TableColumn, TableFormula
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(__file__)
B = os.path.join(HERE, "build")
ROOT = os.path.normpath(os.environ.get("MCI_ROOT") or os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
DATASET = f"{ROOT}/2_Dataset"
B1_DIR = f"{ROOT}/1_Sessions/B1_Nen_tang_Excel_Du_lieu/Practice"
B2_DIR = f"{ROOT}/1_Sessions/B2_Ham_va_KPI/Practice"
B4_DIR = f"{ROOT}/1_Sessions/B4_Power_Query/Practice"
stores = pd.read_pickle(f"{B}/stores.pkl")
products = pd.read_pickle(f"{B}/products.pkl")
customers = pd.read_pickle(f"{B}/customers.pkl")
sales = pd.read_pickle(f"{B}/sales.pkl")
target = pd.read_pickle(f"{B}/target.pkl")
rnd = random.Random(7)

F = "Arial"
HDR_FONT = Font(name=F, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="203864")
BODY = Font(name=F, size=10)
NOTE = Font(name=F, size=10, italic=True, color="585858")
TITLE = Font(name=F, size=14, bold=True, color="C00000")
DATE_FMT = "dd/mm/yyyy"
NUM_FMT = "#,##0"

def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

def write_table(ws, df, name, start_row=1, style="TableStyleMedium2", formats=None):
    cols = list(df.columns)
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=start_row, column=j, value=c)
        cell.font = Font(name=F, bold=True)
    for i, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for j, v in enumerate(row, start=1):
            if isinstance(v, pd.Timestamp):
                v = v.to_pydatetime()
            elif isinstance(v, (np.integer,)):
                v = int(v)
            elif isinstance(v, (np.floating,)):
                v = float(v)
            c = ws.cell(row=i, column=j, value=v)
            c.font = BODY
            if formats and cols[j - 1] in formats:
                c.number_format = formats[cols[j - 1]]
    ref = f"A{start_row}:{get_column_letter(len(cols))}{start_row + len(df)}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(t)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    return t

store_region = stores.set_index("Store_Code")["Region"]
prod_name = products.set_index("Product_Code")["Product_Name"]

# =====================================================================
# 1) Master dataset (instructor reference)
# =====================================================================
os.makedirs(f"{DATASET}", exist_ok=True)
wb = Workbook()
ws = wb.active; ws.title = "README"
readme = [
    ("MCI MART — BỘ DỮ LIỆU CASE (DỮ LIỆU GIẢ LẬP)", TITLE),
    ("Khóa học: Excel for Business Intelligence Analyst — MCI. Dữ liệu được sinh ngẫu nhiên có kiểm soát (seed 2025); mọi tên người, cửa hàng, sản phẩm là hư cấu.", NOTE),
    ("Phạm vi: giao dịch 01/01/2024 – 30/04/2025; 12 cửa hàng, 3 miền; 40 sản phẩm, 6 ngành hàng; 1.500 khách hàng.", NOTE),
    ("", None),
    ("Bảng", Font(name=F, bold=True)),
    ("Sales — 1 dòng = 1 dòng sản phẩm trong 1 đơn hàng (grain: Order_ID × Line_No). Revenue = Quantity × Unit_Price (giá bán tại thời điểm giao dịch).", BODY),
    ("Product — danh mục sản phẩm; Unit_Price = giá niêm yết 2025; Unit_Cost = giá vốn dùng tính biên lợi nhuận.", BODY),
    ("Store — cửa hàng; Region = Bắc/Trung/Nam; ST12 khai trương 01/06/2024.", BODY),
    ("Customer — khách hàng thành viên; Segment = Thường/Thân thiết/VIP.", BODY),
    ("Target_2025 — mục tiêu doanh thu 2025 theo Cửa hàng × Tháng, dạng cột tháng (cần Unpivot ở B4). = doanh thu cùng tháng 2024 × 110%, làm tròn triệu; ST12 dùng trung bình T9–T12/2024.", BODY),
    ("", None),
    ("Các \"câu chuyện\" cài sẵn để khai thác insight", Font(name=F, bold=True)),
    ("1. Tết Ất Tỵ (29/01/2025) và Tết Giáp Thìn (10/02/2024): doanh thu tăng mạnh ~3 tuần trước Tết, giảm sâu trong kỳ nghỉ; Bánh kẹo & Đồ uống tăng tỷ trọng.", BODY),
    ("2. ST07 Nha Trang cải tạo mặt bằng 10–31/03/2025: doanh thu tháng 3/2025 giảm rõ rệt so với target và cùng kỳ.", BODY),
    ("3. ST12 Biên Hòa mới mở 06/2024: tăng dần trong 3 tháng đầu; không có dữ liệu cùng kỳ cho T1–T5 → YoY trống/không so được.", BODY),
    ("4. Giá niêm yết 2025 tăng ~4% so với 2024; tăng trưởng nền ~8%.", BODY),
    ("5. Cuối tuần doanh thu cao hơn ~25%.", BODY),
    ("6. Target 2025 lập theo cùng tháng 2024 nên KHÔNG tính việc Tết dời từ tháng 2 (2024) sang tháng 1 (2025): T1/2025 vượt target (~121%), T2/2025 hụt (~83%) — bài học về cách lập target và đọc % đạt.", BODY),
    ("7. Biên lợi nhuận gộp 2025 (~33–35%) cao hơn 2024 (~31%) vì giá bán tăng trong khi giá vốn dùng chung bảng Product 2025 — cần nêu rõ là giả định của dataset.", BODY),
    ("8. Mức đạt target tháng 3/2025 theo cửa hàng dao động mạnh (ST10 ~145%, ST07 ~45%) — ST07 là trường hợp có nguyên nhân rõ (cải tạo).", BODY),
]
for i, (txt, font) in enumerate(readme, start=1):
    c = ws.cell(row=i, column=1, value=txt)
    if font: c.font = font
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 140

ws = wb.create_sheet("Sales")
write_table(ws, sales, "tblSales_All", formats={"Order_Date": DATE_FMT, "Unit_Price": NUM_FMT, "Revenue": NUM_FMT})
set_widths(ws, [15, 8, 12, 11, 12, 12, 9, 11, 13])
ws = wb.create_sheet("Product")
write_table(ws, products, "tblProduct", formats={"Unit_Price": NUM_FMT, "Unit_Cost": NUM_FMT})
set_widths(ws, [13, 40, 18, 12, 12])
ws = wb.create_sheet("Store")
write_table(ws, stores, "tblStore", formats={"Open_Date": DATE_FMT})
set_widths(ws, [11, 24, 13, 9, 11, 12])
ws = wb.create_sheet("Customer")
write_table(ws, customers, "tblCustomer", formats={"Join_Date": DATE_FMT})
set_widths(ws, [12, 24, 8, 11, 12, 12, 16])
ws = wb.create_sheet("Target_2025")
write_table(ws, target, "tblTarget2025", formats={c: NUM_FMT for c in target.columns if c.startswith("2025")})
set_widths(ws, [11, 24] + [14] * 12)
wb.save(f"{DATASET}/MCI_Mart_Master.xlsx")

# =====================================================================
# 2) B1 — dirty report-style file
# =====================================================================
os.makedirs(f"{B1_DIR}", exist_ok=True)
q1 = sales[(sales.Order_Date >= "2025-01-01") & (sales.Order_Date <= "2025-03-31")].copy()
q1["Region"] = q1["Store_Code"].map(store_region)
q1["Product_Name"] = q1["Product_Code"].map(prod_name)
reg_order = {"Bắc": 0, "Trung": 1, "Nam": 2}
q1["_r"] = q1["Region"].map(reg_order)
q1 = q1.sort_values(["_r", "Order_Date", "Order_ID", "Line_No"]).reset_index(drop=True)
clean_cols = ["Region", "Order_ID", "Line_No", "Order_Date", "Store_Code", "Customer_ID",
              "Product_Code", "Product_Name", "Quantity", "Unit_Price"]
clean = q1[clean_cols].copy()

# choose defects (indices in clean)
idx_all = list(clean.index)
st03 = list(clean.index[clean.Store_Code == "ST03"])
bad_store = rnd.sample(st03, 3)
remaining = [i for i in idx_all if i not in bad_store]
neg_qty = rnd.sample(remaining, 2)
remaining = [i for i in remaining if i not in neg_qty]
txt_qty = rnd.sample(remaining, 15)
remaining = [i for i in remaining if i not in txt_qty]
txt_date = rnd.sample(remaining, 20)

wb = Workbook(); ws = wb.active; ws.title = "BaoCao_Q1_2025"
ws["A1"] = "BÁO CÁO BÁN HÀNG QUÝ 1/2025 — MCI MART"; ws["A1"].font = TITLE
ws.merge_cells("A1:J1")
ws["A2"] = "Xuất từ hệ thống POS ngày 05/04/2025 • Đơn vị tiền: VNĐ"; ws["A2"].font = NOTE
ws.merge_cells("A2:J2")
grp_fill = PatternFill("solid", fgColor="FFC000")
for rng_, label in [("A4:A5", "Vùng"), ("B4:F4", "THÔNG TIN ĐƠN HÀNG"), ("G4:H4", "SẢN PHẨM"), ("I4:J4", "GIÁ TRỊ")]:
    first = rng_.split(":")[0]
    ws[first] = label
    ws.merge_cells(rng_)
    ws[first].font = Font(name=F, bold=True)
    ws[first].fill = grp_fill
    ws[first].alignment = Alignment(horizontal="center", vertical="center")
for j, c in enumerate(clean_cols[1:], start=2):
    cell = ws.cell(row=5, column=j, value=c)
    cell.font = HDR_FONT; cell.fill = HDR_FILL

r = 6
region_fills = {"Bắc": "DDEBF7", "Trung": "E2EFDA", "Nam": "FCE4D6"}
dirty_log = []
blank_row_at = None
for reg in ["Bắc", "Trung", "Nam"]:
    block = clean[clean.Region == reg]
    start = r
    mid = block.index[len(block) // 2]
    for i, row in block.iterrows():
        if reg == "Trung" and i == mid:
            blank_row_at = r
            r += 1                                  # defect: blank row inside data
        vals = row.tolist()
        oid, ln, od, sc, cid, pc, pn, qty, up = vals[1:]
        if i in bad_store:
            sc = "ST3"
        if i in neg_qty:
            qty = -int(qty)
        cells = [None, oid, int(ln), od.to_pydatetime(), sc, cid, pc, pn, int(qty), int(up)]
        if i in txt_qty:
            cells[8] = str(int(qty))
        if i in txt_date:
            cells[3] = od.strftime("%d/%m/%Y")
        for j, v in enumerate(cells, start=1):
            if j == 1: continue
            c = ws.cell(row=r, column=j, value=v)
            c.font = BODY
            c.fill = PatternFill("solid", fgColor=region_fills[reg])
            if j == 4 and not isinstance(v, str): c.number_format = DATE_FMT
            if j == 10: c.number_format = NUM_FMT
        r += 1
    end = r - 1
    ws.cell(row=start, column=1, value=f"Miền {reg}").font = Font(name=F, bold=True)
    ws.merge_cells(start_row=start, start_column=1, end_row=end, end_column=1)
    ws.cell(row=start, column=1).alignment = Alignment(vertical="top")
    # subtotal row (defect)
    ws.cell(row=r, column=2, value=f"Tổng Miền {reg}").font = Font(name=F, bold=True)
    ws.cell(row=r, column=9, value=int(block.Quantity.sum())).font = Font(name=F, bold=True)
    for j in range(2, 11):
        ws.cell(row=r, column=j).fill = PatternFill("solid", fgColor="BFBFBF")
    r += 1
ws.cell(row=r, column=2, value="TỔNG CỘNG").font = Font(name=F, bold=True, color="C00000")
ws.cell(row=r, column=9, value=int(clean.Quantity.sum())).font = Font(name=F, bold=True, color="C00000")
set_widths(ws, [12, 15, 7, 12, 10, 11, 11, 40, 8, 11])
dirty_path = f"{B1_DIR}/B1_MCI_Mart_Q1_2025_Dirty.xlsx"
wb.save(dirty_path)

# =====================================================================
# 3) B1 — answer file (also B2 starting file)
# =====================================================================
wb = Workbook(); ws = wb.active; ws.title = "Sales"
ans = clean.copy()
n = len(ans)
ans_cols = clean_cols + ["Revenue", "Pct_Share"]
for j, c in enumerate(ans_cols, start=1):
    ws.cell(row=1, column=j, value=c).font = Font(name=F, bold=True)
for i, row in enumerate(ans.itertuples(index=False), start=2):
    vals = list(row)
    for j, v in enumerate(vals, start=1):
        if isinstance(v, pd.Timestamp): v = v.to_pydatetime()
        elif isinstance(v, np.integer): v = int(v)
        c = ws.cell(row=i, column=j, value=v); c.font = BODY
        if j == 4: c.number_format = DATE_FMT
        if j == 10: c.number_format = NUM_FMT
    c = ws.cell(row=i, column=11, value="=tblSales[[#This Row],[Quantity]]*tblSales[[#This Row],[Unit_Price]]")
    c.font = BODY; c.number_format = NUM_FMT
    c = ws.cell(row=i, column=12, value="=tblSales[[#This Row],[Revenue]]/'Kiem_tra'!$C$5")
    c.font = BODY; c.number_format = "0.000%"
t = Table(displayName="tblSales", ref=f"A1:L{n + 1}")
t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)

ws.add_table(t)
ws.freeze_panes = "A2"
set_widths(ws, [9, 15, 7, 12, 10, 11, 11, 40, 9, 11, 13, 10])
dv_store = DataValidation(type="list", formula1="=Store!$A$2:$A$13", allow_blank=False,
                          showErrorMessage=True, errorTitle="Mã cửa hàng không hợp lệ",
                          error="Chỉ nhập mã có trong sheet Store (ST01–ST12).")
dv_store.add(f"E2:E{n + 1}")
dv_qty = DataValidation(type="whole", operator="greaterThan", formula1="0", showErrorMessage=True,
                        errorTitle="Số lượng không hợp lệ", error="Quantity phải là số nguyên lớn hơn 0.")
dv_qty.add(f"I2:I{n + 1}")
ws.add_data_validation(dv_store); ws.add_data_validation(dv_qty)

ws2 = wb.create_sheet("Store")
write_table(ws2, stores, "tblStore", formats={"Open_Date": DATE_FMT})
set_widths(ws2, [11, 24, 13, 9, 11, 12])

exp_rev = int((ans.Quantity * ans.Unit_Price).sum())
ws3 = wb.create_sheet("Kiem_tra")
ws3["A1"] = "KIỂM TRA / ĐỐI SOÁT — Brick output B1"; ws3["A1"].font = TITLE
ws3["A2"] = "Cột C = công thức tính trên tblSales; cột D = giá trị đáp án (tính từ bộ sinh dữ liệu, xem README của Dataset); cột E = khớp?"
ws3["A2"].font = NOTE
hdr = ["Chỉ tiêu", "", "Kết quả (công thức)", "Đáp án", "Khớp?"]
for j, h in enumerate(hdr, start=1):
    c = ws3.cell(row=4, column=j, value=h); c.font = HDR_FONT; c.fill = HDR_FILL
checks = [
    ("Tổng doanh thu Q1/2025 (VNĐ)", "=SUM(tblSales[Revenue])", exp_rev, NUM_FMT),
    ("Số dòng dữ liệu", "=ROWS(tblSales[Order_ID])", n, NUM_FMT),
    ("Tổng số lượng (Quantity)", "=SUM(tblSales[Quantity])", int(ans.Quantity.sum()), NUM_FMT),
    ("Tổng cột Pct_Share", "=SUM(tblSales[Pct_Share])", 1, "0.00%"),
    ("Số dòng Store_Code không có trong danh mục", "=SUMPRODUCT(--(COUNTIF(tblStore[Store_Code],tblSales[Store_Code])=0))", 0, "0"),
    ("Số dòng Quantity ≤ 0", '=COUNTIF(tblSales[Quantity],"<=0")', 0, "0"),
    ("Số ô Quantity dạng text", "=SUMPRODUCT(--ISTEXT(tblSales[Quantity]))", 0, "0"),
    ("Số ô Order_Date dạng text", "=SUMPRODUCT(--ISTEXT(tblSales[Order_Date]))", 0, "0"),
]
for k, (label, f, exp, fmt) in enumerate(checks, start=5):
    ws3.cell(row=k, column=1, value=label).font = BODY
    c = ws3.cell(row=k, column=3, value=f); c.number_format = fmt; c.font = BODY
    c = ws3.cell(row=k, column=4, value=exp); c.number_format = fmt; c.font = Font(name=F, size=10, color="0000FF")
    c = ws3.cell(row=k, column=5, value=f'=IF(ROUND(C{k},6)=ROUND(D{k},6),"✔","✘")'); c.font = BODY
ws3["A14"] = "Ghi chú: C5 là ô tổng doanh thu mà cột Pct_Share tham chiếu tuyệt đối ($C$5) — ví dụ minh họa slide 19."
ws3["A14"].font = NOTE
set_widths(ws3, [44, 2, 20, 16, 8])

ws4 = wb.create_sheet("Dap_an_loi")
ws4["A1"] = "ĐÁP ÁN — CÁC LỖI CÀI SẴN TRONG FILE DIRTY VÀ CÁCH SỬA"; ws4["A1"].font = TITLE
rows = [
    ("#", "Loại lỗi", "Số lượng", "Vị trí / nhận biết", "Cách sửa"),
    (1, "Tiêu đề báo cáo (dòng 1–2) và dòng trống (dòng 3)", "3 dòng", "Trên cùng sheet", "Xóa dòng 1–3"),
    (2, "Tiêu đề 2 dòng, ô gộp nhóm (dòng 4)", "1 dòng", "\"THÔNG TIN ĐƠN HÀNG\", \"SẢN PHẨM\", \"GIÁ TRỊ\"", "Unmerge; xóa dòng nhóm; đặt tên cột A = Region"),
    (3, "Ô gộp cột Vùng theo khối miền", "3 khối", "Cột A", "Unmerge → Go To Special > Blanks → =ô trên → Ctrl+Enter → Paste Values; đổi \"Miền Bắc\" → \"Bắc\"…"),
    (4, "Dòng tổng giữa dữ liệu", "3 dòng + 1 dòng TỔNG CỘNG", "\"Tổng Miền …\", \"TỔNG CỘNG\" (cột B)", "Filter cột B chứa \"Tổng\" → xóa"),
    (5, "Dòng trống giữa dữ liệu", "1 dòng", f"Giữa khối Miền Trung (dòng {blank_row_at} trong file Dirty)", "Xóa dòng; kiểm tra bằng Ctrl+↓"),
    (6, "Quantity lưu dạng text", "15 ô", "Tam giác xanh, căn trái", "Error Checking → Convert to Number"),
    (7, "Order_Date lưu dạng text dd/mm/yyyy", "20 ô", "Căn trái, không đổi định dạng được", "DATEVALUE hoặc Text to Columns (DMY)"),
    (8, "Store_Code sai \"ST3\"", "3 ô", "Circle Invalid Data sau khi đặt Validation", "Sửa thành ST03 (cùng Miền Bắc, đối chiếu sheet Store)"),
    (9, "Quantity âm", "2 ô", "Circle Invalid Data (Whole > 0)", "Lỗi dấu khi nhập — đổi thành số dương"),
    ("—", "Ghi chú, không tính là lỗi dữ liệu: định dạng thừa (tô màu theo miền, màu tiêu đề)", "—", "Toàn bảng", "Clear Formats trước khi Ctrl+T (tùy chọn)"),
]
for i, rr in enumerate(rows, start=3):
    for j, v in enumerate(rr, start=1):
        c = ws4.cell(row=i, column=j, value=v)
        c.font = HDR_FONT if i == 3 else BODY
        if i == 3: c.fill = HDR_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
set_widths(ws4, [4, 42, 16, 40, 60])
ws4["A15"] = f"Sau khi sửa: {n:,} dòng dữ liệu; tổng doanh thu {exp_rev:,} VNĐ.".replace(",", ".")
ws4["A15"].font = Font(name=F, bold=True)

ans_path = f"{B1_DIR}/B1_MCI_Mart_Q1_2025_Answer.xlsx"
wb.save(ans_path)

# =====================================================================
# 4) Raw data for later sessions (B4 CSV folder, lookups, target)
# =====================================================================
# B2_Product_DuplicateKey.xlsx (P021 trùng khóa) nay dựng bằng build_b2.py, bản phát học viên không còn ghi chú giảng viên ở G1.

os.makedirs(f"{B4_DIR}/Data_Sales", exist_ok=True)
os.makedirs(f"{B4_DIR}/New_Month", exist_ok=True)
csv_cols = ["Order_ID", "Line_No", "Order_Date", "Store_Code", "Customer_ID", "Product_Code", "Quantity", "Unit_Price"]

def write_csv(df, path, dirty=False, seed=0):
    r_ = random.Random(seed)
    out = df[csv_cols].copy()
    out["Order_Date"] = out["Order_Date"].dt.strftime("%d/%m/%Y")
    out["Quantity"] = out["Quantity"].astype(str)
    out["Unit_Price"] = out["Unit_Price"].astype(int).astype(str)
    log = {}
    if dirty:
        # thousands separator VN style "25.000"
        out["Unit_Price"] = out["Unit_Price"].map(lambda s: f"{int(s):,}".replace(",", "."))
        # trailing spaces in Store_Code
        idx = r_.sample(list(out.index), 25)
        out.loc[idx, "Store_Code"] = out.loc[idx, "Store_Code"] + " "
        log["store_code_trailing_space"] = 25
        # one exact duplicate row
        d = out.loc[[r_.choice(list(out.index))]]
        out = pd.concat([out, d]).sort_index(kind="stable")
        log["duplicate_rows"] = 1
    out.to_csv(path, index=False, encoding="utf-8-sig", quoting=csv.QUOTE_MINIMAL)
    return log

logs = {}
s24 = sales[sales.Order_Date.dt.year == 2024]
logs["Sales_2024.csv"] = write_csv(s24, f"{B4_DIR}/Data_Sales/Sales_2024.csv")
for m in (1, 2, 3):
    part = sales[(sales.Order_Date.dt.year == 2025) & (sales.Order_Date.dt.month == m)]
    logs[f"Sales_2025-{m:02d}.csv"] = write_csv(part, f"{B4_DIR}/Data_Sales/Sales_2025-{m:02d}.csv", dirty=True, seed=m)
part = sales[(sales.Order_Date.dt.year == 2025) & (sales.Order_Date.dt.month == 4)]
logs["Sales_2025-04.csv"] = write_csv(part, f"{B4_DIR}/New_Month/Sales_2025-04.csv", dirty=True, seed=4)

wb = Workbook(); ws = wb.active; ws.title = "Product"
p_b4 = pd.concat([products, products[products.Product_Code == "P015"]]).sort_values("Product_Code", kind="stable")
write_table(ws, p_b4, "Product", formats={"Unit_Price": NUM_FMT, "Unit_Cost": NUM_FMT})
set_widths(ws, [13, 40, 18, 12, 12])
ws = wb.create_sheet("Store"); write_table(ws, stores, "Store", formats={"Open_Date": DATE_FMT}); set_widths(ws, [11, 24, 13, 9, 11, 12])
ws = wb.create_sheet("Customer"); write_table(ws, customers, "Customer", formats={"Join_Date": DATE_FMT}); set_widths(ws, [12, 24, 8, 11, 12, 12, 16])
wb.save(f"{B4_DIR}/B4_Lookups.xlsx")

wb = Workbook(); ws = wb.active; ws.title = "Target_2025"
ws["A1"] = "KẾ HOẠCH DOANH THU 2025 THEO CỬA HÀNG (VNĐ)"; ws["A1"].font = TITLE
tcols = ["Store_Code", "Store_Name"] + [f"Tháng {m}" for m in range(1, 13)]
for j, c in enumerate(tcols, start=1):
    x = ws.cell(row=3, column=j, value=c); x.font = HDR_FONT; x.fill = HDR_FILL
for i, row in enumerate(target.itertuples(index=False), start=4):
    for j, v in enumerate(row, start=1):
        c = ws.cell(row=i, column=j, value=int(v) if isinstance(v, (np.integer,)) else v)
        c.font = BODY
        if j > 2: c.number_format = NUM_FMT
set_widths(ws, [11, 24] + [14] * 12)
wb.save(f"{B4_DIR}/B4_Target_2025.xlsx")

meta = {
    "b1_rows": n, "b1_revenue": exp_rev, "b1_quantity": int(ans.Quantity.sum()),
    "b1_orders": int(ans.Order_ID.nunique()),
    "b1_blank_row_at": blank_row_at,
    "b1_dirty_last_row": r,
    "b1_region_rows": ans.groupby("Region").size().to_dict(),
    "b1_region_revenue": {k: int(v) for k, v in (ans.Quantity * ans.Unit_Price).groupby(ans.Region).sum().items()},
    "b4_csv_logs": logs,
    "b4_rows": {"2024": len(s24), **{f"2025-{m:02d}": int(((sales.Order_Date.dt.year == 2025) & (sales.Order_Date.dt.month == m)).sum()) for m in (1, 2, 3, 4)}},
}
json.dump(meta, open(f"{B}/meta.json", "w"), ensure_ascii=False, indent=1, default=int)
print(json.dumps(meta, ensure_ascii=False, indent=1, default=int))
