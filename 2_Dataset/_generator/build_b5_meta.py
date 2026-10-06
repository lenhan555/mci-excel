"""Buổi 5 (Data Model & DAX): số liệu dùng trong buổi, ghi _generator/build_b5_meta.json.

Chạy được trên Mac lẫn Windows (chỉ cần openpyxl, pandas). Không dựng file nào trong thư mục buổi: file đáp án có
Data Model, relationship và measure thật phải dựng bằng Excel trên Windows (build_b5.py, chưa có).

Nguồn = chính file khởi đầu B5 (Practice/B4_MCI_Mart_Answer.xlsx: bảng Sales trên sheet Sales_PQ, KPI_SUMIFS của B2) và
các file lookup / target B4 mà query của file đó đọc (B4_Lookups.xlsx, B4_Target_2025.xlsx), làm đúng các bước query
(Product bỏ P015 trùng; Target Unpivot). Mỗi measure trong kịch bản được tính lại bằng pandas đúng như DAX sẽ tính:
lọc theo bảng lookup (Store, Product, Calendar), tỷ lệ = tổng ÷ tổng, BLANK khi không có dữ liệu.
Đối chiếu với MCI_Mart_Master.xlsx và KPI_SUMIFS; lệch thì dừng.

    /Users/leopham/anaconda3/bin/python _generator/build_b5_meta.py     # Mac;  Windows: python _generator\\build_b5_meta.py
"""
import json, os, sys
from datetime import date
import pandas as pd
from openpyxl import load_workbook

for s_ in (sys.stdout, sys.stderr):
    if hasattr(s_, "reconfigure"): s_.reconfigure(encoding="utf-8", errors="replace")

ROOT = os.path.normpath(os.environ.get("MCI_ROOT") or os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
GEN = os.path.dirname(os.path.abspath(__file__))
B4P = os.path.join(ROOT, "1_Sessions", "B4_Power_Query", "Practice")
STARTER = os.path.join(B4P, "B4_MCI_Mart_Answer.xlsx")
MASTER = os.path.join(ROOT, "2_Dataset", "MCI_Mart_Master.xlsx")

vn = lambda x: f"{x:,.0f}".replace(",", ".")
pct = lambda x, d=1: "(trống)" if x is None else f"{x * 100:.{d}f}%".replace(".", ",")

# ---------------------------------------------------------------- 6 bảng của Data Model (đúng như query B4)
wb = load_workbook(STARTER, read_only=True, data_only=True)
rows = list(wb["Sales_PQ"].iter_rows(values_only=True))
sales = pd.DataFrame(rows[1:], columns=rows[0])
sales["Order_Date"] = pd.to_datetime(sales.Order_Date)

kpi_rows = list(wb["KPI_SUMIFS"].iter_rows(values_only=True))
wb.close()

lk = pd.ExcelFile(os.path.join(B4P, "B4_Lookups.xlsx"))
product_raw = lk.parse("Product")
product = product_raw.drop_duplicates("Product_Code")
store = lk.parse("Store")
customer = lk.parse("Customer")

tw = pd.read_excel(os.path.join(B4P, "B4_Target_2025.xlsx"), sheet_name="Target_2025", header=2)
target = tw.melt(id_vars=["Store_Code", "Store_Name"], var_name="Attribute", value_name="Target")
target["Thang"] = target.Attribute.str.split(" ").str[-1].astype(int)
target["Ngay_Dau_Thang"] = pd.to_datetime({"year": 2025, "month": target.Thang, "day": 1})
target = target[["Store_Code", "Thang", "Target", "Ngay_Dau_Thang"]]

cal = pd.DataFrame({"Date": pd.date_range("2024-01-01", "2025-12-31")})
cal["Nam"], cal["Thang"], cal["Quy"] = cal.Date.dt.year, cal.Date.dt.month, cal.Date.dt.quarter

# ---------------------------------------------------------------- kiểm tra nguồn
master = pd.read_excel(MASTER, sheet_name="Sales")
assert len(sales) == len(master) == 30247, (len(sales), len(master))
assert int(sales.Revenue.sum()) == int(master.Revenue.sum()) == 4436849500
assert (sales.Revenue == sales.Quantity * sales.Unit_Price).all()
assert len(product) == 40 and len(product_raw) == 41 and len(store) == 12 and len(customer) == 1500
assert len(target) == 144 and len(cal) == 731

# Khóa của lookup duy nhất; mọi khóa ngoại tìm thấy bên lookup (relationship 1-* hợp lệ, không dòng "trống")
keys = {
    "Product": product.Product_Code.is_unique, "Store": store.Store_Code.is_unique,
    "Customer": customer.Customer_ID.is_unique, "Calendar": cal.Date.is_unique,
}
assert all(keys.values()), keys
orphans = {
    "Sales[Product_Code]": int((~sales.Product_Code.isin(product.Product_Code)).sum()),
    "Sales[Store_Code]": int((~sales.Store_Code.isin(store.Store_Code)).sum()),
    "Sales[Customer_ID]": int((~sales.Customer_ID.isin(customer.Customer_ID)).sum()),
    "Sales[Order_Date]": int((~sales.Order_Date.isin(cal.Date)).sum()),
    "Target[Store_Code]": int((~target.Store_Code.isin(store.Store_Code)).sum()),
    "Target[Ngay_Dau_Thang]": int((~target.Ngay_Dau_Thang.isin(cal.Date)).sum()),
}
assert not any(orphans.values()), orphans
p015_bad = int((product_raw.Product_Code == "P015").sum())

# ---------------------------------------------------------------- bảng "đã nối" để tính như measure
s = sales.merge(product[["Product_Code", "Category", "Unit_Cost"]], on="Product_Code", how="left") \
         .merge(store[["Store_Code", "Region", "Open_Date"]], on="Store_Code", how="left") \
         .merge(customer[["Customer_ID", "Segment"]], on="Customer_ID", how="left")
s["COGS"] = s.Quantity * s.Unit_Cost                      # calculated column: Sales[Quantity] * RELATED(Product[Unit_Cost])
s["Y"], s["M"], s["Q"] = s.Order_Date.dt.year, s.Order_Date.dt.month, s.Order_Date.dt.quarter
t = target.merge(store[["Store_Code", "Region"]], on="Store_Code")


def kpi(d):
    """Các measure cơ bản trên một tập dòng Sales đã lọc (None = BLANK)."""
    if len(d) == 0:
        return {"rev": None, "orders": None, "aov": None, "cogs": None, "gm": None, "cust": None, "rows": 0}
    rev, orders, cogs = int(d.Revenue.sum()), int(d.Order_ID.nunique()), int(d.COGS.sum())
    return {"rev": rev, "orders": orders, "aov": rev / orders, "cogs": cogs, "gm": (rev - cogs) / rev,
            "cust": int(d.Customer_ID.nunique()), "rows": len(d)}


def div(a, b):
    return None if a is None or not b else a / b


q1 = s[(s.Y == 2025) & (s.Q == 1)]
REG = ["Bắc", "Trung", "Nam"]
CATS = list(product.Category.drop_duplicates())

# ---------------------------------------------------------------- Phần 1–2: mô hình, grain, Target
merged_cells = len(sales) * ((product.shape[1] - 1) + (store.shape[1] - 1) + (customer.shape[1] - 1))
lookup_cells = product.shape[0] * (product.shape[1] - 1) + store.shape[0] * (store.shape[1] - 1) + customer.shape[0] * (customer.shape[1] - 1)
target_total = int(target.Target.sum())
target_by_region = {r: int(t[t.Region == r].Target.sum()) for r in REG}
target_by_month = {int(m): int(v) for m, v in target.groupby("Thang").Target.sum().items()}

# ---------------------------------------------------------------- Phần 3: calculated column vs measure
q1k = kpi(q1)
q1_row_gm = (q1.Revenue - q1.COGS) / q1.Revenue
gm_wrong_avg = float(q1_row_gm.mean())                     # AVERAGE của cột GM từng dòng (sai)
gm_wrong_sum = float(q1_row_gm.sum())                      # SUM của cột GM từng dòng (vô nghĩa)

# ---------------------------------------------------------------- Phần 4: filter context, đối soát KPI_SUMIFS
reg_month = {r: {m: kpi(q1[(q1.Region == r) & (q1.M == m)]) for m in (1, 2, 3)} for r in REG}
reg_q1 = {r: kpi(q1[q1.Region == r]) for r in REG}
month_q1 = {m: kpi(q1[q1.M == m]) for m in (1, 2, 3)}


def kpi_block(title):
    """Đọc một khối Miền × Tháng của sheet KPI_SUMIFS (file khởi đầu) theo tiêu đề khối."""
    i = next(i for i, r in enumerate(kpi_rows) if r[0] and str(r[0]).startswith(title))
    out = {}
    for r in kpi_rows[i + 1:i + 8]:
        if r[0] in REG + ["Tổng"]:
            out[r[0]] = list(r[1:5])
    return out


sumifs = {"rev": kpi_block("Doanh thu"), "orders": kpi_block("Số đơn"), "aov": kpi_block("Giá trị đơn"), "gm": kpi_block("Gross Margin")}
recon = []
for k in ("rev", "orders", "aov", "gm"):
    for r in REG + ["Tổng"]:
        for j, m in enumerate((1, 2, 3, "Q1")):
            mine = (month_q1[m] if r == "Tổng" else reg_month[r][m]) if m != "Q1" else (q1k if r == "Tổng" else reg_q1[r])
            want = sumifs[k][r][j]
            ok = abs(mine[k] - want) < (1e-9 if k == "gm" else 0.5 if k != "aov" else 1e-6)
            recon.append(ok)
            assert ok, (k, r, m, mine[k], want)

# ---------------------------------------------------------------- Phần 5: CALCULATE, ALL, Target
bac_q1 = reg_q1["Bắc"]["rev"]
share_q1 = {r: reg_q1[r]["rev"] / q1k["rev"] for r in REG}
d25 = s[s.Y == 2025]
actual_2025 = {m: kpi(d25[d25.M == m])["rev"] for m in range(1, 13)}
vs_target_month = {m: div(actual_2025[m], target_by_month[m]) for m in range(1, 13)}
variance_month = {m: (None if actual_2025[m] is None else actual_2025[m] - target_by_month[m]) for m in range(1, 13)}
variance_naive = {m: (actual_2025[m] or 0) - target_by_month[m] for m in range(1, 13)}   # không có IF(ISBLANK…)
t14 = target[target.Thang <= 4]
ytd_target_t4 = int(t14.Target.sum())
store_month = {}
for sc in store.Store_Code:
    store_month[sc] = {}
    for m in range(1, 5):
        a = kpi(d25[(d25.Store_Code == sc) & (d25.M == m)])["rev"]
        tg = int(target[(target.Store_Code == sc) & (target.Thang == m)].Target.iloc[0])
        store_month[sc][m] = {"rev": a, "target": tg, "pct": div(a, tg)}
rag = {m: {"duoi_80": sorted(sc for sc in store_month if store_month[sc][m]["pct"] < 0.8),
           "80_90": sorted(sc for sc in store_month if 0.8 <= store_month[sc][m]["pct"] < 0.9)} for m in range(1, 5)}
region_month_t = {r: {m: {"rev": kpi(d25[(d25.Region == r) & (d25.M == m)])["rev"],
                          "target": int(t[(t.Region == r) & (t.Thang == m)].Target.sum())} for m in range(1, 5)} for r in REG}
for r in REG:
    for m in range(1, 5):
        v = region_month_t[r][m]; v["pct"] = v["rev"] / v["target"]
# Revenue vs Target % khi đặt Category vào hàng (sai grain): doanh thu ngành ÷ target CẢ chuỗi
cat_vs_target_q1 = {c: q1[q1.Category == c].Revenue.sum() / int(target[target.Thang <= 3].Target.sum()) for c in CATS}
q1_target = int(target[target.Thang <= 3].Target.sum())

# ---------------------------------------------------------------- Phần 6: time intelligence
d24 = s[s.Y == 2024]
rev_2024 = {m: int(d24[d24.M == m].Revenue.sum()) for m in range(1, 13)}
yoy_month = {m: div(actual_2025[m] - rev_2024[m], rev_2024[m]) if actual_2025[m] is not None else None for m in range(1, 13)}
yoy_naive_may = div((actual_2025[5] or 0) - rev_2024[5], rev_2024[5])        # không có IF(ISBLANK…): -100%
ytd_2025 = {m: sum(actual_2025[k] for k in range(1, m + 1)) for m in range(1, 5)}
ytd_2024 = {m: sum(rev_2024[k] for k in range(1, m + 1)) for m in range(1, 13)}
ytd_yoy_t4 = (ytd_2025[4] - ytd_2024[4]) / ytd_2024[4]
st12_2024_t1_t4 = int(d24[(d24.Store_Code == "ST12") & (d24.M <= 4)].Revenue.sum())
st12_2025_t1_t4 = int(d25[(d25.Store_Code == "ST12") & (d25.M <= 4)].Revenue.sum())
lfl = s[s.Open_Date <= pd.Timestamp("2024-01-01")]                            # cửa hàng có đủ hai kỳ
lfl_stores = int(lfl.Store_Code.nunique())
lfl_yoy_month = {m: (lfl[(lfl.Y == 2025) & (lfl.M == m)].Revenue.sum() - lfl[(lfl.Y == 2024) & (lfl.M == m)].Revenue.sum())
                 / lfl[(lfl.Y == 2024) & (lfl.M == m)].Revenue.sum() for m in range(1, 5)}
lfl_ytd_2025 = int(lfl[(lfl.Y == 2025) & (lfl.M <= 4)].Revenue.sum())
lfl_ytd_2024 = int(lfl[(lfl.Y == 2024) & (lfl.M <= 4)].Revenue.sum())
store_yoy_t1_t4 = {}
for sc in store.Store_Code:
    a = int(d25[(d25.Store_Code == sc) & (d25.M <= 4)].Revenue.sum())
    b = int(d24[(d24.Store_Code == sc) & (d24.M <= 4)].Revenue.sum())
    store_yoy_t1_t4[sc] = {"rev_2025": a, "rev_2024": b, "yoy": div(a - b, b) if b else None}

# ---------------------------------------------------------------- Phần 7: KPI dictionary, câu chuyện #5
gm_year = {y: kpi(s[s.Y == y])["gm"] for y in (2024, 2025)}
gm_t14 = {y: kpi(s[(s.Y == y) & (s.M <= 4)])["gm"] for y in (2024, 2025)}
price_gap = {}  # giá bán trung bình 2025 so với 2024 trên cùng sản phẩm (T1–T4)
pp = s[s.M <= 4].groupby(["Product_Code", "Y"]).apply(lambda d: d.Revenue.sum() / d.Quantity.sum()).unstack()
price_gap = float((pp[2025] / pp[2024]).mean() - 1)

# ---------------------------------------------------------------- Phần 8: thực hành
kpi_month_2025 = {m: kpi(d25[d25.M == m]) for m in range(1, 5)}
for m in range(1, 5):
    kpi_month_2025[m]["yoy"] = yoy_month[m]
    kpi_month_2025[m]["vs_target"] = vs_target_month[m]
aov_cat_q1 = {c: kpi(q1[q1.Category == c]) for c in CATS}
cat_check = "Bánh kẹo"
avg_line_rev_cat = float(q1[q1.Category == cat_check].Revenue.mean())         # AVERAGE(Revenue) — không phải AOV
cust_sum_regions_q1 = sum(reg_q1[r]["cust"] for r in REG)                    # cộng 3 miền > tổng: DISTINCTCOUNT không cộng dồn
t1_t2_pct = (actual_2025[1] + actual_2025[2]) / (target_by_month[1] + target_by_month[2])
q1_vs_target = q1k["rev"] / q1_target
b4hw = json.load(open(os.path.join(GEN, "build_b4_meta.json"), encoding="utf-8"))["hw"]

# ---------------------------------------------------------------- ghi meta
meta = {
    "source": {"starter": os.path.relpath(STARTER, ROOT), "rows": len(sales), "revenue": int(sales.Revenue.sum())},
    "tables": {"Sales": len(sales), "Product": len(product), "Store": len(store), "Customer": len(customer),
               "Target": len(target), "Calendar": len(cal)},
    "categories": CATS, "n_categories": len(CATS), "p015_rows_raw": p015_bad,
    "orphans": orphans,
    "merged_cells": merged_cells, "lookup_cells": lookup_cells,
    "target_total": target_total, "target_by_region": target_by_region, "target_by_month": target_by_month,
    "q1": q1k, "q1_target": q1_target,
    "gm_wrong_avg_q1": gm_wrong_avg, "gm_wrong_sum_q1": gm_wrong_sum,
    "reg_month_q1": {r: {str(m): reg_month[r][m] for m in reg_month[r]} for r in REG},
    "reg_q1": reg_q1, "month_q1": {str(m): month_q1[m] for m in month_q1},
    "recon_cells": len(recon), "recon_ok": sum(recon),
    "bac_t3_2025": reg_month["Bắc"][3]["rev"], "bac_q1": bac_q1, "share_q1": share_q1,
    "actual_2025": actual_2025, "vs_target_month": vs_target_month, "variance_month": variance_month,
    "variance_naive": variance_naive, "ytd_target_t4": ytd_target_t4,
    "store_month": store_month, "rag": rag, "region_month_target": region_month_t,
    "cat_vs_target_q1": cat_vs_target_q1,
    "rev_2024": rev_2024, "yoy_month": yoy_month, "yoy_naive_may": yoy_naive_may,
    "ytd_2025": ytd_2025, "ytd_2024": ytd_2024, "ytd_yoy_t4": ytd_yoy_t4,
    "st12": {"open": "2024-06-01", "rev_2024_t1_t4": st12_2024_t1_t4, "rev_2025_t1_t4": st12_2025_t1_t4},
    "lfl_stores": lfl_stores, "lfl_yoy_month": lfl_yoy_month, "lfl_ytd_2025": lfl_ytd_2025, "lfl_ytd_2024": lfl_ytd_2024,
    "lfl_ytd_yoy": (lfl_ytd_2025 - lfl_ytd_2024) / lfl_ytd_2024,
    "store_yoy_t1_t4": store_yoy_t1_t4,
    "gm_year": gm_year, "gm_t1_t4": gm_t14, "price_gap_2025_vs_2024": price_gap,
    "kpi_month_2025": kpi_month_2025, "aov_cat_q1": aov_cat_q1, "cat_check": cat_check,
    "avg_line_rev_cat_check": avg_line_rev_cat,
    "cust_sum_regions_q1": cust_sum_regions_q1, "t1_t2_pct": t1_t2_pct, "q1_vs_target": q1_vs_target,
    "b4_hw": {"st07_part1_pct": b4hw["st07_part1_pct"], "st07_pct": b4hw["st07_pct"], "pct_all": b4hw["pct_all"]},
}
# ---------------------------------------------------------------- Bài tập về nhà: "báo cáo cũ" của phòng Kinh doanh
# Bốn lỗi cài sẵn, mỗi lỗi là một nguyên nhân khiến báo cáo cũ lệch số của Tài chính (câu hỏi mở buổi):
#  (1) Target chưa nối Calendar + (2) báo cáo lấy tháng từ Sales[Thang] → target tháng = target cả năm 2025;
#  (3) GM% = AVERAGE cột GM từng dòng; (4) "Số đơn" = đếm dòng (COUNTROWS) → AOV = doanh thu ÷ số dòng.
hw_months = {}
for mth in range(1, 5):
    dm = d25[d25.M == mth]
    k = kpi(dm)
    hw_months[mth] = {
        "tc_rev": k["rev"], "tc_orders": k["orders"], "tc_aov": k["aov"], "tc_gm": k["gm"],
        "tc_target": target_by_month[mth], "tc_vs_target": vs_target_month[mth],
        "cu_target": target_total, "cu_vs_target": k["rev"] / target_total,
        "cu_gm_avg": float(((dm.Revenue - dm.COGS) / dm.Revenue).mean()),
        "cu_orders_rows": k["rows"], "cu_aov_rows": k["rev"] / k["rows"],
        "lfl_yoy": lfl_yoy_month[mth],
    }
meta["hw"] = {
    "months": hw_months,
    "lfl_ytd_yoy": meta["lfl_ytd_yoy"], "ytd_yoy_t4": ytd_yoy_t4, "lfl_stores": lfl_stores,
    "below80": {m: rag[m]["duoi_80"] for m in range(1, 5)},
    "seg_cust_q1": {sg: int(q1[q1.Segment == sg].Customer_ID.nunique()) for sg in ("Thường", "Thân thiết", "VIP")},
    "seg_rev_q1": {sg: int(q1[q1.Segment == sg].Revenue.sum()) for sg in ("Thường", "Thân thiết", "VIP")},
}
st07 = store_month["ST07"][3]
meta["info"] = {  # khóa phẳng cho infographic (build_infographic.py điền {{info.…}})
    "sales_rows": len(sales), "product": len(product), "store": len(store), "customer": len(customer),
    "target_rows": len(target), "calendar_rows": len(cal), "q1_rev": q1k["rev"], "q1_orders": q1k["orders"],
    "q1_aov": round(q1k["aov"]), "q1_gm": q1k["gm"], "q1_cust": q1k["cust"], "bac_t3": reg_month["Bắc"][3]["rev"],
    "t1_pct": vs_target_month[1], "t2_pct": vs_target_month[2], "t3_pct": vs_target_month[3], "t4_pct": vs_target_month[4],
    "st07_t3_pct": st07["pct"], "st07_t3_rev": st07["rev"], "st07_t3_target": st07["target"],
    "ytd_t4": ytd_2025[4], "ytd_yoy_t4": ytd_yoy_t4, "lfl_ytd_yoy": meta["lfl_ytd_yoy"],
    "gm_2024": gm_year[2024], "gm_2025": gm_year[2025], "target_total": target_total,
}
with open(os.path.join(GEN, "build_b5_meta.json"), "w", encoding="utf-8") as f:
    json.dump(meta, f, ensure_ascii=False, indent=1, default=lambda o: int(o) if hasattr(o, "item") else str(o))

# ---------------------------------------------------------------- in tóm tắt
print("Bảng:", meta["tables"], "| khóa ngoại không khớp:", sum(orphans.values()))
print(f"Merge 3 lookup vào Sales: {vn(merged_cells)} ô lặp; giữ lookup riêng: {vn(lookup_cells)} ô")
print(f"Target cả năm {vn(target_total)}; theo miền {', '.join(f'{r} {vn(v)}' for r, v in target_by_region.items())}")
print(f"Q1/2025: doanh thu {vn(q1k['rev'])}, đơn {vn(q1k['orders'])}, AOV {vn(q1k['aov'])}, GM {pct(q1k['gm'], 2)}, KH {vn(q1k['cust'])}, COGS {vn(q1k['cogs'])}")
print(f"  GM từng dòng: AVERAGE {pct(gm_wrong_avg, 2)} · SUM {vn(gm_wrong_sum * 100)}%")
print(f"  Đối soát KPI_SUMIFS: {sum(recon)}/{len(recon)} ô khớp (doanh thu, số đơn, AOV, GM% × 4 dòng × 4 cột)")
print(f"  Bắc × T3 = {vn(reg_month['Bắc'][3]['rev'])}; % tổng: {', '.join(f'{r} {pct(v)}' for r, v in share_q1.items())}")
for m in range(1, 5):
    k = kpi_month_2025[m]
    print(f"T{m}/2025: DT {vn(k['rev'])} · đơn {k['orders']} · AOV {vn(k['aov'])} · GM {pct(k['gm'])} · KH {k['cust']} · "
          f"Target {vn(target_by_month[m])} = {pct(vs_target_month[m])} · chênh {vn(variance_month[m])} · LY {vn(rev_2024[m])} · YoY {pct(yoy_month[m])} · YTD {vn(ytd_2025[m])}")
print(f"T5/2025 không chặn BLANK: Variance {vn(variance_naive[5])}, YoY {pct(yoy_naive_may)}")
print(f"YTD T4: 2025 {vn(ytd_2025[4])} vs 2024 {vn(ytd_2024[4])} = {pct(ytd_yoy_t4)}; like-for-like ({lfl_stores} cửa hàng) {pct(meta['lfl_ytd_yoy'])}")
print(f"ST12 T1–T4: 2024 {vn(st12_2024_t1_t4)} · 2025 {vn(st12_2025_t1_t4)}")
print(f"ST07 T3: {vn(st07['rev'])} / {vn(st07['target'])} = {pct(st07['pct'])}; T4 = {pct(store_month['ST07'][4]['pct'])}")
print("RAG:", {m: v for m, v in rag.items()})
print("Miền × tháng % target:", {r: [pct(region_month_t[r][m]['pct']) for m in range(1, 5)] for r in REG})
print(f"GM năm: 2024 {pct(gm_year[2024], 2)} · 2025 {pct(gm_year[2025], 2)}; T1–T4: {pct(gm_t14[2024], 2)} → {pct(gm_t14[2025], 2)}; giá 2025/2024 +{pct(price_gap)}")
print(f"Category × Revenue vs Target % (Q1, sai grain): {', '.join(f'{c} {pct(v)}' for c, v in cat_vs_target_q1.items())} (target Q1 {vn(q1_target)})")
c = aov_cat_q1[cat_check]
print(f"AOV {cat_check} Q1: {vn(c['rev'])} / {c['orders']} đơn = {vn(c['aov'])} (AVERAGE dòng = {vn(avg_line_rev_cat)})")
print("YoY cửa hàng T1–T4:", {k: pct(v['yoy']) for k, v in store_yoy_t1_t4.items()})
print(f"Active Customers Q1: 3 miền cộng lại {cust_sum_regions_q1} ≠ tổng {q1k['cust']}; T1+T2 đạt {pct(t1_t2_pct)}; Q1 đạt {pct(q1_vs_target)}")
print("BTVN — báo cáo cũ so với Tài chính:")
for m_, h in meta["hw"]["months"].items():
    print(f"  T{m_}: target {vn(h['cu_target'])} → {pct(h['cu_vs_target'])} (đúng {vn(h['tc_target'])} → {pct(h['tc_vs_target'])}) · "
          f"GM {pct(h['cu_gm_avg'], 2)} (đúng {pct(h['tc_gm'], 2)}) · đơn {vn(h['cu_orders_rows'])} → AOV {vn(h['cu_aov_rows'])} "
          f"(đúng {vn(h['tc_orders'])} → {vn(h['tc_aov'])}) · YoY LFL {pct(h['lfl_yoy'])}")
print("  Phân khúc Q1 (KH, doanh thu):", {k: (meta['hw']['seg_cust_q1'][k], vn(v)) for k, v in meta['hw']['seg_rev_q1'].items()})
print("Meta:", os.path.join(GEN, "build_b5_meta.json"))
