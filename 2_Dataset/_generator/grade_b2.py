#!/usr/bin/env python3
"""Chấm bài tập về nhà B2 (Hàm và KPI). Chỉ ĐỌC file học viên.
Dùng:
  python3 grade_b2.py <file_nop.xlsx> <B2_BTVN_..._Answer.xlsx>                      -> JSON một bài
  python3 grade_b2.py --folder <thu_muc_Bai_nop> <..._Answer.xlsx> [--cham-tay f.json] -> JSON cả lớp
f.json (chấm tay việc 5 và 6): {"ten_file.xlsx": {"v5": "Đạt|Một phần|Chưa đạt", "insight": 0-3}}
Không quy điểm khi chưa có thang và trọng số do người dùng đưa.
"""
import sys, json, re, zipfile, collections, glob, os, difflib
from decimal import Decimal, ROUND_HALF_UP
import openpyxl
from openpyxl.formula.translate import Translator

def val(c):
    v = c.value
    return v.text if hasattr(v, "text") else v

def is_formula(v):
    return isinstance(v, str) and v.startswith("=")

def load(path):
    return openpyxl.load_workbook(path), openpyxl.load_workbook(path, data_only=True)

def find_sales(wb):
    for ws in wb:
        for row in range(1, min(10, ws.max_row) + 1):
            h = [c.value for c in ws[row]]
            if "Order_ID" in h and "Line_No" in h:
                ws._grading_header_row = row
                alias = {"UnitCost": "Unit_Cost", "Gross Margin": "Gross_Margin"}
                return ws, {alias.get(n, n): i for i, n in enumerate(h) if n}
    return None, {}

def ref_rows(ws_v, idx):
    d = {}
    for r in ws_v.iter_rows(min_row=getattr(ws_v, '_grading_header_row', 1) + 1, values_only=True):
        if r[idx["Order_ID"]] is None:
            continue
        d[(r[idx["Order_ID"]], r[idx["Line_No"]])] = r
    return d

def find_kpi(wbv, wbf):
    """Tìm hai bảng KPI theo nhãn, không theo vị trí."""
    out = {}
    for name in wbv.sheetnames:
        if name.strip().lower() not in ("kpi", "kpi_sumifs", "sheet1"):
            continue
        wv, wf = wbv[name], wbf[name]
        out["sheet"] = name
        cells = {c.coordinate: c for row in wv.iter_rows() for c in row if c.value is not None}
        # KPI1: hàng có nhãn Bắc/Trung/Nam, cột nhãn ngành, cột Tổng
        regrows = {}
        for c in cells.values():
            if isinstance(c.value, str) and c.value.strip() in ("Bắc", "Trung", "Nam") and c.column == 1:
                regrows.setdefault(c.value.strip(), []).append(c.row)
        out["regrows"] = regrows
        hdr = {}
        for c in cells.values():
            if isinstance(c.value, str):
                hdr.setdefault(c.value.strip(), []).append(c)
        out["hdr"] = {k: [x.coordinate for x in v] for k, v in hdr.items()}
        out["ws_v"], out["ws_f"] = wv, wf
        return out
    return out

def grade(nop, ans):
    """Kiểm toàn cột, nhãn bảng linh hoạt, tách cache/công thức; không quy điểm."""
    import xml.etree.ElementTree as ET
    res = {"file": nop, "kiem_tra": {}, "viec": {}, "ghi_chu": []}
    try:
        if os.path.splitext(nop)[1].lower() != '.xlsx':
            raise ValueError('không phải .xlsx')
        nf, nv = load(nop); af, av = load(ans)
        if 'Kiem_tra' not in av: raise ValueError('Answer thiếu Kiem_tra')
    except Exception as e:
        res['loi_mo_file'] = str(e); return res
    sf, idx = find_sales(nf); sv, _ = find_sales(nv); sa, ai = find_sales(av)
    if sf is None:
        res['viec'] = {str(i): ['Chưa đạt', 'Không tìm thấy Order_ID + Line_No'] for i in range(1, 7)}
        return res
    nr, ar = ref_rows(sv, idx), ref_rows(sa, ai)
    hr = getattr(sf, '_grading_header_row', 1)
    rownums = {(sv.cell(r, idx['Order_ID']+1).value, sv.cell(r,idx['Line_No']+1).value):r for r in range(hr+1,sv.max_row+1) if sv.cell(r,idx['Order_ID']+1).value is not None}
    total_rows = sum(sv.cell(r,idx['Order_ID']+1).value is not None for r in range(hr+1,sv.max_row+1))
    res['bang']={'sheet': sf.title, 'header_row': hr, 'tables': [(t.name,t.ref) for t in sf.tables.values()], 'so_dong':total_rows, 'duplicate_keys':total_rows-len(rownums), 'missing_keys':len(ar.keys()-nr.keys()), 'extra_keys':len(nr.keys()-ar.keys())}
    co={}; errors={}; diffs={}; hands={}; wraps={}; funcs={}; samples={}; label_diffs=0
    for col in ('Region','Category','Unit_Cost','COGS','Gross_Margin','Phan_khuc','Revenue'):
        er=[];df=[];hd=[];wr=[];fc=collections.Counter()
        for key,row in nr.items():
            c=sf.cell(rownums[key],idx[col]+1) if col in idx else None
            f=val(c) if c else None;v=row[idx[col]] if col in idx else None
            if v is None or isinstance(v,str) and v.startswith('#'):er.append(c.coordinate if c else col)
            if not is_formula(f):hd.append(c.coordinate if c else col)
            if is_formula(f):
                fc.update(re.findall(r'XLOOKUP|VLOOKUP|INDEX|MATCH|IFS|IF|SUMIFS',f.upper()))
                if re.search(r'IFERROR|IFNA',f,re.I):wr.append(c.coordinate)
            a=ar.get(key)
            if a and col in ai:
                ex=a[ai[col]]
                same=str(v).strip().casefold()==str(ex).strip().casefold() if col=='Phan_khuc' else v==ex
                if col=='Phan_khuc' and not same and str(v).strip().split(' ')[0].casefold()==str(ex).strip().split(' ')[0].casefold():
                    label_diffs+=1; same=True
                if not same:df.append({'o':c.coordinate if c else col,'key':key,'hoc_vien':v,'dap_an':ex,'formula':f})
        errors[col]=len(er);diffs[col]=len(df);hands[col]={'count':len(hd),'sample':hd[:8]};wraps[col]={'count':len(wr),'sample':wr[:8]};funcs[col]=dict(fc);samples[col]=df[:8]
    res['cot']={'loi_trong':errors,'so_dong_lech':diffs,'go_tay':hands,'boc_loi':wraps,'ham':funcs,'mau_lech':samples}
    trailing=[{'o':sf.cell(rownums[k],idx['Product_Code']+1).coordinate,'value':r[idx['Product_Code']]} for k,r in nr.items() if isinstance(r[idx['Product_Code']],str) and r[idx['Product_Code']]!=r[idx['Product_Code']].strip()]
    fixes=[]
    for r in (53,230,1311,1520):
        key=(sa.cell(r,ai['Order_ID']+1).value,sa.cell(r,ai['Line_No']+1).value);n=nr.get(key)
        fixes.append({'answer_row':r,'key':key,'o':sf.cell(rownums[key],idx['Product_Code']+1).coordinate if key in rownums else None,'got':n[idx['Product_Code']] if n else None,'expected':sa.cell(r,ai['Product_Code']+1).value})
    res['ma_dau_cach']={'con':trailing,'loi_cai_san':fixes}
    bad1=bool(trailing) or any(errors[k] or diffs[k] or wraps[k]['count'] for k in ('Region','Category','Unit_Cost')) or res['bang']['missing_keys'] or res['bang']['extra_keys'] or res['bang']['duplicate_keys']
    partial1=any(hands[k]['count'] for k in ('Region','Category','Unit_Cost')) or sf.title!='Sales' or any(c.value=='UnitCost' for c in sf[hr])
    res['viec']['1']=['Chưa đạt' if bad1 else 'Một phần' if partial1 else 'Đạt',f'{sf.title}: {total_rows} dòng; còn dấu cách={len(trailing)}, lệch tra cứu={sum(diffs[k] for k in ("Region","Category","Unit_Cost"))}, lỗi/trống={sum(errors[k] for k in ("Region","Category","Unit_Cost"))}']
    bad2=any(errors[k] or diffs[k] for k in ('COGS','Gross_Margin','Phan_khuc'))
    partial2=any(hands[k]['count'] for k in ('COGS','Gross_Margin','Phan_khuc')) or any(c.value=='Gross Margin' for c in sf[hr]) or label_diffs>0
    res['phan_khuc_khac_nhan_dung_nguong']=label_diffs
    res['viec']['2']=['Chưa đạt' if bad2 else 'Một phần' if partial2 else 'Đạt',f'{sf.title}: lệch COGS={diffs["COGS"]}, Gross_Margin={diffs["Gross_Margin"]}, Phan_khuc={diffs["Phan_khuc"]}; gõ tay={ {k:hands[k]["count"] for k in ("COGS","Gross_Margin","Phan_khuc")} }; nhãn khác nhưng đúng ngưỡng={label_diffs}']
    if not (nr.keys() & ar.keys()):res['viec']['2']=['Không kiểm được','Cần GV xem: dữ liệu khác kỳ; không có khóa dòng chung để đối chiếu Phan_khuc với bài T4']
    def suma(col):return sum(r[idx[col]] for r in nr.values() if isinstance(r[idx[col]],(int,float))) if col in idx else None
    res['du_lieu_tinh_lai']={c:suma(c) for c in ('Revenue','COGS','Gross_Margin')}
    res['du_lieu_tinh_lai']['so_don']=sum(r[idx['Line_No']]==1 for r in nr.values())
    res['du_lieu_tinh_lai']['COGS_tu_dau_vao']=sum(r[idx['Quantity']]*r[idx['Unit_Cost']] for r in nr.values() if 'Unit_Cost' in idx and isinstance(r[idx['Quantity']],(int,float)) and isinstance(r[idx['Unit_Cost']],(int,float)))
    cats=('Đồ uống','Sữa','Thực phẩm khô','Bánh kẹo','Chăm sóc cá nhân','Gia dụng');regs=('Bắc','Trung','Nam')
    def cs(s):return [c for c in s._cells.values() if c.value is not None]
    k1={}; ks=None
    for s in nv:
        if s.title==sf.title:continue
        cc=cs(s)
        for rh in sorted({c.row for c in cc if isinstance(c.value,str) and c.value in cats}):
            catcol={c.value:c.column for c in cc if c.row==rh and c.value in cats}
            if len(catcol)!=6:continue
            regcells=[c for c in cc if c.value in regs and rh<c.row<=rh+5]
            if not regcells:continue
            lc=collections.Counter(c.column for c in regcells).most_common(1)[0][0]
            rs={c.value:c.row for c in regcells if c.column==lc}
            if len(rs)!=3:continue
            tc=next((c.column for c in cc if c.row==rh and str(c.value).strip().lower() in ('tổng','total')),None)
            tr=next((c.row for c in cc if c.column==lc and max(rs.values())<c.row<=max(rs.values())+3 and str(c.value).strip().lower() in ('tổng','total')),None)
            fs=[nf[s.title].cell(rs[r],catcol[cat]) for r in regs for cat in cats];base=min(fs,key=lambda c:(c.row,c.column));bf=val(base)
            same=is_formula(bf) and all(is_formula(val(c)) and Translator(bf,origin=base.coordinate).translate_formula(c.coordinate).replace(' ','')==val(c).replace(' ','') for c in fs)
            mixed=bool(re.search(r'\$[A-Z]+\d+',str(bf))) and bool(re.search(r'(?<!\$)[A-Z]+\$\d+',str(bf)))
            mismatch=[]
            for reg in regs:
                for cat in cats:
                    c=s.cell(rs[reg],catcol[cat]);ex=sum(r[ai['Revenue']] for r in ar.values() if r[ai['Region']]==reg and r[ai['Category']]==cat)
                    if c.value!=ex:mismatch.append({'o':c.coordinate,'value':c.value,'expected':ex})
            tots=[nf[s.title].cell(tr,c) for c in catcol.values()] if tr else []
            if tc:tots += [nf[s.title].cell(r,tc) for r in rs.values()]
            if tr and tc:tots += [nf[s.title].cell(tr,tc)]
            if tc:
                for reg in regs:
                    c=s.cell(rs[reg],tc);ex=sum(r[ai['Revenue']] for r in ar.values() if r[ai['Region']]==reg)
                    if c.value!=ex:mismatch.append({'o':c.coordinate,'value':c.value,'expected':ex})
            if tr:
                for cat in cats:
                    c=s.cell(tr,catcol[cat]);ex=sum(r[ai['Revenue']] for r in ar.values() if r[ai['Category']]==cat)
                    if c.value!=ex:mismatch.append({'o':c.coordinate,'value':c.value,'expected':ex})
            k1={'sheet':s.title,'anchor':base.coordinate,'formula':bf,'pull_same':same,'mixed_refs':mixed,'sumifs_all':all('SUMIFS' in str(val(c)).upper() for c in fs),'total_cell':s.cell(tr,tc).coordinate if tr and tc else None,'total':s.cell(tr,tc).value if tr and tc else None,'totals_formula':bool(tc and tr) and all(is_formula(val(c)) for c in tots),'mismatch':mismatch};ks=s;break
        if ks:break
    res['kpi1']=k1
    if not k1:res['viec']['3']=['Chưa đạt','Không có bảng 3 miền x 6 ngành hàng']
    else:
        b=bool(k1['mismatch']) or not k1['total_cell'];p=ks.title!='KPI' or not all(k1[k] for k in ('pull_same','mixed_refs','sumifs_all','totals_formula'))
        res['viec']['3']=['Chưa đạt' if b else 'Một phần' if p else 'Đạt',f'{ks.title}!{k1["anchor"]}: kéo đồng nhất={k1["pull_same"]}; $ hỗn hợp={k1["mixed_refs"]}; ô lệch={len(k1["mismatch"])}; tổng {k1["total_cell"]}={k1["total"]}']
    metrics={}; ms=ks or next((s for s in nv if 'kpi' in s.title.lower() and 'card' not in s.title.lower()),None)
    if ms:
        cc=cs(ms);mf=nf[ms.title]
        for name,pat in [('so_don',r'^(số đơn|so don)'),('aov',r'AOV'),('gm_pct',r'(gross.?margin.*%|%.*gross.?margin)')]:
            heads=[c for c in cc if isinstance(c.value,str) and re.search(pat,c.value,re.I)]
            if not heads:continue
            if name=='gm_pct' and any(c.column>1 for c in heads):heads=[c for c in heads if c.column>1]
            hd=min(heads,key=lambda c:c.row)
            regc=[c for c in cc if c.value in regs and hd.row<c.row<=hd.row+6]
            labelcol=min((c.column for c in regc),default=1)
            lr=[c for c in regc if c.column==labelcol]
            if len(lr)!=3:continue
            tot=next((c for c in cc if c.column==labelcol and max(c.row for c in lr)<c.row<=max(c.row for c in lr)+3 and str(c.value).strip().lower() in ('tổng','total')),None)
            col=hd.column if hd.column>labelcol else labelcol+1
            metrics[name]={'values':{c.value:ms.cell(c.row,col).value for c in lr},'formulas':{c.value:val(mf.cell(c.row,col)) for c in lr},'total':ms.cell(tot.row,col).value if tot else None,'total_cell':ms.cell(tot.row,col).coordinate if tot else None,'total_formula':val(mf.cell(tot.row,col)) if tot else None}
    res['kpi2']={'sheet':ms.title if ms else None,**metrics};probs=[];sev=False
    for k,m in metrics.items():
        m['mismatch']=[]
        for reg,v in m['values'].items():
            rr=[r for r in ar.values() if r[ai['Region']]==reg]
            rev=sum(r[ai['Revenue']] for r in rr);orders=sum(r[ai['Line_No']]==1 for r in rr);gm=sum(r[ai['Gross_Margin']] for r in rr)
            ex=orders if k=='so_don' else rev/orders if k=='aov' else gm/rev
            if not isinstance(v,(int,float)) or round(v,6)!=round(ex,6):m['mismatch'].append({'region':reg,'value':v,'expected':ex});sev=True
        if m['mismatch']:probs.append(k+': giá trị miền không khớp')
    for k in ('so_don','aov','gm_pct'):
        if k not in metrics or metrics[k]['total'] is None:probs.append('thiếu '+k+' hoặc dòng Tổng');sev=True
    if 'so_don' in metrics:
        cl=openpyxl.utils.get_column_letter(idx['Line_No']+1)
        for reg,f in metrics['so_don']['formulas'].items():
            ok='COUNTIFS' in str(f).upper() and (bool(re.search(r'Line_No\]\s*,\s*1\b',str(f))) or bool(re.search(r'\$?'+cl+r':\$?'+cl+r'\s*,\s*1\b',str(f))))
            if not ok:probs.append(reg+': chưa dùng COUNTIFS với Line_No=1')
    for k in ('aov','gm_pct'):
        if k in metrics and metrics[k]['total'] is not None:
            f=metrics[k]['total_formula']
            if not is_formula(f) or '/' not in f or 'AVERAGE' in f.upper():probs.append(k+': tổng chưa là tổng chia tổng');sev=True
    res['viec']['4']=['Chưa đạt' if sev else 'Một phần' if probs or ms and ms.title!='KPI' else 'Đạt','; '.join(probs) or 'Đủ COUNTIFS Line_No=1 và tổng chia tổng']
    got=[total_rows,sum(errors[k] for k in ('Region','Category','Unit_Cost')),suma('Revenue'),k1.get('total'),metrics.get('so_don',{}).get('total'),suma('COGS'),metrics.get('gm_pct',{}).get('total'),metrics.get('aov',{}).get('total')]
    for i,r in enumerate(range(5,13)):
        ex=av['Kiem_tra'].cell(r,4).value;g=got[i]
        res['kiem_tra'][av['Kiem_tra'].cell(r,1).value]={'hoc_vien':g,'dap_an':ex,'khop':isinstance(g,(int,float)) and round(g,6)==round(ex,6),'answer_cell':f'Kiem_tra!D{r}'}
    res['card']={};res['insight']=[];res['textbox']=[]
    for s in nv:
        norm=s.title.strip().lower().replace(' ','_')
        if norm=='kpi_card':res['card']={'sheet':s.title,'cells':[[c.coordinate,c.value] for c in cs(s)]}
        if norm=='insight':res['insight'] += [{'sheet':s.title,'o':c.coordinate,'text':str(c.value)} for c in cs(s) if isinstance(c.value,str) and len(c.value)>25]
    # XML drawings and legacy VML textboxes are inspected independently of openpyxl.
    with zipfile.ZipFile(nop) as z:
        for n in z.namelist():
            if n.startswith('xl/drawings/') and (n.endswith('.xml') or n.endswith('.vml')):
                try:
                    e=ET.fromstring(z.read(n));ts=[a.text for a in e.iter() if a.tag.split('}')[-1]=='t' and a.text]
                    if not ts:
                        ts=[' '.join(a.itertext()) for a in e.iter() if a.tag.split('}')[-1]=='textbox']
                    if ts:res['textbox'].append({'xml':n,'text':' '.join(ts)})
                except ET.ParseError:pass
        try:res['core']={e.tag.split('}')[-1]:e.text for e in ET.fromstring(z.read('docProps/core.xml'))}
        except Exception:res['core']={}
    res['viec']['5']=['Không kiểm được','Cần đọc đủ 10 trường và nội dung AOV'] if res['card'] else ['Chưa đạt','Không có KPI_Card']
    res['viec']['6']=['Không kiểm được','Cần đọc insight và đối chiếu số dẫn'] if res['insight'] or res['textbox'] else ['Chưa đạt','Không có sheet Insight chứa câu insight']
    res['sheet_dap_an']=sorted(set(nv.sheetnames)&{'Dap_an_loi','Insight_mau'})
    res['answer_mau']=[[c.coordinate,c.value] for c in cs(av['Insight_mau'])] if 'Insight_mau' in av else []
    return res

def av_get(av, sheet, coord):
    return av[sheet][coord].value

def ap_cham_tay(res, v5, tp):
    """v5: kết quả việc 5 (KPI_Card) do người chấm đọc; tp: số thành phần insight đạt 0-3 (việc 6)."""
    if "loi_mo_file" in res:
        return res
    res["viec"]["5"][0] = v5
    res["insight_thanh_phan_dat"] = tp
    res['viec']['6'][0] = 'Đạt' if tp == 3 else 'Một phần' if res.get('insight') or res.get('textbox') else 'Chưa đạt'
    bad = any(not x['khop'] for x in res['kiem_tra'].values()) or any(x[0]=='Chưa đạt' for x in res['viec'].values())
    partial = any(x[0] in ('Một phần','Không kiểm được') for x in res['viec'].values()) or tp<3
    res['ket_qua'] = 'Chưa đạt - nộp lại' if bad else 'Đạt - cần sửa nhỏ' if partial else 'Đạt'
    return res


def similarity(results):
    items = [(r["file"], ' '.join(x['text'] for x in r.get("insight", []))) for r in results if "loi_mo_file" not in r]
    out = []
    for i in range(len(items)):
        for j in range(i + 1, len(items)):
            a, b = items[i][1], items[j][1]
            if a and b:
                ratio = difflib.SequenceMatcher(None, a.lower(), b.lower()).ratio()
                if ratio >= 0.6:
                    out.append({"a": items[i][0], "b": items[j][0], "ti_le_giong": round(ratio, 2)})
    return out


def main(argv):
    sys.stdout.reconfigure(encoding="utf-8")
    if len(argv) >= 4 and argv[1] == "--folder":
        files = sorted(f for f in glob.glob(os.path.join(argv[2], "**", "*.xlsx"), recursive=True) if not os.path.basename(f).startswith("~$"))
        results = [grade(f, argv[3]) for f in files]
        tay = {}
        if len(argv) >= 6 and argv[4] == "--cham-tay":
            tay = json.load(open(argv[5], encoding="utf8"))
        for r in results:
            t = tay.get(os.path.basename(r["file"]))
            if t:
                ap_cham_tay(r, t["v5"], t["insight"])
        for r in results:
            r["file"] = os.path.basename(r["file"])
        out = {"bai": results, "giong_nhau_insight": similarity(results)}
    elif len(argv) >= 3:
        out = grade(argv[1], argv[2])
    else:
        print(__doc__)
        return 2
    print(json.dumps(out, ensure_ascii=False, indent=1, default=str))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
