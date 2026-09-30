"""Tạo bản slide phát cho học viên: Slide_B<n>.pptx → Slide_B<n>_HocVien.pptx.

Bỏ toàn bộ speaker notes và phần đáp án quiz: các shape tên "Answer …" (khung, chữ cái, nhãn "✓ Đáp án")
và "Dim …" (lớp làm mờ phương án sai), cùng animation hiện đáp án. Học viên chỉ thấy câu hỏi và 4 phương án.

Chạy lại mỗi lần sửa deck giảng viên. Cần python-pptx.
    python3 _generator/build_student_slides.py            # mọi buổi có Slide_B<n>.pptx
    python3 _generator/build_student_slides.py B2         # chỉ các buổi được nêu
"""
import glob, os, sys, zipfile
for s_ in (sys.stdout, sys.stderr):
    if hasattr(s_, "reconfigure"): s_.reconfigure(encoding="utf-8", errors="replace")
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

ROOT = os.environ.get("MCI_ROOT") or os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

ONLY = {a.upper() for a in sys.argv[1:]}
for src in sorted(glob.glob(f"{ROOT}/1_Sessions/*/Slide_B[0-9].pptx")):
    if ONLY and os.path.basename(src)[6:8].upper() not in ONLY: continue
    dst = src.replace(".pptx", "_HocVien.pptx")
    prs = Presentation(src)
    removed = 0
    for slide in prs.slides:
        for rId, rel in list(slide.part.rels.items()):
            if rel.reltype in (RT.NOTES_SLIDE, RT.COMMENTS):
                slide.part.drop_rel(rId)
        answer_shapes = [sh for sh in slide.shapes if sh.name.startswith(("Answer ", "Dim "))]
        for sh in answer_shapes:
            sh._element.getparent().remove(sh._element); removed += 1
        if answer_shapes:
            timing = slide._element.find(qn("p:timing"))
            if timing is not None: slide._element.remove(timing)
    prs.save(dst)
    with zipfile.ZipFile(dst) as z:
        left = [n for n in z.namelist() if "notesSlide" in n or "comments/" in n]
    assert not left, left
    left_ans = [i for i, sl in enumerate(Presentation(dst).slides, 1)
                for sh in sl.shapes if sh.name.startswith(("Answer ", "Dim ")) or (sh.has_text_frame and "Đáp án" in sh.text_frame.text)]
    assert not left_ans, left_ans
    print("ok", os.path.basename(dst), len(prs.slides), f"slide, không còn speaker notes; đã bỏ {removed} shape đáp án")
