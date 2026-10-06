"""Dựng brief capstone B7: Brief/Capstone_Brief_HatViet.docx (phát học viên cuối B6) bằng python-docx.

Theo mẫu "PROJECT BRIEF" của Statistics PDF (p41, 64, 99, 125, 155): khung giới thiệu, thư của khách hàng (From / Subject), Key Objectives.
Số liệu lấy từ build_capstone_meta.json (gen_capstone.py); phần rubric Student-Friendly lấy nguyên văn từ Rubric_B7_Capstone.md.
    python3 _generator/build_brief_b7.py
"""
import json, os, re
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.normpath(os.environ.get("MCI_ROOT") or os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
GEN = os.path.dirname(os.path.abspath(__file__))
B7 = f"{ROOT}/1_Sessions/B7_Capstone"
OUT = f"{B7}/Brief"
os.makedirs(OUT, exist_ok=True)
M = json.load(open(f"{GEN}/build_capstone_meta.json", encoding="utf-8"))
CK = M["checks"]
RED, NAVY, GOLD, GREY = RGBColor(0xC0, 0, 0), RGBColor(0x20, 0x38, 0x64), RGBColor(0x8A, 0x65, 0x00), RGBColor(0x58, 0x58, 0x58)
vn = lambda x: f"{round(x):,}".replace(",", ".")


def base_doc():
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for m in ("left_margin", "right_margin"): setattr(sec, m, Cm(2))
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    for name, size, color in (("Heading 1", 16, NAVY), ("Heading 2", 13, NAVY), ("Heading 3", 11.5, RED)):
        h = d.styles[name]; h.font.name = "Arial"; h.font.size = Pt(size); h.font.color.rgb = color; h.font.bold = True
        rf = h.element.rPr.rFonts
        for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(k)) is not None: del rf.attrib[qn(k)]
        for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"): rf.set(qn("w:" + k.split(":")[1]), "Arial")
    fp = sec.footer.paragraphs[0]
    fp.text = "MCI – Magic Code Institute · Excel for Business Intelligence Analyst · Capstone Buổi 7 · Dữ liệu giả lập, mọi tên đều hư cấu"
    fp.runs[0].font.size = Pt(8); fp.runs[0].font.color.rgb = GREY
    return d


def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_); tcPr.append(shd)


def para(d, text="", bold=False, color=None, size=None, italic=False):
    p = d.add_paragraph()
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic
        if color: r.font.color.rgb = color
        if size: r.font.size = Pt(size)
    return p


def rich(d, parts):
    p = d.add_paragraph()
    for t, f in parts:
        r = p.add_run(t); r.bold = f.get("bold", False); r.italic = f.get("italic", False)
        if f.get("color"): r.font.color.rgb = f["color"]
    return p


def bullet(d, text, bold_prefix=None):
    p = d.add_paragraph(style="List Bullet")
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    p.add_run(text)
    return p


def table(d, rows, widths, header=True, size=9.5, first_bold=False):
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for j, w in enumerate(widths): t.columns[j].width = Cm(w)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.width = Cm(widths[j]); c.text = ""
            r = c.paragraphs[0].add_run(str(v)); r.font.size = Pt(size)
            if i < len(rows) - 1: c.paragraphs[0].paragraph_format.keep_with_next = True
            if header and i == 0:
                r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); shade(c, "203864")
            elif first_bold and j == 0:
                r.bold = True
    d.add_paragraph()
    return t


def box(d, lines, fill="F3F5F9"):
    """Khung một ô, nền xám nhạt (khung giới thiệu và thư của khách hàng)."""
    t = d.add_table(rows=1, cols=1); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0); shade(c, fill); c.text = ""
    first = True
    for text, kw in lines:
        p = c.paragraphs[0] if first else c.add_paragraph(); first = False
        r = p.add_run(text); r.bold = kw.get("bold", False); r.italic = kw.get("italic", False)
        r.font.size = Pt(kw.get("size", 10.5))
        if kw.get("color"): r.font.color.rgb = kw["color"]
    d.add_paragraph()


d = base_doc()
para(d, "MCI – MAGIC CODE INSTITUTE · CAPSTONE BUỔI 7", bold=True, color=GREY, size=9)
d.add_heading("HẠT VIỆT COFFEE | PROJECT BRIEF", level=1)
box(d, [("Bạn là một BI Consultant vừa được Hạt Việt Coffee, một chuỗi cà phê và bánh ngọt gồm 10 chi nhánh trên cả nước, mời hỗ trợ báo cáo kinh doanh. "
         "Dữ liệu và công ty đều giả lập; tên người, chi nhánh và sản phẩm đều hư cấu.", {"italic": True})])

d.add_heading("Thư của khách hàng", level=2)
box(d, [("From: Lan Phương (Giám đốc Vận hành, Hạt Việt Coffee)", {"bold": True}),
        ("Subject: Báo cáo kinh doanh hằng tháng cho Ban điều hành", {"bold": True}),
        ("Chào bạn,", {}),
        ("Cuối năm chúng tôi sắp họp Ban điều hành. Mỗi tháng phòng Vận hành xuất dữ liệu bán hàng từ POS ra một file CSV, nhưng chúng tôi chưa có báo cáo nào "
         "cho Ban điều hành nhìn là hiểu ngay. Tôi chỉ có một cảm nhận chung: doanh thu năm nay có vẻ nhỉnh hơn năm ngoái, nhưng tôi không chắc lợi nhuận có tăng theo. "
         "Năm nay công ty cũng có một chi nhánh mới và vài thay đổi trong menu và cách bán hàng.", {}),
        ("Bạn có thể giúp chúng tôi dựng một dashboard một trang để theo dõi mỗi tháng, kèm vài nhận xét và đề xuất cho Ban điều hành được không? "
         "Từ tháng sau chúng tôi sẽ gửi thêm file dữ liệu mới và muốn bạn chỉ cần bấm Refresh là xong.", {}),
        ("Cảm ơn bạn!", {})])

d.add_heading("Key Objectives", level=2)
for s in ["Tự viết brief cho dashboard: người xem, 3–5 câu hỏi đo lường được, tần suất refresh, quyết định mà dashboard hỗ trợ (ở đây khách hàng chỉ nêu mong muốn chung, chưa có câu hỏi đo lường được).",
          "Nạp và làm sạch dữ liệu bằng Power Query, có thể Refresh khi thêm file tháng mới.",
          "Dựng mô hình ngôi sao trong Data Model và định nghĩa KPI bằng explicit measure.",
          "Dựng dashboard một trang cho Ban điều hành, có mốc so sánh và trạng thái RAG.",
          "Tìm ra những điều Ban điều hành nên biết, kiểm chứng bằng dữ liệu, và đề xuất hành động theo cấu trúc Situation → Complication → Resolution."]:
    bullet(d, s)

d.add_heading("Sản phẩm phải nộp", level=2)
table(d, [["#", "Sản phẩm", "Nội dung"],
          ["1", "File Excel B7_<HoTen>.xlsx",
           "Các query Power Query; Data Model (star schema) và explicit measure; KPI Definition Card cho từng KPI; sheet Brief; dashboard một trang (5–8 visual); sheet Insight; ghi chú đối soát (số dòng và tổng doanh thu của nguồn)."],
          ["2", "Thuyết trình 5–7 phút tại Buổi 7",
           "Cá nhân. Kể bằng dashboard, không kể lại các bước kỹ thuật: tình hình, vấn đề đã kiểm chứng, đề xuất cho Ban điều hành. Có phần hỏi đáp (một câu về filter context, một câu phản biện)."]],
      [0.8, 4.4, 10.8], first_bold=True)
rich(d, [("Hạn nộp: ", {"bold": True}), ("trước Buổi 7, tại https://hocvien.mcivietnam.com/auth#/ (tài khoản như các buổi trước). Làm cá nhân, không làm nhóm. Máy Windows có Power Query và Power Pivot như các buổi trước.", {})])

d.add_heading("Lịch Buổi 7", level=2)
table(d, [["Phút", "Nội dung"],
          ["20", "Ôn tập: 3 chuỗi xương sống, \"3 điều cần nhớ\" mỗi buổi; xếp các câu chuyện dữ liệu vào thang 5 tầng; quiz tổng hợp"],
          ["40", "Hoàn thiện dự án, hỏi đáp, bài kiểm tra Refresh: giảng viên đưa một file dữ liệu tháng mới để bạn thả vào folder rồi Refresh (chấm tiêu chí C2)"],
          ["70", "Thuyết trình 5–7 phút mỗi học viên, nhận xét theo rubric"],
          ["20", "Tổng kết khóa học, lộ trình tiếp theo (Power BI, thống kê suy luận)"]], [1.8, 14.2], first_bold=True)

d.add_heading("Quy tắc làm bài", level=2)
for s in ["Mọi bước làm sạch nằm trong Power Query (Applied Steps); không sửa tay dữ liệu trên sheet. File mới được thả vào folder Sales thì Refresh phải chạy không lỗi và ra số đúng.",
          "Đường dẫn folder nên là tham số (Parameter) để đổi máy không phải sửa query.",
          "Mọi con số trên dashboard lấy từ measure (CUBEVALUE hoặc Pivot). Không viết công thức KPI ở tầng hiển thị.",
          "Mỗi visual trả lời đúng một câu hỏi trong brief của bạn; visual nào không trả lời câu nào thì bỏ.",
          "Ghi lại số đối soát của nguồn (số dòng, tổng doanh thu) và so với kết quả sau khi làm sạch."]:
    bullet(d, s)

# ---------------------------------------------------------------- dữ liệu
d.add_page_break()
d.add_heading("Dữ liệu được cung cấp", level=2)
para(d, "Thư mục Du_lieu_Capstone/ gồm:", )
table(d, [["File", "Nội dung"],
          ["Sales/Sales_2024.csv", "Toàn bộ bán hàng năm 2024 (một file)"],
          ["Sales/Sales_2025-01.csv … Sales_2025-11.csv", "Bán hàng năm 2025, mỗi tháng một file (xuất từ POS). Hằng tháng có thêm một file mới"],
          ["HatViet_Lookups.xlsx", "Bốn Excel Table: tblProduct (sản phẩm), tblBranch (chi nhánh), tblChannel (kênh bán), tblCustomer (khách thành viên)"],
          ["HatViet_Target_2025.xlsx", "Target doanh thu 2025 theo chi nhánh × tháng, do phòng Kế hoạch lập"]], [6.4, 9.6], first_bold=True)
d.add_heading("Mô tả cột", level=3)
table(d, [["Bảng", "Cột", "Ghi chú"],
          ["Sales (CSV)", "Order_ID, Line_No", "Một đơn có một hoặc nhiều dòng; Order_ID + Line_No là khóa của dòng"],
          ["", "Order_Time", "Thời điểm bán (ngày giờ, định dạng của POS xuất ra)"],
          ["", "Branch_Code, Channel_Code, SKU", "Khóa nối sang tblBranch, tblChannel, tblProduct"],
          ["", "Customer_ID", "Mã khách thành viên; để trống khi khách không dùng thẻ thành viên"],
          ["", "Qty, Unit_Price, Discount_Pct", "Số lượng, đơn giá thực tế lúc bán (VNĐ), % giảm giá của đơn"],
          ["tblProduct", "SKU, Product_Name, Category, Unit_Price, Unit_Cost, Launch_Date, Updated_At", "Unit_Price ở đây là giá niêm yết hiện hành; Unit_Cost là giá vốn mỗi đơn vị"],
          ["tblBranch", "Branch_Code, Branch_Name, City, Region, Open_Date, Area_m2, Seats", "10 chi nhánh thuộc 3 miền (Bắc, Trung, Nam)"],
          ["tblChannel", "Channel_Code, Channel_Name, Commission_Rate", "Kênh bán và tỷ lệ hoa hồng trả cho nền tảng (0% với kênh tự bán)"],
          ["tblCustomer", "Customer_ID, Join_Date, Tier, Home_Branch", "Khách thành viên: hạng thẻ và chi nhánh quen thuộc"],
          ["Target_2025", "Branch_Code, Branch_Name, Tháng 1 … Tháng 12", "Target doanh thu từng chi nhánh theo tháng (VNĐ)"]],
      [2.6, 6.4, 7.0], size=9, first_bold=True)
d.add_heading("Thuật ngữ nghiệp vụ", level=3)
table(d, [["Thuật ngữ", "Định nghĩa của khách hàng"],
          ["Doanh thu (Revenue)", "Qty × Unit_Price × (1 − Discount_Pct / 100), tính từ Unit_Price trong Sales"],
          ["Giá vốn (COGS)", "Qty × Unit_Cost của sản phẩm (từ tblProduct)"],
          ["Biên lợi nhuận gộp (Gross Margin %)", "(Doanh thu − COGS) ÷ Doanh thu"],
          ["Hoa hồng nền tảng", "Doanh thu của đơn × Commission_Rate của kênh bán"],
          ["Số đơn / AOV", "Số Order_ID khác nhau; AOV = Doanh thu ÷ Số đơn"],
          ["Đạt target", "Doanh thu ÷ Target của cùng chi nhánh và tháng. ✔ Đạt ≥ 100% · ▲ Cảnh báo 90% đến dưới 100% · ✘ Không đạt dưới 90% (theo Warning 90% của KPI Dictionary)"]],
      [5.2, 10.8], first_bold=True)
d.add_heading("Số tự đối chiếu sau khi làm sạch", level=3)
para(d, f"Ghép {len(CK['sales_files_students'])} file Sales ({', '.join(CK['sales_files_students'][:2])}, … , {CK['sales_files_students'][-1]}), "
        f"làm sạch và bỏ dòng trùng. Bạn phải ra: {vn(CK['total_rows_clean_students'])} dòng bán hàng và tổng doanh thu {vn(CK['total_rev_students'])} VNĐ. "
        "Ra số khác thì quay lại tìm nguyên nhân: đây là thói quen đối soát (D2) mà rubric chấm ở C2. Dữ liệu thô có thể có vấn đề; bạn cần tìm và xử lý chúng bằng Power Query, không sửa tay.")

# ---------------------------------------------------------------- gợi ý
d.add_heading("Gợi ý quy trình (đối chiếu với các buổi đã học)", level=2)
table(d, [["Bước", "Việc cần làm", "Học ở", "Tiêu chí"],
          ["1", "Đọc thư, viết brief: người xem, 3–5 câu hỏi có đủ chỉ số · chiều · thời gian · mốc so sánh", "B6", "C1"],
          ["2", "Query Sales (From Folder), lookups và Target (Unpivot); làm sạch, đặt tên query, tham số folder; ghi số đối soát", "B4", "C2"],
          ["3", "Nạp vào Data Model; fact Sales và Target nối qua lookup chung, Calendar liên tục; ẩn cột kỹ thuật", "B5", "C3"],
          ["4", "Viết explicit measure theo tầng; KPI Definition Card; đối soát ít nhất một KPI bằng cách tính thứ hai", "B2, B5", "C4"],
          ["5", "Dựng dashboard: KPI → Trend → Detail; target và RAG có biểu tượng; Slicer nối đủ; chú thích điểm bất thường; phát hành", "B3, B6", "C5"],
          ["6", "Drill-down để kiểm chứng nguyên nhân; viết insight theo SCR kèm hành động; chuẩn bị 5–7 phút", "B6", "C6"]],
      [1.6, 9.4, 2.0, 3.0], first_bold=True)
rich(d, [("Gợi ý thời gian: ", {"bold": True}), ("khoảng 6–8 giờ làm ở nhà, chia nhiều buổi. Làm theo thứ tự trên, kiểm số sau mỗi bước.", {})])
rich(d, [("Lưu ý: ", {"bold": True, "color": NAVY}),
         ("brief này cố ý chỉ nêu mong muốn chung. Phần việc đầu tiên của bạn là biến nó thành câu hỏi đo lường được. "
          "Dữ liệu có những điều thú vị cần tự khám phá; câu insight tốt là câu bạn đã kiểm chứng trên dashboard, không phải câu đoán.", {})])

# ---------------------------------------------------------------- rubric student-friendly
rub = open(f"{B7}/Rubric_B7_Capstone.md", encoding="utf-8").read()
m = re.search(r"## Student-Friendly Version\n(.*?)\n---\n\n## Calibration Notes", rub, re.S)
d.add_page_break()
d.add_heading("Rubric tự đánh giá (bản Student-Friendly)", level=2)
para(d, "Dùng để tự đánh giá trước khi nộp. Chấm theo 6 tiêu chí, mỗi tiêu chí 4 level. Hoàn thành khóa, đạt Level 3 (Builder): không tiêu chí nào ở Level 1, "
        "và C2 từ Level 3, và C5 từ Level 3. Đạt Level 4 (Modeler): đạt Builder, và C3 = Level 4, và C4 = Level 4. Chưa đạt thì được nộp lại một lần sau khi nhận phản hồi. "
        "Trọng số: C1 10% · C2 20% · C3 15% · C4 20% · C5 20% · C6 15%. Bản đầy đủ: Rubric_B7_Capstone.md.", italic=True, color=GREY)
if m:
    for line in m.group(1).strip().splitlines():
        line = line.strip()
        if not line or line.startswith("*Dùng để"):
            continue
        if line.startswith("**C"):
            d.add_heading(line.strip("*"), level=3).paragraph_format.keep_with_next = True
        elif line.startswith("- L"):
            lvl, _, text = line[2:].partition(":")
            p = d.add_paragraph(style="List Bullet")
            p.add_run(f"L{lvl}: ").bold = True
            # in nghiêng cụm *Thử ngay:* / *Bài kiểm tra 5 giây:*
            for k, part in enumerate(re.split(r"\*(.+?)\*", text.strip())):
                r = p.add_run(part)
                if k % 2: r.italic = True; r.bold = True
else:
    para(d, "(không tìm thấy mục Student-Friendly trong Rubric_B7_Capstone.md)", color=RED)

d.add_heading("Checklist trước khi nộp", level=3).paragraph_format.keep_with_next = True
for s in ["Brief: mỗi câu hỏi có đủ chỉ số, chiều, thời gian, mốc so sánh; mỗi visual trả lời một câu.",
          f"Số đối soát: {vn(CK['total_rows_clean_students'])} dòng và {vn(CK['total_rev_students'])} VNĐ sau khi làm sạch.",
          "Thả thêm một file Sales vào folder rồi Refresh: không lỗi, số dòng và tổng doanh thu khớp.",
          "Mô hình hình ngôi sao, Calendar liên tục, cột kỹ thuật đã ẩn; lọc theo chi nhánh, kênh, sản phẩm đều ra số đúng.",
          "Mỗi KPI có Definition Card và explicit measure; KPI tỷ lệ là tổng chia tổng; đã đối soát một KPI bằng cách tính khác.",
          "Dashboard 5–8 visual; KPI có mốc so sánh và RAG có biểu tượng; Slicer nối mọi visual; qua bài kiểm tra 5 giây.",
          "Insight theo SCR kèm đề xuất cho đúng người xem; biết trả lời \"vậy thì sao?\".",
          "Đã nộp B7_<HoTen>.xlsx trước Buổi 7."]:
    bullet(d, "☐  " + s)
d.save(f"{OUT}/Capstone_Brief_HatViet.docx")
print("ok", f"{OUT}/Capstone_Brief_HatViet.docx")
