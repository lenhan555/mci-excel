r"""Dựng file Excel của Buổi 3 (PivotTable): đáp án thực hành, file demo (bản giảng viên + _HocVien), bài tập về nhà.

Chỉ đọc file đã có: MCI_Mart_Master.xlsx và B1_MCI_Mart_Q1_2025_Answer.xlsx (thứ tự dòng Q1). Không đọc file Answer B2
(đã lưu bằng Excel, openpyxl sẽ làm mất Data Validation) mà dựng lại đúng cấu trúc của nó, rồi thêm phần của Buổi 3.

PivotTable, Slicer, Timeline, PivotChart ghi bằng pivot_xml.py (openpyxl không tạo được). Cuối script, mọi file được
MỞ BẰNG MICROSOFT EXCEL (Mac: AppleScript; Windows: Excel qua COM, gọi bằng PowerShell), Refresh All và lưu lại, rồi
script đọc lại giá trị để kiểm tra kết quả mong đợi (sheet Kiem_tra).
Mọi số liệu tính từ dataset; con số dùng trong kịch bản in ra cuối script và lưu vào build_b3_meta.json cạnh script.
Đường dẫn tự lấy theo vị trí script (thư mục MCI Excel = hai cấp trên _generator), nên chạy được ở máy nào cũng được.

Chạy từ thư mục MCI Excel (cần pandas, numpy, openpyxl):
    Windows:  python 2_Dataset\_generator\build_b3.py
    Mac:      python3 2_Dataset/_generator/build_b3.py
    Thêm --no-excel để chỉ dựng file, bỏ bước mở bằng Excel (khi đó Pivot chưa có số, Kiem_tra chưa kiểm được).
Trước khi chạy: đóng các file B3 (Practice, Demo, Homework) nếu đang mở trong Excel. Script ghi đè toàn bộ các file đó.
"""
import json, os, random, sys
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pivot_xml import Cache, Pivot, Value, Slicer, Timeline, Chart, inject, excel_resave

for s_ in (sys.stdout, sys.stderr):                           # Windows: in được tiếng Việt, dấu ✔ kể cả khi chuyển hướng ra file
    if hasattr(s_, "reconfigure"): s_.reconfigure(encoding="utf-8", errors="replace")
ROOT = os.environ.get("MCI_ROOT") or os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
MASTER = f"{ROOT}/2_Dataset/MCI_Mart_Master.xlsx"
B1_ANS = f"{ROOT}/1_Sessions/B1_Nen_tang_Excel_Du_lieu/Practice/B1_MCI_Mart_Q1_2025_Answer.xlsx"
B3 = f"{ROOT}/1_Sessions/B3_PivotTable"
PRACTICE, DEMO, HW = f"{B3}/Practice", f"{B3}/Demo", f"{B3}/Homework"
for d in (PRACTICE, DEMO, HW): os.makedirs(d, exist_ok=True)
USE_EXCEL = "--no-excel" not in sys.argv

def _locked(p):
    """Windows khóa file đang mở trong Excel: phát hiện trước để không dựng dở dang (file này mới, file kia cũ)."""
    try:
        with open(p, "r+b"): return False
    except PermissionError:
        return True
_busy = [os.path.basename(p) for d in (PRACTICE, DEMO, HW) for p in sorted(os.path.join(d, n) for n in os.listdir(d))
         if p.endswith(".xlsx") and _locked(p)]
if _busy:
    sys.exit("Đang mở trong Excel, hãy đóng rồi chạy lại: " + ", ".join(_busy))

F = "Arial"
HDR_FONT = Font(name=F, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="203864")
TITLE = Font(name=F, size=14, bold=True, color="203864")
BODY = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
INPUT = Font(name=F, size=10, color="0000FF")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
NOTE = Font(name=F, size=10, italic=True, color="585858")
DATE_FMT, NUM_FMT, PCT_FMT = "dd/mm/yyyy", "#,##0", "0.0%"
MONTH_FMT = '"Tháng "0'
vn = lambda x: f"{x:,.0f}".replace(",", ".")
pct = lambda x, d=1: f"{x * 100:.{d}f}%".replace(".", ",")
dec = lambda x, d=1: f"{x:.{d}f}".replace(".", ",")
spct = lambda x, d=1: ("+" if x > 0 else "−" if x < 0 else "") + pct(abs(x), d)

def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

def header(ws, row, labels, col=1):
    for j, h in enumerate(labels, start=col):
        c = ws.cell(row=row, column=j, value=h); c.font = HDR_FONT; c.fill = HDR_FILL

def py(v):
    if isinstance(v, pd.Timestamp): return v.to_pydatetime()
    if isinstance(v, float) and np.isnan(v): return None
    return v.item() if hasattr(v, "item") else v

def write_rows(ws, df, start_row=2, formats=None, font=BODY):
    for i, row in enumerate(df.itertuples(index=False), start=start_row):
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=py(v)); c.font = font
            if formats and df.columns[j - 1] in formats: c.number_format = formats[df.columns[j - 1]]

def add_table(ws, name, ncols, nrows, start_row=1):
    t = Table(displayName=name, ref=f"A{start_row}:{get_column_letter(ncols)}{start_row + nrows}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)

def table_df(ws, df, name, formats=None, widths=None):
    header(ws, 1, list(df.columns)); write_rows(ws, df, 2, formats)
    add_table(ws, name, len(df.columns), len(df)); ws.freeze_panes = "A2"
    if widths: set_widths(ws, widths)

def guide(ws, title, steps):
    ws["A1"] = title; ws["A1"].font = TITLE
    header(ws, 3, ["Slide", "Thao tác", "Kết quả mong đợi"])
    for i, (slide, action, result) in enumerate(steps, start=4):
        for j, v in enumerate((slide, action, result), start=1):
            c = ws.cell(row=i, column=j, value=v); c.font = BODY
            c.alignment = Alignment(wrap_text=True, vertical="top")
    set_widths(ws, [8, 72, 54])

def sheet_title(ws, title, question=None):
    ws["A1"] = title; ws["A1"].font = TITLE
    if question: ws["A2"] = "Câu hỏi: " + question; ws["A2"].font = NOTE

FORBIDDEN_HV = ("trước lớp", "cho lớp", "giảng viên", "demo", "Demo_")

def save_both(wb, name, work_sheet):
    """Bản giảng viên: mở ra ở sheet làm việc, Huong_dan để cuối. Bản _HocVien: Huong_dan đứng đầu, tiêu đề VÍ DỤ,
    không còn chữ dành cho giảng viên (kiểm tra bằng assert). Trả về hai đường dẫn."""
    g = wb["Huong_dan"]
    wb.move_sheet(g, offset=len(wb.sheetnames) - 1 - wb.sheetnames.index("Huong_dan"))
    wb.active = wb.sheetnames.index(work_sheet)
    for ws in wb.worksheets: ws.sheet_view.tabSelected = ws.title == work_sheet
    p1 = f"{DEMO}/{name}.xlsx"; wb.save(p1)
    g["A1"] = g["A1"].value.replace("DEMO B3", "VÍ DỤ B3")
    for row in g.iter_rows(min_row=4):
        for c in row:
            if isinstance(c.value, str):
                low = c.value.lower()
                assert not any(w.lower() in low for w in FORBIDDEN_HV), (name, c.value)
    wb.move_sheet(g, offset=-wb.sheetnames.index("Huong_dan"))
    wb.active = 0
    for ws in wb.worksheets: ws.sheet_view.tabSelected = ws.title == "Huong_dan"
    p2 = f"{DEMO}/{name}_HocVien.xlsx"; wb.save(p2)
    return p1, p2

# ---------------------------------------------------------------------
# Dữ liệu (giống build_b2.py: thứ tự dòng của file khởi đầu, làm giàu bằng Product/Store)
# ---------------------------------------------------------------------
products = pd.read_excel(MASTER, "Product")
stores = pd.read_excel(MASTER, "Store")
m_sales = pd.read_excel(MASTER, "Sales")
b1 = pd.read_excel(B1_ANS, "Sales", usecols="A:J")
N = len(b1)
cost = products.set_index("Product_Code")["Unit_Cost"]
cat = products.set_index("Product_Code")["Category"]
region_of = stores.set_index("Store_Code")["Region"]
q1 = b1.copy()
q1["Revenue"] = q1.Quantity * q1.Unit_Price
q1["Category"] = q1.Product_Code.map(cat)
q1["Unit_Cost"] = q1.Product_Code.map(cost)
q1["COGS"] = q1.Quantity * q1.Unit_Cost
q1["Gross_Margin"] = q1.Revenue - q1.COGS
q1["Thang"] = q1.Order_Date.dt.month
q1["Thu"] = q1.Order_Date.dt.strftime("%a")                     # giá trị Excel hiện trên máy tiếng Anh (TEXT(…,"ddd"))
q1["wd"] = q1.Order_Date.dt.dayofweek                           # 0 = thứ Hai
q1["So_don"] = (q1.Line_No == 1).astype(int)
REV, COGS = int(q1.Revenue.sum()), int(q1.COGS.sum())
GM, ORDERS = REV - COGS, int(q1.So_don.sum())
REGIONS, MONTHS = ["Bắc", "Trung", "Nam"], [1, 2, 3]
CATS_Q1 = list(q1.groupby("Category").Revenue.sum().sort_values(ascending=False).index)
STORES = sorted(stores.Store_Code)
DAYS_EN = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]

meta = {"q1": {"rows": N, "revenue": REV, "orders": ORDERS, "cogs": COGS, "gm_pct": GM / REV, "aov": REV / ORDERS,
               "avg_line": REV / N, "max_line": int(q1.Revenue.max())}}
bym = q1.groupby("Thang").agg(rev=("Revenue", "sum"), orders=("So_don", "sum"), gm=("Gross_Margin", "sum"))
meta["month"] = {int(m): {"rev": int(r.rev), "orders": int(r.orders), "aov": r.rev / r.orders, "gm_pct": r.gm / r.rev} for m, r in bym.iterrows()}
meta["month"][2]["mom"] = bym.rev[2] / bym.rev[1] - 1
meta["month"][3]["mom"] = bym.rev[3] / bym.rev[2] - 1
meta["month"][1]["cum"], meta["month"][2]["cum"], meta["month"][3]["cum"] = [int(x) for x in bym.rev.cumsum()]
grid = q1.pivot_table(index="Region", columns="Thang", values="Revenue", aggfunc="sum")
meta["grid"] = {f"{r}-{m}": int(grid.loc[r, m]) for r in REGIONS for m in MONTHS}
meta["grid"]["bac_t3_lines"] = int(((q1.Region == "Bắc") & (q1.Thang == 3)).sum())
reg = q1.groupby("Region").Revenue.sum()
meta["region"] = {r: {"rev": int(reg[r]), "share": reg[r] / REV} for r in REGIONS}
cs = q1.pivot_table(index="Category", columns="Thang", values="Revenue", aggfunc="sum")
meta["cat_share"] = {c: {**{int(m): cs.loc[c, m] / cs[m].sum() for m in MONTHS}, "q1": cs.loc[c].sum() / REV, "rev": int(cs.loc[c].sum())} for c in CATS_Q1}
# Tết: theo ngày, theo tuần (Group 7 ngày từ 01/01), theo giai đoạn
daily = q1.groupby("Order_Date").Revenue.sum().reindex(pd.date_range("2025-01-01", "2025-03-31"), fill_value=0)
wk = (daily.index - pd.Timestamp("2025-01-01")).days // 7
weekly = daily.groupby(wk).sum()
meta["weekly"] = [{"from": (pd.Timestamp("2025-01-01") + pd.Timedelta(days=7 * k)).strftime("%d/%m"),
                   "to": (pd.Timestamp("2025-01-01") + pd.Timedelta(days=7 * k + 6)).strftime("%d/%m"), "rev": int(v)} for k, v in weekly.items()]
SEG = {"dau_thang": ("2025-01-01", "2025-01-08"), "truoc_tet": ("2025-01-09", "2025-01-28"), "nghi_tet": ("2025-01-29", "2025-02-02"),
       "sau_tet": ("2025-02-03", "2025-03-31")}
meta["tet"] = {k: {"days": len(daily[a:b]), "rev": int(daily[a:b].sum()), "per_day": float(daily[a:b].mean())} for k, (a, b) in SEG.items()}
meta["tet"]["normal_weeks_min"] = int(weekly.iloc[5:].min()); meta["tet"]["normal_weeks_max"] = int(weekly.iloc[5:].max())
seg = np.select([q1.Order_Date.between(*SEG["truoc_tet"]), q1.Order_Date.between(*SEG["nghi_tet"])], ["truoc_tet", "nghi_tet"], "thuong")
sshare = q1.assign(seg=seg).pivot_table(index="Category", columns="seg", values="Revenue", aggfunc="sum")
sshare = sshare / sshare.sum()
meta["tet"]["share_banh_keo"] = {"truoc_tet": sshare.loc["Bánh kẹo", "truoc_tet"], "thuong": sshare.loc["Bánh kẹo", "thuong"]}
meta["tet"]["share_do_uong"] = {"truoc_tet": sshare.loc["Đồ uống", "truoc_tet"], "thuong": sshare.loc["Đồ uống", "thuong"]}
du = q1[q1.Category == "Đồ uống"].groupby(seg[q1.Category == "Đồ uống"]).Revenue.sum()
n_thuong = 90 - meta["tet"]["truoc_tet"]["days"] - meta["tet"]["nghi_tet"]["days"]
meta["tet"]["do_uong_per_day_ratio"] = (du["truoc_tet"] / 20) / (du["thuong"] / n_thuong)
# Thứ trong tuần
thu = q1.groupby("wd").Revenue.sum()
meta["thu"] = {DAYS_EN[k]: int(v) for k, v in thu.items()}
meta["thu_sorted"] = [DAYS_EN[k] for k in thu.sort_values(ascending=False).index]
meta["weekend"] = {"avg_weekend": float(thu[[5, 6]].mean()), "avg_weekday": float(thu[[0, 1, 2, 3, 4]].mean())}
meta["weekend"]["ratio"] = meta["weekend"]["avg_weekend"] / meta["weekend"]["avg_weekday"]
fm = q1[q1.Thang >= 2].groupby("wd").Revenue.sum()
meta["thu_t2t3"] = {DAYS_EN[k]: int(v) for k, v in fm.items()}
meta["thu_t2t3_sorted"] = [DAYS_EN[k] for k in fm.sort_values(ascending=False).index]
# Cửa hàng × Tháng, ST07
sm = q1.pivot_table(index="Store_Code", columns="Thang", values="Revenue", aggfunc="sum")
mom3 = (sm[3] / sm[2] - 1).sort_values()
meta["store"] = {s: {int(m): int(sm.loc[s, m]) for m in MONTHS} for s in STORES}
meta["store_mom3"] = [[s, float(v)] for s, v in mom3.items()]
st7 = q1[(q1.Store_Code == "ST07") & (q1.Thang == 3)]
a7, b7 = st7[st7.Order_Date < "2025-03-10"], st7[st7.Order_Date >= "2025-03-10"]
st7f = q1[(q1.Store_Code == "ST07") & (q1.Thang == 2)]
meta["st07"] = {"t3_rev": int(st7.Revenue.sum()), "t3_rows": len(st7), "t3_orders": int(st7.So_don.sum()),
                "mom3": float(mom3["ST07"]), "rank_drop": 1 + list(mom3.index).index("ST07"),
                "t2_orders": int(st7f.So_don.sum()), "t2_per_day": st7f.So_don.sum() / 28,
                "before": {"rows": len(a7), "orders": int(a7.So_don.sum()), "days_sold": a7.Order_Date.nunique(), "days": 9},
                "after": {"rows": len(b7), "orders": int(b7.So_don.sum()), "days_sold": b7.Order_Date.nunique(), "days": 22}}
meta["st07"]["before"]["per_day"] = meta["st07"]["before"]["orders"] / 9
meta["st07"]["after"]["per_day"] = meta["st07"]["after"]["orders"] / 22
st4 = q1[(q1.Store_Code == "ST04") & (q1.Thang == 3)]
meta["st04"] = {"mom3": float(mom3["ST04"]), "t3_orders": int(st4.So_don.sum()), "days_sold": st4.Order_Date.nunique(),
                "before_orders": int(st4[st4.Order_Date < "2025-03-10"].So_don.sum()), "after_orders": int(st4[st4.Order_Date >= "2025-03-10"].So_don.sum())}
meta["hcm"] = int(q1[q1.Store_Code.isin(["ST08", "ST09", "ST10"])].Revenue.sum())
tp = q1.groupby(["Product_Code", "Product_Name"]).Revenue.sum().sort_values(ascending=False).head(5)
meta["top5"] = [[c, n, int(v)] for (c, n), v in tp.items()]
pb = q1.groupby((q1.Unit_Price // 100000) * 100000).agg(rows=("Revenue", "size"), rev=("Revenue", "sum"))
meta["price_band"] = [[int(k), int(r.rows), int(r.rev)] for k, r in pb.iterrows()]
byc = q1.groupby("Category").agg(rev=("Revenue", "sum"), gm=("Gross_Margin", "sum"))
meta["cat_gm"] = {c: r.gm / r.rev for c, r in byc.iterrows()}

# =====================================================================
# Bảng tblSales giống file Answer B2 (file khởi đầu B3)
# =====================================================================
BASE = ["Region", "Order_ID", "Line_No", "Order_Date", "Store_Code", "Customer_ID", "Product_Code", "Product_Name", "Quantity", "Unit_Price"]
NEW_B2 = ["Revenue", "Pct_Share", "Category", "Unit_Cost", "COGS", "Gross_Margin", "Thang", "Quy", "Thu", "Phan_khuc"]
TR = lambda c: f"tblSales[[#This Row],[{c}]]"
def b2_formulas(pct_share):
    return {
        "Revenue": (f"={TR('Quantity')}*{TR('Unit_Price')}", NUM_FMT),
        "Pct_Share": (pct_share, "0.000%"),
        "Category": (f"=_xlfn.XLOOKUP({TR('Product_Code')},tblProduct[Product_Code],tblProduct[Category])", None),
        "Unit_Cost": (f"=_xlfn.XLOOKUP({TR('Product_Code')},tblProduct[Product_Code],tblProduct[Unit_Cost])", NUM_FMT),
        "COGS": (f"={TR('Quantity')}*{TR('Unit_Cost')}", NUM_FMT),
        "Gross_Margin": (f"={TR('Revenue')}-{TR('COGS')}", NUM_FMT),
        "Thang": (f"=MONTH({TR('Order_Date')})", "0"),
        "Quy": (f'="Q"&ROUNDUP(MONTH({TR("Order_Date")})/3,0)', None),
        "Thu": (f'=TEXT({TR("Order_Date")},"ddd")', None),
        "Phan_khuc": (f'=_xlfn.IFS({TR("Unit_Price")}>=100000,"Cao",{TR("Unit_Price")}>=30000,"Trung",TRUE,"Thấp")', None),
        "So_don": (f"=IF({TR('Line_No')}=1,1,0)", "0"),
    }

def sales_sheet(wb, pct_share, with_so_don):
    ws = wb.active; ws.title = "Sales"
    cols = BASE + NEW_B2 + (["So_don"] if with_so_don else [])
    fm = b2_formulas(pct_share)
    header(ws, 1, cols)
    for i, row in enumerate(b1.itertuples(index=False), start=2):
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=py(v)); c.font = BODY
            if j == 4: c.number_format = DATE_FMT
            if j == 10: c.number_format = NUM_FMT
        for j, col in enumerate(cols[len(BASE):], start=len(BASE) + 1):
            f, fmt = fm[col]
            c = ws.cell(row=i, column=j, value=f); c.font = BODY
            if fmt: c.number_format = fmt
    add_table(ws, "tblSales", len(cols), N)
    ws.freeze_panes = "A2"
    set_widths(ws, [9, 15, 7, 12, 10, 11, 11, 36, 9, 11, 12, 9, 16, 10, 12, 13, 7, 6, 6, 10, 8])
    dv = DataValidation(type="list", formula1="=Store!$A$2:$A$13", showErrorMessage=True,
                        errorTitle="Mã cửa hàng không hợp lệ", error="Chỉ nhập mã có trong sheet Store (ST01–ST12).")
    dv.add(f"E2:E{N + 1}"); ws.add_data_validation(dv)
    dv = DataValidation(type="whole", operator="greaterThan", formula1="0", showErrorMessage=True,
                        errorTitle="Số lượng không hợp lệ", error="Quantity phải là số nguyên lớn hơn 0.")
    dv.add(f"I2:I{N + 1}"); ws.add_data_validation(dv)
    return cols

def kpi_sumifs_sheet(wb, blocks=("rev", "orders", "aov", "gm")):
    """Giống sheet KPI_SUMIFS của đáp án B2: Doanh thu ở B6:D8 (Bắc–T3 = D6), tổng ở E9."""
    k = wb.create_sheet("KPI_SUMIFS")
    k["A1"] = "BẢNG KPI THEO MIỀN × THÁNG — QUÝ 1/2025 (SUMIFS)"; k["A1"].font = TITLE
    k["A2"] = "Mỗi ô = một bộ lọc: Miền (cột A, khóa cột $A) × Tháng (dòng tiêu đề, khóa dòng). Bảng này là “đáp án” để đối soát ở Buổi 3 và Buổi 5."
    k["A2"].font = NOTE
    titles = {"rev": "Doanh thu (VNĐ)", "orders": "Số đơn", "aov": "Giá trị đơn trung bình – AOV (VNĐ)", "gm": "Gross Margin %"}
    anchor, r0 = {}, 4
    for key in blocks:
        k.cell(row=r0, column=1, value=titles[key]).font = BOLD
        h = r0 + 1
        header(k, h, ["Miền", 1, 2, 3, "Quý 1"])
        for j in range(2, 5): k.cell(row=h, column=j).number_format = MONTH_FMT
        anchor[key] = h
        for i, reg in enumerate(REGIONS + ["Tổng"], start=h + 1):
            k.cell(row=i, column=1, value=reg).font = BOLD if reg == "Tổng" else BODY
            for j in range(2, 6):
                col = get_column_letter(j)
                tot = reg == "Tổng"
                if key == "rev":
                    f = (f"=SUMIFS(tblSales[Revenue],tblSales[Region],$A{i},tblSales[Thang],{col}${h})" if not tot and j < 5
                         else (f"=SUM(B{i}:D{i})" if j == 5 else f"=SUM({col}{h + 1}:{col}{h + 3})")); fmt = NUM_FMT
                elif key == "orders":
                    f = (f"=COUNTIFS(tblSales[Region],$A{i},tblSales[Thang],{col}${h},tblSales[Line_No],1)" if not tot and j < 5
                         else (f"=SUM(B{i}:D{i})" if j == 5 else f"=SUM({col}{h + 1}:{col}{h + 3})")); fmt = NUM_FMT
                elif key == "aov":
                    f = f"={col}{anchor['rev'] + (i - h)}/{col}{anchor['orders'] + (i - h)}"; fmt = NUM_FMT
                else:
                    ri = anchor["rev"] + (i - h)
                    if not tot and j < 5: f = f"=SUMIFS(tblSales[Gross_Margin],tblSales[Region],$A{i},tblSales[Thang],{col}${h})/{col}{ri}"
                    elif not tot: f = f"=SUMIFS(tblSales[Gross_Margin],tblSales[Region],$A{i})/{col}{ri}"
                    elif j < 5: f = f"=SUMIFS(tblSales[Gross_Margin],tblSales[Thang],{col}${h})/{col}{ri}"
                    else: f = f"=SUM(tblSales[Gross_Margin])/{col}{ri}"
                    fmt = PCT_FMT
                c = k.cell(row=i, column=j, value=f); c.font = BOLD if tot else BODY; c.number_format = fmt
        r0 = h + 6
    k.cell(row=r0, column=1, value="AOV và Gross Margin % của dòng Tổng tính bằng tổng ÷ tổng, KHÔNG lấy trung bình của các miền.").font = NOTE
    set_widths(k, [34, 15, 15, 15, 16])
    return k

CARD_FIELDS = ["Tên KPI", "Mục đích (trả lời câu hỏi gì)", "Công thức", "Nguồn dữ liệu", "Grain",
               "Dimensions", "Owner", "Target", "Ngưỡng Warning / Critical", "Ghi chú giới hạn"]
CARDS = [
    ["Total Revenue (Tổng doanh thu)", "Tháng này bán được bao nhiêu tiền, ở miền / cửa hàng nào?",
     "SUM(Quantity × Unit_Price) = SUM(tblSales[Revenue])", "tblSales (giá bán thực tế trong giao dịch)",
     "Dữ liệu: 1 dòng = 1 sản phẩm trong 1 đơn. KPI báo cáo theo tháng", "Miền, cửa hàng, ngành hàng, tháng",
     "Giám đốc bán hàng", "Theo bảng Target (Tháng × Cửa hàng) — dùng từ Buổi 4", "90% / 80% target",
     "Cộng dồn được theo mọi chiều. Chưa trừ trả hàng, khuyến mại."],
    ["Total Orders (Số đơn)", "Có bao nhiêu lượt mua hàng?",
     "Số Order_ID khác nhau = COUNTIFS(…, tblSales[Line_No], 1)", "tblSales (cột Order_ID, Line_No)",
     "Dữ liệu: 1 dòng = 1 sản phẩm trong 1 đơn, nên KHÔNG đếm số dòng. KPI báo cáo theo tháng", "Miền, cửa hàng, tháng",
     "Giám đốc bán hàng", "1.000 đơn/tháng", "90% / 80%",
     "Một đơn có nhiều dòng: đếm dòng Line_No = 1. Không lọc theo ngành hàng được (một đơn có nhiều ngành)."],
    ["Gross Margin % (Biên lợi nhuận gộp)", "Mỗi đồng doanh thu còn lại bao nhiêu sau giá vốn?",
     "(SUM(Revenue) − SUM(COGS)) ÷ SUM(Revenue); COGS = Quantity × Unit_Cost", "tblSales + tblProduct (Unit_Cost qua XLOOKUP)",
     "KPI tỷ lệ, báo cáo theo tháng", "Ngành hàng, miền, tháng", "Finance (Tài chính)", "33%", "90% / 80% (29,7% / 26,4%)",
     "Là tỷ lệ: tổng ÷ tổng, không lấy trung bình các tỷ lệ. Giá vốn lấy một bảng Product duy nhất (giả định)."],
]

def kpi_card_sheet(wb):
    kc = wb.create_sheet("KPI_Card")
    kc["A1"] = "KPI DEFINITION CARD — 3 KPI của Buổi 2"; kc["A1"].font = TITLE
    kc["A2"] = "Viết card TRƯỚC khi viết công thức. Mẫu đầy đủ: ../B2_Ham_va_KPI/Handout/B2_KPI_Definition_Card.docx"; kc["A2"].font = NOTE
    header(kc, 4, ["Trường"] + [f"Card {i}" for i in (1, 2, 3)])
    for i, fld in enumerate(CARD_FIELDS, start=5):
        kc.cell(row=i, column=1, value=fld).font = BOLD
        for j, card in enumerate(CARDS, start=2):
            c = kc.cell(row=i, column=j, value=card[i - 5]); c.font = BODY; c.alignment = Alignment(wrap_text=True, vertical="top")
    kc.cell(row=16, column=1, value="Giá trị Q1/2025").font = BOLD
    for j, f in enumerate(["=SUM(tblSales[Revenue])", "=KPI_SUMIFS!E16", "=SUM(tblSales[Gross_Margin])/SUM(tblSales[Revenue])"], start=2):
        c = kc.cell(row=16, column=j, value=f); c.font = BOLD; c.number_format = PCT_FMT if j == 4 else NUM_FMT
    set_widths(kc, [26, 42, 42, 42])

# =====================================================================
# 1) PRACTICE — B3_MCI_Mart_Q1_2025_Answer.xlsx (brick output B3; file khởi đầu B4)
# =====================================================================
ANS = f"{PRACTICE}/B3_MCI_Mart_Q1_2025_Answer.xlsx"
wb = Workbook()
SALES_COLS = sales_sheet(wb, "=tblSales[[#This Row],[Revenue]]/Kiem_tra!$C$5", with_so_don=True)
ws = wb.create_sheet("Product"); table_df(ws, products, "tblProduct", {"Unit_Price": NUM_FMT, "Unit_Cost": NUM_FMT}, [13, 40, 18, 12, 12])
ws = wb.create_sheet("Store"); table_df(ws, stores, "tblStore", {"Open_Date": DATE_FMT}, [11, 24, 13, 9, 11, 12])
kpi_sumifs_sheet(wb)
kpi_card_sheet(wb)
PV_SHEETS = {
    "PV_MienThang": ("DOANH THU THEO MIỀN × THÁNG — PIVOT (đối chiếu với KPI_SUMIFS của Buổi 2)",
                     "Doanh thu từng miền, từng tháng Quý 1/2025 là bao nhiêu? Có khớp bảng SUMIFS Buổi 2 không?"),
    "PV_NganhHang": ("TỶ TRỌNG NGÀNH HÀNG TRONG TỪNG THÁNG (% of Column Total)",
                     "Trong doanh thu mỗi tháng, mỗi ngành hàng chiếm bao nhiêu phần trăm?"),
    "PV_Thang": ("THEO THÁNG: MoM, LŨY KẾ, AOV, GROSS MARGIN % (Show Values As + Calculated Field)",
                 "Tháng sau tăng/giảm bao nhiêu so với tháng trước? Mỗi đơn trung bình bao nhiêu tiền?"),
    "PV_Ngay": ("DOANH THU THEO NGÀY — QUÝ 1/2025", "Vì sao tháng 1 cao hơn hẳn tháng 2? Doanh thu thay đổi thế nào quanh Tết (29/01/2025)?"),
    "PV_Thu": ("DOANH THU THEO THỨ TRONG TUẦN", "Ngày nào trong tuần bán được nhiều nhất?"),
    "PV_CuaHang": ("DOANH THU THEO CỬA HÀNG × THÁNG (trái) · % SO VỚI THÁNG TRƯỚC (phải)",
                   "Cửa hàng nào giảm bất thường trong tháng 3? Double-click ô ST07 – tháng 3 để xem các dòng gốc."),
}
for nm, (t, qn) in PV_SHEETS.items():
    sheet_title(wb.create_sheet(nm), t, qn)
pvm = wb["PV_MienThang"]
pvm["A11"] = "ĐỐI SOÁT PIVOT ↔ SUMIFS (Buổi 2)"; pvm["A11"].font = BOLD
header(pvm, 12, ["Ô cần đối chiếu", "Pivot (GETPIVOTDATA)", "SUMIFS B2", "Chênh lệch"])
recon = [("Bắc – Tháng 3", '=GETPIVOTDATA("Doanh thu",$A$4,"Region","Bắc","Thang",3)', "=KPI_SUMIFS!D6"),
         ("Nam – Tháng 1", '=GETPIVOTDATA("Doanh thu",$A$4,"Region","Nam","Thang",1)', "=KPI_SUMIFS!B8"),
         ("Tổng Quý 1", '=GETPIVOTDATA("Doanh thu",$A$4)', "=KPI_SUMIFS!E9")]
for i, (lab, f1, f2) in enumerate(recon, start=13):
    pvm.cell(row=i, column=1, value=lab).font = BODY
    for j, f in ((2, f1), (3, f2), (4, f"=B{i}-C{i}")):
        c = pvm.cell(row=i, column=j, value=f); c.font = BOLD if j == 4 else BODY; c.number_format = NUM_FMT
pvm["A17"] = "Chênh lệch phải bằng 0. Bỏ chọn mọi mục trên Slicer trước khi đối chiếu (Slicer lọc cả Pivot)."; pvm["A17"].font = NOTE
set_widths(pvm, [18, 16, 16, 16, 16, 14, 3])
set_widths(wb["PV_CuaHang"], [12, 14, 14, 14, 14, 3, 3, 12, 11, 11, 11])
# Kết quả Show Details (double-click ô ST07 – tháng 3)
det = wb.create_sheet("ST07_T3_ChiTiet")
det["A1"] = "KẾT QUẢ SHOW DETAILS: double-click ô ST07 × Tháng 3 ở PV_CuaHang (Excel tạo sheet mới như thế này)"; det["A1"].font = TITLE
det["A2"] = (f"{len(st7)} dòng, {meta['st07']['t3_orders']} đơn. Từ 01 đến 09/03: {meta['st07']['before']['orders']} đơn trong 9 ngày; "
             f"từ 10/03: {meta['st07']['after']['orders']} đơn trong 22 ngày, {22 - meta['st07']['after']['days_sold']} ngày không có đơn nào.")
det["A2"].font = NOTE
dcols = ["Region", "Order_ID", "Line_No", "Order_Date", "Store_Code", "Product_Name", "Category", "Quantity", "Unit_Price", "Revenue"]
st7s = st7.sort_values(["Order_Date", "Order_ID", "Line_No"])[dcols]
header(det, 4, dcols); write_rows(det, st7s, 5, {"Order_Date": DATE_FMT, "Unit_Price": NUM_FMT, "Revenue": NUM_FMT})
t = Table(displayName="tblST07_T3", ref=f"A4:{get_column_letter(len(dcols))}{4 + len(st7s)}"); t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
det.add_table(t); set_widths(det, [8, 15, 7, 12, 10, 36, 16, 9, 11, 12])
for i, dd in enumerate(st7s.Order_Date, start=5):
    if dd >= pd.Timestamp("2025-03-10"):
        for j in range(1, len(dcols) + 1): det.cell(row=i, column=j).fill = PatternFill("solid", fgColor="FDE9E7")
# Kiem_tra
kt = wb.create_sheet("Kiem_tra")
kt["A1"] = "KIỂM TRA / ĐỐI SOÁT — Brick output B3"; kt["A1"].font = TITLE
kt["A2"] = ("Cột C = công thức (GETPIVOTDATA đọc thẳng từ Pivot); cột D = đáp án tính từ dataset; cột E = khớp? "
            "C5 là ô tổng doanh thu mà cột Pct_Share tham chiếu ($C$5). Bỏ chọn Slicer/Timeline trước khi kiểm tra.")
kt["A2"].font = NOTE
header(kt, 4, ["Chỉ tiêu", "", "Kết quả (công thức)", "Đáp án", "Khớp?"])
m = meta
checks = [
    ("Tổng doanh thu Q1/2025 (VNĐ)", "=SUM(tblSales[Revenue])", REV, NUM_FMT),
    ("Số dòng dữ liệu", "=ROWS(tblSales[Order_ID])", N, NUM_FMT),
    ("Số đơn = tổng cột So_don", "=SUM(tblSales[So_don])", ORDERS, NUM_FMT),
    ("Tổng Pivot Miền × Tháng = tổng cột Revenue", '=GETPIVOTDATA("Doanh thu",PV_MienThang!$A$4)', REV, NUM_FMT),
    ("Ô Pivot Bắc–T3 − ô SUMIFS Bắc–T3 (Buổi 2)", '=GETPIVOTDATA("Doanh thu",PV_MienThang!$A$4,"Region","Bắc","Thang",3)-KPI_SUMIFS!D6', 0, NUM_FMT),
    ("Tỷ trọng Bánh kẹo trong tháng 1 (% of Column Total)", '=GETPIVOTDATA("Tỷ trọng",PV_NganhHang!$A$4,"Category","Bánh kẹo","Thang",1)', m["cat_share"]["Bánh kẹo"][1], "0.0%"),
    ("MoM tháng 2 so với tháng 1 (% Difference From)", '=GETPIVOTDATA("MoM %",PV_Thang!$A$4,"Thang",2)', m["month"][2]["mom"], "0.0%"),
    ("Lũy kế đến tháng 2 (Running Total)", '=GETPIVOTDATA("Lũy kế",PV_Thang!$A$4,"Thang",2)', m["month"][2]["cum"], NUM_FMT),
    ("AOV Quý 1 (Calculated Field = Revenue / So_don)", '=GETPIVOTDATA("AOV (VNĐ/đơn)",PV_Thang!$A$4)', REV / ORDERS, NUM_FMT),
    ("Gross Margin % Quý 1 (Calculated Field)", '=GETPIVOTDATA("Gross Margin %",PV_Thang!$A$4)', GM / REV, "0.0%"),
    ("Doanh thu ST07 tháng 3", '=GETPIVOTDATA("Doanh thu",PV_CuaHang!$A$4,"Store_Code","ST07","Thang",3)', m["st07"]["t3_rev"], NUM_FMT),
    ("ST07 tháng 3 so với tháng 2 (% Difference From)", '=GETPIVOTDATA("% so tháng trước",PV_CuaHang!$H$4,"Store_Code","ST07","Thang",3)', m["st07"]["mom3"], "0.0%"),
    ("Dòng đầu Pivot theo Thứ (sắp xếp giảm dần)", "=PV_Thu!B5", m["thu"][m["thu_sorted"][0]], NUM_FMT),
    ("Số dòng Show Details ST07 – tháng 3", "=ROWS(tblST07_T3[Order_ID])", len(st7), NUM_FMT),
]
for r, (label, f, exp, fmt) in enumerate(checks, start=5):
    kt.cell(row=r, column=1, value=label).font = BODY
    c = kt.cell(row=r, column=3, value=f); c.number_format = fmt; c.font = BODY
    c = kt.cell(row=r, column=4, value=exp); c.number_format = fmt; c.font = INPUT
    kt.cell(row=r, column=5, value=f'=IF(ROUND(C{r},6)=ROUND(D{r},6),"✔","✘")').font = BODY
KT_ROWS = (5, 4 + len(checks))
set_widths(kt, [50, 2, 20, 16, 8])
wb.active = wb.sheetnames.index("PV_MienThang")
for w_ in wb.worksheets: w_.sheet_view.tabSelected = w_.title == "PV_MienThang"
wb.save(ANS)

ITEMS = {"Region": REGIONS, "Thang": MONTHS, "Category": CATS_Q1, "Store_Code": STORES}
c_ans = Cache("sales", "tblSales", SALES_COLS, calc_fields={"AOV": "=Revenue/So_don", "GM_pct": "=Gross_Margin/Revenue"},
              date_fields=("Order_Date",), items=ITEMS)
ANS_PIVOTS = [
    Pivot("PV_MienThang", "A4", "pvtMienThang", "sales", rows=["Region"], cols=["Thang"], values=[Value("Revenue", "Doanh thu")]),
    Pivot("PV_NganhHang", "A4", "pvtNganhHang", "sales", rows=["Category"], cols=["Thang"], values=[Value("Revenue", "Tỷ trọng", fmt="0.0%", show_as="pct_col")]),
    Pivot("PV_Thang", "A4", "pvtThang", "sales", rows=["Thang"],
          values=[Value("Revenue", "Doanh thu"), Value("Revenue", "MoM %", fmt="0.0%", show_as="pct_diff", base_field="Thang"),
                  Value("Revenue", "Lũy kế", show_as="run_total", base_field="Thang"), Value("So_don", "Số đơn"),
                  Value("AOV", "AOV (VNĐ/đơn)"), Value("GM_pct", "Gross Margin %", fmt="0.0%")]),
    Pivot("PV_Ngay", "A4", "pvtNgay", "sales", rows=["Order_Date"], values=[Value("Revenue", "Doanh thu")]),
    Pivot("PV_Thu", "A4", "pvtThu", "sales", rows=["Thu"], values=[Value("Revenue", "Doanh thu")], sort_desc_by="Doanh thu"),
    Pivot("PV_CuaHang", "A4", "pvtCuaHang", "sales", rows=["Store_Code"], cols=["Thang"], values=[Value("Revenue", "Doanh thu")]),
    Pivot("PV_CuaHang", "H4", "pvtCuaHangPct", "sales", rows=["Store_Code"], cols=["Thang"],
          values=[Value("Revenue", "% so tháng trước", fmt="0.0%", show_as="pct_diff", base_field="Thang")], total_col=False),
]
ANS_SLICERS = [Slicer("PV_MienThang", "Store_Code", ["pvtMienThang", "pvtNganhHang", "pvtThang"], "G4", "I20", columns=2)]
ANS_TIMELINES = [Timeline("PV_Ngay", "Order_Date", ["pvtNgay", "pvtThu"], "D4", "N11")]
tt = m["tet"]
ANS_CHARTS = [
    Chart("PV_Ngay", "pvtNgay", "line", f"3 tuần trước Tết mỗi ngày bán ~{dec(tt['truoc_tet']['per_day'] / 1e6)} triệu, nghỉ Tết còn ~{tt['nghi_tet']['per_day'] / 1e6:.0f} triệu",
          "D13", "P33", "PV_Ngay!$A$5:$A$94", [("PV_Ngay!$B$4", "PV_Ngay!$B$5:$B$94")]),
    Chart("PV_Thu", "pvtThu", "col", f"Thứ Bảy, Chủ nhật bán nhiều hơn ngày thường khoảng {m['weekend']['ratio'] * 100 - 100:.0f}%",
          "D4", "L22", "PV_Thu!$A$5:$A$11", [("PV_Thu!$B$4", "PV_Thu!$B$5:$B$11")], labels=True),
    Chart("PV_NganhHang", "pvtNganhHang", "col", f"Bánh kẹo chiếm {pct(m['cat_share']['Bánh kẹo'][1])} doanh thu tháng 1, còn {pct(m['cat_share']['Bánh kẹo'][2])} tháng 2",
          "G4", "Q22", "PV_NganhHang!$A$6:$A$11", [(f"PV_NganhHang!${c}$5", f"PV_NganhHang!${c}$6:${c}$11") for c in "BCD"], legend=True, val_fmt="0%"),
]

# =====================================================================
# 2) DEMO
# =====================================================================
# ---- B3_Demo_Pivot.xlsx — file demo chính: dựng Pivot trực tiếp trên tblSales (giống file khởi đầu)
wb = Workbook()
sales_sheet(wb, "=tblSales[[#This Row],[Revenue]]/KPI_SUMIFS!$E$9", with_so_don=False)
ws = wb.create_sheet("Product"); table_df(ws, products, "tblProduct", {"Unit_Price": NUM_FMT, "Unit_Cost": NUM_FMT}, [13, 40, 18, 12, 12])
ws = wb.create_sheet("Store"); table_df(ws, stores, "tblStore", {"Open_Date": DATE_FMT}, [11, 24, 13, 9, 11, 12])
kpi_sumifs_sheet(wb, blocks=("rev", "orders"))
wk_ = m["weekly"]; s7 = m["st07"]; s4 = m["st04"]; we = m["weekend"]; tp5 = m["top5"]; cs_ = m["cat_share"]
g = wb.create_sheet("Huong_dan")
NEWPV = "Chọn một ô trong tblSales (sheet Sales) → Insert → PivotTable → From Table/Range → New Worksheet; đổi tên sheet mới thành "
guide(g, "DEMO B3 — DỰNG VÀ PHÂN TÍCH BẰNG PIVOTTABLE (slide 9–29)", [
    (9, NEWPV + "PV_MienThang.", "Hộp Create PivotTable tự ghi Table/Range: tblSales. Sheet mới có Pivot trống (ô A3) và khung PivotTable Fields (20 trường)."),
    (10, "Kéo Region vào Rows, Thang vào Columns, Revenue vào Values.",
     f"Bảng Miền × Tháng: Bắc–T1 {vn(m['grid']['Bắc-1'])} … Grand Total {vn(REV)}; hàng tổng T1 {vn(m['month'][1]['rev'])}, T2 {vn(m['month'][2]['rev'])}, T3 {vn(m['month'][3]['rev'])}."),
    (11, "View → New Window → Arrange → Vertical: một cửa sổ sheet PV_MienThang, một cửa sổ sheet KPI_SUMIFS. So ô Bắc × 3 của Pivot với ô D6.",
     f"Cả hai = {vn(m['grid']['Bắc-3'])}: tọa độ (Region = Bắc, Thang = 3) chính là hai tiêu chí của SUMIFS."),
    (11, "Double-click ô Bắc × 3 trong Pivot (Show Details). Xem xong thì xóa sheet chi tiết.", f"Sheet mới có {m['grid']['bac_t3_lines']} dòng, đúng các dòng Bắc tháng 3."),
    (11, "Ô G4 của PV_MienThang: gõ = rồi bấm vào ô Bắc × 3 của Pivot, Enter.",
     "Excel tự viết =GETPIVOTDATA(…,$A$3,\"Region\",\"Bắc\",\"Thang\",3). Tắt: PivotTable Analyze → Options ▾ → bỏ tick Generate GetPivotData."),
    (14, "Chuột phải một số trong Pivot → Value Field Settings → Custom Name: Doanh thu → Number Format → Number, 0 chữ số thập phân, tick Use 1000 Separator.",
     "Số hiện dạng 145.814.000 (hoặc 145,814,000 tùy cài đặt vùng); tiêu đề đổi thành Doanh thu."),
    (14, "PivotTable Design → PivotTable Styles: chọn một style; Grand Totals / Subtotals để bật tắt dòng tổng.", "Định dạng giữ nguyên khi Refresh."),
    (15, "Kéo Store_Code vào Rows dưới Region. PivotTable Design → Report Layout: lần lượt Compact, Outline, Tabular; thử Repeat All Item Labels. Xong thì bỏ Store_Code, chọn lại Outline.",
     "Compact: một cột nhãn thụt lề. Outline: mỗi trường một cột, tổng phụ ở trên. Tabular: mỗi trường một cột, tổng phụ ở dưới."),
    (18, NEWPV + "PV_Thu. Thu vào Rows, Revenue vào Values. Chuột phải một số → Sort → Sort Largest to Smallest.",
     "Thứ tự: " + ", ".join(f"{d} {vn(m['thu'][d])}" for d in m["thu_sorted"]) + " (máy tiếng Việt hiện T2…CN). "
     f"Trung bình T7–CN {vn(we['avg_weekend'])}, T2–T6 {vn(we['avg_weekday'])}: cao hơn khoảng {we['ratio'] * 100 - 100:.0f}%."),
    (18, NEWPV + "PV_Top5. Product_Name vào Rows, Revenue vào Values. Mũi tên ở ô tiêu đề hàng → Value Filters → Top 10 → sửa thành 5 → OK.",
     "Top 5: " + "; ".join(f"{n} {vn(v)}" for _, n, v in tp5) + "."),
    (19, NEWPV + "PV_Tuan. Order_Date vào Rows (Excel 365 tự nhóm theo tháng), Revenue vào Values. Chuột phải một ngày → Group → bỏ chọn Months, chọn Days, Number of days: 7 → OK.",
     "13 nhóm 7 ngày từ 01/01: " + "; ".join(f"{w['from']}–{w['to']} {vn(w['rev'])}" for w in wk_[:6]) + " …; các tuần sau "
     f"{vn(tt['normal_weeks_min'])}–{vn(tt['normal_weeks_max'])}."),
    (19, "Pivot Store_Code × Revenue (có thể dùng lại PV_MienThang): giữ Ctrl chọn ST08, ST09, ST10 → chuột phải → Group; đổi tên Group1 thành TP.HCM. Xong thì Ungroup.",
     f"Xuất hiện trường Store_Code2; nhóm TP.HCM = {vn(m['hcm'])}."),
    (20, NEWPV + "PV_CuaHang. Store_Code vào Rows, Thang vào Columns, Revenue vào Values. Double-click ô ST07 × 3, sắp xếp Order_Date tăng dần.",
     f"ST07 tháng 3 = {vn(s7['t3_rev'])} ({spct(s7['mom3'])} so với tháng 2). Sheet chi tiết có {s7['t3_rows']} dòng ({s7['t3_orders']} đơn): 01–09/03 có {s7['before']['orders']} đơn; "
     f"từ 10/03 chỉ {s7['after']['orders']} đơn trong 22 ngày ({22 - s7['after']['days_sold']} ngày không có đơn). Tháng 2: {s7['t2_orders']} đơn."),
    (20, "Xóa sheet chi tiết. Double-click ô ST04 × 3 để so sánh, rồi xóa sheet chi tiết.",
     f"ST04 tháng 3 ({spct(s4['mom3'])}) có {s4['t3_orders']} đơn rải đều cả tháng ({s4['days_sold']} ngày có bán): không có quãng trống như ST07."),
    (21, "Sheet PV_Thu → PivotTable Analyze → Insert Slicer → tick Region. Chuột phải Slicer → Report Connections → tick thêm Pivot của sheet PV_CuaHang.",
     "Bấm Nam: cả hai Pivot chỉ còn số của Miền Nam. Bỏ lọc: nút phễu có dấu X."),
    (21, "PivotTable Analyze → Insert Timeline → tick Order_Date. Kéo chọn tháng 2 và tháng 3.",
     "Pivot Thu chỉ còn tháng 2–3: " + ", ".join(f"{d} {vn(m['thu_t2t3'][d])}" for d in m["thu_t2t3_sorted"][:3]) + ": bỏ tháng Tết, cuối tuần vẫn dẫn đầu. Bỏ lọc Timeline."),
    (23, "Sheet PV_MienThang: chuột phải một số → Summarize Values By → Count, rồi Average, rồi Max; cuối cùng trả về Sum.",
     f"Count = {vn(N)} dòng (không phải {vn(ORDERS)} đơn). Average = {vn(REV / N)} / dòng. Max = {vn(m['q1']['max_line'])}."),
    (24, NEWPV + "PV_NganhHang. Category vào Rows, Thang vào Columns, Revenue vào Values. Chuột phải một số → Show Values As → % of Column Total.",
     "Bánh kẹo: " + " · ".join(f"T{k} {pct(cs_['Bánh kẹo'][k])}" for k in MONTHS) + "; Đồ uống: " + " · ".join(f"T{k} {pct(cs_['Đồ uống'][k])}" for k in MONTHS) + "."),
    (24, "Tạm đổi: bỏ Thang khỏi Columns, thay Category bằng Region → Show Values As → % of Grand Total. Xong thì trả lại Category × Thang, % of Column Total.",
     "Bắc " + pct(m["region"]["Bắc"]["share"]) + " · Trung " + pct(m["region"]["Trung"]["share"]) + " · Nam " + pct(m["region"]["Nam"]["share"]) + "."),
    (25, NEWPV + "PV_Thang. Thang vào Rows, kéo Revenue vào Values 3 lần. Cột thứ 2: Show Values As → % Difference From → Base field Thang, Base item (previous). Cột thứ 3: Running Total In → Thang.",
     f"MoM: T2 {spct(m['month'][2]['mom'])}, T3 {spct(m['month'][3]['mom'])}. Lũy kế: {vn(m['month'][1]['cum'])} · {vn(m['month'][2]['cum'])} · {vn(m['month'][3]['cum'])}."),
    (26, "Sheet Sales: thêm cột So_don (ô U1), ô U2 gõ =IF([@Line_No]=1,1,0). Quay lại PV_Thang, Refresh (Alt+F5), kéo So_don vào Values.",
     f"Sum of So_don = {vn(ORDERS)} đơn (Count of Order_ID vẫn ra {vn(N)} dòng)."),
    (26, "PivotTable Analyze → Fields, Items & Sets → Calculated Field. Name: AOV, Formula: =Revenue/So_don → Add → OK. Làm tiếp GM_pct = Gross_Margin/Revenue, định dạng %.",
     f"AOV: T1 {vn(m['month'][1]['aov'])} · T2 {vn(m['month'][2]['aov'])} · T3 {vn(m['month'][3]['aov'])} · Grand Total {vn(REV / ORDERS)} (tổng ÷ tổng). GM_pct: {pct(GM / REV)}."),
    (28, "Sheet PV_Thang: chọn Pivot → PivotTable Analyze → PivotChart → Pie. Sau đó Change Chart Type → Line.",
     "Pie 3 miếng không cho thấy xu hướng; Line cho thấy T1 cao, T2 sụt, T3 nhích lên."),
    (29, NEWPV + "PV_Ngay. Order_Date vào Rows → chuột phải → Group → chỉ chọn Days (Number of days: 1); Revenue vào Values; PivotTable Analyze → PivotChart → Line. Gõ tiêu đề là câu kết luận.",
     f"Đường tăng từ 09/01 (~{dec(tt['truoc_tet']['per_day'] / 1e6)} triệu/ngày), sụt còn ~{dec(tt['nghi_tet']['per_day'] / 1e6)} triệu/ngày 29/01–02/02, sau đó quanh {dec(tt['sau_tet']['per_day'] / 1e6)} triệu/ngày."),
])
DEMO_PIVOT = save_both(wb, "B3_Demo_Pivot", "Sales")

# ---- B3_Demo_Refresh.xlsx — slide 12: Pivot đọc từ cache; nguồn là vùng thường vs Table
t4_all = m_sales[m_sales.Order_ID.str.startswith("HD2504")].copy()
t4_all["Region"] = t4_all.Store_Code.map(region_of); t4_all["Category"] = t4_all.Product_Code.map(cat)
add_rows = t4_all.sort_values(["Order_Date", "Order_ID", "Line_No"]).head(10)
RC = ["Region", "Order_ID", "Line_No", "Order_Date", "Store_Code", "Category", "Quantity", "Unit_Price"]
src = q1[RC].copy()
wb = Workbook(); ws = wb.active; ws.title = "Sales_Vung"
header(ws, 1, RC + ["Revenue"]); write_rows(ws, src, 2, {"Order_Date": DATE_FMT, "Unit_Price": NUM_FMT})
for i in range(2, N + 2):
    c = ws.cell(row=i, column=9, value=f"=G{i}*H{i}"); c.font = BODY; c.number_format = NUM_FMT
ws.freeze_panes = "A2"; set_widths(ws, [9, 15, 7, 12, 10, 16, 9, 11, 12])
ws2 = wb.create_sheet("Sales_Table")
header(ws2, 1, RC + ["Revenue"]); write_rows(ws2, src, 2, {"Order_Date": DATE_FMT, "Unit_Price": NUM_FMT})
for i in range(2, N + 2):
    c = ws2.cell(row=i, column=9, value="=tblSales[[#This Row],[Quantity]]*tblSales[[#This Row],[Unit_Price]]"); c.font = BODY; c.number_format = NUM_FMT
add_table(ws2, "tblSales", 9, N); ws2.freeze_panes = "A2"; set_widths(ws2, [9, 15, 7, 12, 10, 16, 9, 11, 12])
for nm, t_ in (("PV_Vung", "PIVOT TỪ VÙNG THƯỜNG — nguồn: Sales_Vung!$A$1:$I$6188 (cố định)"), ("PV_Table", "PIVOT TỪ EXCEL TABLE — nguồn: tblSales (tự mở rộng)")):
    sheet_title(wb.create_sheet(nm), t_)
td = wb.create_sheet("Them_dong")
td["A1"] = "10 DÒNG THÁNG 4/2025 ĐỂ DÁN THÊM — copy A4:H13, dán vào dòng trống đầu tiên dưới dữ liệu của Sales_Vung và Sales_Table"; td["A1"].font = BOLD
header(td, 3, RC); write_rows(td, add_rows[RC], 4, {"Order_Date": DATE_FMT, "Unit_Price": NUM_FMT})
set_widths(td, [9, 15, 7, 12, 10, 16, 9, 11])
first = q1.iloc[0]
add_rev = int((add_rows.Quantity * add_rows.Unit_Price).sum())
add_by_reg = (add_rows.assign(r=add_rows.Quantity * add_rows.Unit_Price).groupby("Region").r.sum()).to_dict()
meta["refresh"] = {"row2": {"order": first.Order_ID, "product": first.Product_Code, "qty": int(first.Quantity), "price": int(first.Unit_Price), "region": first.Region},
                   "bump": 100 * int(first.Unit_Price), "add_rows": len(add_rows), "add_rev": add_rev,
                   "add_by_region": {k: int(v) for k, v in add_by_reg.items()}}
mr = meta["refresh"]
g = wb.create_sheet("Huong_dan")
guide(g, "DEMO B3 — PIVOT ĐỌC TỪ CACHE: REFRESH, NGUỒN LÀ VÙNG THƯỜNG HAY TABLE (slide 12)", [
    (12, "Mở PV_Table và PV_Vung: hai Pivot Region × Revenue dựng từ cùng một dữ liệu, một từ vùng thường, một từ tblSales.",
     f"Cả hai: Bắc {vn(meta['region']['Bắc']['rev'])} · Trung {vn(meta['region']['Trung']['rev'])} · Nam {vn(meta['region']['Nam']['rev'])} · tổng {vn(REV)}."),
    (12, f"Sheet Sales_Table, ô G2 (Quantity của {mr['row2']['order']}, {mr['row2']['product']}): đổi {mr['row2']['qty']} thành {mr['row2']['qty'] + 100}. Quay lại PV_Table.",
     "Pivot KHÔNG đổi: Pivot đọc từ bộ nhớ đệm (cache), không đọc trực tiếp ô."),
    (12, "Chuột phải trong PV_Table → Refresh (Alt+F5).",
     f"Bắc tăng thêm {vn(mr['bump'])} thành {vn(meta['region']['Bắc']['rev'] + mr['bump'])}. Ctrl+Z ở Sales_Table để trả lại {mr['row2']['qty']}, Refresh lại."),
    (12, "Sheet Them_dong: copy A4:H13. Dán vào ô A6189 của Sales_Table (dòng trống ngay dưới bảng) và ô A6189 của Sales_Vung. Ở Sales_Vung kéo công thức I6188 xuống I6198.",
     "Sales_Table tự mở rộng thành 6.197 dòng, cột Revenue tự điền công thức. Sales_Vung có thêm 10 dòng nhưng không có gì thay đổi về nguồn Pivot."),
    (12, "Data → Refresh All (Ctrl+Alt+F5).",
     f"PV_Table: tổng {vn(REV + mr['add_rev'])} (thêm {vn(mr['add_rev'])} của {mr['add_rows']} dòng tháng 4). PV_Vung: vẫn {vn(REV)} vì nguồn cố định $A$1:$I$6188."),
    (12, "Chọn PV_Vung → PivotTable Analyze → Change Data Source → sửa thành Sales_Vung!$A$1:$I$6198 → OK.",
     f"PV_Vung = {vn(REV + mr['add_rev'])}. Với Table thì không phải làm bước này mỗi lần có dữ liệu mới."),
])
DEMO_REFRESH = save_both(wb, "B3_Demo_Refresh", "PV_Table")
c_vung = Cache("vung", None, RC + ["Revenue"], items={"Region": REGIONS}, sheet="Sales_Vung", ref=f"A1:I{N + 1}")
c_tab = Cache("tab", "tblSales", RC + ["Revenue"], items={"Region": REGIONS})
REFRESH_PIVOTS = [Pivot("PV_Vung", "A3", "pvtVung", "vung", rows=["Region"], values=[Value("Revenue", "Doanh thu")]),
                  Pivot("PV_Table", "A3", "pvtTable", "tab", rows=["Region"], values=[Value("Revenue", "Doanh thu")])]

# ---- B3_Demo_Loi_Du_lieu.xlsx — slide 19 (Cannot group that selection), 23 (Pivot mặc định Count)
rnd = random.Random(33)
LC = ["Region", "Order_ID", "Line_No", "Order_Date", "Store_Code", "Category", "Quantity", "Unit_Price", "Revenue"]
lo = q1[LC].copy().astype(object)
idx_text_rev = sorted(rnd.sample(range(N), 2)); idx_blank_rev = sorted(rnd.sample([i for i in range(N) if i not in idx_text_rev], 1))
idx_text_date = sorted(rnd.sample([i for i in range(N) if i not in idx_text_rev + idx_blank_rev], 3))
bad_rev_text = {i: int(q1.Revenue.iloc[i]) for i in idx_text_rev}
bad_rev_blank = {i: int(q1.Revenue.iloc[i]) for i in idx_blank_rev}
for i in idx_text_rev: lo.iat[i, LC.index("Revenue")] = str(int(q1.Revenue.iloc[i]))
for i in idx_blank_rev: lo.iat[i, LC.index("Revenue")] = None
for i in idx_text_date: lo.iat[i, LC.index("Order_Date")] = q1.Order_Date.iloc[i].strftime("%d/%m/%Y")
wb = Workbook(); ws = wb.active; ws.title = "Sales_Loi"
header(ws, 1, LC)
for i, row in enumerate(lo.itertuples(index=False), start=2):
    for j, v in enumerate(row, start=1):
        c = ws.cell(row=i, column=j, value=py(v) if not isinstance(v, str) else v); c.font = BODY
        if j == 4 and not isinstance(v, str): c.number_format = DATE_FMT
        if j in (8, 9) and not isinstance(v, str): c.number_format = NUM_FMT
add_table(ws, "tblSalesLoi", len(LC), N); ws.freeze_panes = "A2"; set_widths(ws, [9, 15, 7, 12, 10, 16, 9, 11, 12])
lost = sum(bad_rev_text.values()) + sum(bad_rev_blank.values())
meta["loi"] = {"text_rev_rows": [i + 2 for i in idx_text_rev], "text_rev_values": list(bad_rev_text.values()),
               "blank_rev_row": [i + 2 for i in idx_blank_rev], "blank_rev_value": list(bad_rev_blank.values()),
               "text_date_rows": [i + 2 for i in idx_text_date], "text_dates": [q1.Order_Date.iloc[i].strftime("%d/%m/%Y") for i in idx_text_date],
               "count_default": N - 1, "sum_before_fix": REV - lost, "lost": lost}
ml = meta["loi"]
g = wb.create_sheet("Huong_dan")
guide(g, "DEMO B3 — HAI LỖI DỮ LIỆU LÀM PIVOT SAI (slide 19, 23)", [
    (23, "Chọn một ô trong tblSalesLoi → Insert → PivotTable → New Worksheet. Kéo Region vào Rows, Revenue vào Values.",
     f"Pivot hiện “Count of Revenue” (không phải Sum): cột Revenue có ô chữ và ô trống. Tổng Count = {vn(ml['count_default'])}."),
    (23, "Chuột phải một số → Summarize Values By → Sum.",
     f"Tổng = {vn(ml['sum_before_fix'])}, thiếu {vn(ml['lost'])} so với {vn(REV)}: Sum bỏ qua ô chữ và ô trống, không báo lỗi."),
    (23, f"Tìm ô lỗi: sheet Sales_Loi, lọc cột Revenue. Ô chữ ở dòng {', '.join(map(str, ml['text_rev_rows']))} (căn trái, tam giác xanh) → Convert to Number. "
         f"Ô trống ở dòng {ml['blank_rev_row'][0]} → gõ lại = Quantity × Unit_Price. Quay lại Pivot, Refresh.",
     f"Tổng = {vn(REV)}. Kéo Revenue vào Values lần nữa: Excel mặc định Sum."),
    (19, "Pivot mới từ tblSalesLoi: Order_Date vào Rows. Chuột phải một ngày → Group.",
     "Excel báo “Cannot group that selection.” Đầu danh sách ngày có 3 dòng chữ, đứng riêng trước các ngày thật: " + ", ".join(sorted(ml["text_dates"])) + "."),   # Excel 365 Windows xếp mục chữ trước ngày
    (19, f"Tìm ô ngày dạng chữ: lọc cột Order_Date (các giá trị chữ nằm cuối danh sách lọc), dòng {', '.join(map(str, ml['text_date_rows']))}. Gõ lại thành ngày thật (hoặc =DATEVALUE(…)). Refresh.",
     "Group chạy được: chọn Months hoặc Days."),
])
DEMO_LOI = save_both(wb, "B3_Demo_Loi_Du_lieu", "Sales_Loi")

# =====================================================================
# 3) HOMEWORK — B3_BTVN_MCI_Mart_T1_T4_2025.xlsx (+ _Answer)
# =====================================================================
t4 = m_sales[m_sales.Order_ID.str.startswith("HD2504")].copy()
t4["Product_Name"] = t4.Product_Code.map(products.set_index("Product_Code")["Product_Name"])
t4["Region"] = t4.Store_Code.map(region_of)
t4 = t4.sort_values(["Store_Code", "Order_Date", "Order_ID", "Line_No"]).reset_index(drop=True)
HC = ["Region", "Order_ID", "Line_No", "Order_Date", "Store_Code", "Customer_ID", "Product_Code", "Product_Name", "Category", "Quantity", "Unit_Price", "Unit_Cost"]
h4 = pd.concat([b1.assign(Category=b1.Product_Code.map(cat)), t4.assign(Category=t4.Product_Code.map(cat))], ignore_index=True)
h4["Unit_Cost"] = h4.Product_Code.map(cost)
h4 = h4[HC].copy()
NH = len(h4)
h4["Revenue"] = h4.Quantity * h4.Unit_Price; h4["COGS"] = h4.Quantity * h4.Unit_Cost
h4["Thang"] = h4.Order_Date.dt.month; h4["So_don"] = (h4.Line_No == 1).astype(int)
rnd = random.Random(43)
apr_idx = list(h4.index[h4.Thang == 4])
qty_text = sorted(rnd.sample([i for i in apr_idx if h4.Quantity[i] >= 2], 6))
HF = ["Revenue", "COGS", "Gross_Margin", "Thang", "Thu", "So_don"]
TT = lambda c_: f"tblSales4T[[#This Row],[{c_}]]"
HFORM = {"Revenue": (f"={TT('Quantity')}*{TT('Unit_Price')}", NUM_FMT), "COGS": (f"={TT('Quantity')}*{TT('Unit_Cost')}", NUM_FMT),
         "Gross_Margin": (f"={TT('Revenue')}-{TT('COGS')}", NUM_FMT), "Thang": (f"=MONTH({TT('Order_Date')})", "0"),
         "Thu": (f'=TEXT({TT("Order_Date")},"ddd")', None), "So_don": (f"=IF({TT('Line_No')}=1,1,0)", "0")}
bym4 = h4.groupby("Thang").agg(rev=("Revenue", "sum"), rows=("Revenue", "size"), orders=("So_don", "sum"), qty=("Quantity", "sum"), cogs=("COGS", "sum"))
reg4 = h4.pivot_table(index="Region", columns="Thang", values="Revenue", aggfunc="sum")
reg_mom4 = (reg4[4] / reg4[3] - 1)
cat4 = h4.pivot_table(index="Category", columns="Thang", values="Revenue", aggfunc="sum"); cat4s = cat4 / cat4.sum()
st4m = h4.pivot_table(index="Store_Code", columns="Thang", values="Revenue", aggfunc="sum")
st07_4 = h4[(h4.Store_Code == "ST07") & (h4.Thang == 4)]
qty_lost = int(h4.Quantity[qty_text].sum())
meta["hw"] = {"rows": NH, "revenue": int(h4.Revenue.sum()), "by_month": {int(k): {"rev": int(r.rev), "rows": int(r.rows), "orders": int(r.orders), "qty": int(r.qty),
                                                                                      "aov": r.rev / r.orders, "gm_pct": 1 - r.cogs / r.rev} for k, r in bym4.iterrows()},
              "mom4": bym4.rev[4] / bym4.rev[3] - 1, "reg_mom4": {r: float(reg_mom4[r]) for r in REGIONS},
              "best_region": reg_mom4.idxmax(), "banh_keo_share": {int(k): float(cat4s.loc["Bánh kẹo", k]) for k in cat4s.columns},
              "st07": {int(k): int(st4m.loc["ST07", k]) for k in st4m.columns}, "st07_mom4": float(st4m.loc["ST07", 4] / st4m.loc["ST07", 3] - 1),
              "st07_t4_orders": int(st07_4.So_don.sum()), "st07_t4_days": st07_4.Order_Date.nunique(),
              "qty_text_rows": [i + 2 for i in qty_text], "qty_text_values": [int(h4.Quantity[i]) for i in qty_text], "qty_lost": qty_lost,
              "qty_t4_wrong": int(bym4.qty[4]) - qty_lost, "count_qty_default": NH}
mh = meta["hw"]

def hw_sales(wb, dirty):
    ws = wb.active; ws.title = "Sales"
    header(ws, 1, HC + HF)
    for i, row in enumerate(h4[HC].itertuples(index=False), start=2):
        for j, v in enumerate(row, start=1):
            val = py(v)
            if dirty and j == HC.index("Quantity") + 1 and (i - 2) in qty_text: val = str(val)
            c = ws.cell(row=i, column=j, value=val); c.font = BODY
            if j == 4: c.number_format = DATE_FMT
            if j in (11, 12) and not isinstance(val, str): c.number_format = NUM_FMT
        for j, col in enumerate(HF, start=len(HC) + 1):
            f, fmt = HFORM[col]
            c = ws.cell(row=i, column=j, value=f); c.font = BODY
            if fmt: c.number_format = fmt
    add_table(ws, "tblSales4T", len(HC) + len(HF), NH); ws.freeze_panes = "A2"
    set_widths(ws, [9, 15, 7, 12, 10, 11, 11, 36, 16, 9, 11, 11, 12, 12, 13, 7, 6, 8])
    return ws

def control_sheet(wb):
    cs_ = wb.create_sheet("So_lieu_he_thong")
    cs_["A1"] = "SỐ LIỆU ĐỐI SOÁT TỪ HỆ THỐNG BÁN HÀNG (tổng theo tháng, xuất cùng file dữ liệu)"; cs_["A1"].font = TITLE
    cs_["A2"] = "Pivot của bạn phải khớp đúng các con số này trước khi phân tích tiếp."; cs_["A2"].font = NOTE
    header(cs_, 4, ["Tháng", "Số dòng", "Số đơn", "Tổng Quantity", "Doanh thu (VNĐ)"])
    for i, (k, r) in enumerate(bym4.iterrows(), start=5):
        for j, v in enumerate([int(k), int(r.rows), int(r.orders), int(r.qty), int(r.rev)], start=1):
            c = cs_.cell(row=i, column=j, value=v); c.font = BODY; c.number_format = MONTH_FMT if j == 1 else NUM_FMT
    i = 5 + len(bym4)
    cs_.cell(row=i, column=1, value="Tổng").font = BOLD
    for j, v in enumerate([NH, int(bym4.orders.sum()), int(bym4.qty.sum()), int(bym4.rev.sum())], start=2):
        c = cs_.cell(row=i, column=j, value=v); c.font = BOLD; c.number_format = NUM_FMT
    set_widths(cs_, [10, 12, 12, 16, 18])

wb = Workbook(); hw_sales(wb, dirty=True)
control_sheet(wb)
s_ = wb.create_sheet("Store"); table_df(s_, stores, "tblStore", {"Open_Date": DATE_FMT}, [11, 24, 13, 9, 11, 12])
de = wb.create_sheet("De_bai", 0)
tasks = [
    ("1", "Đối soát trước khi phân tích", "Dựng Pivot: Thang vào Rows; Revenue, So_don, Quantity vào Values (Sum). So từng tháng với sheet So_lieu_he_thong. "
     "Có cột nào lệch? Tìm nguyên nhân trong dữ liệu, sửa ở tblSales4T, Refresh cho đến khi khớp. Ghi lại lỗi đã sửa vào sheet Insight."),
    ("2", "Miền nào tăng mạnh nhất?", "Pivot Region × Thang, Revenue. Thêm Revenue lần hai, Show Values As → % Difference From (Thang, previous). "
     "Trả lời: miền nào tăng mạnh nhất tháng 4 so với tháng 3?"),
    ("3", "Cơ cấu ngành hàng", "Pivot Category × Thang, Show Values As → % of Column Total. Tỷ trọng Bánh kẹo thay đổi thế nào từ tháng 1 đến tháng 4?"),
    ("4", "ST07 đã hồi phục chưa?", "Pivot Store_Code × Thang. Thêm cột % so với tháng trước. Double-click ô ST07 × 4 và ST07 × 3 để so số đơn, số ngày có bán."),
    ("5", "AOV bằng Calculated Field", "Pivot theo Thang: Calculated Field AOV = Revenue / So_don. So AOV tháng 4 với Quý 1."),
    ("6", "PivotChart + Slicer", "Mỗi câu 2–4 có một PivotChart, tiêu đề là câu kết luận. Chèn Slicer Region, nối (Report Connections) với các Pivot của câu 2–4."),
    ("7", "Viết insight", "Sheet Insight: 1 câu có đủ con số (có so sánh), nguyên nhân khả dĩ, hàm ý."),
]
de["A1"] = "BÀI TẬP VỀ NHÀ BUỔI 3 (BẮT BUỘC) — BÁO CÁO PIVOT 4 THÁNG ĐẦU NĂM 2025"; de["A1"].font = TITLE
de["A2"] = ("Nộp trước Buổi 4. Đặt tên file: B3_BTVN_<HoTen>.xlsx. Dữ liệu tháng 1–4/2025 đã làm giàu (tblSales4T, "
            f"{vn(NH)} dòng). Phòng IT vừa ghép thêm tháng 4 vào, hãy kiểm tra trước khi tin các con số.")
de["A2"].font = NOTE
de["A3"] = "Nộp bài tại https://hocvien.mcivietnam.com/auth#/ — Tài khoản: email bạn đăng ký tại MCI · Mật khẩu: hv12345678"; de["A3"].font = BOLD
header(de, 5, ["#", "Việc cần làm", "Yêu cầu"])
for i, row in enumerate(tasks, start=6):
    for j, val in enumerate(row, start=1):
        c = de.cell(row=i, column=j, value=val); c.font = BODY; c.alignment = Alignment(wrap_text=True, vertical="top")
kr = 6 + len(tasks) + 1
de.cell(row=kr, column=1, value="Tự kiểm tra trước khi nộp").font = BOLD
for i, tx in enumerate([f"Pivot theo tháng khớp So_lieu_he_thong ở cả 3 cột Revenue, So_don, Quantity (tổng Quantity tháng 4 = {vn(mh['by_month'][4]['qty'])})",
                        f"Tổng doanh thu 4 tháng = {vn(mh['revenue'])} VNĐ; tháng 4 = {vn(mh['by_month'][4]['rev'])} VNĐ",
                        f"Tháng 4 so với tháng 3: {spct(mh['mom4'])}",
                        "Mỗi PivotChart có tiêu đề là câu kết luận; Slicer Region lọc được các Pivot của câu 2–4"], start=kr + 1):
    de.cell(row=i, column=2, value="☐ " + tx).font = BODY
set_widths(de, [5, 34, 100])
wb.active = 0
for w_ in wb.worksheets: w_.sheet_view.tabSelected = w_.title == "De_bai"
HW_DE = f"{HW}/B3_BTVN_MCI_Mart_T1_T4_2025.xlsx"
wb.save(HW_DE)

# --- Đáp án
wb = Workbook(); hw_sales(wb, dirty=False)
control_sheet(wb)
s_ = wb.create_sheet("Store"); table_df(s_, stores, "tblStore", {"Open_Date": DATE_FMT}, [11, 24, 13, 9, 11, 12])
HWPV = {"PV_DoiSoat": ("CÂU 1 — ĐỐI SOÁT PIVOT VỚI SỐ LIỆU HỆ THỐNG", "Pivot theo tháng có khớp số liệu hệ thống không?"),
        "PV_Mien": ("CÂU 2 — MIỀN × THÁNG, % SO VỚI THÁNG TRƯỚC", f"Miền nào tăng mạnh nhất tháng 4 so với tháng 3?"),
        "PV_Nganh": ("CÂU 3 — TỶ TRỌNG NGÀNH HÀNG THEO THÁNG", "Tỷ trọng Bánh kẹo thay đổi thế nào từ tháng 1 đến tháng 4?"),
        "PV_CuaHang": ("CÂU 4 — CỬA HÀNG × THÁNG, % SO VỚI THÁNG TRƯỚC", "ST07 đã hồi phục sau đợt cải tạo tháng 3 chưa?"),
        "PV_AOV": ("CÂU 5 — AOV THEO THÁNG (Calculated Field)", "Mỗi đơn tháng 4 trung bình bao nhiêu tiền?")}
for nm, (t_, qn) in HWPV.items(): sheet_title(wb.create_sheet(nm), t_, qn)
dso = wb["PV_DoiSoat"]
# Khối đối soát đặt từ cột J: vùng giữ chỗ của Pivot 3 giá trị ở A4 là A4:H8 (pivot_xml), Excel xóa vùng này khi Refresh.
dso["J4"] = "Đối soát với So_lieu_he_thong (chênh lệch phải = 0)"; dso["J4"].font = BOLD
header(dso, 5, ["Tháng", "Δ Doanh thu", "Δ Số đơn", "Δ Quantity"], col=10)
for i, k in enumerate(sorted(bym4.index), start=6):
    dso.cell(row=i, column=10, value=int(k)).number_format = MONTH_FMT
    r_ = 5 + (i - 6)
    for j, (cap, col) in enumerate((("Doanh thu", "E"), ("Số đơn", "C"), ("Tổng Quantity", "D")), start=11):
        c = dso.cell(row=i, column=j, value=f'=GETPIVOTDATA("{cap}",$A$4,"Thang",{int(k)})-So_lieu_he_thong!{col}{r_}'); c.number_format = NUM_FMT; c.font = BODY
set_widths(dso, [10, 16, 12, 14, 3, 3, 3, 3, 3, 10, 14, 12, 12])
kt = wb.create_sheet("Kiem_tra")
kt["A1"] = "KIỂM TRA / ĐỐI SOÁT — Bài tập về nhà B3"; kt["A1"].font = TITLE
kt["A2"] = "Chấm bài: so từng chỉ tiêu của học viên với cột Đáp án. Bỏ chọn Slicer trước khi so."; kt["A2"].font = NOTE
header(kt, 4, ["Chỉ tiêu", "", "Kết quả (công thức)", "Đáp án", "Khớp?"])
hchecks = [("Số dòng dữ liệu", "=ROWS(tblSales4T[Order_ID])", NH, NUM_FMT),
           ("Số ô Quantity còn dạng chữ", "=SUMPRODUCT(--ISTEXT(tblSales4T[Quantity]))", 0, "0"),
           ("Tổng Quantity tháng 4 (Pivot)", '=GETPIVOTDATA("Tổng Quantity",PV_DoiSoat!$A$4,"Thang",4)', mh["by_month"][4]["qty"], NUM_FMT),
           ("Tổng doanh thu 4 tháng (VNĐ)", '=GETPIVOTDATA("Doanh thu",PV_DoiSoat!$A$4)', mh["revenue"], NUM_FMT),
           ("Doanh thu tháng 4 (VNĐ)", '=GETPIVOTDATA("Doanh thu",PV_DoiSoat!$A$4,"Thang",4)', mh["by_month"][4]["rev"], NUM_FMT),
           (f"Miền tăng mạnh nhất T4 so với T3: {mh['best_region']}", f'=GETPIVOTDATA("% so tháng trước",PV_Mien!$A$4,"Region","{mh["best_region"]}","Thang",4)', mh["reg_mom4"][mh["best_region"]], "0.0%"),
           ("Tỷ trọng Bánh kẹo tháng 4", '=GETPIVOTDATA("Tỷ trọng",PV_Nganh!$A$4,"Category","Bánh kẹo","Thang",4)', mh["banh_keo_share"][4], "0.0%"),
           ("ST07 tháng 4 so với tháng 3", '=GETPIVOTDATA("% so tháng trước",PV_CuaHang!$H$4,"Store_Code","ST07","Thang",4)', mh["st07_mom4"], "0.0%"),
           ("AOV tháng 4 (Calculated Field)", '=GETPIVOTDATA("AOV (VNĐ/đơn)",PV_AOV!$A$4,"Thang",4)', mh["by_month"][4]["aov"], NUM_FMT)]
for r, (label, f, exp, fmt) in enumerate(hchecks, start=5):
    kt.cell(row=r, column=1, value=label).font = BODY
    c = kt.cell(row=r, column=3, value=f); c.number_format = fmt; c.font = BODY
    c = kt.cell(row=r, column=4, value=exp); c.number_format = fmt; c.font = INPUT
    kt.cell(row=r, column=5, value=f'=IF(ROUND(C{r},6)=ROUND(D{r},6),"✔","✘")').font = BODY
HW_KT_ROWS = (5, 4 + len(hchecks))
set_widths(kt, [50, 2, 20, 16, 8])
ea = wb.create_sheet("Dap_an_loi")
ea["A1"] = "ĐÁP ÁN — LỖI CÀI SẴN VÀ LỖI HAY GẶP"; ea["A1"].font = TITLE
rows = [("#", "Lỗi", "Vị trí / dấu hiệu", "Cách sửa"),
        (1, f"{len(qty_text)} ô Quantity của tháng 4 lưu dạng chữ", "Sheet Sales, dòng " + ", ".join(map(str, mh["qty_text_rows"])) +
         f". Kéo Quantity vào Values ra “Count of Quantity”; đổi sang Sum thì tháng 4 ra {vn(mh['qty_t4_wrong'])}, thiếu {vn(qty_lost)} so với hệ thống ({vn(mh['by_month'][4]['qty'])}). "
         "Doanh thu vẫn khớp vì phép nhân tự đổi chữ thành số.",
         "Lọc Quantity / ô căn trái, tam giác xanh → Convert to Number; Refresh Pivot; đổi lại Summarize Values By → Sum"),
        (2, "Đếm số đơn bằng Count of Order_ID", f"Tháng 4 ra {vn(mh['by_month'][4]['rows'])} (số dòng) thay vì {vn(mh['by_month'][4]['orders'])} đơn", "Sum of So_don"),
        (3, "AOV = AVERAGE của AOV các tháng / cửa hàng", "Trung bình của trung bình", "Calculated Field Revenue / So_don (tổng ÷ tổng)"),
        (4, "Quên Refresh sau khi sửa dữ liệu", "Pivot vẫn Count hoặc vẫn thiếu Quantity", "Data → Refresh All (Ctrl+Alt+F5)")]
for i, rr in enumerate(rows, start=3):
    for j, val in enumerate(rr, start=1):
        c = ea.cell(row=i, column=j, value=val); c.font = HDR_FONT if i == 3 else BODY
        if i == 3: c.fill = HDR_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
set_widths(ea, [4, 36, 64, 44])
ins = wb.create_sheet("Insight_mau")
ins["A1"] = "INSIGHT MẪU (tham khảo khi chấm; học viên viết theo cách của mình)"; ins["A1"].font = TITLE
ins["A3"] = (f"\"Tháng 4/2025 doanh thu {vn(mh['by_month'][4]['rev'])} VNĐ, tăng {pct(mh['mom4'])} so với tháng 3; tăng mạnh nhất là Miền {mh['best_region']} "
             f"({spct(mh['reg_mom4'][mh['best_region']])}). ST07 Nha Trang từ {vn(mh['st07'][3])} (tháng 3) lên {vn(mh['st07'][4])} (tháng 4, {mh['st07_t4_orders']} đơn trong "
             f"{mh['st07_t4_days']} ngày có bán), nên đợt sụt tháng 3 là tạm thời, nhiều khả năng do cửa hàng cải tạo; khi so cửa hàng theo tháng 3 cần ghi chú ST07 "
             f"để không kết luận sai về hiệu quả cửa hàng.\"")
ins["A3"].font = BODY; ins["A3"].alignment = Alignment(wrap_text=True, vertical="top")
ins["A5"] = "Chấm: đủ 3 thành phần (con số có so sánh · nguyên nhân khả dĩ · hàm ý), con số khớp dữ liệu, và đã nêu lỗi Quantity dạng chữ ở câu 1."; ins["A5"].font = NOTE
set_widths(ins, [120])
wb.active = wb.sheetnames.index("Kiem_tra")
for w_ in wb.worksheets: w_.sheet_view.tabSelected = w_.title == "Kiem_tra"
HW_ANS = f"{HW}/B3_BTVN_MCI_Mart_T1_T4_2025_Answer.xlsx"
wb.save(HW_ANS)
HW_COLS = HC + HF
c_hw = Cache("hw", "tblSales4T", HW_COLS, calc_fields={"AOV": "=Revenue/So_don"}, date_fields=("Order_Date",),
             items={"Region": REGIONS, "Thang": [1, 2, 3, 4], "Category": CATS_Q1, "Store_Code": STORES})
HW_PIVOTS = [
    Pivot("PV_DoiSoat", "A4", "pvtDoiSoat", "hw", rows=["Thang"], values=[Value("Revenue", "Doanh thu"), Value("So_don", "Số đơn"), Value("Quantity", "Tổng Quantity")]),
    Pivot("PV_Mien", "A4", "pvtMien", "hw", rows=["Region"], cols=["Thang"], values=[Value("Revenue", "% so tháng trước", fmt="0.0%", show_as="pct_diff", base_field="Thang")], total_col=False),
    Pivot("PV_Mien", "A12", "pvtMienDT", "hw", rows=["Region"], cols=["Thang"], values=[Value("Revenue", "Doanh thu")]),
    Pivot("PV_Nganh", "A4", "pvtNganh", "hw", rows=["Category"], cols=["Thang"], values=[Value("Revenue", "Tỷ trọng", fmt="0.0%", show_as="pct_col")]),
    Pivot("PV_CuaHang", "A4", "pvtCuaHangDT", "hw", rows=["Store_Code"], cols=["Thang"], values=[Value("Revenue", "Doanh thu")]),
    Pivot("PV_CuaHang", "H4", "pvtCuaHangPct", "hw", rows=["Store_Code"], cols=["Thang"], values=[Value("Revenue", "% so tháng trước", fmt="0.0%", show_as="pct_diff", base_field="Thang")], total_col=False),
    Pivot("PV_AOV", "A4", "pvtAOV", "hw", rows=["Thang"], values=[Value("Revenue", "Doanh thu"), Value("So_don", "Số đơn"), Value("AOV", "AOV (VNĐ/đơn)")]),
]
best = mh["best_region"]
HW_SLICERS = [Slicer("PV_Mien", "Region", ["pvtMien", "pvtMienDT", "pvtNganh", "pvtCuaHangDT", "pvtCuaHangPct"], "H4", "J11")]
HW_CHARTS = [
    Chart("PV_Mien", "pvtMienDT", "col", f"Tháng 4: Miền {best} tăng mạnh nhất ({spct(mh['reg_mom4'][best])} so với tháng 3)", "H13", "R31",
          "PV_Mien!$A$14:$A$16", [(f"PV_Mien!${c}$13", f"PV_Mien!${c}$14:${c}$16") for c in "BCDE"], legend=True),
    Chart("PV_Nganh", "pvtNganh", "col", f"Bánh kẹo từ {pct(mh['banh_keo_share'][1])} (tháng 1) còn {pct(mh['banh_keo_share'][4])} (tháng 4)", "H4", "R22",
          "PV_Nganh!$A$6:$A$11", [(f"PV_Nganh!${c}$5", f"PV_Nganh!${c}$6:${c}$11") for c in "BCDE"], legend=True, val_fmt="0%"),
    Chart("PV_CuaHang", "pvtCuaHangDT", "col", f"ST07 hồi phục: tháng 4 đạt {dec(mh['st07'][4] / 1e6)} triệu, {spct(mh['st07_mom4'])} so với tháng 3 ({dec(mh['st07'][3] / 1e6)} triệu)",
          "A20", "L40", "PV_CuaHang!$A$6:$A$17", [(f"PV_CuaHang!${c}$5", f"PV_CuaHang!${c}$6:${c}$17") for c in "BCDE"], legend=True),
]

# =====================================================================
# 4) Ghi Pivot / Slicer / Timeline / PivotChart, mở bằng Excel, lưu lại, kiểm tra
# =====================================================================
inject(ANS, [c_ans], ANS_PIVOTS, ANS_SLICERS, ANS_TIMELINES, ANS_CHARTS)
for p_ in DEMO_REFRESH:
    inject(p_, [c_vung, c_tab], REFRESH_PIVOTS)
inject(HW_ANS, [c_hw], HW_PIVOTS, HW_SLICERS, (), HW_CHARTS)

# Khóa phẳng cho infographic (build_infographic.py chỉ đọc khóa dạng chữ-số-gạch dưới; số thập phân viết sẵn kiểu Việt Nam)
wk_rev = [w["rev"] for w in m["weekly"]]
meta["info"] = {"bac_t3": m["grid"]["Bắc-3"], "bac_t3_lines": m["grid"]["bac_t3_lines"], "rows": N, "orders": ORDERS, "revenue": REV,
                "aov": REV / ORDERS, "gm_pct": GM / REV, "mom2": spct(m["month"][2]["mom"]), "mom3": spct(m["month"][3]["mom"]),
                "banh_keo_t1": m["cat_share"]["Bánh kẹo"][1], "banh_keo_t2": m["cat_share"]["Bánh kẹo"][2],
                "weekend_pct": m["weekend"]["ratio"] - 1, "weekend_round": f"{(m['weekend']['ratio'] - 1) * 100:.0f}%", "pre_tet_per_day": dec(m["tet"]["truoc_tet"]["per_day"] / 1e6),
                "tet_per_day": f"{m['tet']['nghi_tet']['per_day'] / 1e6:.0f}", "week_tet": dec(wk_rev[4] / 1e6),
                "week_pre_min": dec(min(wk_rev[1:4]) / 1e6), "week_pre_max": dec(max(wk_rev[1:4]) / 1e6),
                "st07_t3": m["st07"]["t3_rev"], "st07_mom3": spct(m["st07"]["mom3"]), "st07_after_orders": m["st07"]["after"]["orders"],
                "st07_zero_days": 22 - m["st07"]["after"]["days_sold"], "count_lines": N,
                "hw_rows": m["hw"]["rows"], "hw_rev": m["hw"]["revenue"], "hw_t4": m["hw"]["by_month"][4]["rev"], "hw_mom4": spct(m["hw"]["mom4"])}
JS = lambda o: o.item() if hasattr(o, "item") else str(o)
with open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "build_b3_meta.json"), "w", encoding="utf-8", newline="\n") as f_:
    json.dump(meta, f_, ensure_ascii=False, indent=1, default=JS)

if USE_EXCEL:
    def check(path, sheet, rows):
        wb_ = load_workbook(path, data_only=True)
        ws_ = wb_[sheet]
        res = [(ws_.cell(row=r, column=1).value, ws_.cell(row=r, column=3).value, ws_.cell(row=r, column=5).value) for r in range(rows[0], rows[1] + 1)]
        bad = [x for x in res if x[2] != "✔"]
        print(os.path.basename(path), sheet, f"{len(res) - len(bad)}/{len(res)} ✔", "LỖI:" if bad else "", bad or "")
        return not bad
    ok = True
    for p_ in [ANS, HW_ANS, HW_DE] + list(DEMO_PIVOT) + list(DEMO_REFRESH) + list(DEMO_LOI):
        excel_resave(p_)
    ok &= check(ANS, "Kiem_tra", KT_ROWS)
    ok &= check(HW_ANS, "Kiem_tra", HW_KT_ROWS)
    for p_ in DEMO_REFRESH:
        wb_ = load_workbook(p_, data_only=True)
        v1, v2 = wb_["PV_Vung"]["B7"].value, wb_["PV_Table"]["B7"].value
        print(os.path.basename(p_), "PV_Vung tổng", v1, "PV_Table tổng", v2)
        ok &= v1 == v2 == REV
    print("KIỂM TRA:", "tất cả khớp" if ok else "CÓ LỖI")

print(json.dumps({k: meta[k] for k in ("q1", "month", "tet", "weekend", "st07", "loi", "refresh", "hw")}, ensure_ascii=False, indent=1, default=JS))
