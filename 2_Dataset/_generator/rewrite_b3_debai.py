"""Viết lại riêng sheet De_bai trong file đề bài tập về nhà Buổi 3, không dựng lại các file khác của Buổi 3.

Nội dung đề nằm ở b3_debai.py (build_b3.py dùng chung). Script này: sao lưu file đề, thay sheet De_bai bằng nội dung mới
bằng openpyxl, rồi mở bằng Excel lưu lại để các ô công thức của Sales có giá trị tính sẵn.

Chạy từ thư mục MCI Excel, sau khi đóng file đề trong Excel:
    python 2_Dataset\\_generator\\rewrite_b3_debai.py
"""
import json, os, shutil, sys, warnings
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from b3_debai import fill_de_bai
from pivot_xml import excel_resave

ROOT = os.environ.get("MCI_ROOT") or os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
HW = f"{ROOT}/1_Sessions/B3_PivotTable/Homework"
DE = f"{HW}/B3_BTVN_MCI_Mart_T1_T4_2025.xlsx"
META = os.path.join(os.path.dirname(os.path.abspath(__file__)), "build_b3_meta.json")

vn = lambda x: f"{x:,.0f}".replace(",", ".")
spct = lambda x, d=1: ("+" if x > 0 else "-" if x < 0 else "") + f"{abs(x) * 100:.{d}f}%".replace(".", ",")

mh = json.load(open(META, encoding="utf-8"))["hw"]
facts = {"rows": vn(mh["rows"]), "qty4": vn(mh["by_month"]["4"]["qty"]), "rev": vn(mh["revenue"]),
         "rev4": vn(mh["by_month"]["4"]["rev"]), "mom4": spct(mh["mom4"])}

shutil.copyfile(DE, DE + ".bak")
warnings.simplefilter("ignore")
wb = load_workbook(DE)
old = wb["De_bai"]
idx = wb.sheetnames.index("De_bai")
wb.remove(old)
de = wb.create_sheet("De_bai", idx)
fill_de_bai(de, facts)
wb.active = idx
for w in wb.worksheets:
    w.sheet_view.tabSelected = w.title == "De_bai"
wb.save(DE)
excel_resave(DE)
os.remove(DE + ".bak")
print("Đã viết lại De_bai:", DE)
