"""Dựng tài liệu phát cho Buổi 3 (.docx) bằng python-docx: handout tự học B3_Handout.docx.

Số liệu ví dụ tính từ dataset (xem build_b3.py, build_b3_meta.json). Ảnh chụp thật lấy từ Materials/B3/Images (H<nn>_…png);
chưa có ảnh thì để khung [CHỤP]. Xuất PDF bằng Microsoft Word (xem 2_Dataset/README.md).
    python3 _generator/build_docs_b3.py
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.image.image import Image as DocxImage

ROOT = os.path.normpath(os.environ.get("MCI_ROOT") or os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
OUT = f"{ROOT}/1_Sessions/B3_PivotTable/Handout"
IMG = f"{ROOT}/Materials/B3/Images"   # ảnh chụp thật, danh sách: _DANH_SACH_ANH.md
MISSING = []
os.makedirs(OUT, exist_ok=True)
RED, NAVY, GOLD, GREY = RGBColor(0xC0, 0, 0), RGBColor(0x20, 0x38, 0x64), RGBColor(0x8A, 0x65, 0x00), RGBColor(0x58, 0x58, 0x58)

def base_doc(footer="MCI – Magic Code Institute · Excel for Business Intelligence Analyst · Buổi 3"):
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for m in ("left_margin", "right_margin"): setattr(sec, m, Cm(2))
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    for name, size, color in (("Heading 1", 16, NAVY), ("Heading 2", 13, NAVY), ("Heading 3", 11.5, RED)):
        h = d.styles[name]; h.font.name = "Arial"; h.font.size = Pt(size); h.font.color.rgb = color; h.font.bold = True
        rf = h.element.rPr.rFonts
        for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(k)) is not None: del rf.attrib[qn(k)]
        for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"): rf.set(qn(k), "Arial")
    fp = sec.footer.paragraphs[0]
    fp.text = footer
    fp.runs[0].font.size = Pt(8); fp.runs[0].font.color.rgb = GREY
    return d

def para(d, text="", bold=False, color=None, size=None, italic=False):
    p = d.add_paragraph()
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic
        if color: r.font.color.rgb = color
        if size: r.font.size = Pt(size)
    return p

def rich(d, parts):
    p = d.add_paragraph()
    for t, f in parts:
        r = p.add_run(t); r.bold = f.get("bold", False); r.italic = f.get("italic", False)
        if f.get("color"): r.font.color.rgb = f["color"]
    return p

def bullet(d, text):
    p = d.add_paragraph(style="List Bullet"); p.add_run(text); return p

def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_)
    tcPr.append(shd)

def table(d, rows, widths, header=True, size=9.5, first_bold=False, min_h=None):
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for j, w in enumerate(widths): t.columns[j].width = Cm(w)
    for i, row in enumerate(rows):
        if min_h and i: t.rows[i].height = Cm(min_h)
        for j, v in enumerate(row):
            c = t.cell(i, j); c.width = Cm(widths[j]); c.text = ""
            r = c.paragraphs[0].add_run(str(v)); r.font.size = Pt(size)
            if i < len(rows) - 1: c.paragraphs[0].paragraph_format.keep_with_next = True   # không tách bảng qua trang
            if header and i == 0:
                r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); shade(c, "203864")
            elif first_bold and j == 0:
                r.bold = True
    d.add_paragraph()
    return t

def figure(d, what, where, img, width=15):
    """Chèn ảnh chụp thật từ Materials/B3/Images nếu đã có; chưa có thì để khung chờ [CHỤP] (CLAUDE.md, mục Ảnh minh họa)."""
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(f"{IMG}/{img}"):
        # Không phóng ảnh quá cỡ thật (zoom 100%, 96 dpi) và không cao quá 10 cm, để ảnh nhỏ không bị vỡ, ảnh vuông không chiếm cả trang.
        px = DocxImage.from_file(f"{IMG}/{img}")
        w = min(width, px.px_width * 2.54 / 96)
        if w * px.px_height / px.px_width > 10: w = 10 * px.px_width / px.px_height
        p.add_run().add_picture(f"{IMG}/{img}", width=Cm(w)); return p
    MISSING.append(img)
    r = p.add_run(f"[CHỤP] {what}"); r.italic = True; r.font.color.rgb = GREY; r.font.size = Pt(9)
    r = p.add_run(f"\n{where} → {img}"); r.font.color.rgb = GREY; r.font.size = Pt(8)
    pPr = p._p.get_or_add_pPr(); bdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "dashed"); e.set(qn("w:sz"), "6"); e.set(qn("w:color"), "A6A6A6"); e.set(qn("w:space"), "8"); bdr.append(e)
    pPr.append(bdr)
    return p

def tool(d, name, purpose, how=None, example=None, tip=None, protip=None, shot=None):
    """Khối công cụ theo style khóa học: tên (đỏ) → mục đích → cách dùng → bảng ví dụ trên case → ảnh chụp → TIP / PRO TIP.
    shot = (nội dung ảnh, file · sheet · vùng, tên file ảnh)."""
    d.add_heading(name, level=3).paragraph_format.keep_with_next = True
    para(d, purpose).paragraph_format.keep_with_next = True
    for h in how or []: bullet(d, h)
    if example: table(d, example[0], example[1])
    if shot: figure(d, *shot)
    if tip: rich(d, [("TIP: ", {"bold": True, "color": NAVY}), (tip, {})])
    if protip: rich(d, [("PRO TIP: ", {"bold": True, "color": GOLD}), (protip, {})])

# =====================================================================
# HANDOUT B3
# =====================================================================
d = base_doc()
para(d, "MCI – MAGIC CODE INSTITUTE", bold=True, color=GREY, size=9)
d.add_heading("Handout Buổi 3 — Tư duy dữ liệu & PivotTable", level=1)
para(d, "Tài liệu tự học đi kèm Buổi 3. Các công cụ ở mục 2–7 KHÔNG giảng chi tiết trên lớp; đọc sau buổi học rồi làm bài tập về nhà (bắt buộc) ở mục 9. "
        "Ví dụ viết trên file đáp án B3_MCI_Mart_Q1_2025_Answer.xlsx (tblSales có cột So_don, các sheet PV_…). Tên trường và lệnh Excel giữ tiếng Anh như trên Ribbon.", italic=True, color=GREY)

d.add_heading("Mục lục", level=2)
for s in ["1. Tóm tắt: Pivot = SUMIFS theo tọa độ · quy trình dựng · phím tắt", "2. Show Values As mở rộng: % of Parent Row Total, Rank, Index",
          "3. Calculated Item, Solve Order, List Formulas", "4. Report Filter Pages: mỗi cửa hàng một sheet",
          "5. PivotChart: Field Buttons, di chuyển, đổi kiểu", "6. Data → Subtotal và Group / Outline (ngoài Pivot)",
          "7. PRO TIP: Conditional Formatting trong Pivot, Custom List, PivotTable Options", "8. GETPIVOTDATA, khi nào vẫn dùng SUMIFS, lỗi Pivot hay gặp",
          "9. Bài tập về nhà (bắt buộc)", "10. Checklist trước Buổi 4"]:
    bullet(d, s)

# 1 Tóm tắt
d.add_heading("1. Tóm tắt: Pivot = SUMIFS theo tọa độ", level=2)
para(d, "Mỗi ô giá trị của Pivot là một SUMIFS mà Excel viết giúp: tọa độ hàng và cột của ô chính là các tiêu chí lọc. "
        "Ô Bắc × Tháng 3 = SUMIFS(tblSales[Revenue], tblSales[Region], \"Bắc\", tblSales[Thang], 3) = 90.210.000. "
        "Buổi 2 gọi đó là criteria của SUMIFS, Buổi 3 gọi là tọa độ ô Pivot, Buổi 5 sẽ gọi là filter context.")
table(d, [["Bước", "Thao tác", "Ghi nhớ"],
          ["1. Nguồn", "Chọn một ô trong tblSales → Insert → PivotTable → From Table/Range → New Worksheet", "Nguồn là Table thì dòng mới tự vào khi Refresh"],
          ["2. Kéo trường", "Dimension vào Rows / Columns / Filters; Measure vào Values", "Kéo nhầm thì kéo ra khỏi vùng"],
          ["3. Định dạng", "Value Field Settings → Custom Name, Number Format; Report Layout → Outline", "Không bôi đen ô để định dạng"],
          ["4. Đối soát", "=GETPIVOTDATA(…) so với SUMIFS / tổng cột", "Tổng Pivot phải bằng 941.454.000"],
          ["5. Làm mới", "Refresh (Alt+F5) · Refresh All (Ctrl+Alt+F5)", "Pivot đọc từ cache, sửa nguồn xong phải Refresh"]], [2.4, 8.2, 5.4], first_bold=True)
table(d, [["Phím tắt / thao tác", "Tác dụng"],
          ["Alt, N, V, T", "Insert → PivotTable → From Table/Range"],
          ["Alt+F5 · Ctrl+Alt+F5", "Refresh Pivot đang chọn · Refresh All"],
          ["Double-click ô giá trị", "Show Details: sheet mới chứa các dòng gốc"],
          ["Alt+Shift+→ · Alt+Shift+←", "Group · Ungroup mục đang chọn"],
          ["Alt+F1 (khi đang chọn Pivot)", "Chèn nhanh PivotChart kiểu mặc định"]], [6.0, 10.0], first_bold=True)
figure(d, "Tab PivotTable Analyze: các nhóm PivotTable, Active Field, Group, Filter (Insert Slicer / Timeline), Data (Refresh, Change Data Source), Calculations",
       "Practice/B3_MCI_Mart_Q1_2025_Answer.xlsx · sheet PV_MienThang · Ribbon khi chọn một ô trong Pivot", "H01_analyze_tab.png")

# 2 Show Values As mở rộng
d.add_heading("2. Show Values As mở rộng", level=2)
para(d, "Trên lớp đã dùng % of Column Total, % of Grand Total, % Difference From, Running Total In. Ba lựa chọn dưới đây hay dùng khi phân tích nhiều tầng.")
tool(d, "% of Parent Row Total", "Tỷ trọng của mục con trong mục cha ngay trên nó (Region → Store_Code).",
     how=["Rows: Region rồi Store_Code; Values: Revenue → Show Values As → % of Parent Row Total."],
     example=([["Cửa hàng", "Doanh thu Q1", "% trong miền"],
               ["ST08 Quận 1 (Nam)", "130.448.000", "30,5%"], ["ST05 Hải Châu (Trung)", "86.993.000", "45,8%"], ["ST04 Hạ Long (Bắc)", "45.105.000", "13,9%"]], [7, 4.5, 4.5]),
     tip="Mỗi miền cộng lại 100%. Dùng để trả lời “trong miền này, cửa hàng nào gánh doanh thu?”.")
tool(d, "Rank Largest to Smallest", "Xếp hạng mục theo giá trị, không cần sắp xếp lại bảng (giống RANK.EQ của Buổi 2).",
     how=["Values: Revenue lần hai → Show Values As → Rank Largest to Smallest → Base field: Store_Code."],
     example=([["Hạng", "Cửa hàng"], ["1 · 2 · 3", "ST08 · ST01 · ST02"], ["10 · 11 · 12", "ST11 · ST06 · ST04"]], [4, 12]))
tool(d, "Index", "Mức “nặng nhẹ” của một ô so với cả bảng: (Giá trị ô × Tổng chung) ÷ (Tổng hàng × Tổng cột). Lớn hơn 1 = ô đó chiếm phần lớn hơn mức trung bình.",
     how=["Rows: Category; Columns: Region; Values: Revenue → Show Values As → Index."],
     example=([["Ngành × Miền", "Index", "Đọc là"],
               ["Chăm sóc cá nhân × Nam", "1,17", "Miền Nam mua nhóm này nhiều hơn mức trung bình 17%"],
               ["Chăm sóc cá nhân × Trung", "0,79", "Miền Trung mua ít hơn mức trung bình 21%"],
               ["Bánh kẹo × Bắc", "1,00", "Đúng bằng mức trung bình"]], [5, 2, 9]),
     shot=("Pivot Category × Region, Show Values As → Index (Chăm sóc cá nhân × Nam = 1,17)", "Practice/B3_MCI_Mart_Q1_2025_Answer.xlsx · tạo Pivot mới từ tblSales · A3:E10", "H02_show_values_index.png"))

# 3 Calculated Item
d.add_heading("3. Calculated Item, Solve Order, List Formulas", level=2)
tool(d, "Calculated Item", "Tạo MỤC mới trong một trường (khác Calculated Field tạo TRƯỜNG mới). Ví dụ: mục “Hàng FMCG” trong Category = Đồ uống + Sữa + Bánh kẹo.",
     how=["Chọn một ô nhãn của Category → PivotTable Analyze → Fields, Items & Sets → Calculated Item.",
          "Name: Hàng FMCG; Formula: ='Đồ uống' + 'Sữa' + 'Bánh kẹo' → Add → OK."],
     example=([["Dòng", "Doanh thu Q1"], ["Hàng FMCG (mục mới)", "465.501.000"], ["Grand Total khi còn đủ 6 ngành + FMCG", "1.406.955.000 ✘ (cộng hai lần)"], ["Grand Total đúng", "941.454.000"]], [9, 7]),
     shot=("Hộp Insert Calculated Item và Pivot có dòng Hàng FMCG, Grand Total 1.406.955.000", "Practice/B3_MCI_Mart_Q1_2025_Answer.xlsx · Pivot mới Category × Revenue", "H03_calculated_item.png"),
     tip="Grand Total cộng cả mục mới lẫn các mục thành phần → sai. Ẩn các mục thành phần, hoặc tốt hơn: dùng Group (Buổi 3, slide 19) hay thêm cột phân nhóm trong nguồn.",
     protip="Tài liệu Pivot khuyên: KHÔNG dùng Calculated Item nếu không bắt buộc. Trường có Calculated Item thì không Group được, và Values không đổi được sang Average, StdDev, Var.")
tool(d, "Solve Order · List Formulas", "Solve Order: khi có nhiều Calculated Item cùng tác động lên một ô, quyết định công thức nào tính sau cùng. "
     "List Formulas: tạo sheet mới liệt kê mọi Calculated Field / Item của Pivot.",
     how=["PivotTable Analyze → Fields, Items & Sets → Solve Order… / List Formulas."],
     tip="Trên file đáp án, List Formulas liệt kê hai Calculated Field: AOV = Revenue/So_don và GM_pct = Gross_Margin/Revenue. Dùng để kiểm tra và bàn giao báo cáo.")

# 4 Report Filter Pages
d.add_heading("4. Report Filter Pages", level=2)
tool(d, "Show Report Filter Pages", "Tách Pivot thành nhiều sheet, mỗi giá trị của trường ở vùng Filters một sheet: ví dụ mỗi cửa hàng một báo cáo.",
     how=["Kéo Store_Code vào Filters (Pivot Category × Thang).", "PivotTable Analyze → Options ▾ → Show Report Filter Pages → chọn Store_Code → OK."],
     shot=("12 sheet mới ST01 … ST12, mỗi sheet một Pivot đã lọc sẵn theo cửa hàng", "Practice/B3_MCI_Mart_Q1_2025_Answer.xlsx · Pivot Category × Thang, Store_Code ở Filters", "H04_report_filter_pages.png"),
     tip="Tiện khi phải gửi mỗi cửa hàng trưởng một bảng riêng. Các sheet vẫn dùng chung cache, Refresh All cập nhật tất cả.")

# 5 PivotChart
d.add_heading("5. PivotChart: Field Buttons, di chuyển, đổi kiểu", level=2)
tool(d, "PivotChart Analyze", "Các lệnh riêng của PivotChart.",
     example=([["Lệnh", "Tác dụng"],
               ["Field Buttons → Hide All", "Ẩn các nút lọc xám trên biểu đồ khi đưa vào báo cáo"],
               ["Move Chart", "Chuyển biểu đồ sang sheet riêng (chart sheet) hoặc sheet báo cáo"],
               ["Change Chart Type", "Đổi kiểu mà không phải vẽ lại: Pie → Line, Column → Bar…"],
               ["Insert Slicer / Timeline", "Lọc cả Pivot lẫn biểu đồ"]], [6, 10]),
     shot=("PivotChart doanh thu theo thứ với menu Field Buttons đang mở", "Practice/B3_MCI_Mart_Q1_2025_Answer.xlsx · sheet PV_Thu · D4:L22", "H05_field_buttons.png"),
     tip="PivotChart luôn đi theo bố cục Pivot: trường ở Rows thành trục ngang, trường ở Columns thành các chuỗi (legend). Muốn đổi biểu đồ thì đổi Pivot.")

# 6 Subtotal, Group/Outline
d.add_heading("6. Data → Subtotal và Group / Outline", level=2)
para(d, "Hai công cụ cũ hơn Pivot nhưng vẫn gặp trong file báo cáo của doanh nghiệp. Chúng chạy trên VÙNG THƯỜNG (không chạy trên Excel Table).")
tool(d, "Subtotal (Data → Subtotal)", "Chèn dòng tổng phụ mỗi khi giá trị của một cột thay đổi, kèm các nút thu gọn 1 · 2 · 3.",
     how=["Dữ liệu phải sắp xếp theo cột cần tổng phụ (Sales_Vung đã xếp theo Region).",
          "Chọn một ô trong vùng → Data → Subtotal → At each change in: Region · Use function: Sum · Add subtotal to: Revenue → OK.",
          "Bấm nút 2 bên trái để chỉ xem dòng tổng: Bắc 323.978.000 · Trung 189.813.000 · Nam 427.663.000 · Grand Total 941.454.000.",
          "Bỏ: Data → Subtotal → Remove All."],
     shot=("Hộp Subtotal và vùng Sales_Vung thu gọn ở mức 2 (ba dòng tổng miền)", "Demo/B3_Demo_Refresh_HocVien.xlsx · sheet Sales_Vung · A1:I6192 thu gọn", "H06_subtotal.png"),
     tip="Dòng tổng phụ nằm LẪN trong dữ liệu: đúng loại “dòng tổng giữa bảng” mà Buổi 1 đã dọn. Dùng Subtotal để xem nhanh rồi Remove All; đừng lưu dữ liệu gốc có dòng tổng.")
tool(d, "Group / Outline (Data → Group)", "Nhóm các dòng hoặc cột để thu gọn / mở rộng bằng dấu − / +, không thay đổi dữ liệu.",
     how=["Chọn các dòng chi tiết → Data → Group (Alt+Shift+→). Bỏ nhóm: Data → Ungroup (Alt+Shift+←)."],
     tip="Hay dùng để giấu cột phụ trợ trong báo cáo mà không xóa. Khác Group trong Pivot: Group của Pivot tạo mục mới để tính toán.")

# 7 PRO TIP
d.add_heading("7. PRO TIP: Conditional Formatting trong Pivot, Custom List, PivotTable Options", level=2)
tool(d, "Conditional Formatting trong Pivot", "Tô màu theo giá trị ngay trong Pivot, ví dụ Color Scale trên Pivot Cửa hàng × Tháng để thấy ô ST07 × 3 đỏ nhất.",
     how=["Chọn vùng số của Pivot → Home → Conditional Formatting → Color Scales.",
          "Bấm biểu tượng Formatting Options cạnh vùng vừa tô → chọn “All cells showing ‘Doanh thu’ values for ‘Store_Code’ and ‘Thang’”."],
     shot=("Pivot PV_CuaHang tô Color Scale và lựa chọn “All cells showing ‘Doanh thu’ values…”", "Practice/B3_MCI_Mart_Q1_2025_Answer.xlsx · sheet PV_CuaHang · A4:E18", "H07_cf_pivot.png"),
     tip="Chọn “All cells showing…” thì quy tắc tự áp cho ô mới khi Pivot thêm cửa hàng hay tháng. Nếu chỉ áp cho vùng đang chọn, Pivot đổi bố cục là mất.")
tool(d, "Custom List (thứ tự tự đặt)", "Sắp xếp theo thứ tự nghiệp vụ thay vì A → Z, ví dụ thứ trong tuần T2 … CN, hoặc Bắc → Trung → Nam.",
     how=["File → Options → Advanced → General → Edit Custom Lists → gõ Bắc, Trung, Nam (mỗi dòng một mục) → Add.",
          "Trong Pivot: chuột phải → PivotTable Options → Totals & Filters → giữ tick Use Custom Lists when sorting; rồi Sort → Sort A to Z trên trường Region."],
     tip="Excel có sẵn danh sách Mon … Sun, nên máy tiếng Anh tự xếp thứ đúng thứ tự. Máy tiếng Việt hiện T2 … CN và xếp theo chữ cái (CN đứng đầu): thêm Custom List T2, T3, …, T7, CN.")
tool(d, "PivotTable Options → Data", "Các thiết lập về bộ nhớ đệm (cache) của Pivot.",
     example=([["Thiết lập", "Nên dùng khi"],
               ["Save source data with file (bỏ tick)", "File nặng vì cache; bỏ tick thì file nhẹ hơn nhưng phải Refresh khi mở"],
               ["Refresh data when opening the file", "Báo cáo phát cho người khác: luôn thấy số mới nhất"],
               ["Number of items to retain per field: None", "Đã xóa / đổi tên mục ở nguồn mà bộ lọc vẫn còn mục cũ"]], [7, 9]),
     shot=("Hộp PivotTable Options, tab Data", "Practice/B3_MCI_Mart_Q1_2025_Answer.xlsx · PV_MienThang → chuột phải → PivotTable Options → Data", "H08_pivot_options.png"))

# 8 GETPIVOTDATA, SUMIFS, lỗi
d.add_heading("8. GETPIVOTDATA, khi nào vẫn dùng SUMIFS, lỗi Pivot hay gặp", level=2)
tool(d, "GETPIVOTDATA", "Lấy một con số từ Pivot theo tên trường và tên mục, không phụ thuộc vị trí ô.",
     example=([["Công thức", "Kết quả"],
               ['=GETPIVOTDATA("Doanh thu",PV_MienThang!$A$4,"Region","Bắc","Thang",3)', "90.210.000"],
               ['=GETPIVOTDATA("Doanh thu",PV_MienThang!$A$4,"Region",$A13,"Thang",3)', "Tham chiếu ô chứa tên miền: kéo được"],
               ['=GETPIVOTDATA("AOV (VNĐ/đơn)",PV_Thang!$A$4)', "299.445 (Grand Total)"]], [11, 5]),
     tip="Muốn gõ =D6 thay vì GETPIVOTDATA: PivotTable Analyze → Options ▾ → bỏ tick Generate GetPivotData. GETPIVOTDATA đổi theo Slicer đang chọn.")
rich(d, [("Khi nào vẫn dùng SUMIFS thay vì Pivot? ", {"bold": True, "color": NAVY}),
         ("Khi báo cáo có bố cục cố định và ô kết quả được công thức khác tham chiếu: thẻ KPI, bảng so với target, bảng in hằng tháng. "
          "Pivot hợp để khám phá, đổi chiều liên tục; SUMIFS hợp cho “khung báo cáo đã chốt”. Buổi 6 sẽ dùng CUBEVALUE cho việc này.", {})])
table(d, [["Lỗi", "Dấu hiệu", "Xử lý"],
          ["Nguồn không phải Table", "Dòng mới không vào Pivot sau Refresh", "Nguồn Pivot = Table; hoặc Change Data Source"],
          ["Quên Refresh", "Pivot không khớp dữ liệu", "Refresh All (Ctrl+Alt+F5); bật Refresh data when opening the file"],
          ["Cột số có ô trống / chữ", "Pivot mặc định Count thay vì Sum; Sum thiếu tiền", "Sửa dữ liệu nguồn (Convert to Number), Refresh, đổi Summarize Values By"],
          ["Không group được theo ngày", "“Cannot group that selection”", "Cột ngày có ô chữ / trống → sửa nguồn"],
          ["GETPIVOTDATA tự sinh", "Kéo công thức không đổi kết quả", "Tắt Generate GetPivotData hoặc tham chiếu ô trong GETPIVOTDATA"],
          ["Pie cho xu hướng theo tháng", "Không đọc được tăng / giảm", "3 câu hỏi chọn chart → Line"],
          ["Count of Order_ID để đếm đơn", "Ra 6.187 (số dòng) thay vì 3.144 đơn", "Cột So_don =IF([@Line_No]=1,1,0) rồi Sum"]], [4.2, 5.6, 6.2], first_bold=True)

# 9 Homework
d.add_heading("9. Bài tập về nhà (bắt buộc)", level=2)
rich(d, [("File đề: ", {"bold": True}), ("B3_BTVN_MCI_Mart_T1_T4_2025.xlsx — dữ liệu tháng 1–4/2025 đã làm giàu (tblSales4T, 8.106 dòng) và sheet So_lieu_he_thong (số đối soát). "
                                         "Nộp trước Buổi 4, tên file B3_BTVN_<HoTen>.xlsx. Sheet De_bai ghi đủ yêu cầu.", {})])
rich(d, [("Nộp bài: ", {"bold": True, "color": NAVY}), ("https://hocvien.mcivietnam.com/auth#/ — Tài khoản: email bạn đăng ký tại MCI · Mật khẩu: hv12345678", {})])
for s in ["Đối soát trước khi phân tích: Pivot theo Thang với Revenue, So_don, Quantity; so với So_lieu_he_thong. Có cột lệch thì tìm nguyên nhân trong dữ liệu, sửa ở tblSales4T, Refresh cho đến khi khớp.",
          "Pivot Region × Thang, thêm % Difference From (previous): miền nào tăng mạnh nhất tháng 4 so với tháng 3?",
          "Pivot Category × Thang, % of Column Total: tỷ trọng Bánh kẹo thay đổi thế nào từ tháng 1 đến tháng 4?",
          "Pivot Store_Code × Thang và % so với tháng trước; double-click ST07 × 3 và ST07 × 4: ST07 đã hồi phục chưa?",
          "Calculated Field AOV = Revenue / So_don theo tháng.",
          "Mỗi câu 2–4 một PivotChart, tiêu đề là câu kết luận; Slicer Region nối các Pivot. Sheet Insight: 1 câu insight và lỗi dữ liệu bạn đã sửa."]:
    bullet(d, s)
rich(d, [("Tự kiểm tra trước khi nộp: ", {"bold": True, "color": NAVY}),
         ("Pivot khớp So_lieu_he_thong ở cả Revenue, So_don, Quantity (Quantity tháng 4 = 3.448); doanh thu 4 tháng 1.238.419.000 VNĐ; tháng 4 296.965.000 VNĐ, tăng 10,1% so với tháng 3.", {})])

# 10 Checklist
h = d.add_heading("10. Checklist trước Buổi 4", level=2); h.paragraph_format.keep_with_next = True
for s in ["File B3 của bạn: 6 sheet Pivot; tổng Pivot = tổng SUMIFS Buổi 2 = 941.454.000 (khối đối soát chênh lệch 0).",
          "PV_NganhHang, PV_Thang dùng Show Values As; AOV bằng Calculated Field = 299.445.",
          "Slicer Store_Code lọc được cả 3 Pivot; mỗi PivotChart có tiêu đề là câu kết luận và dòng “Câu hỏi:”.",
          "Đã viết câu insight cuối Buổi 3 (Tết, ST07 hoặc cuối tuần).",
          "Đã nộp bài tập về nhà (B3_BTVN_<HoTen>.xlsx) và tìm ra cột lệch ở bước đối soát.",
          "Nếu file chưa khớp: Buổi 4 lớp phát file khởi đầu (đáp án B3).",
          "BẮT BUỘC từ Buổi 4: máy Windows có Excel 365 / 2016+, Data → Get Data mở được Power Query, COM Add-in Microsoft Power Pivot đã bật (checklist Buổi 1). "
          "Dùng Mac: máy ảo Parallels Desktop / VMware Fusion hoặc laptop Windows khác."]:
    bullet(d, "☐  " + s).paragraph_format.keep_with_next = True
d.paragraphs[-1].paragraph_format.keep_with_next = False
d.save(f"{OUT}/B3_Handout.docx")
print("ok;", f"còn {len(MISSING)} khung [CHỤP] chưa có ảnh:" if MISSING else "đủ ảnh.", " ".join(MISSING))
