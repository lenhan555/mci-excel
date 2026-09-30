"""Build B1 homework (bài tập về nhà bắt buộc): đề (file bẩn tháng 4/2025) + đáp án.

Đọc MCI_Mart_Master.xlsx đã có, không cần thư mục build/. Số liệu tháng 4/2025 là dữ liệu thật của dataset.
Đề khác bài trên lớp về hình thức lỗi (nhóm theo cửa hàng, dòng tiêu đề lặp lại, mã sản phẩm thiếu số 0)
để học viên áp dụng lại kỹ năng B1 chứ không làm lại y hệt.
"""
import os, random
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = os.path.normpath(os.environ.get("MCI_ROOT") or os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
MASTER = f"{ROOT}/2_Dataset/MCI_Mart_Master.xlsx"
HW = f"{ROOT}/1_Sessions/B1_Nen_tang_Excel_Du_lieu/Homework"
os.makedirs(HW, exist_ok=True)
rnd = random.Random(41)

F = "Arial"
HDR_FONT = Font(name=F, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="203864")
TITLE = Font(name=F, size=14, bold=True, color="203864")
BODY = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
NOTE = Font(name=F, size=10, italic=True, color="585858")
DATE_FMT = "dd/mm/yyyy"
NUM_FMT = "#,##0"

def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

def header(ws, row, labels):
    for j, h in enumerate(labels, start=1):
        c = ws.cell(row=row, column=j, value=h); c.font = HDR_FONT; c.fill = HDR_FILL

def py(v):
    if isinstance(v, pd.Timestamp): return v.to_pydatetime()
    return v.item() if hasattr(v, "item") else v

def write_df(ws, df, start_row=1, formats=None):
    header(ws, start_row, list(df.columns))
    for i, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=py(v)); c.font = BODY
            if formats and df.columns[j - 1] in formats: c.number_format = formats[df.columns[j - 1]]

# ---------------------------------------------------------------------
# Dữ liệu tháng 4/2025
# ---------------------------------------------------------------------
sales = pd.read_excel(MASTER, "Sales")
products = pd.read_excel(MASTER, "Product")
stores = pd.read_excel(MASTER, "Store")
t4 = sales[sales.Order_ID.str.startswith("HD2504")].copy()
t4["Product_Name"] = t4.Product_Code.map(products.set_index("Product_Code")["Product_Name"])
t4 = t4.sort_values(["Store_Code", "Order_Date", "Order_ID", "Line_No"]).reset_index(drop=True)
cols = ["Order_ID", "Line_No", "Order_Date", "Store_Code", "Customer_ID", "Product_Code", "Product_Name", "Quantity", "Unit_Price"]
clean = t4[cols].copy()
n = len(clean)
exp_rev = int((clean.Quantity * clean.Unit_Price).sum())
exp_qty = int(clean.Quantity.sum())

# Chọn lỗi cài sẵn (chỉ số trong clean)
idx = list(clean.index)
bad_store = rnd.sample(list(clean.index[clean.Store_Code == "ST05"]), 2)          # "ST5"
rest = [i for i in idx if i not in bad_store]
bad_prod = rnd.sample(rest, 3); rest = [i for i in rest if i not in bad_prod]     # "P25" thay vì "P025"
neg_qty = rnd.sample(rest, 2); rest = [i for i in rest if i not in neg_qty]
txt_qty = rnd.sample(rest, 10); rest = [i for i in rest if i not in txt_qty]
txt_date = rnd.sample(rest, 12)
store_list = sorted(clean.Store_Code.unique())
repeat_hdr_after = "ST06"                                                          # tiêu đề lặp lại sau khối ST06
blank_store = "ST09"                                                               # dòng trống giữa khối ST09

# ---------------------------------------------------------------------
# ĐỀ: B1_BTVN_MCI_Mart_T4_2025.xlsx
# ---------------------------------------------------------------------
wb = Workbook()
de = wb.active; de.title = "De_bai"
ws = wb.create_sheet("BaoCao_T4_2025")
ws["A1"] = "BÁO CÁO BÁN HÀNG THÁNG 4/2025 — MCI MART (THEO CỬA HÀNG)"; ws["A1"].font = Font(name=F, size=14, bold=True, color="C00000")
ws.merge_cells("A1:I1")
ws["A2"] = "Xuất từ hệ thống POS ngày 02/05/2025 • Đơn vị tiền: VNĐ"; ws["A2"].font = NOTE
ws.merge_cells("A2:I2")

def write_header_row(r):
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=r, column=j, value=c); cell.font = HDR_FONT; cell.fill = HDR_FILL

write_header_row(4)
r = 5
blk_fill = PatternFill("solid", fgColor="FFF2CC")
sub_fill = PatternFill("solid", fgColor="BFBFBF")
store_name = stores.set_index("Store_Code")["Store_Name"]
log = {"sub_rows": [], "blank_row": None, "repeat_hdr_row": None}
for sc in store_list:
    block = clean[clean.Store_Code == sc]
    c = ws.cell(row=r, column=1, value=f"Cửa hàng {sc} — {store_name[sc]}"); c.font = BOLD
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
    for j in range(1, 10): ws.cell(row=r, column=j).fill = blk_fill
    log["sub_rows"].append(r); r += 1
    mid = block.index[len(block) // 2]
    for i, row in block.iterrows():
        if sc == blank_store and i == mid:
            log["blank_row"] = r; r += 1
        oid, ln, od, s, cid, pc, pn, qty, up = row.tolist()
        if i in bad_store: s = "ST5"
        if i in bad_prod: pc = "P" + str(int(pc[1:]))
        if i in neg_qty: qty = -int(qty)
        vals = [oid, int(ln), od.to_pydatetime(), s, cid, pc, pn, int(qty), int(up)]
        if i in txt_qty: vals[7] = str(int(qty))
        if i in txt_date: vals[2] = od.strftime("%d/%m/%Y")
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=j, value=v); cell.font = BODY
            if j == 3 and not isinstance(v, str): cell.number_format = DATE_FMT
            if j == 9: cell.number_format = NUM_FMT
        r += 1
    ws.cell(row=r, column=1, value=f"Cộng {sc}").font = BOLD
    ws.cell(row=r, column=8, value=int(block.Quantity.sum())).font = BOLD
    for j in range(1, 10): ws.cell(row=r, column=j).fill = sub_fill
    log["sub_rows"].append(r); r += 1
    if sc == repeat_hdr_after:
        log["repeat_hdr_row"] = r; write_header_row(r); r += 1
ws.cell(row=r, column=1, value="TỔNG CỘNG").font = Font(name=F, bold=True, color="C00000")
ws.cell(row=r, column=8, value=exp_qty).font = Font(name=F, bold=True, color="C00000")
grand_row = r
ws.freeze_panes = "A5"
set_widths(ws, [15, 7, 12, 10, 11, 11, 40, 9, 11])

# Tổng cột Quantity nếu SUM nguyên cột khi chưa làm sạch (dùng cho insight mẫu)
dirty_sum = sum(int(q) for i, q in clean.Quantity.items() if i not in txt_qty and i not in neg_qty) \
            - sum(int(clean.Quantity[i]) for i in neg_qty) + exp_qty * 2

s = wb.create_sheet("Store"); write_df(s, stores[["Store_Code", "Store_Name", "City", "Region"]]); set_widths(s, [11, 24, 13, 9])
p = wb.create_sheet("Product"); write_df(p, products[["Product_Code", "Product_Name", "Category", "Unit_Price"]], formats={"Unit_Price": NUM_FMT}); set_widths(p, [13, 40, 18, 12])

tasks = [
    ("1", "Làm sạch cấu trúc sheet BaoCao_T4_2025",
     "Xóa tiêu đề báo cáo, các dòng \"Cửa hàng …\", \"Cộng …\", \"TỔNG CỘNG\", dòng tiêu đề bị lặp lại và dòng trống. Kết quả: 1 dòng tiêu đề ở dòng 1, ngay bên dưới là dữ liệu."),
    ("2", "Sửa kiểu dữ liệu", "Chuyển các ô Quantity và Order_Date đang ở dạng text thành số và ngày."),
    ("3", "Tạo Excel Table", "Ctrl + T, đặt tên bảng tblSales_T4. Đổi tên sheet thành Sales."),
    ("4", "Cột Revenue", "Thêm cột Revenue = [@Quantity]*[@Unit_Price]."),
    ("5", "Tổng doanh thu và % đóng góp",
     "Tạo sheet Kiem_tra, ô C5 = tổng doanh thu. Thêm cột Pct_Share = Revenue ÷ ô tổng, dùng tham chiếu $ để kéo/tự điền không bị lệch."),
    ("6", "Data Validation",
     "Store_Code: List lấy từ sheet Store. Product_Code: List lấy từ sheet Product. Quantity: Whole number > 0. Viết thông báo lỗi tiếng Việt."),
    ("7", "Circle Invalid Data",
     "Khoanh các ô sai sẵn có, đối chiếu sheet Store / Product để sửa về giá trị đúng. Làm lại đến khi không còn ô nào bị khoanh."),
    ("8", "Viết insight", "Tạo sheet Insight, viết 1 câu có đủ: con số, nguyên nhân, hàm ý (theo mẫu slide 40)."),
]
de["A1"] = "BÀI TẬP VỀ NHÀ BUỔI 1 (BẮT BUỘC) — LÀM SẠCH BÁO CÁO BÁN HÀNG THÁNG 4/2025"; de["A1"].font = TITLE
de["A2"] = "Nộp trước Buổi 2. Đặt tên file: B1_BTVN_<HoTen>.xlsx. Dùng đúng các kỹ năng của Buổi 1; chưa cần hàm SUMIFS, XLOOKUP."; de["A2"].font = NOTE
de["A3"] = "Nộp bài tại https://hocvien.mcivietnam.com/auth#/ — Tài khoản: email bạn đăng ký tại MCI · Mật khẩu: hv12345678"; de["A3"].font = BOLD
header(de, 5, ["#", "Việc cần làm", "Yêu cầu"])
for i, row in enumerate(tasks, start=6):
    for j, v in enumerate(row, start=1):
        c = de.cell(row=i, column=j, value=v); c.font = BODY; c.alignment = Alignment(wrap_text=True, vertical="top")
k = 6 + len(tasks) + 1
de.cell(row=k, column=1, value="Tự kiểm tra trước khi nộp").font = BOLD
checks_hv = [
    f"Ctrl + ↓ từ dòng dữ liệu đầu chạy thẳng tới dòng cuối; chọn ô Order_ID đầu tiên → Ctrl + Shift + ↓ (không gồm tiêu đề): Count trên Status Bar = {n:,}".replace(",", "."),
    f"Tổng doanh thu = {exp_rev:,} VNĐ".replace(",", "."),
    f"Tổng Quantity = {exp_qty:,}".replace(",", "."),
    "Tổng cột Pct_Share = 100%",
    "Circle Invalid Data: không còn ô nào bị khoanh",
]
for i, t in enumerate(checks_hv, start=k + 1):
    de.cell(row=i, column=2, value="☐ " + t).font = BODY
set_widths(de, [5, 34, 90])
wb.active = 0
wb.save(f"{HW}/B1_BTVN_MCI_Mart_T4_2025.xlsx")

# ---------------------------------------------------------------------
# ĐÁP ÁN: B1_BTVN_MCI_Mart_T4_2025_Answer.xlsx (chỉ giảng viên / trợ giảng chấm bài)
# ---------------------------------------------------------------------
wb = Workbook(); ws = wb.active; ws.title = "Sales"
ans_cols = cols + ["Revenue", "Pct_Share"]
header(ws, 1, ans_cols)
for i, row in enumerate(clean.itertuples(index=False), start=2):
    for j, v in enumerate(row, start=1):
        c = ws.cell(row=i, column=j, value=py(v)); c.font = BODY
        if j == 3: c.number_format = DATE_FMT
        if j == 9: c.number_format = NUM_FMT
    c = ws.cell(row=i, column=10, value="=tblSales_T4[[#This Row],[Quantity]]*tblSales_T4[[#This Row],[Unit_Price]]"); c.font = BODY; c.number_format = NUM_FMT
    c = ws.cell(row=i, column=11, value="=tblSales_T4[[#This Row],[Revenue]]/Kiem_tra!$C$5"); c.font = BODY; c.number_format = "0.000%"
t = Table(displayName="tblSales_T4", ref=f"A1:K{n + 1}")
t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(t); ws.freeze_panes = "A2"
set_widths(ws, [15, 7, 12, 10, 11, 11, 40, 9, 11, 13, 10])
for dv in [
    DataValidation(type="list", formula1="=Store!$A$2:$A$13", showErrorMessage=True,
                   errorTitle="Mã cửa hàng không hợp lệ", error="Chỉ nhập mã có trong sheet Store (ST01–ST12)."),
    DataValidation(type="list", formula1=f"=Product!$A$2:$A${len(products) + 1}", showErrorMessage=True,
                   errorTitle="Mã sản phẩm không hợp lệ", error="Chỉ nhập mã có trong sheet Product (P001–P040)."),
    DataValidation(type="whole", operator="greaterThan", formula1="0", showErrorMessage=True,
                   errorTitle="Số lượng không hợp lệ", error="Quantity phải là số nguyên lớn hơn 0."),
]:
    ws.add_data_validation(dv)
ws.data_validations.dataValidation[0].add(f"D2:D{n + 1}")
ws.data_validations.dataValidation[1].add(f"F2:F{n + 1}")
ws.data_validations.dataValidation[2].add(f"H2:H{n + 1}")

s = wb.create_sheet("Store"); write_df(s, stores[["Store_Code", "Store_Name", "City", "Region"]]); set_widths(s, [11, 24, 13, 9])
p = wb.create_sheet("Product"); write_df(p, products[["Product_Code", "Product_Name", "Category", "Unit_Price"]], formats={"Unit_Price": NUM_FMT}); set_widths(p, [13, 40, 18, 12])

k3 = wb.create_sheet("Kiem_tra")
k3["A1"] = "KIỂM TRA / ĐỐI SOÁT — Bài tập về nhà B1"; k3["A1"].font = TITLE
k3["A2"] = "Chấm bài: so từng dòng của học viên với cột Đáp án. C5 là ô tổng doanh thu mà cột Pct_Share tham chiếu tuyệt đối ($C$5)."; k3["A2"].font = NOTE
header(k3, 4, ["Chỉ tiêu", "", "Kết quả (công thức)", "Đáp án", "Khớp?"])
checks = [
    ("Tổng doanh thu tháng 4/2025 (VNĐ)", "=SUM(tblSales_T4[Revenue])", exp_rev, NUM_FMT),
    ("Số dòng dữ liệu", "=ROWS(tblSales_T4[Order_ID])", n, NUM_FMT),
    ("Tổng số lượng (Quantity)", "=SUM(tblSales_T4[Quantity])", exp_qty, NUM_FMT),
    ("Tổng cột Pct_Share", "=SUM(tblSales_T4[Pct_Share])", 1, "0.00%"),
    ("Số dòng Store_Code không có trong danh mục", "=SUMPRODUCT(--(COUNTIF(Store!$A$2:$A$13,tblSales_T4[Store_Code])=0))", 0, "0"),
    ("Số dòng Product_Code không có trong danh mục", f"=SUMPRODUCT(--(COUNTIF(Product!$A$2:$A${len(products) + 1},tblSales_T4[Product_Code])=0))", 0, "0"),
    ("Số dòng Quantity ≤ 0", '=COUNTIF(tblSales_T4[Quantity],"<=0")', 0, "0"),
    ("Số ô Quantity dạng text", "=SUMPRODUCT(--ISTEXT(tblSales_T4[Quantity]))", 0, "0"),
    ("Số ô Order_Date dạng text", "=SUMPRODUCT(--ISTEXT(tblSales_T4[Order_Date]))", 0, "0"),
]
for k, (label, f, exp, fmt) in enumerate(checks, start=5):
    k3.cell(row=k, column=1, value=label).font = BODY
    c = k3.cell(row=k, column=3, value=f); c.number_format = fmt; c.font = BODY
    c = k3.cell(row=k, column=4, value=exp); c.number_format = fmt; c.font = Font(name=F, size=10, color="0000FF")
    k3.cell(row=k, column=5, value=f'=IF(ROUND(C{k},6)=ROUND(D{k},6),"✔","✘")').font = BODY
set_widths(k3, [46, 2, 20, 16, 8])

e = wb.create_sheet("Dap_an_loi")
e["A1"] = "ĐÁP ÁN — CÁC LỖI CÀI SẴN TRONG ĐỀ VÀ CÁCH SỬA"; e["A1"].font = TITLE
rows = [
    ("#", "Loại lỗi", "Số lượng", "Vị trí / nhận biết", "Cách sửa"),
    (1, "Tiêu đề báo cáo (dòng 1–2) và dòng trống (dòng 3)", "3 dòng", "Trên cùng sheet", "Xóa dòng 1–3"),
    (2, "Dòng nhóm \"Cửa hàng STxx — …\" (ô gộp A:I)", f"{len(store_list)} dòng", "Đầu mỗi khối cửa hàng, nền vàng", "Filter cột A chứa \"Cửa hàng\" → xóa"),
    (3, "Dòng cộng \"Cộng STxx\" và \"TỔNG CỘNG\"", f"{len(store_list)} + 1 dòng", f"Cuối mỗi khối, nền xám; TỔNG CỘNG ở dòng {grand_row}", "Filter cột A chứa \"Cộng\" → xóa"),
    (4, "Dòng tiêu đề lặp lại giữa dữ liệu", "1 dòng", f"Dòng {log['repeat_hdr_row']} (sau khối {repeat_hdr_after})", "Xóa dòng"),
    (5, "Dòng trống giữa dữ liệu", "1 dòng", f"Dòng {log['blank_row']} (giữa khối {blank_store})", "Xóa dòng; kiểm tra bằng Ctrl + ↓"),
    (6, "Quantity lưu dạng text", f"{len(txt_qty)} ô", "Tam giác xanh, căn trái", "Error Checking → Convert to Number"),
    (7, "Order_Date lưu dạng text dd/mm/yyyy", f"{len(txt_date)} ô", "Căn trái", "DATEVALUE hoặc Text to Columns (DMY)"),
    (8, "Store_Code sai \"ST5\"", f"{len(bad_store)} ô", "Circle Invalid Data (List từ Store)", "Sửa thành ST05 (nằm trong khối ST05)"),
    (9, "Product_Code thiếu số 0 (\"P25\" thay vì \"P025\")", f"{len(bad_prod)} ô", "Circle Invalid Data (List từ Product)", "Đối chiếu Product_Name với sheet Product, sửa về mã 3 chữ số"),
    (10, "Quantity âm", f"{len(neg_qty)} ô", "Circle Invalid Data (Whole > 0)", "Lỗi dấu khi nhập — đổi thành số dương"),
]
for i, rr in enumerate(rows, start=3):
    for j, v in enumerate(rr, start=1):
        c = e.cell(row=i, column=j, value=v); c.font = HDR_FONT if i == 3 else BODY
        if i == 3: c.fill = HDR_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
set_widths(e, [4, 44, 14, 40, 54])
e.cell(row=len(rows) + 4, column=1, value=f"Sau khi sửa: {n:,} dòng dữ liệu; tổng doanh thu {exp_rev:,} VNĐ; tổng Quantity {exp_qty:,}.".replace(",", ".")).font = BOLD

ins = wb.create_sheet("Insight_mau")
ins["A1"] = "INSIGHT MẪU (tham khảo khi chấm; học viên viết theo cách của mình)"; ins["A1"].font = TITLE
vn = lambda x: f"{x:,}".replace(",", ".")                      # chỉ đổi dấu nghìn của con số, không đụng dấu câu
ins["A3"] = (f"\"Báo cáo tháng 4/2025 có {vn(n)} dòng giao dịch và tôi tìm thấy 10 loại lỗi. Lỗi nguy hiểm nhất là {len(store_list) + 1} dòng cộng "
             f"(\"Cộng STxx\" và \"TỔNG CỘNG\") nằm lẫn trong dữ liệu, nên nếu không sửa thì SUM cột Quantity ra {vn(dirty_sum)} thay vì {vn(exp_qty)}, "
             f"gần gấp ba, và mọi báo cáo theo cửa hàng đều sai theo.\"")
ins["A3"].font = BODY; ins["A3"].alignment = Alignment(wrap_text=True, vertical="top")
ins["A5"] = "Chấm: đủ 3 thành phần (con số · nguyên nhân · hàm ý) và con số khớp dữ liệu."; ins["A5"].font = NOTE
set_widths(ins, [120])
wb.save(f"{HW}/B1_BTVN_MCI_Mart_T4_2025_Answer.xlsx")
print("ok", n, exp_rev, exp_qty, "dirty SUM Quantity:", dirty_sum, log)
