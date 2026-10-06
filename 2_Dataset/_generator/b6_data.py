"""Nạp dữ liệu nguồn dùng chung cho các script Buổi 6 (cùng nguồn build_b5_meta.py / build_b6_meta.py).

Trả về: s (Sales đã nối Product/Store/Customer, thêm COGS, Y, M, Q), product, store, customer, target, daily (doanh thu theo ngày).
"""
import os
import pandas as pd
from openpyxl import load_workbook

ROOT = os.path.normpath(os.environ.get("MCI_ROOT") or os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
B4P = os.path.join(ROOT, "1_Sessions", "B4_Power_Query", "Practice")
STARTER = os.path.join(B4P, "B4_MCI_Mart_Answer.xlsx")
MASTER = os.path.join(ROOT, "2_Dataset", "MCI_Mart_Master.xlsx")

wb = load_workbook(STARTER, read_only=True, data_only=True)
rows = list(wb["Sales_PQ"].iter_rows(values_only=True))
wb.close()
sales = pd.DataFrame(rows[1:], columns=rows[0])
sales["Order_Date"] = pd.to_datetime(sales.Order_Date)
lk = pd.ExcelFile(os.path.join(B4P, "B4_Lookups.xlsx"))
product = lk.parse("Product").drop_duplicates("Product_Code")
store = lk.parse("Store")
customer = lk.parse("Customer")
tw = pd.read_excel(os.path.join(B4P, "B4_Target_2025.xlsx"), sheet_name="Target_2025", header=2)
target = tw.melt(id_vars=["Store_Code", "Store_Name"], var_name="Attribute", value_name="Target")
target["Thang"] = target.Attribute.str.split(" ").str[-1].astype(int)
target = target[["Store_Code", "Thang", "Target"]]

master = pd.read_excel(MASTER, sheet_name="Sales")
assert len(sales) == len(master) == 30247 and int(sales.Revenue.sum()) == 4436849500

s = sales.merge(product[["Product_Code", "Product_Name", "Category", "Unit_Cost", "Unit_Price"]], on="Product_Code", how="left") \
         .merge(store[["Store_Code", "Store_Name", "City", "Region", "Store_Type", "Open_Date"]], on="Store_Code", how="left") \
         .merge(customer[["Customer_ID", "Segment"]], on="Customer_ID", how="left")
s["COGS"] = s.Quantity * s.Unit_Cost
s["Y"], s["M"], s["Q"] = s.Order_Date.dt.year, s.Order_Date.dt.month, s.Order_Date.dt.quarter
daily = s.groupby("Order_Date").Revenue.sum().astype(int)
assert len(daily) == 486
