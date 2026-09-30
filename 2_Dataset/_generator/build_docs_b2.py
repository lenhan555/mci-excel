"""Dựng tài liệu phát cho Buổi 2 (.docx) bằng python-docx: handout tự học và mẫu KPI Definition Card.

Số liệu ví dụ tính từ dataset (xem build_b2.py). Xuất PDF bằng Microsoft Word (xem 2_Dataset/README.md).
    python3 _generator/build_docs_b2.py
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.normpath(os.environ.get("MCI_ROOT") or os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
OUT = f"{ROOT}/1_Sessions/B2_Ham_va_KPI/Handout"
IMG = f"{ROOT}/Materials/B2/Images"   # ảnh chụp thật, danh sách: _DANH_SACH_ANH.md
MISSING = []
os.makedirs(OUT, exist_ok=True)
RED, NAVY, GOLD, GREY = RGBColor(0xC0, 0, 0), RGBColor(0x20, 0x38, 0x64), RGBColor(0x8A, 0x65, 0x00), RGBColor(0x58, 0x58, 0x58)

def base_doc(footer="MCI – Magic Code Institute · Excel for Business Intelligence Analyst · Buổi 2"):
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for m in ("left_margin", "right_margin"): setattr(sec, m, Cm(2))
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    for name, size, color in (("Heading 1", 16, NAVY), ("Heading 2", 13, NAVY), ("Heading 3", 11.5, RED)):
        h = d.styles[name]; h.font.name = "Arial"; h.font.size = Pt(size); h.font.color.rgb = color; h.font.bold = True
        rf = h.element.rPr.rFonts
        for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(k)) is not None: del rf.attrib[qn(k)]
        for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"): rf.set(qn(k), "Arial")
    fp = sec.footer.paragraphs[0]
    fp.text = footer
    fp.runs[0].font.size = Pt(8); fp.runs[0].font.color.rgb = GREY
    return d

def para(d, text="", bold=False, color=None, size=None, italic=False):
    p = d.add_paragraph()
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic
        if color: r.font.color.rgb = color
        if size: r.font.size = Pt(size)
    return p

def rich(d, parts):
    p = d.add_paragraph()
    for t, f in parts:
        r = p.add_run(t); r.bold = f.get("bold", False); r.italic = f.get("italic", False)
        if f.get("color"): r.font.color.rgb = f["color"]
    return p

def bullet(d, text):
    p = d.add_paragraph(style="List Bullet"); p.add_run(text); return p

def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_)
    tcPr.append(shd)

def table(d, rows, widths, header=True, size=9.5, first_bold=False, min_h=None):
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for j, w in enumerate(widths): t.columns[j].width = Cm(w)
    for i, row in enumerate(rows):
        if min_h and i: t.rows[i].height = Cm(min_h)
        for j, v in enumerate(row):
            c = t.cell(i, j); c.width = Cm(widths[j]); c.text = ""
            r = c.paragraphs[0].add_run(str(v)); r.font.size = Pt(size)
            if i < len(rows) - 1: c.paragraphs[0].paragraph_format.keep_with_next = True   # không tách bảng qua trang
            if header and i == 0:
                r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); shade(c, "203864")
            elif first_bold and j == 0:
                r.bold = True
    d.add_paragraph()
    return t

def figure(d, what, where, img, width=15):
    """Chèn ảnh chụp thật từ Materials/B2/Images nếu đã có; chưa có thì để khung chờ [CHỤP] (CLAUDE.md, mục Ảnh minh họa)."""
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(f"{IMG}/{img}"):
        p.add_run().add_picture(f"{IMG}/{img}", width=Cm(width)); return p
    MISSING.append(img)
    r = p.add_run(f"[CHỤP] {what}"); r.italic = True; r.font.color.rgb = GREY; r.font.size = Pt(9)
    r = p.add_run(f"\n{where} → {img}"); r.font.color.rgb = GREY; r.font.size = Pt(8)
    pPr = p._p.get_or_add_pPr(); bdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "dashed"); e.set(qn("w:sz"), "6"); e.set(qn("w:color"), "A6A6A6"); e.set(qn("w:space"), "8"); bdr.append(e)
    pPr.append(bdr)
    return p

def tool(d, name, purpose, how=None, example=None, tip=None, protip=None, shot=None):
    """Khối công cụ theo style khóa học: tên (đỏ) → mục đích → cách dùng → bảng ví dụ trên case → ảnh chụp → TIP / PRO TIP.
    shot = (nội dung ảnh, file · sheet · vùng, tên file ảnh)."""
    d.add_heading(name, level=3).paragraph_format.keep_with_next = True
    para(d, purpose).paragraph_format.keep_with_next = True
    for h in how or []: bullet(d, h)
    if example: table(d, example[0], example[1])
    if shot: figure(d, *shot)
    if tip: rich(d, [("TIP: ", {"bold": True, "color": NAVY}), (tip, {})])
    if protip: rich(d, [("PRO TIP: ", {"bold": True, "color": GOLD}), (protip, {})])

# =====================================================================
# HANDOUT B2
# =====================================================================
d = base_doc()
para(d, "MCI – MAGIC CODE INSTITUTE", bold=True, color=GREY, size=9)
d.add_heading("Handout Buổi 2 — Hàm cho phân tích: từ dữ liệu giao dịch đến KPI", level=1)
para(d, "Tài liệu tự học đi kèm Buổi 2. Các hàm ở đây KHÔNG giảng chi tiết trên lớp; đọc sau buổi học rồi làm bài tập về nhà (bắt buộc) ở mục 9. "
        "Ví dụ viết trên file đáp án B2_MCI_Mart_Q1_2025_Answer.xlsx (bảng tblSales đã làm giàu, tblProduct, tblStore, sheet KPI_SUMIFS).", italic=True, color=GREY)
para(d, "Dấu phân cách: máy cài định dạng vùng Việt Nam dùng dấu chấm phẩy ( ; ) giữa các đối số. Tài liệu viết theo dấu phẩy ( , ); hãy đổi nếu máy bạn dùng ( ; ).", italic=True, color=GREY)

d.add_heading("Mục lục", level=2)
for s in ["1. KPI Definition Card · Dimension, Measure, Grain", "2. Hàm logic bổ sung: SWITCH, IF lồng", "3. Tổng hợp có điều kiện nâng cao: MAXIFS/MINIFS, tiêu chí động, SUMPRODUCT",
          "4. Tra cứu bổ sung: HLOOKUP, CHOOSE, INDEX/MATCH hai chiều, OFFSET, INDIRECT, ROW/COLUMN",
          "5. Text & Date bổ sung: SEARCH/FIND, HYPERLINK, TODAY/NOW, WEEKDAY, WORKDAY, YEARFRAC, DATEDIF",
          "6. Số ngẫu nhiên: RAND, RANDBETWEEN, RANDARRAY", "7. Conditional Formatting bằng công thức",
          "8. PRO TIP: Dynamic Arrays (FILTER, SORT, UNIQUE, LET), Top 5, tra cứu gần đúng theo mẫu", "9. Bài tập về nhà (bắt buộc)", "10. Checklist trước Buổi 3"]:
    bullet(d, s)

# 1 KPI card
d.add_heading("1. KPI Definition Card · Dimension, Measure, Grain", level=2)
para(d, "Viết card TRƯỚC khi viết công thức. Card trả lời: KPI tên gì, tính thế nào, lấy dữ liệu ở đâu, xem theo chiều nào, ai chịu trách nhiệm, mốc nào là đạt. "
        "Mẫu in 1 trang: B2_KPI_Definition_Card.pdf.")
table(d, [["Khái niệm", "Câu hỏi", "Trên tblSales"],
          ["Dimension", "Theo cái gì?", "Region, Store_Code, Thang, Category, Phan_khuc"],
          ["Measure", "Đo cái gì?", "Quantity, Revenue, COGS, Gross_Margin"],
          ["Grain", "Mỗi dòng là gì?", "1 sản phẩm trong 1 đơn (Order_ID × Line_No): 6.187 dòng, 3.144 đơn"]], [3, 3.4, 9.6], first_bold=True)
table(d, [["KPI (Q1/2025)", "Công thức", "Giá trị", "Ghi chú"],
          ["Total Revenue", "SUM(Revenue)", "941.454.000", "Cộng dồn được theo mọi chiều"],
          ["Total Orders", "COUNTIFS(…, Line_No, 1)", "3.144", "Không đếm số dòng (6.187)"],
          ["Avg Order Value", "Revenue ÷ Orders", "299.445", "Tổng ÷ tổng; trung vị giá trị đơn chỉ 198.000"],
          ["Gross Margin %", "(Revenue − COGS) ÷ Revenue", "34,2%", "AVERAGE 6 ngành ra 31,2% — sai"]], [3.4, 4.4, 2.4, 5.8], first_bold=True)
figure(d, "Sheet KPI_Card với 3 card đã điền và dòng giá trị Q1/2025", "Practice/B2_MCI_Mart_Q1_2025_Answer.xlsx · sheet KPI_Card · A4:D16", "H01_kpi_card.png")

# 2 Logic
d.add_heading("2. Hàm logic bổ sung", level=2)
tool(d, "SWITCH", "So một giá trị với danh sách trường hợp, trả về kết quả tương ứng. Gọn hơn IFS khi so bằng (=).",
     how=["Cú pháp: =SWITCH(giá_trị, trường_hợp_1, kết_quả_1, …, [mặc_định])."],
     example=([["Công thức (trong tblSales)", "Kết quả"],
               ['=SWITCH([@Region],"Bắc","North","Trung","Central","Nam","South")', "North / Central / South"],
               ['=SWITCH([@Quy],"Q1","Quý 1","Q2","Quý 2","Khác")', "Quý 1"]], [11, 5]),
     tip="SWITCH chỉ so bằng. Cần so lớn hơn / nhỏ hơn (ngưỡng giá) thì dùng IFS.")
tool(d, "IF lồng (Excel 2016/2019)", "Máy chưa có IFS thì lồng IF: điều kiện sau nằm trong nhánh “ngược lại” của điều kiện trước.",
     example=([["IFS", "IF lồng tương đương"],
               ['=IFS([@Unit_Price]>=100000,"Cao",[@Unit_Price]>=30000,"Trung",TRUE,"Thấp")',
                '=IF([@Unit_Price]>=100000,"Cao",IF([@Unit_Price]>=30000,"Trung","Thấp"))']], [8, 8]))

# 3 Aggregation
d.add_heading("3. Tổng hợp có điều kiện nâng cao", level=2)
tool(d, "MAXIFS / MINIFS", "Giá trị lớn nhất / nhỏ nhất có điều kiện, cùng cấu trúc với SUMIFS.",
     example=([["Công thức", "Kết quả Q1/2025"],
               ['=MAXIFS(tblSales[Revenue],tblSales[Region],"Nam")', "1.110.000 (dòng lớn nhất Miền Nam)"],
               ['=MAXIFS(tblSales[Revenue],tblSales[Region],"Bắc")', "1.470.000 (HD2501-00009: 6 chảo chống dính)"],
               ['=MINIFS(tblSales[Revenue],tblSales[Region],"Bắc")', "19.000"]], [10, 6]))
tool(d, "Tiêu chí động trong SUMIFS / COUNTIFS", "Tiêu chí có thể là phép so sánh ghép với ô, hoặc ký tự đại diện * ? .",
     example=([["Công thức", "Ý nghĩa"],
               ['=COUNTIFS(tblSales[Revenue],">="&H1)', "Số dòng có doanh thu ≥ giá trị ở ô H1 (H1 = 1000000 → 13 dòng)"],
               ['=SUMIFS(tblSales[Revenue],tblSales[Order_Date],">="&DATE(2025,1,1),tblSales[Order_Date],"<="&EOMONTH(DATE(2025,1,1),0))', "Doanh thu tháng 1 lọc theo khoảng ngày: 425.850.000"],
               ['=COUNTIFS(tblSales[Product_Name],"*hộp*")', "Số dòng có tên sản phẩm chứa “hộp” (không phân biệt hoa thường)"]], [10, 6]),
     tip="Dấu so sánh phải nằm trong ngoặc kép và nối với ô bằng &: \">=\"&H1, không viết >=H1.")
tool(d, "SUMPRODUCT", "Nhân từng cặp phần tử rồi cộng lại; dùng được như một “SUMIFS tự do” với điều kiện bất kỳ.",
     example=([["Công thức", "Kết quả"],
               ["=SUMPRODUCT(tblSales[Quantity],tblSales[Unit_Price])", "941.454.000 (= tổng Revenue, không cần cột phụ)"],
               ['=SUMPRODUCT((tblSales[Region]="Bắc")*(tblSales[Thang]=3)*tblSales[Revenue])', "90.210.000 (giống SUMIFS)"]], [10, 6]),
     protip="Mỗi điều kiện (…=…) là một mảng TRUE/FALSE; nhân với nhau = AND. Chậm hơn SUMIFS trên bảng lớn — chỉ dùng khi SUMIFS không làm được (ví dụ điều kiện là phép tính trên cột).")

# 4 Lookup
d.add_heading("4. Tra cứu bổ sung", level=2)
tool(d, "HLOOKUP", "Như VLOOKUP nhưng tìm ở DÒNG ĐẦU và trả về dòng thứ n bên dưới — dùng khi bảng xếp tháng theo cột.",
     example=([["Công thức (sheet KPI_SUMIFS)", "Kết quả"],
               ["=HLOOKUP(2,KPI_SUMIFS!$B$5:$E$9,2,FALSE)", "87.954.000 (Bắc, tháng 2)"]], [10, 6]),
     tip="Luôn ghi FALSE như VLOOKUP. XLOOKUP làm được cả chiều ngang: =XLOOKUP(2,$B$5:$D$5,$B$6:$D$6).")
tool(d, "INDEX / MATCH hai chiều", "Tra một ô theo cả dòng và cột: MATCH cho dòng, MATCH cho cột.",
     example=([["Công thức (sheet KPI_SUMIFS)", "Kết quả"],
               ['=INDEX($B$6:$D$8,MATCH("Nam",$A$6:$A$8,0),MATCH(3,$B$5:$D$5,0))', "129.438.000 (Nam, tháng 3)"]], [10, 6]),
     shot=("Gõ công thức INDEX/MATCH hai chiều ở G6, nhấn F2: vùng B6:D8, A6:A8, B5:D5 được tô màu; kết quả 129.438.000", "Practice/B2_MCI_Mart_Q1_2025_Answer.xlsx · sheet KPI_SUMIFS · A4:G9", "H02_index_match_2d.png"))
tool(d, "CHOOSE", "Chọn giá trị thứ n trong danh sách.",
     example=([["Công thức", "Kết quả"], ['=CHOOSE(ROUNDUP(MONTH([@Order_Date])/3,0),"Q1","Q2","Q3","Q4")', "Q1 cho mọi dòng Quý 1"]], [10, 6]))
tool(d, "OFFSET · INDIRECT", "OFFSET trả về vùng dịch đi n dòng/cột từ một ô; INDIRECT biến chuỗi chữ thành tham chiếu.",
     example=([["Công thức", "Kết quả"],
               ["=SUM(OFFSET(KPI_SUMIFS!$B$9,0,1,1,2))", "515.604.000 (tổng tháng 2 + 3, dòng Tổng)"],
               ['=INDIRECT("KPI_SUMIFS!E9")', "941.454.000"]], [10, 6]),
     tip="Cả hai là hàm “volatile”: tính lại mỗi khi sửa bất kỳ ô nào, làm file chậm và khó kiểm tra. Ưu tiên INDEX/XLOOKUP; khóa học không dùng hai hàm này ở phần lõi.")
tool(d, "ROW / COLUMN", "Số dòng / số cột của một ô. Hay dùng để đánh số thứ tự tự động.",
     example=([["Công thức", "Kết quả"], ["=ROW()-1 (ở dòng 2 của bảng)", "1, 2, 3… khi kéo xuống"], ["=ROWS(tblSales[Order_ID])", "6.187 (số dòng của cột)"]], [10, 6]))

# 5 Text & Date
d.add_heading("5. Text & Date bổ sung", level=2)
tool(d, "SEARCH / FIND", "Tìm vị trí một chuỗi con. SEARCH không phân biệt hoa thường và nhận ký tự đại diện; FIND phân biệt hoa thường.",
     example=([["Công thức (trên tblProduct)", "Kết quả trên 40 sản phẩm"],
               ['=ISNUMBER(SEARCH("hộp",[@Product_Name]))', "8 sản phẩm TRUE (gồm “Hộp quà Tết…”, “Hộp đựng…”)"],
               ['=ISNUMBER(FIND("hộp",[@Product_Name]))', "6 sản phẩm TRUE (bỏ qua chữ “Hộp” viết hoa)"]], [10, 6]),
     shot=("Thêm tạm 2 cột SEARCH / FIND vào tblProduct, lọc TRUE; dòng “Hộp quà Tết cao cấp” khác nhau giữa hai cột", "Practice/B2_MCI_Mart_Q1_2025_Answer.xlsx · sheet Product · A1:G41", "H03_search_find.png"))
tool(d, "HYPERLINK", "Tạo liên kết bấm được tới trang web, file hoặc ô trong file.",
     example=([["Công thức", "Kết quả"], ['=HYPERLINK("#Kiem_tra!A1","Mở sheet Kiểm tra")', "Bấm để nhảy tới sheet Kiem_tra"]], [10, 6]))
tool(d, "TODAY / NOW · WEEKDAY · WORKDAY", "Ngày hôm nay / ngày giờ hiện tại; thứ trong tuần dạng số; ngày làm việc thứ n sau một ngày (bỏ qua cuối tuần và ngày lễ).",
     example=([["Công thức", "Kết quả"],
               ["=TODAY()  ·  =NOW()", "Ngày hôm nay · ngày giờ hiện tại (tự đổi mỗi lần mở file)"],
               ["=WEEKDAY(DATE(2025,1,29),2)", "3 (thứ Tư, mùng 1 Tết Ất Tỵ)"],
               ["=WORKDAY(DATE(2025,1,24),3,H10:H15)", "05/02/2025 (với danh sách ngày lễ 01/01 và 27–31/01/2025)"]], [10, 6]),
     shot=("Khối NETWORKDAYS (tháng 1: 23 → 17), danh sách ngày lễ H10:H15 và một ô =WORKDAY(DATE(2025,1,24),3,H10:H15) ra 05/02/2025", "Demo/B2_Demo_Text_Date_HocVien.xlsx · sheet Vi_du_Date · A8:H15", "H04_workday.png"))
tool(d, "YEARFRAC · DATEDIF", "Số năm (có phần lẻ) / số tháng tròn giữa hai ngày — ví dụ tuổi cửa hàng.",
     example=([["Công thức (tblStore, ST12 mở 01/06/2024)", "Kết quả"],
               ["=YEARFRAC([@Open_Date],DATE(2025,3,31))", "≈ 0,83 năm"],
               ['=DATEDIF([@Open_Date],DATE(2025,3,31),"m")', "9 tháng"]], [10, 6]),
     tip="DATEDIF không hiện trong danh sách gợi ý của Excel nhưng vẫn dùng được; ngày bắt đầu phải nhỏ hơn ngày kết thúc, nếu không ra #NUM!.")

# 6 Random
d.add_heading("6. Số ngẫu nhiên", level=2)
tool(d, "RAND · RANDBETWEEN · RANDARRAY", "Sinh số ngẫu nhiên: lấy mẫu kiểm tra, tạo dữ liệu thử.",
     example=([["Công thức", "Kết quả"], ["=RAND()", "Số thực trong [0, 1)"], ["=RANDBETWEEN(1,6187)", "Một số dòng ngẫu nhiên để kiểm tra tay"],
               ["=RANDARRAY(5,1,1,6187,TRUE)", "5 số dòng ngẫu nhiên (Excel 365, tràn xuống 5 ô)"]], [10, 6]),
     tip="Kết quả đổi mỗi lần Excel tính lại. Muốn giữ cố định: copy → Paste Values (Alt H V V).")

# 7 CF formula
d.add_heading("7. Conditional Formatting bằng công thức", level=2)
tool(d, "New Rule → Use a formula to determine which cells to format",
     "Tô màu cả dòng theo điều kiện: công thức trả về TRUE thì ô được định dạng.",
     how=["Chọn vùng dữ liệu của tblSales (không gồm tiêu đề), ô hiện hành ở dòng 2.",
          "Home → Conditional Formatting → New Rule → Use a formula… → gõ =$T2=\"Cao\" (cột T = Phan_khuc) → Format → chọn màu nền.",
          "Khóa CỘT ($T) để mọi ô trong dòng cùng nhìn vào cột Phan_khuc; KHÔNG khóa dòng (2) để quy tắc trôi theo từng dòng."],
     shot=("Hộp New Formatting Rule với =$T2=\"Cao\" và bảng tblSales có các dòng phân khúc Cao được tô màu", "Practice/B2_MCI_Mart_Q1_2025_Answer.xlsx · sheet Sales · A1:T21", "H05_cf_formula.png"),
     tip="Đây chính là tham chiếu hỗn hợp của Buổi 1. Quản lý các quy tắc: Conditional Formatting → Manage Rules.")

# 8 PRO TIP dynamic arrays
d.add_heading("8. PRO TIP: Dynamic Arrays và tra cứu theo mẫu (Excel 365)", level=2)
para(d, "Một công thức trả về nhiều ô, kết quả tự “tràn” (spill) xuống dưới. Tham chiếu cả vùng tràn bằng dấu #, ví dụ A4#. File thực hành: Demo/B2_Demo_SUMIFS_HocVien.xlsx, sheet Vi_du_DynamicArray.")
table(d, [["Hàm", "Công thức", "Kết quả"],
          ["UNIQUE + SORT", "=SORT(UNIQUE(tblSales[Store_Code]))", "12 mã ST01…ST12"],
          ["FILTER", '=FILTER(tblSales[[Order_ID]:[Revenue]],(tblSales[Region]="Trung")*(tblSales[Revenue]>=1000000))', "Các dòng Miền Trung có doanh thu ≥ 1.000.000"],
          ["SUMIFS với vùng tràn", "=SUMIFS(tblSales[Revenue],tblSales[Store_Code],A4#)", "Doanh thu 12 cửa hàng, tự dài theo danh sách"],
          ["Top 5 (SORTBY + TAKE)", "=TAKE(SORTBY(UNIQUE(tblSales[Product_Code]),SUMIFS(tblSales[Revenue],tblSales[Product_Code],UNIQUE(tblSales[Product_Code])),-1),5)",
           "P027, P039, P013, P028, P023"],
          ["LET", "=LET(r,SUM(tblSales[Revenue]),o,COUNTIFS(tblSales[Line_No],1),r/o)", "299.445 (AOV) — đặt tên cho kết quả trung gian"],
          ["XLOOKUP theo mẫu", '=XLOOKUP("*quà*",tblProduct[Product_Name],tblProduct[Product_Code],,2)', "P027 (match_mode 2 = ký tự đại diện)"]],
      [3.2, 8.6, 4.2], size=9)
figure(d, "Vùng tràn của =SORT(UNIQUE(…)) (viền xanh) và kết quả FILTER; ô A21 dùng A4#", "Demo/B2_Demo_SUMIFS_HocVien.xlsx · sheet Vi_du_DynamicArray · A3:I22", "H06_dynamic_array.png")
rich(d, [("Tìm giá trị trùng / không trùng: ", {"bold": True}), ("=COUNTIF(tblProduct[Product_Code],[@Product_Code])>1 đánh dấu mã trùng (bài P021); Home → Conditional Formatting → Highlight Cells Rules → Duplicate Values để tô màu; "
                                                               "Data → Remove Duplicates để xóa (cẩn thận: xóa luôn, hãy kiểm tra trước — Buổi 4 làm việc này bằng Power Query).", {})])
figure(d, "Duplicate Values tô màu hai dòng P021 (dòng 22–23)", "Practice/B2_Product_DuplicateKey.xlsx · sheet Product · A18:E26", "H07_duplicate_values.png", width=12)

# 9 Homework
d.add_heading("9. Bài tập về nhà (bắt buộc)", level=2)
rich(d, [("File đề: ", {"bold": True}), ("B2_BTVN_MCI_Mart_T4_2025.xlsx — dữ liệu tháng 4/2025 đã làm sạch (tblSales_T4, 1.919 dòng) cùng bảng Store, Product. "
                                         "Nộp trước Buổi 3, tên file B2_BTVN_<HoTen>.xlsx. Sheet De_bai ghi đủ yêu cầu.", {})])
rich(d, [("Nộp bài: ", {"bold": True, "color": NAVY}), ("https://hocvien.mcivietnam.com/auth#/ — Tài khoản: email bạn đăng ký tại MCI · Mật khẩu: hv12345678", {})])
for s in ["XLOOKUP Region (tblStore), Category và Unit_Cost (tblProduct). Có ô #N/A thì tìm nguyên nhân (gợi ý: =LEN của mã) và sửa, không dùng IFERROR để che.",
          "Thêm COGS, Gross_Margin, Phan_khuc (IFS theo Unit_Price, ngưỡng như trên lớp).",
          "Sheet KPI: bảng Doanh thu Miền × Ngành hàng bằng MỘT công thức SUMIFS kéo ra cả bảng ($ hỗn hợp), có dòng và cột Tổng.",
          "Cùng sheet: bảng theo Miền gồm Doanh thu, Số đơn (Line_No = 1), AOV, Gross Margin %; dòng Tổng tính tổng ÷ tổng.",
          "Sheet KPI_Card: card cho Avg Order Value (10 trường). Sheet Insight: 1 câu có con số có so sánh, nguyên nhân khả dĩ, hàm ý."]:
    bullet(d, s)
rich(d, [("Tự kiểm tra trước khi nộp: ", {"bold": True, "color": NAVY}),
         ("không còn #N/A; doanh thu 296.965.000 VNĐ; 966 đơn; COGS 196.595.000 VNĐ; Gross Margin 33,8%.", {})])

# 10 Checklist
h = d.add_heading("10. Checklist trước Buổi 3", level=2); h.paragraph_format.keep_with_next = True
for s in ["File B2 của bạn: tblSales có Category, Unit_Cost, COGS, Gross_Margin, Thang, Quy, Thu, Phan_khuc; không còn #N/A.",
          "tblProduct không còn mã trùng (P021 chỉ còn 1 dòng); tổng COGS Q1 = 619.663.000.",
          "Sheet KPI_SUMIFS: tổng doanh thu 941.454.000; 3.144 đơn. GIỮ sheet này: Buổi 3 và Buổi 5 dùng để đối soát.",
          "Đủ 3 KPI Definition Card (Total Revenue, Total Orders, Gross Margin %).",
          "Đã viết câu insight cuối Buổi 2.",
          "Đã nộp bài tập về nhà (B2_BTVN_<HoTen>.xlsx), tự kiểm tra khớp 296.965.000 VNĐ và 966 đơn.",
          "Nếu file chưa khớp: Buổi 3 lớp phát file khởi đầu (đáp án B2).",
          "Dùng Mac: đã có máy Windows (máy ảo hoặc laptop Windows khác) — hạn chót trước Buổi 4."]:
    bullet(d, "☐  " + s).paragraph_format.keep_with_next = True
d.paragraphs[-1].paragraph_format.keep_with_next = False
d.save(f"{OUT}/B2_Handout.docx")

# =====================================================================
# MẪU KPI DEFINITION CARD (1 trang, in 3 tờ / học viên)
# =====================================================================
d = base_doc()
para(d, "MCI – MAGIC CODE INSTITUTE", bold=True, color=GREY, size=9)
d.add_heading("KPI Definition Card", level=1)
para(d, "Viết card TRƯỚC khi viết công thức. Mỗi KPI một tờ. Chữ xám là gợi ý — xóa hoặc viết đè.", italic=True, color=GREY)
rows = [["Trường", "Nội dung"],
        ["Tên KPI", ""], ["Mục đích (trả lời câu hỏi gì, cho ai)", ""], ["Công thức (khái niệm, rồi công thức Excel)", ""],
        ["Nguồn dữ liệu (bảng, cột)", ""], ["Grain (mỗi dòng dữ liệu là gì · KPI báo cáo theo mức nào)", ""],
        ["Dimensions (xem theo chiều nào; chiều nào KHÔNG xem được)", ""], ["Owner (ai chịu trách nhiệm)", ""],
        ["Target (con số hoặc bảng Target)", ""], ["Ngưỡng Warning / Critical", "Warning: ____ %   ·   Critical: ____ %"],
        ["Ghi chú giới hạn (giả định, dữ liệu còn thiếu)", ""]]
table(d, rows, [5.2, 11.8], size=10.5, first_bold=True, min_h=1.75)
rich(d, [("Ví dụ (Card 1): ", {"bold": True, "color": NAVY}),
         ("Total Revenue · Tháng này bán được bao nhiêu tiền, ở đâu? · SUM(Quantity × Unit_Price) = SUM(tblSales[Revenue]) · tblSales · "
          "dữ liệu: 1 sản phẩm trong 1 đơn, KPI: tháng · Miền, cửa hàng, ngành hàng, tháng · Giám đốc bán hàng · bảng Target (từ Buổi 4) · 90% / 80% · chưa trừ trả hàng, khuyến mại.", {"color": GREY})])
d.save(f"{OUT}/B2_KPI_Definition_Card.docx")
print("ok;", f"còn {len(MISSING)} khung [CHỤP] chưa có ảnh:" if MISSING else "đủ ảnh.", " ".join(MISSING))
