"""Dựng deck Buổi 3: Slide_B3.pptx từ Slide_Script_B3.md, theo style của Slide_B1.pptx (cùng khung với build_slides_b2.py).

- Master, layout, theme và slide bìa lấy từ Slide_B1.pptx (không sửa file B1; chỉ đọc làm khuôn).
- Chữ trên slide khai báo trong script này, khớp phần "Trên slide" của kịch bản.
- Speaker notes lấy NGUYÊN VĂN từ kịch bản (Lời giảng, Ghi chú giảng viên), bỏ định dạng markdown,
  nên sửa lời giảng thì sửa Slide_Script_B3.md rồi chạy lại script này.
- Quiz: shape đáp án tên "Answer …", lớp mờ "Dim …", hiện khi bấm (build_student_slides.py sẽ xóa ở bản _HocVien).
- Ảnh chụp thật lấy từ Materials/B3/Images (tên S<số slide>_…png); chưa có thì để khung [CHỤP].

    python3 _generator/build_slides_b3.py     # cần python-pptx
    python3 _generator/build_student_slides.py B3
"""
import copy, os, re
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

ROOT = os.path.normpath(os.environ.get("MCI_ROOT") or os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
TEMPLATE = f"{ROOT}/1_Sessions/B1_Nen_tang_Excel_Du_lieu/Slide_B1.pptx"
B3 = f"{ROOT}/1_Sessions/B3_PivotTable"
SCRIPT = f"{B3}/Slide_Script_B3.md"
OUT = f"{B3}/Slide_B3.pptx"
IMG = f"{ROOT}/Materials/B3/Images"        # ảnh chụp thật: S<số slide>_<chủ đề>.png (danh sách: _DANH_SACH_ANH.md)
MISSING = []

NAVY, INK, INK2, MUTED = "203864", "1B2230", "585858", "6B6F76"
RED, GOLD, ALT, HAIR, WHITE = "C00000", "FFC000", "F3F5F9", "D9DEE8", "FFFFFF"
GOOD, GOOD_BG, BAD_TXT, BAD_BG, ACCENT = "2E7D32", "E8F5E9", "9C0006", "FFC7CE", "3398DA"
L, R, W = 1.67, 12.74, 11.07          # lề trái, lề phải, bề rộng vùng nội dung (in)

# ---------------------------------------------------------------------
# Khuôn: mở B1, giữ slide bìa, xóa các slide còn lại
# ---------------------------------------------------------------------
prs = Presentation(TEMPLATE)
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst)[1:]:
    prs.part.drop_rel(sldId.rId); sldIdLst.remove(sldId)
LAYOUT = {l.name: l for l in prs.slide_layouts}

def rgb(h): return RGBColor.from_string(h)

def _runs(p, parts, size, color, bold):
    for part in parts:
        if isinstance(part, str): part = (part, {})
        text, f = part
        r = p.add_run(); r.text = text
        r.font.size = Pt(f.get("size", size)); r.font.bold = f.get("b", bold)
        r.font.color.rgb = rgb(f.get("color", color)); r.font.name = "Arial"
        if f.get("i"): r.font.italic = True

def tb(s, x, y, w, h, paras, size=18, color=INK, bold=False, anchor="t", align=None, name="TextBox", bullet=False, space=6):
    """paras: str | list[para]; para: str | list[run]; run: str | (text, {b, color, size, i})."""
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); sh.name = name
    tf = sh.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "ctr": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str): paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align: p.alignment = {"l": PP_ALIGN.LEFT, "ctr": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        if i: p.space_before = Pt(space)
        if isinstance(para, str): para = [para]
        if bullet: para = [("•  ", {"color": NAVY, "b": True})] + list(para)
        _runs(p, para, size, color, bold)
    return sh

def box(s, x, y, w, h, fill=None, line=None, lw=1, shape=MSO_SHAPE.RECTANGLE, radius=None, name="Box", alpha=None):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h)); sh.name = name
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0] = radius
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    else: sh.fill.background()
    if line: sh.line.color.rgb = rgb(line); sh.line.width = Pt(lw)
    else: sh.line.fill.background()
    sh.shadow.inherit = False
    if alpha is not None:
        clr = sh.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        a = etree.SubElement(clr, qn("a:alpha")); a.set("val", str(alpha))
    return sh

def chip(s, x, y, w, text, fill=None, line=NAVY, color=NAVY, size=14, name="Chip", h=0.4):
    sh = box(s, x, y, w, h, fill=fill, line=line, lw=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2, name=name)
    tf = sh.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; _runs(p, [text], size, color, True)
    return sh

def vline(s, x, y, h):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x), Inches(y + h))
    ln.line.color.rgb = rgb(HAIR); ln.line.width = Pt(1); return ln

def content(title, rail):
    s = prs.slides.add_slide(LAYOUT["Title Only"])
    s.shapes.title.text = title
    tb(s, 0, 1.5, 1.07, 5.25, rail, size=14, color=WHITE, bold=True, anchor="ctr", align="ctr", name="Rail label")
    s.shapes[-1].text_frame._bodyPr.set("vert", "vert270")
    return s

def section(title, body=None):
    s = prs.slides.add_slide(LAYOUT["Section Header"])
    s.shapes.title.text = title
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx == 1:
            if body: ph.text = body
            else: ph._element.getparent().remove(ph._element)
    return s

def tip(s, x, y, w, h, text, pro=False):
    box(s, x, y, w, h, fill=ALT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, name=("PRO TIP" if pro else "TIP") + " callout")
    if pro:
        chip(s, x + 0.12, y + 0.12, 1.1, "PRO TIP", fill=GOLD, line=None, color=INK, size=12, name="PRO TIP chip", h=0.35)
        tb(s, x + 1.32, y + 0.1, w - 1.44, h - 0.2, text, size=16, name="PRO TIP text")
    else:
        tb(s, x + 0.12, y + 0.12, 0.6, 0.35, "TIP", size=14, color=NAVY, bold=True, anchor="ctr")
        tb(s, x + 0.75, y + 0.1, w - 0.87, h - 0.2, text, size=16, name="TIP text")

def tool(s, x, y, w, name, purpose, h_purpose=0.7, size=24):
    tb(s, x, y, w, 0.5, name, size=size, color=RED, bold=True, name="Tool name")
    tb(s, x, y + 0.5, w, h_purpose, purpose, size=16, color=INK2, name="Tool purpose")

def table(s, x, y, colw, rows, size=16, head_size=None, rowh=0.45, bold_first=False, name="Table",
          header=True, fills=None, colors=None, align=None):
    """rows[0] là tiêu đề (nền navy). fills/colors: {(r,c): hex}. Chữ trong ô có thể là list run."""
    nr, nc = len(rows), len(rows[0])
    gf = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(sum(colw)), Inches(rowh * nr)); gf.name = name
    tbl = gf.table
    tbl._tbl.tblPr.set("firstRow", "1"); tbl._tbl.tblPr.set("bandRow", "0")
    style = tbl._tbl.tblPr.find(qn("a:tableStyleId"))
    if style is not None: tbl._tbl.tblPr.remove(style)
    for j, cw in enumerate(colw): tbl.columns[j].width = Inches(cw)
    for i in range(nr):
        tbl.rows[i].height = Inches(rowh)
        for j in range(nc):
            c = tbl.cell(i, j); v = rows[i][j]
            c.margin_left = c.margin_right = Inches(0.08); c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]
            if align and align[j]: p.alignment = {"l": PP_ALIGN.LEFT, "ctr": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align[j]]
            head = header and i == 0
            col = WHITE if head else (colors or {}).get((i, j), INK)
            parts = v if isinstance(v, list) else [str(v)]
            _runs(p, parts, head_size or size, col, head or (bold_first and j == 0))
            tcPr = c._tc.get_or_add_tcPr()
            for side in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                ln = etree.SubElement(tcPr, qn(side))
                if side in ("a:lnL", "a:lnR"):
                    ln.set("w", "0"); etree.SubElement(ln, qn("a:noFill"))
                else:
                    ln.set("w", "12700"); sf = etree.SubElement(ln, qn("a:solidFill"))
                    etree.SubElement(sf, qn("a:srgbClr")).set("val", HAIR)
            f = NAVY if head else (fills or {}).get((i, j))
            if f:
                sf = etree.SubElement(tcPr, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr")).set("val", f)
            else:
                etree.SubElement(tcPr, qn("a:noFill"))
    return gf

def fx(s, x, y, w, formula, size=16, name="Formula bar"):
    """Thanh công thức mô phỏng Excel: ô 'fx' + công thức."""
    box(s, x, y, w, 0.45, fill=WHITE, line=HAIR, name=name)
    tb(s, x + 0.1, y, 0.45, 0.45, [[("fx", {"i": True})]], size=14, color=MUTED, anchor="ctr")
    tb(s, x + 0.6, y, w - 0.7, 0.45, formula, size=size, color=INK, anchor="ctr", name="Formula")

def stat(s, x, y, w, h, label, value, sub=None, name="Stat card", vcolor=NAVY, vsize=32):
    box(s, x, y, w, h, fill=WHITE, line=HAIR, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, name=name)
    paras = [[(label, {"size": 14, "color": INK2})], [(value, {"size": vsize, "b": True, "color": vcolor})]]
    if sub: paras.append([(sub, {"size": 14, "color": INK})])
    tb(s, x + 0.2, y + 0.1, w - 0.4, h - 0.2, paras, anchor="ctr", space=2)

def shot(s, x, y, w, h, what, where, img, name="Image placeholder [CHỤP]"):
    """Ảnh chụp thật (Excel 365 Windows, zoom 100%) — xem CLAUDE.md, mục Ảnh minh họa.
    Có file Materials/B3/Images/<img> thì chèn ảnh vào khung (giữ tỉ lệ, canh giữa, viền hairline);
    chưa có thì để khung chờ [CHỤP] ghi nội dung, file · sheet · vùng và tên file ảnh cần lưu."""
    path = f"{IMG}/{img}"
    if os.path.exists(path):
        pic = s.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))
        if pic.height > Inches(h):
            ratio = Inches(h) / pic.height; pic.height = Inches(h); pic.width = int(pic.width * ratio)
        pic.left = Inches(x) + (Inches(w) - pic.width) // 2; pic.top = Inches(y) + (Inches(h) - pic.height) // 2
        pic.line.color.rgb = rgb(HAIR); pic.line.width = Pt(0.75); pic.name = f"Screenshot {img[:-4]}"
        return pic
    sh = box(s, x, y, w, h, fill=ALT, line="A6A6A6", lw=1.25, name=name)
    sh.line.dash_style = MSO_LINE.DASH
    tb(s, x + 0.2, y + 0.1, w - 0.4, h - 0.2, [[("[CHỤP] ", {"b": True, "color": NAVY}), (what, {"color": INK2})],
                                               [(f"{where} → {img}", {"color": MUTED, "size": 12})]], size=14, anchor="ctr", align="ctr", space=4)
    MISSING.append(img)
    return sh

def footnote(s, text):
    tb(s, L, 6.8, W, 0.35, text, size=12, color=MUTED, anchor="b", name="Footnote")

# ---------------------------------------------------------------------
# Speaker notes từ kịch bản
# ---------------------------------------------------------------------
md = open(SCRIPT, encoding="utf-8").read()
blocks = re.split(r"^### Slide (\d+) · ", md, flags=re.M)
NOTES = {}

def clean(t):
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t); t = re.sub(r"(?<!\w)\*(.+?)\*(?!\w)", r"\1", t)
    return t.replace("`", "").strip()

def fmt_min(m):
    return (f"{m:g}").replace(".", ",")

for k in range(1, len(blocks), 2):
    n = int(blocks[k]); body = blocks[k + 1]
    head, rest = body.split("\n", 1)
    body = re.split(r"^(?:---|## )", rest, flags=re.M)[0]
    m = re.search(r"· ([\d.]+)'\s*$", head) or re.search(r"\((\d+) phút\)", head)
    out = []
    ans = re.search(r"\*\*([A-D])\)", body)
    if head.startswith("QUESTION"):
        out.append(f"Khoảng 40 giây. Đáp án: {ans.group(1)}.")
    elif m:
        out.append(f"Thời lượng: {fmt_min(float(m.group(1)))} phút.")
    for sec in re.finditer(r"^\*\*(Lời giảng[^*]*|Ghi chú giảng viên[^*]*):\*\*\n(.*?)(?=^\*\*(?:Lời giảng|Ghi chú giảng viên)|\Z)", body, flags=re.M | re.S):
        label, txt = sec.group(1), sec.group(2)
        head_, _, tail_ = label.partition(" (")
        out.append(head_.upper() + (f" ({tail_}" if tail_ else ""))
        for line in txt.splitlines():
            line = line.rstrip()
            if line.startswith(">"): line = line[1:].strip()
            elif line.lstrip().startswith("- "): line = line.lstrip()[2:]
            else: line = line.strip()
            if line: out.append(clean(line))
    NOTES[n] = "\n".join(out)
assert len(NOTES) == 42, sorted(NOTES)

def notes(s, n):
    s.notes_slide.notes_text_frame.text = NOTES[n]

# ---------------------------------------------------------------------
# Quiz: 4 phương án, đáp án hiện khi bấm
# ---------------------------------------------------------------------
def quiz(n, topic, question, options, answer, rail="PHẦN 8 · QUIZ & TỔNG KẾT", qsize=22, osize=18):
    s = content(f"Question {n - 33}: {topic}", rail)
    tb(s, L, 1.65, W, 1.0, question, size=qsize, color=INK)
    pos = {"A": (L, 2.85), "B": (7.33, 2.85), "C": (L, 4.35), "D": (7.33, 4.35)}
    ow, oh = 5.41, 1.3
    for letter, (x, y) in pos.items():
        box(s, x, y, ow, oh, fill=WHITE, line=HAIR, name=f"Option {letter}")
        if letter == answer:
            box(s, x, y, ow, oh, fill=GOOD_BG, line=GOOD, lw=3, name=f"Answer highlight {letter}")
        lt = box(s, x + 0.2, y + (oh - 0.6) / 2, 0.6, 0.6, fill=NAVY, shape=MSO_SHAPE.OVAL, name=f"Letter {letter}")
        _center_text(lt, letter)
        if letter == answer:
            al = box(s, x + 0.2, y + (oh - 0.6) / 2, 0.6, 0.6, fill=GOOD, shape=MSO_SHAPE.OVAL, name=f"Answer letter {letter}")
            _center_text(al, letter)
        tb(s, x + 1.0, y, ow - 1.1, oh, options[letter], size=osize, anchor="ctr")
    for letter, (x, y) in pos.items():
        if letter != answer: box(s, x, y, ow, oh, fill=WHITE, name=f"Dim {letter}", alpha=60000)
    ax, ay = pos[answer]
    chip(s, ax + ow - 1.6, ay - 0.2, 1.45, "✓ Đáp án", fill=GOOD, line=None, color=WHITE, size=14, name="Answer chip", h=0.32)
    _reveal(s, [sh for sh in s.shapes if sh.name.startswith(("Answer", "Dim"))])
    notes(s, n)

def _center_text(sh, text):
    tf = sh.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; _runs(p, [text], 20, WHITE, True)

def _reveal(s, shapes):
    """Một lần bấm: mọi shape đáp án/làm mờ hiện dần (fade), giống deck B1."""
    ids = [sh.shape_id for sh in shapes]
    cid = [3]
    def nid(): cid[0] += 1; return cid[0]
    pars = []
    for k, spid in enumerate(ids):
        a, b = nid(), nid(); c = nid()
        node = "clickEffect" if k == 0 else "withEffect"
        pars.append(f'<p:par><p:cTn id="{a}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="{node}">'
                    f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:set><p:cBhvr><p:cTn id="{b}" dur="1" fill="hold">'
                    f'<p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
                    f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
                    f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{c}" dur="400"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
                    f'</p:childTnLst></p:cTn></p:par>')
    bld = "".join(f'<p:bldP spid="{i}" grpId="0" animBg="1"/>' for i in ids)
    xml = ('<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par>'
           '<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek">'
           '<p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst><p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>'
           f'<p:childTnLst><p:par><p:cTn id="{nid() + 900}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{"".join(pars)}</p:childTnLst></p:cTn></p:par>'
           '</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>'
           f'<p:bldLst>{bld}</p:bldLst></p:timing>')
    s._element.append(etree.fromstring(xml))

# =====================================================================
# SLIDES
# =====================================================================
BLUE, ORANGE = "2A78D6", "EB6834"
R0 = "PHẦN 0 · MỞ ĐẦU"; R1 = "PHẦN 1 · VÌ SAO PIVOT"; R2 = "PHẦN 2 · TẠO PIVOT"; R3 = "PHẦN 3 · ĐỊNH DẠNG"
R4 = "PHẦN 4 · SẮP XẾP · NHÓM"; R5 = "PHẦN 5 · TÍNH TOÁN"; R6 = "PHẦN 6 · PIVOTCHART"; R7 = "PHẦN 7 · THỰC HÀNH"

def card(s, x, y, w, h, title, body, size=15, fill=WHITE, line=HAIR, title_color=NAVY, name="Card"):
    box(s, x, y, w, h, fill=fill, line=line, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05, name=name)
    paras = [[(title, {"b": True, "color": title_color, "size": size + 1})]] + [p if isinstance(p, list) else [p] for p in body]
    tb(s, x + 0.18, y + 0.12, w - 0.36, h - 0.24, paras, size=size, space=5)

def bricks(s, done, nxt):
    box(s, 9.07, 4.6, 3.67, 1.0, fill=NAVY, name=f"Brick {done[:2]}")
    tb(s, 9.17, 4.6, 3.47, 1.0, "✓  " + done, size=18, color=WHITE, bold=True, anchor="ctr")
    box(s, 9.07, 3.3, 3.67, 1.0, fill=WHITE, line=BLUE, lw=2.25, name=f"Brick {nxt[:2]}")
    tb(s, 9.17, 3.3, 3.47, 1.0, nxt, size=18, color=NAVY, bold=True, anchor="ctr")

# 1 · Bìa
s1 = prs.slides[0]
for sh in s1.shapes:
    if sh.name.startswith("subTitle"):
        sh.text_frame.paragraphs[0].runs[0].text = "Buổi 3: Tư duy dữ liệu & PivotTable"
        for r in sh.text_frame.paragraphs[0].runs[1:]: r.text = ""
    elif sh.has_text_frame and sh.text_frame.text.startswith("Yêu cầu máy"):
        p = sh.text_frame.paragraphs[0]
        p.runs[0].text = "Mở sẵn: file khởi đầu B2_MCI_Mart_Q1_2025_Answer.xlsx (đáp án Buổi 2)"
        for r in p.runs[1:]: r.text = ""
    elif sh.has_text_frame and sh.text_frame.text.startswith("Mở sẵn"):
        p = sh.text_frame.paragraphs[0]
        p.runs[0].text = "Mở sẵn: file khởi đầu B2_MCI_Mart_Q1_2025_Answer.xlsx (đáp án Buổi 2)"
        for r in p.runs[1:]: r.text = ""
s1.notes_slide.notes_text_frame.text = NOTES[1]

# 2 · Câu hỏi mở buổi
s = content("Câu hỏi mở buổi", R0)
chip(s, 10.03, 0.6, 2.7, "Tầng 2 — Dimensional", name="Tier chip")
tb(s, L, 1.8, 5.2, 2.4, [[("“Làm sao trả lời 20 câu hỏi ‘theo chiều nào’ trong 20 phút?”", {"b": True, "color": NAVY})]], size=28)
tb(s, L, 4.35, 5.2, 0.4, "— Giám đốc bán hàng, MCI Mart", size=16, color=INK2)
tb(s, L, 5.0, 5.2, 1.2, "SUMIFS trả lời một câu hỏi bằng một bảng. Hai mươi câu hỏi là hai mươi bảng.", size=16, color=INK2)
qs = [("theo cửa hàng?", "Doanh thu 12 cửa hàng × 3 tháng"), ("theo thứ trong tuần?", "Ngày nào khách đi siêu thị nhiều?"),
      ("theo tuần?", "Tuần nào bán mạnh nhất quý?"), ("theo ngành hàng?", "Mỗi ngành chiếm bao nhiêu %?"),
      ("cửa hàng nào bất thường?", "Ai giảm mạnh trong tháng 3?"), ("vì sao tháng 1 > tháng 2?", "Câu hỏi còn treo từ Buổi 2")]
for k, (a, b) in enumerate(qs):
    x, y = 7.25 + (k % 2) * 2.8, 1.8 + (k // 2) * 1.5
    box(s, x, y, 2.65, 1.3, fill=ALT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, name=f"Question card {k + 1}")
    tb(s, x + 0.15, y + 0.12, 2.35, 1.1, [[(a, {"b": True, "color": NAVY, "size": 16})], [(b, {"size": 13, "color": INK2})]], space=4)
notes(s, 2)

# 3 · Recap B2
s = content("Recap Buổi 2 & file khởi đầu", R0)
tb(s, L, 1.75, 7.0, 4.6, [
    [("Viên gạch B2: ", {"b": True, "color": NAVY}), "tblSales làm giàu (Category, COGS, Gross_Margin, Thang, Thu…) · KPI_SUMIFS: tổng 941.454.000 · T1 425.850.000 · T2 245.996.000."],
    [("Bài về nhà B2 (tháng 4): ", {"b": True, "color": NAVY}), "296.965.000 · 966 đơn. Lỗi cài sẵn: 4 mã Product_Code dính dấu cách cuối (XLOOKUP ra #N/A → TRIM, không che bằng IFERROR)."],
    [("File khởi đầu hôm nay ", {"b": True, "color": NAVY}), "= đáp án B2 (B2_MCI_Mart_Q1_2025_Answer.xlsx). Lưu thành B3_<HoTen>.xlsx. Giữ sheet KPI_SUMIFS để đối soát Pivot."],
], size=18, bullet=True, space=14)
bricks(s, "B2 · Bảng KPI (SUMIFS)", "B3 · Báo cáo Pivot")
footnote(s, "Demo: ../B2_Ham_va_KPI/Practice/B2_MCI_Mart_Q1_2025_Answer.xlsx (sheet KPI_SUMIFS) · ../B2_Ham_va_KPI/Homework/B2_BTVN_MCI_Mart_T4_2025_Answer.xlsx (Kiem_tra)")
notes(s, 3)

# 4 · Thang 5 tầng
s = content("Hôm nay leo lên tầng 2", R0)
tiers = [("5", "Sắp tới sẽ thế nào?", "Predictive", "B6"), ("4", "Vậy thì sao, nên làm gì?", "Prescriptive", "B6"),
         ("3", "Vì sao?", "Diagnostic", "B3 drill-down · B5"), ("2", "Ở đâu, khi nào, nhóm nào?", "Dimensional", "B3"),
         ("1", "Chuyện gì đã xảy ra?", "Descriptive", "B1–B2 ✓")]
for k, (n_, q_, t_, b_) in enumerate(tiers):
    y = 1.65 + k * 0.92
    on = n_ == "2"; part = n_ == "3"
    box(s, L, y, 6.3, 0.8, fill=NAVY if on else ("DCE3EF" if part else ALT), line=BLUE if part else None, lw=1.5, name=f"Tier {n_}")
    tb(s, L + 0.2, y, 6.0, 0.8, [[(f"Tầng {n_} · ", {"b": True, "color": WHITE if on else NAVY}), (q_, {"b": on, "color": WHITE if on else INK}),
                                  (f"   {t_} · {b_}", {"size": 12, "color": "CDE2FB" if on else INK2})]], size=17, anchor="ctr")
tb(s, 8.3, 1.65, 4.44, 0.45, "Ba câu hỏi tự tìm đáp án hôm nay", size=17, color=NAVY, bold=True)
for k, t in enumerate(["Vì sao tháng 1 cao hơn hẳn tháng 2?", "Cửa hàng nào giảm bất thường trong tháng 3?", "Ngày nào trong tuần bán chạy nhất?"]):
    y = 2.25 + k * 1.25
    lt = box(s, 8.3, y + 0.2, 0.6, 0.6, fill=NAVY, shape=MSO_SHAPE.OVAL, name=f"Story {k + 1}"); _center_text(lt, str(k + 1))
    tb(s, 9.05, y, 3.69, 1.0, t, size=17, anchor="ctr")
tb(s, 8.3, 6.1, 4.44, 0.5, "Viên gạch B3: báo cáo Pivot tương tác + PivotChart", size=14, color=INK2)
notes(s, 4)

# 5 · Chia phần 1
notes(section("1. Vì sao cần PivotTable"), 5)

# 6 · Giới hạn của SUMIFS
s = content("Giới hạn của SUMIFS", R1)
card(s, L, 1.65, 5.3, 4.35, "SUMIFS: mỗi câu hỏi = một bảng mới", [
    "• Cửa hàng × Tháng: gõ 12 mã, 3 tháng, kéo 36 ô.", "• Theo tuần: tblSales chưa có cột Tuần → thêm cột trước.",
    "• Thêm một chiều = viết lại công thức.", [("≈ 5 phút / câu hỏi", {"b": True, "color": BAD_TXT, "size": 22})]], size=16, fill=ALT, line=None)
card(s, 7.44, 1.65, 5.3, 4.35, "PivotTable: kéo thả", [
    "• Kéo trường vào Rows / Columns / Values.", "• Đổi chiều phân tích trong vài giây.", "• Bấm đúp để xem dòng gốc.",
    [("≈ 10 giây / câu hỏi", {"b": True, "color": GOOD, "size": 22})]], size=16, fill=ALT, line=None)
for k, z in enumerate(["Filters", "Columns", "Rows", "Values"]):
    chip(s, 7.62 + k * 1.25, 5.3, 1.15, z, fill=WHITE, line=BLUE, color=NAVY, size=13, name=f"Zone {z}", h=0.42)
tb(s, L, 6.15, W, 0.5, "5 lợi ích của Pivot: Powerful · Beautiful · Fast · Accurate · Flexible. Nhưng mỗi ô Pivot vẫn là một SUMIFS.", size=15, color=INK2)
notes(s, 6)

# 7 · Dữ liệu cho Pivot
s = content("Dữ liệu cho Pivot: tidy + Table", R1)
tb(s, L, 1.65, W, 0.9, [[("Nguồn Pivot: ", {"b": True, "color": NAVY}), "hình chữ nhật (mỗi cột một trường, mỗi dòng một bản ghi) · tiêu đề một dòng · không dòng tổng, không ô gộp · nên là Excel Table (B1)."]], size=16)
table(s, L, 2.65, [2.6, 5.4, 3.07], [
    ["Vai trò", "Ví dụ trong tblSales", "Kéo vào"],
    [[("Dimension", {"b": True, "color": BLUE})], "Region, Store_Code, Thang, Thu, Category, Order_Date", "Rows · Columns · Filters"],
    [[("Measure", {"b": True, "color": ORANGE})], "Revenue, Quantity, COGS, Gross_Margin", "Values"],
    [[("Grain", {"b": True, "color": NAVY})], "1 sản phẩm trong 1 đơn (6.187 dòng ≠ 3.144 đơn)", "Quyết định Pivot đếm cái gì"],
], size=16, rowh=0.62)
tip(s, L, 5.35, W, 0.95, "Dữ liệu đã tidy ở B1 là lý do hôm nay chỉ cần kéo thả. Còn dòng tổng hay ô gộp thì Pivot cộng sai (tổng gấp đôi).")
notes(s, 7)

# 8 · Chia phần 2
notes(section("2. Tạo PivotTable: mỗi ô là một bộ lọc"), 8)

# 9 · Insert PivotTable
s = content("Insert PivotTable", R2)
tool(s, L, 1.6, 5.3, "Insert PivotTable", "Tạo Pivot trống từ một Table / vùng dữ liệu.", 0.45)
vline(s, 7.2, 1.6, 1.0)
tool(s, 7.4, 1.6, 5.34, "Recommended PivotTables", "Excel gợi ý sẵn vài Pivot để bắt đầu.", 0.45)
shot(s, L, 2.75, 7.2, 3.9, "Hộp Create PivotTable: Table/Range = tblSales, chọn New Worksheet", "Demo/B3_Demo_Pivot.xlsx · sheet Sales · chọn A2 → Insert → PivotTable → From Table/Range", "S09_create_pivot.png")
tip(s, 9.05, 2.75, 3.69, 2.2, "Chọn MỘT ô trong Table rồi Insert → PivotTable: Excel tự điền tên Table. Phím tắt: Alt, N, V, T.")
tb(s, 9.05, 5.15, 3.69, 1.4, "Hai câu hỏi của hộp thoại: dữ liệu ở đâu? đặt Pivot ở đâu?", size=15, color=INK2)
footnote(s, "Ảnh: Demo/B3_Demo_Pivot.xlsx, sheet Sales → Insert → PivotTable")
notes(s, 9)

# 10 · PivotTable Fields
s = content("PivotTable Fields: 4 vùng", R2)
tool(s, L, 1.6, W, "PivotTable Fields", "Danh sách trường của nguồn và 4 vùng: Filters (lọc cả Pivot) · Columns · Rows · Values (số cần tính).", 0.45)
shot(s, L, 2.6, 7.6, 3.95, "Pivot Region × Thang (3 × 3 + Grand Total) và khung PivotTable Fields: Region ở Rows, Thang ở Columns, Sum of Revenue ở Values", "Demo/B3_Demo_Pivot.xlsx · sheet PV_MienThang (sheet vừa tạo) · A3:E8 + khung Fields", "S10_field_list.png")
stat(s, 9.45, 2.6, 3.29, 1.6, "Miền × Tháng, 10 giây", "941.454.000", "Grand Total = tổng Buổi 2", vsize=26)
tip(s, 9.45, 4.35, 3.29, 2.2, "Dimension vào Rows / Columns, Measure vào Values. Kéo nhầm thì kéo trường ra khỏi vùng.")
footnote(s, "Ảnh: Demo/B3_Demo_Pivot.xlsx, sheet PV_MienThang")
notes(s, 10)

# 11 · Pivot hoạt động thế nào
s = content("Pivot hoạt động thế nào: tọa độ = bộ lọc", R2)
box(s, L, 1.75, 2.1, 2.3, fill=ALT, line=HAIR, name="Source column")
tb(s, L, 1.75, 2.1, 2.3, [[("tblSales[Revenue]", {"b": True, "color": NAVY, "size": 15})], "6.187 dòng", [("∑ phép tổng", {"b": True})]], size=16, anchor="ctr", align="ctr", space=8)
for k, t in enumerate(["Region = Bắc", "Thang = 3"]):
    y = 1.8 + k * 1.2
    box(s, 4.1, y, 2.95, 0.9, fill=WHITE, line=BLUE, lw=2, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.15, name=f"Coordinate {k + 1}")
    tb(s, 4.1, y, 2.95, 0.9, [[(("Tọa độ hàng  " if k == 0 else "Tọa độ cột  "), {"color": INK2, "size": 12}), (t, {"b": True, "color": NAVY})]], size=16, anchor="ctr", align="ctr")
for y in (2.25, 3.45):
    a = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(3.77), Inches(2.9), Inches(4.1), Inches(y)); a.line.color.rgb = rgb(ACCENT); a.line.width = Pt(2)
    b = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(7.05), Inches(y), Inches(7.35), Inches(2.9)); b.line.color.rgb = rgb(ACCENT); b.line.width = Pt(2)
box(s, 7.35, 2.2, 1.85, 1.4, fill=NAVY, name="Result cell")
tb(s, 7.35, 2.2, 1.85, 1.4, [[("Ô Pivot Bắc × 3", {"size": 12, "color": "CDE2FB"})], [("90.210.000", {"b": True, "size": 19, "color": WHITE})]], anchor="ctr", align="ctr", space=4)
table(s, L, 4.3, [2.5, 2.5, 2.5], [["B2 · SUMIFS", "B3 · PivotTable", "B5 · DAX"], ["criteria", "tọa độ ô Pivot", "filter context"]],
      size=15, rowh=0.42, align=["ctr", "ctr", "ctr"], name="Chain C", fills={(1, 1): "DCE3EF"}, colors={(1, 1): NAVY})
shot(s, 9.4, 1.75, 3.34, 3.0, "Pivot Miền × Tháng và khối đối soát: GETPIVOTDATA Bắc–T3 = SUMIFS B2 = 90.210.000, chênh lệch 0", "Practice/B3_MCI_Mart_Q1_2025_Answer.xlsx · sheet PV_MienThang · A4:D15", "S11_pivot_vs_sumifs.png")
tip(s, L, 5.35, W, 1.25, "Gõ = rồi bấm vào ô Pivot: Excel tự viết GETPIVOTDATA (đúng ô dù Pivot đổi bố cục, nhưng kéo công thức không đổi kết quả). Tắt: PivotTable Analyze → Options ▾ → bỏ tick Generate GetPivotData. Double-click ô Bắc × 3 → 635 dòng gốc.")
footnote(s, "Hình xương sống Chuỗi C dùng lại từ Buổi 2 · Ảnh: Practice/B3_MCI_Mart_Q1_2025_Answer.xlsx, sheet PV_MienThang")
notes(s, 11)

# 12 · Refresh
s = content("Refresh · Change Data Source", R2)
tool(s, L, 1.6, 5.3, "Refresh · Refresh All", "Cập nhật Pivot sau khi nguồn thay đổi (Alt+F5 · Ctrl+Alt+F5).", 0.7)
vline(s, 7.2, 1.6, 1.2)
tool(s, 7.4, 1.6, 5.34, "Change Data Source", "Đổi / mở rộng vùng nguồn của Pivot.", 0.7)
shot(s, L, 2.95, 6.2, 2.35, "Sau khi dán 10 dòng tháng 4 và Refresh All: PV_Table tổng 942.937.000, PV_Vung vẫn 941.454.000", "Demo/B3_Demo_Refresh.xlsx · sheet PV_Table, PV_Vung · A3:B7", "S12a_refresh_table_vs_range.png")
shot(s, 8.05, 2.95, 4.69, 2.35, "Hộp Change PivotTable Data Source: Sales_Vung!$A$1:$I$6198", "Demo/B3_Demo_Refresh.xlsx · sheet PV_Vung", "S12b_change_source.png")
tip(s, L, 5.45, 5.4, 1.2, "Pivot đọc từ BỘ NHỚ ĐỆM (cache), không đọc thẳng ô: sửa dữ liệu xong phải Refresh.")
tip(s, 7.25, 5.45, 5.49, 1.2, "Nguồn là Table → dòng mới tự vào Pivot khi Refresh. Nguồn là vùng $A$1:$I$6188 → không.", pro=True)
footnote(s, "Ảnh: Demo/B3_Demo_Refresh.xlsx, sheet PV_Table và PV_Vung")
notes(s, 12)

# 13 · Chia phần 3
notes(section("3. Định dạng: Pivot dễ đọc cho người xem"), 13)

# 14 · Value Field Settings
s = content("Value Field Settings · PivotTable Styles", R3)
tool(s, L, 1.6, 5.3, "Value Field Settings", "Đặt tên và định dạng số cho cả trường giá trị (Number Format).", 0.7)
vline(s, 7.2, 1.6, 1.2)
tool(s, 7.4, 1.6, 5.34, "PivotTable Styles", "Kiểu bảng; Grand Totals, Subtotals: bật / tắt dòng tổng.", 0.7)
shot(s, L, 2.95, 6.9, 3.65, "Hộp Value Field Settings (Custom Name: Doanh thu) và Number Format: Number, 0 chữ số, Use 1000 Separator", "Demo/B3_Demo_Pivot.xlsx · sheet PV_MienThang", "S14_value_field_settings.png")
tip(s, 8.75, 2.95, 3.99, 1.75, "Định dạng qua Value Field Settings, KHÔNG bôi đen ô: Pivot đổi bố cục là mất định dạng.")
tip(s, 8.75, 4.85, 3.99, 1.75, "Conditional Formatting trong Pivot: chọn “All cells showing ‘Doanh thu’ values…” để quy tắc tự áp cho ô mới.", pro=True)
footnote(s, "Ảnh: Demo/B3_Demo_Pivot.xlsx, sheet PV_MienThang")
notes(s, 14)

# 15 · Report Layout
s = content("Compact · Outline · Tabular", R3)
tool(s, L, 1.6, W, "Report Layout", "Cách xếp nhiều trường trên hàng. Outline: mỗi trường một cột, nên dùng khi phân tích.", 0.45)
for k, (lab, img, rng) in enumerate([("Compact (mặc định)", "S15a_compact.png", "A3:F20"), ("Outline", "S15b_outline.png", "A3:G20"), ("Tabular + Repeat labels", "S15c_tabular.png", "A3:G20")]):
    x = L + k * 3.75
    tb(s, x, 2.55, 3.55, 0.35, lab, size=15, color=NAVY, bold=True)
    shot(s, x, 2.95, 3.55, 2.6, f"Region + Store_Code trên hàng, {lab}", f"Demo/B3_Demo_Pivot.xlsx · PV_MienThang · {rng}", img)
tip(s, L, 5.7, W, 0.9, "Muốn chép kết quả Pivot thành bảng dữ liệu mới: Tabular + Repeat All Item Labels + tắt Subtotals và Grand Totals.")
footnote(s, "Ảnh: Demo/B3_Demo_Pivot.xlsx, sheet PV_MienThang (Region + Store_Code trên hàng)")
notes(s, 15)

# 16 · Nghỉ
notes(section("Nghỉ 10 phút", "Trước khi quay lại: mở Demo/B3_Demo_Loi_Du_lieu_HocVien.xlsx (trợ giảng đã gửi)"), 16)

# 17 · Chia phần 4
notes(section("4. Sắp xếp, lọc, nhóm, đào sâu: đi tìm câu chuyện"), 17)

# 18 · Sort & Filter
s = content("Sort · Value Filters (Top 5)", R4)
tool(s, L, 1.6, 5.3, "Sort", "Sắp xếp mục theo tên hoặc theo giá trị.", 0.45)
vline(s, 7.2, 1.6, 1.0)
tool(s, 7.4, 1.6, 5.34, "Label / Value Filters", "Lọc theo tên (chứa…) hoặc theo giá trị (Top 10…).", 0.45)
shot(s, L, 2.65, 3.3, 3.0, "Pivot Thu × Doanh thu sắp xếp giảm dần: Sun 159.845.000 … Thu 111.408.000", "Practice/B3_MCI_Mart_Q1_2025_Answer.xlsx · PV_Thu · A4:B12", "S18a_sort_thu.png")
stat(s, 5.1, 2.65, 2.0, 3.0, "T7–CN so với T2–T6", "+25%", "157,1 so với 125,4 triệu / thứ", vsize=30)
shot(s, 7.4, 2.65, 5.34, 3.0, "Hộp Top 10 Filter (Top 5 Items by Sum of Revenue) và 5 sản phẩm; Hộp quà Tết cao cấp 135.850.000 đứng đầu", "Demo/B3_Demo_Pivot.xlsx · PV_Top5 (sheet vừa tạo) · A3:B9", "S18b_top5.png")
tip(s, L, 5.8, W, 0.85, "Thứ tự T2 → CN cố định: tạo Custom List (File → Options → Advanced → Edit Custom Lists), rồi Sort A to Z.", pro=True)
footnote(s, "Ảnh: Practice/B3_MCI_Mart_Q1_2025_Answer.xlsx, sheet PV_Thu · Demo/B3_Demo_Pivot.xlsx, sheet PV_Top5")
notes(s, 18)

# 19 · Group
s = content("Group: theo ngày, theo nhóm tự chọn", R4)
tool(s, L, 1.6, W, "Group", "Gộp mục thành nhóm mới: ngày → tuần / tháng / quý · số → khoảng · chữ → nhóm tự đặt tên.", 0.45)
shot(s, L, 2.55, 5.6, 3.0, "Hộp Grouping (Days, Number of days: 7) và Pivot 13 tuần", "Demo/B3_Demo_Pivot.xlsx · PV_Tuan (sheet vừa tạo) · A3:B17", "S19a_group_7days.png")
weeks = [("08–14/01", "118,8"), ("15–21/01", "111,7"), ("22–28/01", "112,9"), ("29/01–04/02", "31,6"), ("tuần thường", "53–74")]
for k, (lab, v) in enumerate(weeks):
    x = 7.45 + (k % 3) * 1.78; y = 2.55 + (k // 3) * 0.95
    tet = k == 3
    box(s, x, y, 1.66, 0.85, fill=BAD_BG if tet else WHITE, line=HAIR, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08, name=f"Week {k + 1}")
    tb(s, x + 0.06, y + 0.02, 1.54, 0.81, [[(v, {"b": True, "size": 18, "color": BAD_TXT if tet else NAVY})], [(lab, {"size": 11, "color": INK2})]], anchor="ctr", align="ctr", space=2)
tb(s, 11.01, 3.5, 1.73, 0.85, "triệu VNĐ / tuần; mùng 1 Tết: 29/01", size=12, color=MUTED, anchor="ctr")
shot(s, 7.45, 4.5, 5.29, 1.05, "Thông báo “Cannot group that selection.” và 3 mục ngày dạng chữ", "Demo/B3_Demo_Loi_Du_lieu.xlsx · Pivot mới từ tblSalesLoi", "S19b_cannot_group.png")
tip(s, L, 5.7, W, 0.95, "“Cannot group that selection” = cột ngày có ô CHỮ hoặc ô trống → sửa ở nguồn, Refresh. Excel 365 tự nhóm ngày theo tháng; muốn theo ngày: Group → chỉ chọn Days.")
footnote(s, "Ảnh: Demo/B3_Demo_Pivot.xlsx, sheet PV_Tuan · Demo/B3_Demo_Loi_Du_lieu.xlsx, sheet Sales_Loi")
notes(s, 19)

# 20 · Show Details
s = content("Show Details (drill-down)", R4)
tool(s, L, 1.6, W, "Show Details", "Double-click một ô giá trị → sheet mới chứa các dòng gốc tạo nên con số: từ “bao nhiêu” sang “vì sao”.", 0.45)
shot(s, L, 2.55, 4.4, 3.1, "Pivot Store_Code × Thang, ô ST07 × 3 = 8.539.000 được chọn", "Practice/B3_MCI_Mart_Q1_2025_Answer.xlsx · PV_CuaHang · A4:E18", "S20a_store_month.png")
shot(s, 6.2, 2.55, 3.6, 3.1, "Sheet chi tiết sau khi double-click ô D12, sắp xếp Order_Date: 55 dòng, từ 10/03 thưa hẳn", "Practice/B3_MCI_Mart_Q1_2025_Answer.xlsx · double-click PV_CuaHang!D12", "S20b_show_details.png")
box(s, 9.95, 2.55, 2.79, 3.1, fill=WHITE, line=HAIR, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05, name="Stat card")
tb(s, 10.1, 2.65, 2.49, 2.9, [[("ST07 Nha Trang", {"b": True, "color": NAVY})], [("T3: −52,4%", {"b": True, "size": 22, "color": BAD_TXT})],
                              "T2: 51 đơn", "01–09/03: 13 đơn", [("10–31/03: 15 đơn,", {"b": True})], [("12/22 ngày không có đơn", {"b": True})]], size=14, space=4)
tip(s, L, 5.8, W, 0.85, "Drill-down trước khi kết luận: ST04 cũng giảm 50,5% nhưng 41 đơn rải đều cả tháng. Sheet chi tiết là bản sao tĩnh, xem xong thì xóa.")
footnote(s, "Ảnh: Practice/B3_MCI_Mart_Q1_2025_Answer.xlsx, sheet PV_CuaHang")
notes(s, 20)

# 21 · Slicer & Timeline
s = content("Slicer · Timeline · Report Connections", R4)
for k, (n_, p_) in enumerate([("Slicer", "Nút bấm lọc Pivot, nhìn là biết đang lọc gì."), ("Timeline", "Thanh thời gian: lọc ngày / tháng / quý / năm."),
                              ("Report Connections", "Nối một Slicer với nhiều Pivot cùng nguồn.")]):
    x = L + k * 3.75
    tool(s, x, 1.6, 3.5, n_, p_, 0.7, size=22)
    if k < 2: vline(s, x + 3.62, 1.6, 1.2)
shot(s, L, 2.95, 6.3, 2.65, "Slicer Store_Code (2 cột) cạnh Pivot Miền × Tháng và hộp Report Connections tick 3 Pivot", "Practice/B3_MCI_Mart_Q1_2025_Answer.xlsx · PV_MienThang · A4:I20", "S21a_slicer_connections.png")
shot(s, 8.15, 2.95, 4.59, 2.65, "Timeline Order_Date (mức Months) phía trên PivotChart", "Practice/B3_MCI_Mart_Q1_2025_Answer.xlsx · PV_Ngay · D4:N11", "S21b_timeline.png")
tip(s, L, 5.75, W, 0.9, "Slicer chỉ nối được các Pivot dùng CHUNG nguồn (chung cache). Chọn tháng 2–3 trên Timeline: Chủ nhật, thứ Bảy vẫn đứng đầu → câu chuyện cuối tuần không phải do Tết.")
footnote(s, "Ảnh: Practice/B3_MCI_Mart_Q1_2025_Answer.xlsx, sheet PV_MienThang và PV_Ngay")
notes(s, 21)

# 22 · Chia phần 5
notes(section("5. Tính toán trong Pivot: tính trên kết quả đã tổng hợp"), 22)

# 23 · Summarize Values By
s = content("Summarize Values By", R5)
tool(s, L, 1.6, W, "Summarize Values By", "Chọn phép tổng hợp: Sum · Count · Average · Max · Min…", 0.45)
for k, (lab, v, sub, col) in enumerate([("Sum", "941.454.000", "doanh thu", NAVY), ("Count", "6.187", "dòng ≠ 3.144 đơn", BAD_TXT),
                                        ("Average", "152.166", "/ dòng", NAVY), ("Max", "1.470.000", "dòng lớn nhất", NAVY)]):
    stat(s, L + k * 2.8, 2.55, 2.62, 1.25, lab, v, sub, vcolor=col, vsize=22)
shot(s, L, 3.95, 6.6, 2.65, "Pivot từ tblSalesLoi: Region × Count of Revenue (tổng 6.186) thay vì Sum", "Demo/B3_Demo_Loi_Du_lieu.xlsx · Pivot mới · A3:B7 + khung Values", "S23_count_default.png")
tip(s, 8.45, 3.95, 4.29, 2.65, "Kéo cột số vào Values mà Pivot hiện COUNT → cột có ô trống hoặc chữ. Đổi sang Sum thì ra 940.480.000: Excel lặng lẽ bỏ qua 3 ô lỗi (thiếu 974.000).")
footnote(s, "Ảnh: Demo/B3_Demo_Loi_Du_lieu.xlsx, sheet Sales_Loi → PivotTable · thẻ số: Demo/B3_Demo_Pivot.xlsx")
notes(s, 23)

# 24 · % of Column Total
s = content("Show Values As: % of Column Total", R5)
tool(s, L, 1.6, W, "Show Values As", "Hiển thị con số dưới dạng khác (tỷ trọng, chênh lệch, lũy kế…) mà KHÔNG tạo dữ liệu mới.", 0.45)
shot(s, L, 2.55, 6.4, 3.0, "Category × Thang, % of Column Total: Bánh kẹo 45,6% → 30,9% → 25,5%", "Practice/B3_MCI_Mart_Q1_2025_Answer.xlsx · PV_NganhHang · A4:E12", "S24_pct_column.png")
stat(s, 8.25, 2.55, 2.2, 1.45, "Bánh kẹo · 20 ngày trước Tết", "50,8%", "ngày thường 27,9%", vsize=26)
stat(s, 10.55, 2.55, 2.19, 1.45, "Đồ uống · trước Tết", "11,7%", "ngày thường 11,5%", vsize=26)
tb(s, 8.25, 4.1, 4.49, 1.45, "Đồ uống bán gấp ~1,8 lần mỗi ngày trước Tết, nhưng cả cửa hàng cũng tăng: tỷ trọng đứng yên ≠ doanh thu đứng yên.", size=14, color=INK2)
tip(s, L, 5.7, W, 0.95, "% of COLUMN Total: mỗi cột cộng lại 100%. Tháng ở cột → so cơ cấu TRONG từng tháng. % of Grand Total: Nam 45,4% · Bắc 34,4% · Trung 20,2%.")
footnote(s, "Ảnh: Practice/B3_MCI_Mart_Q1_2025_Answer.xlsx, sheet PV_NganhHang")
notes(s, 24)

# 25 · % Difference From · Running Total
s = content("% Difference From · Running Total", R5)
table(s, L, 1.6, [3.3, 4.2, 3.57], [
    ["Tên", "Mục đích", "Quý 1/2025"],
    [[("% Difference From (previous)", {"b": True, "color": RED})], "Tăng trưởng so với mục liền trước (MoM)", [("T2 −42,2% · T3 +9,6%", {"b": True, "color": NAVY})]],
    [[("Running Total In", {"b": True, "color": RED})], "Lũy kế theo thứ tự mục", "425,9 · 671,8 · 941,5 triệu"],
], size=15, rowh=0.55, name="Functions")
shot(s, L, 3.45, 7.8, 2.2, "Pivot theo Thang: Doanh thu, MoM %, Lũy kế, Số đơn, AOV, Gross Margin %", "Practice/B3_MCI_Mart_Q1_2025_Answer.xlsx · PV_Thang · A4:G9", "S25_mom_running.png")
tip(s, 9.65, 3.45, 3.09, 2.2, "Base item (previous) = so với tháng liền trước; tháng đầu để trống.")
tb(s, L, 5.8, W, 0.8, "Ba cột cùng lấy từ MỘT trường Revenue: Show Values As chỉ đổi cách hiển thị.", size=16, color=NAVY, bold=True)
footnote(s, "Ảnh: Practice/B3_MCI_Mart_Q1_2025_Answer.xlsx, sheet PV_Thang")
notes(s, 25)

# 26 · Calculated Field
s = content("Calculated Field", R5)
tool(s, L, 1.6, W, "Calculated Field", "Tạo trường mới bằng công thức trên TỔNG của các trường khác.", 0.45)
fx(s, L, 2.5, 5.6, "AOV  =Revenue/So_don", size=16)
tb(s, L, 3.05, 5.6, 0.45, [[("Cột đếm trong tblSales:  ", {"color": INK2}), ("So_don =IF([@Line_No]=1,1,0)", {"b": True})]], size=14)
shot(s, L, 3.55, 5.6, 2.1, "Hộp Insert Calculated Field: Name AOV, Formula =Revenue/So_don", "Demo/B3_Demo_Pivot.xlsx · PV_Thang (sau khi thêm cột So_don)", "S26_calc_field.png")
for k, (lab, v, col) in enumerate([("Count of Order_ID", "6.187  ✘", BAD_TXT), ("Sum of So_don", "3.144  ✔", GOOD), ("AOV Grand Total", "299.445", NAVY)]):
    stat(s, 7.45, 2.5 + k * 1.07, 5.29, 0.95, lab, v, None, vcolor=col, vsize=22)
tip(s, L, 5.8, W, 0.85, "Calculated Field luôn tính trên TỔNG: AOV Grand Total = 941.454.000 ÷ 3.144 (tổng ÷ tổng), không phải trung bình AOV các tháng. Cần đếm thì thêm cột đếm vào nguồn.", pro=True)
footnote(s, "Ảnh: Demo/B3_Demo_Pivot.xlsx, sheet PV_Thang")
notes(s, 26)

# 27 · Chia phần 6
notes(section("6. PivotChart: biểu đồ trả lời câu hỏi"), 27)

# 28 · 3 câu hỏi chọn chart
s = content("3 câu hỏi chọn chart", R6)
tb(s, L, 1.6, 5.4, 2.1, [[("1 · ", {"b": True, "color": NAVY}), "Dữ liệu loại gì? (số, nhóm, thời gian…)"],
                         [("2 · ", {"b": True, "color": NAVY}), "Muốn truyền đạt gì? (xu hướng, so sánh, cơ cấu…)"],
                         [("3 · ", {"b": True, "color": NAVY}), "Ai xem? (nhà phân tích, giám đốc…)"]], size=17, space=10)
table(s, L, 3.75, [3.1, 2.3], [["Câu hỏi dữ liệu", "Chart"], ["Xu hướng theo thời gian", [("Line", {"b": True, "color": NAVY})]],
                               ["So sánh giữa các nhóm", [("Column / Bar", {"b": True, "color": NAVY})]], ["Cơ cấu / tỷ trọng", [("Stacked column", {"b": True, "color": NAVY})]]],
      size=15, rowh=0.45, name="Chart chooser")
tb(s, 7.3, 1.6, 2.6, 0.35, [[("✘ Pie cho xu hướng", {"b": True, "color": BAD_TXT})]], size=14)
shot(s, 7.3, 1.95, 2.6, 2.6, "PivotChart Pie doanh thu 3 tháng", "Demo/B3_Demo_Pivot.xlsx · PV_Thang", "S28a_pie_wrong.png")
tb(s, 10.14, 1.6, 2.6, 0.35, [[("✔ Line", {"b": True, "color": GOOD})]], size=14)
shot(s, 10.14, 1.95, 2.6, 2.6, "Cùng dữ liệu, PivotChart Line", "Demo/B3_Demo_Pivot.xlsx · PV_Thang", "S28b_line_right.png")
tip(s, 7.3, 4.75, 5.44, 1.85, "Tiêu đề biểu đồ là CÂU KẾT LUẬN: “Doanh thu sụt 42% trong tháng Tết”, không phải “Doanh thu theo tháng”.")
footnote(s, "Ảnh: Demo/B3_Demo_Pivot.xlsx, sheet PV_Thang · bảng chọn chart đủ 8 loại ở Buổi 6")
notes(s, 28)

# 29 · PivotChart
s = content("PivotChart", R6)
tool(s, L, 1.6, W, "PivotChart", "Biểu đồ gắn với Pivot: đổi trường, lọc Slicer / Timeline thì biểu đồ đổi theo.", 0.45)
shot(s, L, 2.55, 8.0, 4.1, "PivotChart đường doanh thu theo ngày Quý 1 (đỉnh 3 tuần trước Tết, đáy 29/01–02/02) và Timeline phía trên", "Practice/B3_MCI_Mart_Q1_2025_Answer.xlsx · PV_Ngay · D4:P33", "S29_pivotchart_ngay.png")
for k, (lab, v) in enumerate([("01–08/01", "10,7"), ("09–28/01 · trước Tết", "16,5"), ("29/01–02/02 · nghỉ Tết", "3,0"), ("từ 03/02", "9,0")]):
    stat(s, 9.85, 2.55 + k * 0.84, 2.89, 0.76, lab, v + " tr/ngày", None, vcolor=BAD_TXT if k == 2 else NAVY, vsize=17)
tip(s, 9.85, 5.95, 2.89, 0.7, "Ẩn nút lọc: Field Buttons → Hide All.")
footnote(s, "Ảnh: Practice/B3_MCI_Mart_Q1_2025_Answer.xlsx, sheet PV_Ngay")
notes(s, 29)

# 30 · Thực hành 1
s = content("Thực hành 1: Báo cáo Pivot & đối soát", R7)
chip(s, L, 1.6, 1.53, "15 phút", name="Time chip")
tb(s, L, 2.2, 6.9, 4.5, [
    [("1. ", {"b": True, "color": NAVY}), "File B3_<HoTen>.xlsx: thêm cột So_don =IF([@Line_No]=1,1,0) vào tblSales."],
    [("2. ", {"b": True, "color": NAVY}), "PV_MienThang: Region × Thang, Doanh thu (#,##0, Outline). Đối soát GETPIVOTDATA Bắc–T3 − KPI_SUMIFS!D6 = 0."],
    [("3. ", {"b": True, "color": NAVY}), "PV_NganhHang: Category × Thang, % of Column Total."],
    [("4. ", {"b": True, "color": NAVY}), "PV_Thang: Doanh thu, MoM % (% Difference From, previous), Lũy kế, Số đơn (Sum So_don), AOV (Calculated Field)."],
    [("5. ", {"b": True, "color": NAVY}), "Slicer Store_Code, Report Connections nối 3 Pivot."],
], size=15, space=9)
box(s, 8.87, 1.6, 3.87, 5.0, fill=ALT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03, name="Check box")
tb(s, 9.07, 1.75, 3.47, 4.7, [[("Kiểm tra", {"b": True, "color": NAVY, "size": 16})], [("941.454.000", {"b": True, "color": NAVY, "size": 28})],
                              "Tổng Pivot = tổng SUMIFS B2", "Bắc–T3: 90.210.000", "Bánh kẹo T1: 45,6%", "MoM T2: −42,2%", "3.144 đơn · AOV 299.445"], size=15, space=8)
footnote(s, "File khởi đầu: ../B2_Ham_va_KPI/Practice/B2_MCI_Mart_Q1_2025_Answer.xlsx")
notes(s, 30)

# 31 · Thực hành 2
s = content("Thực hành 2: Tìm câu chuyện", R7)
chip(s, L, 1.6, 1.53, "12 phút", name="Time chip")
tb(s, L, 2.2, 6.9, 4.5, [
    [("1. ", {"b": True, "color": NAVY}), "PV_Ngay: Order_Date (Group → Days) × Doanh thu → PivotChart Line; thêm Timeline Order_Date."],
    [("2. ", {"b": True, "color": NAVY}), "PV_Thu: Thu × Doanh thu, sắp xếp giảm dần → PivotChart Column."],
    [("3. ", {"b": True, "color": NAVY}), "PV_CuaHang: Store_Code × Thang + cột % so tháng trước; double-click ST07 × 3, ghi số đơn trước / sau 10/03."],
    [("4. ", {"b": True, "color": NAVY}), "Mỗi biểu đồ: tiêu đề là câu kết luận; dòng “Câu hỏi:” ở ô A2."],
], size=15, space=10)
box(s, 8.87, 1.6, 3.87, 5.0, fill=ALT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03, name="Check box")
tb(s, 9.07, 1.75, 3.47, 4.7, [[("Kiểm tra", {"b": True, "color": NAVY, "size": 16})], [("Chủ nhật", {"b": True, "color": NAVY})], "159.845.000 đứng đầu; thứ Năm 111.408.000 cuối",
                              [("ST07 tháng 3", {"b": True, "color": NAVY})], "8.539.000 (−52,4%), 55 dòng chi tiết",
                              [("Trước Tết", {"b": True, "color": NAVY})], "~16,5 triệu / ngày (09–28/01)"], size=15, space=6)
footnote(s, "Kết quả mẫu: Practice/B3_MCI_Mart_Q1_2025_Answer.xlsx (sheet PV_Ngay, PV_Thu, PV_CuaHang, ST07_T3_ChiTiet)")
notes(s, 31)

# 32 · Chữa bài
s = content("Chữa bài nhanh", R7)
items = ["Tổng Pivot = tổng SUMIFS B2 = 941.454.000 (khối đối soát chênh lệch 0)",
         "PV_NganhHang, PV_Thang dùng Show Values As; AOV bằng Calculated Field (299.445)",
         "Slicer Store_Code lọc được cả 3 Pivot",
         "PivotChart theo ngày, theo thứ; tiêu đề là câu kết luận, có dòng “Câu hỏi:”",
         "Đã drill-down ST07 tháng 3"]
for k, t in enumerate(items):
    y = 1.8 + k * 0.9
    box(s, L, y + 0.05, 0.35, 0.35, fill=WHITE, line=NAVY, lw=1.5, name=f"Checkbox {k + 1}")
    tb(s, L + 0.6, y, W - 0.6, 0.5, t, size=20, anchor="ctr")
    if k < 4:
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(L), Inches(y + 0.72), Inches(R), Inches(y + 0.72)); ln.line.color.rgb = rgb(HAIR)
footnote(s, "File đáp án Practice/B3_MCI_Mart_Q1_2025_Answer.xlsx (sheet Kiem_tra: 14 dòng đối soát) là file khởi đầu Buổi 4.")
notes(s, 32)

# 33 · Quiz
notes(section("Excel Quiz Challenge", "7 câu hỏi nhanh · câu 6–7 là tình huống"), 33)
quiz(34, "Pivot", "Ô Pivot Bắc × Tháng 3 = 90.210.000 tương đương công thức nào?",
     {"A": "=SUM(tblSales[Revenue])", "B": "=SUMIFS(tblSales[Revenue], tblSales[Region],\"Bắc\", tblSales[Thang],3)",
      "C": "=COUNTIFS(tblSales[Region],\"Bắc\", tblSales[Thang],3)", "D": "=AVERAGEIFS(tblSales[Revenue], tblSales[Region],\"Bắc\")"}, "B", osize=14)
quiz(35, "Refresh", "Bạn sửa Quantity trong tblSales nhưng Pivot không đổi. Vì sao?",
     {"A": "Pivot bị khóa", "B": "Pivot đọc từ bộ nhớ đệm, cần Refresh", "C": "Excel đang tính thủ công", "D": "Pivot chỉ đọc cột Revenue"}, "B")
quiz(36, "Summarize Values By", "Kéo Revenue vào Values, Pivot hiện “Count of Revenue”. Nguyên nhân khả dĩ nhất?",
     {"A": "Revenue là cột công thức", "B": "Cột Revenue có ô trống hoặc ô chữ", "C": "Pivot luôn mặc định Count", "D": "Chưa Refresh"}, "B")
quiz(37, "Show Values As", "Tháng ở cột, ngành hàng ở hàng. Muốn biết mỗi ngành chiếm bao nhiêu % doanh thu CỦA TỪNG THÁNG?",
     {"A": "% of Grand Total", "B": "% of Row Total", "C": "% of Column Total", "D": "Running Total In"}, "C")
quiz(38, "Calculated Field", "Calculated Field AOV = Revenue / So_don. Ô Grand Total tính thế nào?",
     {"A": "Trung bình AOV của các tháng", "B": "Tổng Revenue ÷ tổng So_don", "C": "AOV tháng lớn nhất", "D": "Tổng AOV các tháng"}, "B")
quiz(39, "Tình huống — Chọn chart", "Sếp muốn nhìn doanh thu 90 ngày Quý 1 để thấy ảnh hưởng của Tết. Chọn loại biểu đồ nào?",
     {"A": "Pie", "B": "Line", "C": "Stacked column 100%", "D": "Scatter"}, "B", qsize=20)
quiz(40, "Tình huống — Chọn công cụ", "Khi nào VẪN NÊN dùng SUMIFS thay vì Pivot?",
     {"A": "Khi cần đổi chiều phân tích liên tục", "B": "Khi báo cáo có bố cục cố định và ô kết quả được công thức khác tham chiếu (thẻ KPI, so với target)",
      "C": "Khi dữ liệu có nhiều dòng", "D": "Khi cần xem các dòng gốc"}, "B", osize=15)

# 41 · Tổng kết + insight
s = content("Tổng kết & insight của bạn", "PHẦN 8 · QUIZ & TỔNG KẾT")
tb(s, L, 1.65, 5.0, 0.45, "Điểm chính cần nhớ", size=18, color=NAVY, bold=True)
tb(s, L, 2.2, 5.0, 4.4, ["(1) Mỗi ô Pivot = một SUMIFS: tọa độ là bộ lọc.", "(2) Nguồn Pivot là Table; Pivot đọc cache → Refresh.",
                        "(3) Con số bất thường → drill-down trước khi kết luận.", "(4) Show Values As, Calculated Field tính trên kết quả đã tổng hợp: tỷ lệ luôn tổng ÷ tổng.",
                        "(5) Chọn chart theo câu hỏi; tiêu đề là câu kết luận."], size=17, space=10)
box(s, 6.97, 1.65, 5.77, 3.8, fill=ALT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03, name="Insight box")
tb(s, 7.17, 1.8, 5.37, 0.45, "Insight của bạn (1 câu)", size=18, color=NAVY, bold=True)
tb(s, 7.17, 2.35, 5.37, 2.2, "“Doanh thu ___ là ___ (con số, có so sánh), vì ___ (nguyên nhân khả dĩ), nên ___ (hàm ý).”", size=22)
for k, t in enumerate(["Con số", "Nguyên nhân", "Hàm ý"]):
    chip(s, 7.17 + k * 1.84, 4.8, 1.69, t, fill=WHITE, line=HAIR, name=f"Tag {t}", h=0.45)
tb(s, 6.97, 5.6, 5.77, 0.9, "Ba câu chuyện hôm nay: Tết (tháng 1 > tháng 2) · ST07 từ 10/03 · cuối tuần +25%.", size=14, color=INK2)
notes(s, 41)

# 42 · Preview B4
s = content("Buổi sau: Power Query", "PHẦN 8 · QUIZ & TỔNG KẾT")
tb(s, L, 1.65, 7.0, 1.1, "“Làm sao để việc làm sạch dữ liệu chỉ làm MỘT LẦN, từ đó chỉ cần bấm Refresh?”", size=19, color=NAVY, bold=True)
tb(s, L, 2.8, 7.0, 1.0, [[("Giới hạn hôm nay: ", {"b": True}), "tháng mới về dạng CSV bẩn, tách nhiều file → lại copy-paste, sửa tay rồi mới Refresh được Pivot."]], size=16)
tb(s, L, 3.85, 7.0, 2.8, [[("Về nhà", {"b": True, "color": NAVY})], "Đọc handout B3",
                          "Bài tập về nhà (bắt buộc): B3_BTVN_MCI_Mart_T1_T4_2025.xlsx — nộp trước Buổi 4 tại hocvien.mcivietnam.com, tên file B3_BTVN_<HoTen>.xlsx",
                          "Tick “Checklist trước Buổi 4” (mục 10 handout)", [("Buổi 4 bắt buộc máy Windows có Power Query (Data → Get Data)", {"b": True, "color": BAD_TXT})]], size=15, space=6)
bricks(s, "B3 · Báo cáo Pivot", "B4 · Pipeline Power Query")
tb(s, 9.07, 5.7, 3.67, 0.5, "Giữ file đáp án B3: file khởi đầu Buổi 4.", size=13, color=INK2)
notes(s, 42)

assert len(prs.slides) == 42, len(prs.slides)
prs.save(OUT)
print("ok", OUT, len(prs.slides), "slide;", f"còn {len(MISSING)} khung [CHỤP] chưa có ảnh:" if MISSING else "đủ ảnh.", " ".join(MISSING))
