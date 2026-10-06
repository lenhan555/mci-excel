"""Buổi 4 (Power Query): số liệu dùng trong buổi (build_b4_meta.json) và bài tập về nhà bắt buộc.

Chạy được trên Mac lẫn Windows (không cần Power Query). File đáp án của lớp (Practice/B4_MCI_Mart_Answer.xlsx, có
query thật) dựng riêng bằng build_b4.py trên Windows.

1. Đọc các file Practice/ của B4 (CSV bẩn, B4_Lookups.xlsx, B4_Target_2025.xlsx), làm sạch bằng pandas đúng như các
   bước Power Query trong kịch bản (Locale vi-VN, Trim, Remove Duplicates), tính mọi con số lời giảng / deck / handout /
   infographic dùng tới, ghi _generator/build_b4_meta.json.
2. Dựng bài tập về nhà: Homework/B4_BTVN_Du_lieu/ (12 CSV tháng 4/2025 theo cửa hàng, file bổ sung của ST07),
   đề B4_BTVN_MCI_Mart_T4_2025.xlsx (De_bai, So_lieu_he_thong, Product, Target_2025) và đáp án _Answer.xlsx
   (bảng kết quả của các query, Dap_an_M, Kiem_tra, Dap_an_loi, Insight_mau). Mở và lưu lại bằng Excel để có giá
   trị công thức (--no-excel để bỏ bước này).

    python3 _generator/build_b4_homework.py            # Mac (venv có openpyxl, pandas) hoặc Windows: python ...
"""
import csv, json, os, random, shutil, sys
from datetime import date, timedelta
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

for s_ in (sys.stdout, sys.stderr):
    if hasattr(s_, "reconfigure"): s_.reconfigure(encoding="utf-8", errors="replace")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

ROOT = os.path.normpath(os.environ.get("MCI_ROOT") or os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
GEN = os.path.dirname(os.path.abspath(__file__))
B4 = os.path.join(ROOT, "1_Sessions", "B4_Power_Query")
PRACTICE = os.path.join(B4, "Practice")
DATA_SALES = os.path.join(PRACTICE, "Data_Sales")
NEW_MONTH = os.path.join(PRACTICE, "New_Month")
HW = os.path.join(B4, "Homework")
HW_DATA = os.path.join(HW, "B4_BTVN_Du_lieu")
HW_DE = os.path.join(HW, "B4_BTVN_MCI_Mart_T4_2025.xlsx")
HW_ANS = os.path.join(HW, "B4_BTVN_MCI_Mart_T4_2025_Answer.xlsx")
USE_EXCEL = "--no-excel" not in sys.argv

vn = lambda x: f"{x:,.0f}".replace(",", ".")
pct = lambda x, d=1: f"{x * 100:.{d}f}%".replace(".", ",")

# ---------------------------------------------------------------------
# 1. Số liệu trong buổi (từ chính các file phát cho học viên)
# ---------------------------------------------------------------------
FILES = [("Sales_2024.csv", DATA_SALES), ("Sales_2025-01.csv", DATA_SALES), ("Sales_2025-02.csv", DATA_SALES),
         ("Sales_2025-03.csv", DATA_SALES), ("Sales_2025-04.csv", NEW_MONTH)]
COLS = ["Order_ID", "Line_No", "Order_Date", "Store_Code", "Customer_ID", "Product_Code", "Quantity", "Unit_Price"]


def read_raw(name, folder):
    return pd.read_csv(os.path.join(folder, name), dtype=str, encoding="utf-8-sig", keep_default_na=False)


def clean(raw):
    """Đúng các bước của query Sales: Trim Store_Code, Using Locale vi-VN (ngày dd/mm/yyyy, số có dấu chấm ngăn nghìn),
    Remove Duplicates trên mọi cột, Revenue/Nam/Thang."""
    d = raw.copy()
    d["Store_Code"] = d.Store_Code.str.strip()
    d["Order_Date"] = pd.to_datetime(d.Order_Date, format="%d/%m/%Y")
    d["Unit_Price"] = d.Unit_Price.str.replace(".", "", regex=False).astype(int)
    d["Quantity"] = d.Quantity.astype(int)
    d["Line_No"] = d.Line_No.astype(int)
    d = d.drop_duplicates().reset_index(drop=True)
    d["Revenue"] = d.Quantity * d.Unit_Price
    d["Nam"] = d.Order_Date.dt.year
    d["Thang"] = d.Order_Date.dt.month
    return d


meta = {"files": {}}
parts = []
for name, folder in FILES:
    raw = read_raw(name, folder)
    c = clean(raw)
    ws_rows = int((raw.Store_Code != raw.Store_Code.str.strip()).sum())
    day13 = int((raw.Order_Date.str[:2].astype(int) >= 13).sum())
    meta["files"][name] = {"raw": len(raw), "rows": len(c), "revenue": int(c.Revenue.sum()), "ws_store": ws_rows,
                           "dup": len(raw) - len(raw.drop_duplicates()), "day13": day13,
                           "day_le12": len(raw) - day13, "distinct_store_raw": int(raw.Store_Code.nunique())}
    c["File"] = name
    parts.append(c)
sales_all = pd.concat(parts, ignore_index=True)
sales_q = sales_all[sales_all.File != "Sales_2025-04.csv"]
apr = sales_all[sales_all.File == "Sales_2025-04.csv"].reset_index(drop=True)
jan_raw = read_raw("Sales_2025-01.csv", DATA_SALES)
jan = parts[1]

# Giá khi đổi kiểu KHÔNG dùng Locale trên máy en-US: "19.000" thành 19
jan_price_us = jan_raw.Unit_Price.astype(float)

lk = os.path.join(PRACTICE, "B4_Lookups.xlsx")
product_raw = pd.read_excel(lk, "Product")
product = product_raw.drop_duplicates("Product_Code").reset_index(drop=True)
store = pd.read_excel(lk, "Store")
customer = pd.read_excel(lk, "Customer")
p015 = int((sales_q.Product_Code == "P015").sum())
orphans = int((~sales_all.Product_Code.isin(product.Product_Code)).sum())

tw = load_workbook(os.path.join(PRACTICE, "B4_Target_2025.xlsx"), data_only=True)["Target_2025"]
target_rows = [r for r in tw.iter_rows(min_row=4, values_only=True) if r[0]]
target = pd.DataFrame([(r[0], m + 1, int(r[2 + m])) for r in target_rows for m in range(12)], columns=["Store_Code", "Thang", "Target"])
t_month = target.groupby("Thang").Target.sum()

cal_start, cal_end = date(2024, 1, 1), date(2025, 12, 31)
cal_rows = (cal_end - cal_start).days + 1

bym = sales_all.groupby(["Nam", "Thang"]).agg(rows=("Order_ID", "size"), rev=("Revenue", "sum"))
st12 = sales_all[sales_all.Store_Code == "ST12"].groupby(["Nam", "Thang"]).size()
jan_store = jan.groupby("Store_Code").agg(rows=("Order_ID", "size"), rev=("Revenue", "sum"))
act_2025 = sales_all[sales_all.Nam == 2025].groupby("Thang").Revenue.sum()

meta.update({
    "q_rows": len(sales_q), "q_rev": int(sales_q.Revenue.sum()),
    "all_rows": len(sales_all), "all_rev": int(sales_all.Revenue.sum()),
    "q1_rows": int((sales_q.Nam == 2025).sum()), "q1_rev": int(sales_q[sales_q.Nam == 2025].Revenue.sum()),
    "p015_rows": p015, "merge_trap_rows": len(sales_q) + p015, "left_anti_rows": orphans,
    "product_raw": len(product_raw), "product": len(product), "store": len(store), "customer": len(customer),
    "target_rows": len(target), "target_total": int(target.Target.sum()),
    "target_month": {int(k): int(v) for k, v in t_month.items()},
    "actual_2025": {int(k): int(v) for k, v in act_2025.items()},
    "vs_target": {int(k): float(act_2025[k] / t_month[k]) for k in act_2025.index},
    "calendar_rows": cal_rows, "calendar_start": "01/01/2024", "calendar_end": "31/12/2025",
    "tet_2024": "10/02/2024", "tet_2025": "29/01/2025",
    "by_month": [{"nam": int(n), "thang": int(t), "rows": int(r.rows), "rev": int(r.rev)} for (n, t), r in bym.iterrows()],
    "st12_first": f"{int(st12.index[0][1]):02d}/{int(st12.index[0][0])}", "st12_first_rows": int(st12.iloc[0]),
    "st12_before_jun": int(((sales_all.Store_Code == "ST12") & (sales_all.Nam == 2024) & (sales_all.Thang < 6)).sum()),
    "jan_store": {s: {"rows": int(r.rows), "rev": int(r.rev)} for s, r in jan_store.iterrows()},
    "jan_price_us_min": float(jan_price_us.min()), "jan_price_us_max": float(jan_price_us.max()),
    "jan_price_min": int(jan.Unit_Price.min()), "jan_price_max": int(jan.Unit_Price.max()),
})
m_ = meta
assert m_["q_rows"] == 28328 and m_["q_rev"] == 4139884500, (m_["q_rows"], m_["q_rev"])
assert m_["all_rows"] == 30247 and m_["all_rev"] == 4436849500
assert m_["q1_rows"] == 6187 and m_["q1_rev"] == 941454000            # khớp KPI_SUMIFS B2 và Pivot B3
assert m_["left_anti_rows"] == 0 and m_["st12_before_jun"] == 0

# ---------------------------------------------------------------------
# 2. Bài tập về nhà: tháng 4/2025, mỗi cửa hàng một file CSV
# ---------------------------------------------------------------------
rnd = random.Random(2025_04)
STORES = sorted(apr.Store_Code.unique())
SPLIT_STORE, SPLIT_DAY = "ST07", 16                 # ST07 gửi thiếu nửa sau tháng: phần 2 nằm ở Bo_sung/
DUP_STORES = ["ST02", "ST05", "ST10"]               # mỗi file 1 dòng trùng tuyệt đối
LOWER_STORES = {"ST03": 2, "ST08": 2}               # Product_Code viết thường (p018)
SPACE_STORES = {"ST09": 2, "ST11": 2}               # Product_Code dính dấu cách cuối (P025 )
HW_COLS = ["Order_ID", "Line_No", "Order_Date", "Customer_ID", "Product_Code", "Quantity", "Unit_Price"]


def price_vn(v):
    return f"{v:,}".replace(",", ".")


def write_store_file(path, df, issues):
    rows = []
    for _, r in df.iterrows():
        rows.append([r.Order_ID, str(r.Line_No), r.Order_Date.strftime("%d/%m/%Y"), r.Customer_ID, r.Product_Code,
                     str(r.Quantity), price_vn(r.Unit_Price)])
    for kind, idx in issues:
        if kind == "lower": rows[idx][4] = rows[idx][4].lower()
        if kind == "space": rows[idx][4] = rows[idx][4] + " "
    for kind, idx in sorted([i for i in issues if i[0] == "dup"], key=lambda x: -x[1]):
        rows.insert(idx + 1, list(rows[idx]))
    rows.append(["Tổng cộng", "", "", "", "", str(int(df.Quantity.sum())), ""])
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, lineterminator="\r\n"); w.writerow(HW_COLS); w.writerows(rows)
    return rows


if os.path.isdir(HW_DATA): shutil.rmtree(HW_DATA)
os.makedirs(os.path.join(HW_DATA, "Cua_hang")); os.makedirs(os.path.join(HW_DATA, "Bo_sung"))
issues_log = []                                     # (file, dòng trong CSV kể cả dòng tiêu đề, loại, giá trị)
file_rows = {}
for s in STORES:
    d = apr[apr.Store_Code == s].sort_values(["Order_ID", "Line_No"]).reset_index(drop=True)
    chunks = [(f"Cua_hang/{s}_2025-04.csv", d)]
    if s == SPLIT_STORE:
        chunks = [(f"Cua_hang/{s}_2025-04.csv", d[d.Order_Date.dt.day < SPLIT_DAY].reset_index(drop=True)),
                  (f"Bo_sung/{s}_2025-04_phan2.csv", d[d.Order_Date.dt.day >= SPLIT_DAY].reset_index(drop=True))]
    for rel, dd in chunks:
        issues = []
        if s in DUP_STORES and rel.startswith("Cua_hang"):
            issues.append(("dup", rnd.randrange(5, len(dd) - 5)))
        for kind, pool in (("lower", LOWER_STORES), ("space", SPACE_STORES)):
            if s in pool:
                for i in rnd.sample(range(len(dd)), pool[s]): issues.append((kind, i))
        rows = write_store_file(os.path.join(HW_DATA, rel), dd, issues)
        file_rows[rel] = {"data": len(dd), "lines": len(rows)}
        for i, r in enumerate(rows, start=2):
            if r[0] == "Tổng cộng": issues_log.append((rel, i, "footer", r[5]))
            elif r[4] != r[4].strip() or r[4] != r[4].upper(): issues_log.append((rel, i, "code", r[4]))
        seen = {}
        for i, r in enumerate(rows, start=2):
            k = tuple(r)
            if k in seen and r[0] != "Tổng cộng": issues_log.append((rel, i, "dup", f"trùng dòng {seen[k]}"))
            seen.setdefault(k, i)

n_footer = sum(1 for x in issues_log if x[2] == "footer")
n_code = sum(1 for x in issues_log if x[2] == "code")
n_dup = sum(1 for x in issues_log if x[2] == "dup")
n_lower = sum(1 for x in issues_log if x[2] == "code" and x[3].strip() != x[3].strip().upper())
n_space = n_code - n_lower
lines_main = sum(v["lines"] for k, v in file_rows.items() if k.startswith("Cua_hang"))

# Kết quả đúng (đủ file bổ sung) và kết quả khi quên file bổ sung
hw = apr.copy()
g = hw.groupby("Store_Code").agg(So_dong=("Order_ID", "size"), Tong_SL=("Quantity", "sum"), Doanh_thu=("Revenue", "sum")).reset_index()
t4 = target[target.Thang == 4].set_index("Store_Code").Target
g["Target"] = g.Store_Code.map(t4)
g["Pct_Dat"] = g.Doanh_thu / g.Target
missing = hw[(hw.Store_Code == SPLIT_STORE) & (hw.Order_Date.dt.day >= SPLIT_DAY)]
st07 = g.set_index("Store_Code").loc[SPLIT_STORE]
st07_part1_rev = int(st07.Doanh_thu - missing.Revenue.sum())
below90 = g[g.Pct_Dat < 0.9].Store_Code.tolist()
meta["hw"] = {
    "rows": len(hw), "revenue": int(hw.Revenue.sum()), "qty": int(hw.Quantity.sum()), "stores": len(g),
    "target": int(g.Target.sum()), "pct_all": float(g.Doanh_thu.sum() / g.Target.sum()),
    "files_main": sum(1 for k in file_rows if k.startswith("Cua_hang")), "lines_main": lines_main,
    "footer": n_footer, "code_issues": n_code, "code_lower": n_lower, "code_space": n_space, "dup": n_dup,
    "missing_rows": len(missing), "missing_rev": int(missing.Revenue.sum()),
    "rows_without_bosung": len(hw) - len(missing), "rev_without_bosung": int(hw.Revenue.sum() - missing.Revenue.sum()),
    "st07_rev": int(st07.Doanh_thu), "st07_pct": float(st07.Pct_Dat), "st07_target": int(st07.Target),
    "st07_part1_rev": st07_part1_rev, "st07_part1_pct": st07_part1_rev / float(st07.Target),
    "st07_part1_rows": int(st07.So_dong - len(missing)),
    "below90": below90, "below90_n": len(below90),
    "min_store": g.loc[g.Pct_Dat.idxmin(), "Store_Code"], "min_pct": float(g.Pct_Dat.min()),
    "by_store": {r.Store_Code: {"rows": int(r.So_dong), "qty": int(r.Tong_SL), "rev": int(r.Doanh_thu), "target": int(r.Target), "pct": float(r.Pct_Dat)} for r in g.itertuples()},
}
mh = meta["hw"]
assert mh["rows"] == meta["files"]["Sales_2025-04.csv"]["rows"] and mh["revenue"] == meta["files"]["Sales_2025-04.csv"]["revenue"]

# Khóa phẳng cho infographic (build_infographic.py điền {{info.…}})
meta["info"] = {
    "raw_jan": meta["files"]["Sales_2025-01.csv"]["raw"], "rows_jan": meta["files"]["Sales_2025-01.csv"]["rows"],
    "day13_jan": meta["files"]["Sales_2025-01.csv"]["day13"], "distinct_jan": meta["files"]["Sales_2025-01.csv"]["distinct_store_raw"],
    "q_rows": m_["q_rows"], "q_rev": m_["q_rev"], "all_rows": m_["all_rows"], "all_rev": m_["all_rev"],
    "apr_rows": mh["rows"], "apr_rev": mh["revenue"], "merge_trap_rows": m_["merge_trap_rows"],
    "product_raw": m_["product_raw"], "product": m_["product"], "store": m_["store"], "customer": m_["customer"],
    "target_rows": m_["target_rows"], "calendar_rows": m_["calendar_rows"],
    "t1_target": m_["target_month"][1], "t1_actual": m_["actual_2025"][1], "t1_pct": m_["vs_target"][1],
    "t2_target": m_["target_month"][2], "t2_actual": m_["actual_2025"][2], "t2_pct": m_["vs_target"][2],
    "st12_first_rows": m_["st12_first_rows"], "hw_st07_pct": mh["st07_pct"], "hw_st07_part1_pct": mh["st07_part1_pct"],
}
with open(os.path.join(GEN, "build_b4_meta.json"), "w", encoding="utf-8") as f:
    json.dump(meta, f, ensure_ascii=False, indent=1)

# ---------------------------------------------------------------------
# Workbook đề và đáp án
# ---------------------------------------------------------------------
F = "Arial"
HDR_FONT = Font(name=F, bold=True, color="FFFFFF"); HDR_FILL = PatternFill("solid", fgColor="203864")
TITLE = Font(name=F, size=14, bold=True, color="203864"); BODY = Font(name=F, size=10); BOLD = Font(name=F, size=10, bold=True)
INPUT = Font(name=F, size=10, color="0000FF"); NOTE = Font(name=F, size=10, italic=True, color="585858")
MONO = Font(name="Consolas", size=10)
DATE_FMT, NUM_FMT, PCT_FMT = "dd/mm/yyyy", "#,##0", "0.0%"


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1): ws.column_dimensions[get_column_letter(i)].width = w


def header(ws, row, labels, col=1):
    for j, h in enumerate(labels, start=col):
        c = ws.cell(row=row, column=j, value=h); c.font = HDR_FONT; c.fill = HDR_FILL


def table_df(ws, df, name, start_row=1, formats=None, widths=None):
    header(ws, start_row, list(df.columns))
    for i, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for j, v in enumerate(row, start=1):
            v = v.to_pydatetime() if isinstance(v, pd.Timestamp) else (v.item() if hasattr(v, "item") else v)
            c = ws.cell(row=i, column=j, value=v); c.font = BODY
            if formats and df.columns[j - 1] in formats: c.number_format = formats[df.columns[j - 1]]
    t = Table(displayName=name, ref=f"A{start_row}:{get_column_letter(len(df.columns))}{start_row + len(df)}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True); ws.add_table(t)
    if widths: set_widths(ws, widths)


def wrap_rows(ws, start, rows, widths=None, head=True):
    for i, rr in enumerate(rows, start=start):
        for j, val in enumerate(rr, start=1):
            c = ws.cell(row=i, column=j, value=val); c.alignment = Alignment(wrap_text=True, vertical="top")
            c.font = HDR_FONT if (head and i == start) else BODY
            if head and i == start: c.fill = HDR_FILL
    if widths: set_widths(ws, widths)


def common_sheets(wb):
    so = wb.create_sheet("So_lieu_he_thong")
    so["A1"] = "SỐ LIỆU ĐỐI SOÁT TỪ HỆ THỐNG BÁN HÀNG — THÁNG 4/2025 THEO CỬA HÀNG"; so["A1"].font = TITLE
    so["A2"] = "Bảng Group By của bạn (sau khi làm sạch) phải khớp đúng các con số này, từng cửa hàng."; so["A2"].font = NOTE
    header(so, 4, ["Store_Code", "Số dòng", "Tổng Quantity", "Doanh thu (VNĐ)"])
    for i, r in enumerate(g.itertuples(), start=5):
        for j, v in enumerate([r.Store_Code, int(r.So_dong), int(r.Tong_SL), int(r.Doanh_thu)], start=1):
            c = so.cell(row=i, column=j, value=v); c.font = BODY
            if j > 1: c.number_format = NUM_FMT
    i = 5 + len(g)
    so.cell(row=i, column=1, value="Tổng").font = BOLD
    for j, v in enumerate([mh["rows"], mh["qty"], mh["revenue"]], start=2):
        c = so.cell(row=i, column=j, value=v); c.font = BOLD; c.number_format = NUM_FMT
    set_widths(so, [12, 12, 15, 18])
    pr = wb.create_sheet("Product")
    table_df(pr, product[["Product_Code", "Product_Name", "Category", "Unit_Price", "Unit_Cost"]], "Product",
             formats={"Unit_Price": NUM_FMT, "Unit_Cost": NUM_FMT}, widths=[13, 38, 18, 11, 11])
    tg = wb.create_sheet("Target_2025")
    src = load_workbook(os.path.join(PRACTICE, "B4_Target_2025.xlsx"))["Target_2025"]
    for row in src.iter_rows():
        for c in row:
            n = tg.cell(row=c.row, column=c.column, value=c.value)
            n.font = TITLE if c.row == 1 else (HDR_FONT if c.row == 3 else BODY)
            if c.row == 3: n.fill = HDR_FILL
            if c.row > 3 and c.column > 2: n.number_format = NUM_FMT
    set_widths(tg, [11, 22] + [12] * 12)


TASKS = [
    ("1", "Kết nối cả thư mục", "Data → Get Data → From File → From Folder → chọn thư mục B4_BTVN_Du_lieu\\Cua_hang → Combine & Transform Data "
     "(file mẫu: file đầu tiên). Đổi tên query thành Sales_T4. Chưa đụng tới thư mục Bo_sung."),
    ("2", "Store_Code lấy từ tên file", "Các file không có cột Store_Code: mã cửa hàng nằm trong tên file (ST01_2025-04.csv). Giữ cột Source.Name → "
     "Transform → Extract → Text Before Delimiter, dấu \"_\" → đổi tên cột thành Store_Code."),
    ("3", "Làm sạch", "Xóa bước Changed Type tự động. Order_Date → Using Locale: Date, Vietnamese (Vietnam); Unit_Price → Using Locale: Whole Number, "
     "Vietnamese (Vietnam); Line_No, Quantity → Whole Number. Bỏ dòng \"Tổng cộng\" cuối mỗi file. Product_Code: Trim rồi UPPERCASE. "
     "Remove Duplicates (cả bảng). Thêm cột Revenue = Quantity × Unit_Price."),
    ("4", "Kiểm tra mã sản phẩm", "Nạp sheet Product (Data → From Table/Range). Merge Queries as New: Sales_T4 × Product theo Product_Code, "
     "Join Kind = Left Anti. Ghi lại số dòng trước khi sửa mã (bước 3) và sau khi sửa: phải còn 0 dòng."),
    ("5", "Đối soát với hệ thống", "Tạo query Reference từ Sales_T4 → Group By Store_Code: So_dong = Count Rows, Tong_SL = Sum Quantity, "
     "Doanh_thu = Sum Revenue. So từng cửa hàng với sheet So_lieu_he_thong. Cửa hàng nào lệch, vì sao? Xử lý bằng thư mục Bo_sung "
     "(không gõ tay số liệu), Refresh All đến khi khớp cả 12 cửa hàng."),
    ("6", "% đạt target tháng 4", "Nạp Target_2025 (From Table/Range, vùng A3:N15), Unpivot về Store_Code | Thang | Target, Thang là số, lọc Thang = 4. "
     "Merge bảng Group By (câu 5) với Target theo Store_Code, thêm cột Pct_Dat = Doanh_thu / Target. Đặt tên query KetQua_T4."),
    ("7", "Load và insight", "Close & Load: Sales_T4 và KetQua_T4 ra sheet (Table); các query còn lại Only Create Connection. "
     "Sheet Insight: ghi các lỗi đã sửa (loại lỗi, số dòng) và 1 câu insight có con số so sánh, nguyên nhân khả dĩ, hàm ý."),
]
SELF_CHECK = [
    f"Sales_T4: {vn(mh['rows'])} dòng, doanh thu {vn(mh['revenue'])} VNĐ (khớp số tháng 4 sau Refresh trên lớp)",
    "Merge Left Anti Sales_T4 × Product: 0 dòng",
    f"Group By khớp So_lieu_he_thong ở cả 12 cửa hàng (Tổng Quantity {vn(mh['qty'])})",
    f"Cả chuỗi đạt {pct(mh['pct_all'])} target tháng 4; ST07 đạt {pct(mh['st07_pct'])}",
]

# --- Đề
wb = Workbook(); de = wb.active; de.title = "De_bai"
de["A1"] = "BÀI TẬP VỀ NHÀ BUỔI 4 (BẮT BUỘC) — PIPELINE POWER QUERY CHO DỮ LIỆU THÁNG 4/2025 THEO CỬA HÀNG"; de["A1"].font = TITLE
de["A2"] = ("Nộp trước Buổi 5. Đặt tên file: B4_BTVN_<HoTen>.xlsx (làm ngay trên file này). Tháng 4/2025, hệ thống POS mới xuất mỗi cửa hàng một file CSV "
            f"vào thư mục B4_BTVN_Du_lieu\\Cua_hang ({mh['files_main']} file). Kế toán cần doanh thu và % đạt target theo cửa hàng — nhưng trước hết phải khớp số liệu hệ thống.")
de["A2"].font = NOTE; de["A2"].alignment = Alignment(wrap_text=True, vertical="top"); de.row_dimensions[2].height = 42
de["A3"] = "Nộp bài tại https://hocvien.mcivietnam.com/auth#/ — Tài khoản: email bạn đăng ký tại MCI · Mật khẩu: hv12345678"; de["A3"].font = BOLD
de.merge_cells("A2:C2")
header(de, 5, ["#", "Việc cần làm", "Yêu cầu"])
for i, row in enumerate(TASKS, start=6):
    for j, val in enumerate(row, start=1):
        c = de.cell(row=i, column=j, value=val); c.font = BODY; c.alignment = Alignment(wrap_text=True, vertical="top")
kr = 6 + len(TASKS) + 1
de.cell(row=kr, column=1, value="Tự kiểm tra trước khi nộp").font = BOLD
for i, tx in enumerate(SELF_CHECK, start=kr + 1):
    de.cell(row=i, column=2, value="☐ " + tx).font = BODY
de.cell(row=kr + len(SELF_CHECK) + 2, column=1, value="Gợi ý: xem handout B4, mục 7 (lấy thông tin từ tên file) và mục 3 (đối soát trong Power Query).").font = NOTE
set_widths(de, [5, 30, 110])
common_sheets(wb)
ins_ = wb.create_sheet("Insight")
ins_["A1"] = "INSIGHT CỦA BẠN"; ins_["A1"].font = TITLE
ins_["A3"] = "Lỗi dữ liệu đã sửa (loại lỗi · số dòng · cách sửa):"; ins_["A3"].font = BOLD
ins_["A8"] = "Câu insight (con số có so sánh · nguyên nhân khả dĩ · hàm ý):"; ins_["A8"].font = BOLD
set_widths(ins_, [120])
wb.active = 0
for w_ in wb.worksheets: w_.sheet_view.tabSelected = w_.title == "De_bai"
wb.save(HW_DE)

# --- Đáp án
M_SALES = r'''let
    // Tham số: thư mục chứa các file cửa hàng (sửa theo máy của bạn)
    Source = Folder.Files("C:\...\B4_BTVN_Du_lieu\Cua_hang"),
    CsvOnly = Table.SelectRows(Source, each [Extension] = ".csv"),
    Parsed = Table.AddColumn(CsvOnly, "Data", each Table.PromoteHeaders(Csv.Document([Content], [Delimiter = ",", Encoding = 65001, QuoteStyle = QuoteStyle.None]), [PromoteAllScalars = true])),
    Kept = Table.SelectColumns(Parsed, {"Name", "Data"}),
    Expanded = Table.ExpandTableColumn(Kept, "Data", {"Order_ID", "Line_No", "Order_Date", "Customer_ID", "Product_Code", "Quantity", "Unit_Price"}),
    #"Extracted Text Before Delimiter" = Table.TransformColumns(Expanded, {{"Name", each Text.BeforeDelimiter(_, "_"), type text}}),
    #"Renamed Columns" = Table.RenameColumns(#"Extracted Text Before Delimiter", {{"Name", "Store_Code"}}),
    #"Filtered Rows" = Table.SelectRows(#"Renamed Columns", each [Order_ID] <> "Tổng cộng"),
    #"Changed Type" = Table.TransformColumnTypes(#"Filtered Rows", {{"Order_ID", type text}, {"Line_No", Int64.Type}, {"Customer_ID", type text}, {"Product_Code", type text}, {"Quantity", Int64.Type}}),
    #"Changed Type with Locale" = Table.TransformColumnTypes(#"Changed Type", {{"Order_Date", type date}, {"Unit_Price", Int64.Type}}, "vi-VN"),
    #"Trimmed Text" = Table.TransformColumns(#"Changed Type with Locale", {{"Product_Code", Text.Trim, type text}}),
    #"Uppercased Text" = Table.TransformColumns(#"Trimmed Text", {{"Product_Code", Text.Upper, type text}}),
    #"Removed Duplicates" = Table.Distinct(#"Uppercased Text"),
    #"Added Custom" = Table.AddColumn(#"Removed Duplicates", "Revenue", each [Quantity] * [Unit_Price], Int64.Type)
in
    #"Added Custom"'''
M_PRODUCT = '''let
    Source = Excel.CurrentWorkbook(){[Name = "Product"]}[Content],
    #"Changed Type" = Table.TransformColumnTypes(Source, {{"Product_Code", type text}, {"Category", type text}})
in
    #"Changed Type"'''
M_ANTI = '''let
    Source = Table.NestedJoin(Sales_T4, {"Product_Code"}, Product, {"Product_Code"}, "Product", JoinKind.LeftAnti)
in
    Source    // phải ra 0 dòng'''
M_GROUP = '''let
    Source = Sales_T4,
    #"Grouped Rows" = Table.Group(Source, {"Store_Code"}, {{"So_dong", each Table.RowCount(_), Int64.Type}, {"Tong_SL", each List.Sum([Quantity]), Int64.Type}, {"Doanh_thu", each List.Sum([Revenue]), Int64.Type}})
in
    #"Grouped Rows"'''
M_TARGET = '''let
    Source = Excel.CurrentWorkbook(){[Name = "Table1"]}[Content],   // bảng tạo bởi From Table/Range trên A3:N15
    #"Unpivoted Other Columns" = Table.UnpivotOtherColumns(Source, {"Store_Code", "Store_Name"}, "Attribute", "Value"),
    #"Extracted Text After Delimiter" = Table.TransformColumns(#"Unpivoted Other Columns", {{"Attribute", each Text.AfterDelimiter(_, " "), type text}}),
    #"Renamed Columns" = Table.RenameColumns(#"Extracted Text After Delimiter", {{"Attribute", "Thang"}, {"Value", "Target"}}),
    #"Changed Type" = Table.TransformColumnTypes(#"Renamed Columns", {{"Thang", Int64.Type}, {"Target", Int64.Type}}),
    #"Filtered Rows" = Table.SelectRows(#"Changed Type", each [Thang] = 4)
in
    #"Filtered Rows"'''
M_KETQUA = '''let
    Source = Table.NestedJoin(DoiSoat_T4, {"Store_Code"}, Target_T4, {"Store_Code"}, "Target_T4", JoinKind.LeftOuter),
    #"Expanded Target_T4" = Table.ExpandTableColumn(Source, "Target_T4", {"Target"}),
    #"Added Custom" = Table.AddColumn(#"Expanded Target_T4", "Pct_Dat", each [Doanh_thu] / [Target], Percentage.Type),
    #"Sorted Rows" = Table.Sort(#"Added Custom", {{"Store_Code", Order.Ascending}})
in
    #"Sorted Rows"'''

wb = Workbook(); ws = wb.active; ws.title = "Sales_T4"
out = hw.sort_values(["Store_Code", "Order_ID", "Line_No"])[["Order_ID", "Line_No", "Order_Date", "Customer_ID", "Product_Code", "Quantity", "Unit_Price", "Store_Code", "Revenue"]]
table_df(ws, out, "Sales_T4", formats={"Order_Date": DATE_FMT, "Unit_Price": NUM_FMT, "Revenue": NUM_FMT}, widths=[14, 8, 12, 12, 13, 9, 11, 11, 12])
ws.freeze_panes = "A2"
kq = wb.create_sheet("KetQua_T4")
table_df(kq, g[["Store_Code", "So_dong", "Tong_SL", "Doanh_thu", "Target", "Pct_Dat"]], "KetQua_T4",
         formats={"So_dong": NUM_FMT, "Tong_SL": NUM_FMT, "Doanh_thu": NUM_FMT, "Target": NUM_FMT, "Pct_Dat": PCT_FMT}, widths=[12, 10, 10, 14, 14, 10])
common_sheets(wb)
dm = wb.create_sheet("Dap_an_M")
dm["A1"] = "ĐÁP ÁN — MÃ M CỦA CÁC QUERY (Home → Advanced Editor). Hai sheet Sales_T4, KetQua_T4 là kết quả Load của các query này"; dm["A1"].font = TITLE
dm["A2"] = ("Học viên làm bằng giao diện nên tên bước và thứ tự có thể khác; chấm theo kết quả (Kiem_tra). Folder.Files: sửa đường dẫn theo máy. "
            "From Folder bằng giao diện tạo thêm nhóm query trợ giúp (Sample File, Transform File…), kết quả như nhau.")
dm["A2"].font = NOTE
r = 4
for nm, code, load in (("Sales_T4", M_SALES, "Table (sheet Sales_T4)"), ("Product", M_PRODUCT, "Only Create Connection"),
                       ("KiemTra_Ma (Left Anti)", M_ANTI, "Only Create Connection"), ("DoiSoat_T4 (Reference của Sales_T4)", M_GROUP, "Only Create Connection"),
                       ("Target_T4", M_TARGET, "Only Create Connection"), ("KetQua_T4", M_KETQUA, "Table (sheet KetQua_T4)")):
    dm.cell(row=r, column=1, value=f"Query {nm} — Load: {load}").font = BOLD; r += 1
    for line in code.splitlines():
        dm.cell(row=r, column=1, value=line).font = MONO; r += 1
    r += 1
set_widths(dm, [160])

kt = wb.create_sheet("Kiem_tra")
kt["A1"] = "KIỂM TRA / ĐỐI SOÁT — Bài tập về nhà B4"; kt["A1"].font = TITLE
kt["A2"] = "Chấm bài: tính lại từng chỉ tiêu trên bảng Load của học viên (Sales_T4, KetQua_T4), so với cột Đáp án."; kt["A2"].font = NOTE
header(kt, 4, ["Chỉ tiêu", "", "Kết quả (công thức)", "Đáp án", "Khớp?"])
st07_row = 2 + list(g.Store_Code).index(SPLIT_STORE)
checks = [
    ("Số dòng Sales_T4 (đã bỏ Tổng cộng, bỏ trùng, đủ file bổ sung)", "=ROWS(Sales_T4[Order_ID])", mh["rows"], NUM_FMT),
    ("Tổng doanh thu tháng 4 (VNĐ)", "=SUM(Sales_T4[Revenue])", mh["revenue"], NUM_FMT),
    ("Tổng Quantity tháng 4", "=SUM(Sales_T4[Quantity])", mh["qty"], NUM_FMT),
    ("Dòng \"Tổng cộng\" còn lại", '=COUNTIF(Sales_T4[Order_ID],"Tổng cộng")', 0, "0"),
    ("Dòng có Product_Code không có trong Product (Left Anti)", "=SUMPRODUCT(--(COUNTIF(Product[Product_Code],Sales_T4[Product_Code])=0))", 0, "0"),
    ("Số cửa hàng trong KetQua_T4", "=ROWS(KetQua_T4[Store_Code])", mh["stores"], "0"),
    ("Tổng Target tháng 4 (VNĐ)", "=SUM(KetQua_T4[Target])", mh["target"], NUM_FMT),
    ("ST07: doanh thu tháng 4 (đủ 2 file)", '=SUMIFS(KetQua_T4[Doanh_thu],KetQua_T4[Store_Code],"ST07")', mh["st07_rev"], NUM_FMT),
    ("ST07: % đạt target", '=SUMIFS(KetQua_T4[Pct_Dat],KetQua_T4[Store_Code],"ST07")', mh["st07_pct"], PCT_FMT),
    ("Cả chuỗi: % đạt target tháng 4", "=SUM(KetQua_T4[Doanh_thu])/SUM(KetQua_T4[Target])", mh["pct_all"], PCT_FMT),
    ("Số cửa hàng đạt dưới 90% target", '=COUNTIF(KetQua_T4[Pct_Dat],"<0.9")', mh["below90_n"], "0"),
]
for rr, (label, f, exp, fmt) in enumerate(checks, start=5):
    kt.cell(row=rr, column=1, value=label).font = BODY
    c = kt.cell(row=rr, column=3, value=f); c.number_format = fmt; c.font = BODY
    c = kt.cell(row=rr, column=4, value=exp); c.number_format = fmt; c.font = INPUT
    kt.cell(row=rr, column=5, value=f'=IF(ROUND(C{rr},6)=ROUND(D{rr},6),"✔","✘")').font = BODY
KT_LAST = 4 + len(checks)
set_widths(kt, [58, 2, 20, 16, 8])

ea = wb.create_sheet("Dap_an_loi")
ea["A1"] = "ĐÁP ÁN — LỖI CÀI SẴN VÀ LỖI HAY GẶP"; ea["A1"].font = TITLE
foot_files = ", ".join(sorted({x[0].split("/")[-1][:4] for x in issues_log if x[2] == "footer"}))
code_where = "; ".join(f"{x[0].split('/')[-1]} dòng {x[1]} (\"{x[3]}\")" for x in issues_log if x[2] == "code")
dup_where = "; ".join(f"{x[0].split('/')[-1]} dòng {x[1]} ({x[3]})" for x in issues_log if x[2] == "dup")
rows = [("#", "Lỗi", "Vị trí / dấu hiệu", "Cách sửa"),
        (1, f"Thiếu dữ liệu: ST07 gửi 2 file, file thứ hai (từ ngày {SPLIT_DAY}/04) nằm ở thư mục Bo_sung",
         f"Group By: ST07 chỉ {mh['st07_part1_rows']} dòng, {vn(mh['st07_part1_rev'])} VNĐ ({pct(mh['st07_part1_pct'])} target) so với hệ thống "
         f"{mh['by_store'][SPLIT_STORE]['rows']} dòng, {vn(mh['st07_rev'])}. Cả chuỗi thiếu {mh['missing_rows']} dòng, {vn(mh['missing_rev'])} VNĐ.",
         "Copy Bo_sung/ST07_2025-04_phan2.csv vào thư mục Cua_hang → Refresh All. Không gõ tay, không sửa query. Tên file vẫn bắt đầu bằng ST07_ nên Store_Code đúng."),
        (2, f"Dòng \"Tổng cộng\" cuối mỗi file ({n_footer} dòng: {mh['files_main']} file chính + 1 file bổ sung)",
         f"Order_ID = \"Tổng cộng\", Order_Date và Unit_Price trống; Quantity là tổng cả file. Chưa bỏ thì số dòng thừa {n_footer} và Tổng Quantity gấp đôi (Revenue trống nên doanh thu không đổi).",
         "Filter Order_ID ≠ \"Tổng cộng\" (hoặc bỏ dòng Order_Date null). Column profile: Order_Date có Empty là dấu hiệu."),
        (3, f"Product_Code sai dạng: {n_lower} mã viết thường, {n_space} mã dính dấu cách cuối", code_where + ". Left Anti ra "
         f"{n_code} dòng trước khi sửa; chỉ Trim thì còn {n_lower} dòng.", "Transform → Format → Trim rồi UPPERCASE trên Product_Code, trước Merge"),
        (4, f"{n_dup} dòng trùng tuyệt đối", dup_where, "Remove Duplicates trên cả bảng (chọn mọi cột), sau khi đã Trim / UPPERCASE"),
        (5, "Ngày dd/mm/yyyy và giá kiểu Việt Nam (\"19.000\")", "Máy locale English (United States): Order_Date ngày ≥ 13 thành Error, ngày ≤ 12 đảo ngày–tháng; "
         "Unit_Price \"19.000\" thành 19 (không báo lỗi). Doanh thu nhỏ đi 1.000 lần.", "Xóa bước Changed Type tự động; Change Type → Using Locale → Vietnamese (Vietnam)"),
        (6, "Không có Store_Code", "Cột Source.Name (tên file) bị xóa khi Combine hoặc chưa tách", "Extract → Text Before Delimiter \"_\" trên Source.Name; đổi tên Store_Code"),
        (7, "Target: cột Thang vẫn là chữ \"Tháng 4\"", "Lọc Thang = 4 không ra dòng nào, hoặc Merge thiếu Target", "Extract → Text After Delimiter \" \" → Whole Number trước khi lọc"),
        (8, "Pct_Dat tính bằng trung bình / gõ tay", "Cả chuỗi = AVERAGE các cửa hàng", "Cả chuỗi = tổng doanh thu ÷ tổng target (tổng ÷ tổng)")]
wrap_rows(ea, 3, rows, [4, 40, 70, 48])

ins = wb.create_sheet("Insight_mau")
ins["A1"] = "INSIGHT MẪU (tham khảo khi chấm; học viên viết theo cách của mình)"; ins["A1"].font = TITLE
ins["A3"] = (f"\"Tháng 4/2025 cả chuỗi đạt {vn(mh['revenue'])} VNĐ, {pct(mh['pct_all'])} target; nếu chỉ nhìn {mh['files_main']} file ban đầu thì ST07 chỉ đạt "
             f"{pct(mh['st07_part1_pct'])}, nhưng đối soát với hệ thống cho thấy thiếu file nửa sau tháng của ST07 — đủ dữ liệu thì ST07 đạt {pct(mh['st07_pct'])}, "
             f"tức đã hồi phục sau đợt cải tạo tháng 3. Cần theo dõi {', '.join(mh['below90'])} (dưới 90% target), thấp nhất là {mh['min_store']} ({pct(mh['min_pct'])}); "
             "và phải đối soát số dòng từng cửa hàng mỗi tháng trước khi báo cáo, vì thiếu một file là kết luận sai về cả một cửa hàng.\"")
ins["A3"].font = BODY; ins["A3"].alignment = Alignment(wrap_text=True, vertical="top"); ins.row_dimensions[3].height = 90
ins["A5"] = ("Chấm: đủ 3 thành phần (con số có so sánh · nguyên nhân khả dĩ · hàm ý), con số khớp dữ liệu; sheet Insight liệt kê được lỗi thiếu file ST07 "
             "(phát hiện nhờ đối soát) và ít nhất 2 lỗi dữ liệu khác (Tổng cộng, mã sản phẩm, dòng trùng, Locale).")
ins["A5"].font = NOTE; ins["A5"].alignment = Alignment(wrap_text=True)
set_widths(ins, [120])
wb.active = wb.sheetnames.index("Kiem_tra")
for w_ in wb.worksheets: w_.sheet_view.tabSelected = w_.title == "Kiem_tra"
wb.save(HW_ANS)

ok = True
if USE_EXCEL:
    from pivot_xml import excel_resave
    excel_resave(HW_DE)
    vals = excel_resave(HW_ANS, [f"Kiem_tra!E{r}" for r in range(5, KT_LAST + 1)])
    ok = all(v == "✔" for v in vals)
    print("Kiem_tra BTVN:", " ".join(vals))
print(f"Meta: {os.path.join(GEN, 'build_b4_meta.json')}")
print(f"BTVN: {mh['files_main']} file + 1 bổ sung, {lines_main} dòng thô ở Cua_hang; {n_footer} dòng Tổng cộng, {n_code} mã sai dạng, {n_dup} dòng trùng; "
      f"đáp án {vn(mh['rows'])} dòng / {vn(mh['revenue'])} VNĐ; ST07 {pct(mh['st07_part1_pct'])} → {pct(mh['st07_pct'])}; dưới 90%: {mh['below90']}")
if not ok: sys.exit("Kiem_tra BTVN có dòng ✘")
