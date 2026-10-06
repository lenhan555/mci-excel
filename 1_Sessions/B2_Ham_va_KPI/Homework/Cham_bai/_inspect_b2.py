import openpyxl, json, zipfile, xml.etree.ElementTree as ET, pathlib, warnings
warnings.filterwarnings('ignore')
root=pathlib.Path(__file__).parent.parent
out={}
for p in list(root.glob('*.xlsx'))+list((root/'Bai_nop').glob('*.xlsx')):
 w=openpyxl.load_workbook(p,data_only=True)
 f=openpyxl.load_workbook(p,data_only=False)
 z=zipfile.ZipFile(p)
 o={'sheets':w.sheetnames,'cells':{},'formulas':{},'drawings':{}}
 for s in w:
  if 'Order_ID' in [c.value for c in s[1]]:
   o['headers']=[c.value for c in s[1]]
   o['rows']=s.max_row
   o['sample']=[[c.coordinate,c.value,f[s.title][c.coordinate].value] for r in [2,100,s.max_row] for c in s[r]]
  elif s.title not in ('Store','Product','Calendar','Target'):
   o['cells'][s.title]=[[c.coordinate,c.value] for row in s for c in row if c.value is not None]
   o['formulas'][s.title]=[[c.coordinate,c.value] for row in f[s.title] for c in row if c.data_type=='f']
 for n in z.namelist():
  if n.startswith('xl/drawings/') and n.endswith('.xml') or n.startswith('xl/drawings/vmlDrawing'):
   try:
    e=ET.fromstring(z.read(n)); ts=[x.text for x in e.iter() if x.tag.split('}')[-1]=='t' and x.text]
    if ts:o['drawings'][n]=ts
   except Exception:pass
 out[p.name]=o
(root/'Cham_bai'/'_inspection.json').write_text(json.dumps(out,ensure_ascii=False,indent=1,default=str),encoding='utf8')
for n,o in out.items():
 print(n,json.dumps(o,ensure_ascii=False,default=str))
