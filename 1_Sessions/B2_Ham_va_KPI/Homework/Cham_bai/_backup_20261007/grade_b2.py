#!/usr/bin/env python3
"""Chấm bài tập về nhà B2 (Hàm và KPI). Chỉ ĐỌC file học viên.
Dùng:
  python3 grade_b2.py <file_nop.xlsx> <B2_BTVN_..._Answer.xlsx>                      -> JSON một bài
  python3 grade_b2.py --folder <thu_muc_Bai_nop> <..._Answer.xlsx> [--cham-tay f.json] -> JSON cả lớp
f.json (chấm tay việc 5 và 6): {"ten_file.xlsx": {"v5": "Đạt|Một phần|Chưa đạt", "insight": 0-3}}
Thang điểm TẠM (cần GV xác nhận): trọng số nguyên: việc 1-2 mỗi việc 2 điểm, việc 3-5 mỗi việc 1 điểm, việc 6 (insight) 3 điểm = 1 điểm mỗi thành phần (tổng 10); Đạt 100%, Một phần 50%, Chưa đạt 0;
Sửa TRONG_SO và diem_viec() nếu GV đổi thang.
"""
import sys, json, re, zipfile, collections, glob, os, difflib
from decimal import Decimal, ROUND_HALF_UP
import openpyxl
from openpyxl.formula.translate import Translator

def val(c):
    v = c.value
    return v.text if hasattr(v, "text") else v

def is_formula(v):
    return isinstance(v, str) and v.startswith("=")

def load(path):
    return openpyxl.load_workbook(path), openpyxl.load_workbook(path, data_only=True)

def find_sales(wb):
    for ws in wb:
        h = [c.value for c in ws[1]]
        if "Order_ID" in h and "Line_No" in h:
            return ws, {n: i for i, n in enumerate(h) if n}
    return None, {}

def ref_rows(ws_v, idx):
    d = {}
    for r in ws_v.iter_rows(min_row=2, values_only=True):
        if r[idx["Order_ID"]] is None:
            continue
        d[(r[idx["Order_ID"]], r[idx["Line_No"]])] = r
    return d

def find_kpi(wbv, wbf):
    """Tìm hai bảng KPI theo nhãn, không theo vị trí."""
    out = {}
    for name in wbv.sheetnames:
        if name.strip().lower() != "kpi":
            continue
        wv, wf = wbv[name], wbf[name]
        out["sheet"] = name
        cells = {c.coordinate: c for row in wv.iter_rows() for c in row if c.value is not None}
        # KPI1: hàng có nhãn Bắc/Trung/Nam, cột nhãn ngành, cột Tổng
        regrows = {}
        for c in cells.values():
            if isinstance(c.value, str) and c.value.strip() in ("Bắc", "Trung", "Nam") and c.column == 1:
                regrows.setdefault(c.value.strip(), []).append(c.row)
        out["regrows"] = regrows
        hdr = {}
        for c in cells.values():
            if isinstance(c.value, str):
                hdr.setdefault(c.value.strip(), []).append(c)
        out["hdr"] = {k: [x.coordinate for x in v] for k, v in hdr.items()}
        out["ws_v"], out["ws_f"] = wv, wf
        return out
    return out

def grade(nop, ans):
    nv, nf = None, None
    res = {"file": nop, "kiem_tra": {}, "viec": {}, "ghi_chu": []}
    try:
        nf, nv = load(nop)
        af, av = load(ans)
    except Exception as e:
        res["loi_mo_file"] = str(e)
        return res
    ws_n_f, idx = find_sales(nf)
    ws_n_v, _ = find_sales(nv)
    ws_a_v, aidx = find_sales(av)
    if ws_n_f is None:
        res["viec"]["1"] = ["Chưa đạt", "Không tìm thấy bảng dữ liệu có Order_ID/Line_No"]
        return res
    tbl = list(ws_n_f.tables.values())
    res["bang"] = [(t.name, t.ref) for t in tbl]
    n_rows = ref_rows(ws_n_v, idx)
    a_rows = ref_rows(ws_a_v, aidx)
    res["so_dong"] = len(n_rows)

    def col(name):
        return idx.get(name)

    # ---------- Việc 1: tra cứu ----------
    ev = []
    need = ["Region", "Category", "Unit_Cost"]
    miss = [c for c in need if c not in idx]
    if miss:
        v1 = ("Chưa đạt", f"Thiếu cột {miss}")
    else:
        n_na = 0
        for r in n_rows.values():
            for c in need:
                x = r[col(c)]
                if isinstance(x, str) and x.startswith("#") or x is None:
                    n_na += 1
        f2 = {c: val(ws_n_f.cell(2, col(c) + 1)) for c in need}
        funcs = {c: re.findall(r"_xlfn\.XLOOKUP|XLOOKUP|INDEX|VLOOKUP", str(f)) for c, f in f2.items()}
        wrap = [c for c, f in f2.items() if re.search(r"IFERROR|IFNA", str(f))]
        hard = [c for c, f in f2.items() if not is_formula(f)]
        trailing = sum(1 for r in n_rows.values() if isinstance(r[col("Product_Code")], str) and r[col("Product_Code")] != r[col("Product_Code")].strip())
        res["tra_cuu"] = {"o_loi_hoac_trong": n_na, "ham": funcs, "bocIFERROR": wrap, "go_tay": hard, "ma_con_dau_cach": trailing}
        ok = n_na == 0 and not wrap and not hard and trailing == 0
        v1 = ("Đạt" if ok else "Chưa đạt", f"#N/A/trống={n_na}; IFERROR={wrap}; gõ tay={hard}; mã còn dấu cách={trailing}; hàm={funcs}")
    res["viec"]["1"] = list(v1)

    # ---------- Việc 2: cột tính toán ----------
    comp = {}
    for c in ["COGS", "Gross_Margin", "Phan_khuc"]:
        if c not in idx:
            comp[c] = "thiếu cột"
            continue
        f = val(ws_n_f.cell(2, col(c) + 1))
        comp[c] = "công thức" if is_formula(f) else "gõ tay"
    diff_exact = diff_ci = 0
    sample = collections.Counter()
    if "Phan_khuc" in idx:
        for k, r in n_rows.items():
            a = a_rows.get(k)
            if a is None:
                continue
            x, y = r[col("Phan_khuc")], a[aidx["Phan_khuc"]]
            if x != y:
                diff_exact += 1
                sample[(x, y)] += 1
                if str(x).strip().split(" ")[0].casefold() != str(y).strip().split(" ")[0].casefold():
                    diff_ci += 1
    cg = {}
    for c in ["COGS", "Gross_Margin"]:
        if c in idx:
            cg[c] = sum(r[col(c)] for r in n_rows.values() if isinstance(r[col(c)], (int, float)))
    res["cot_tinh"] = {"loai": comp, "tong": cg, "phan_khuc_lech_ky_tu": diff_exact, "phan_khuc_lech_y_nghia": diff_ci, "mau_lech": [[list(k), v] for k, v in sample.items()]}
    all_f = all(v == "công thức" for v in comp.values())
    if not all_f:
        v2 = ("Chưa đạt", f"loại={comp}")
    elif diff_ci:
        v2 = ("Chưa đạt", f"Phan_khuc sai ngưỡng {diff_ci} dòng")
    elif diff_exact:
        v2 = ("Một phần", f"Phan_khuc đúng ngưỡng nhưng khác nhãn đề ('Trung'/'Thấp') ở {diff_exact} dòng: {dict((str(k), v) for k, v in sample.items())}")
    else:
        v2 = ("Đạt", "COGS, Gross_Margin, Phan_khuc là công thức, khớp đáp án")
    res["viec"]["2"] = list(v2)

    # ---------- KPI ----------
    k = find_kpi(nv, nf)
    kv = {}
    if "sheet" not in k:
        res["viec"]["3"] = ["Chưa đạt", "Không có sheet KPI"]
        res["viec"]["4"] = ["Chưa đạt", "Không có sheet KPI"]
    else:
        wv, wf = k["ws_v"], k["ws_f"]
        hd = k["hdr"]
        allreg = k["regrows"]
        def reg_after(h):
            return {r: min(x for x in v if x > h) for r, v in allreg.items() if any(x > h for x in v)}
        tot_all = sorted(openpyxl.utils.cell.coordinate_from_string(c)[1] for c in hd.get("Tổng", []) if openpyxl.utils.cell.coordinate_from_string(c)[0] == "A")
        def tot_after(h):
            c = [x for x in tot_all if x > h]
            return c[0] if c else None
        regrows = reg_after(0)
        # hàng/cột Tổng
        tot_cells = [openpyxl.utils.cell.coordinate_from_string(c) for c in hd.get("Tổng", [])]
        tot_row = [rw for (cl, rw) in tot_cells if openpyxl.utils.column_index_from_string(cl) == 1]
        tot_col = [openpyxl.utils.column_index_from_string(cl) for (cl, rw) in tot_cells if openpyxl.utils.column_index_from_string(cl) > 1]
        tot_row = None
        # KPI1: tìm hàng tiêu đề có nhãn ngành
        cat_cols = {}
        hdr_row1 = None
        for c in wv[1:wv.max_row]:
            pass
        cats = ["Bánh kẹo", "Chăm sóc cá nhân", "Đồ uống", "Gia dụng", "Sữa", "Thực phẩm khô"]
        for cat in cats:
            for coord in hd.get(cat, []):
                cl, rw = openpyxl.utils.cell.coordinate_from_string(coord)
                cat_cols[cat] = openpyxl.utils.column_index_from_string(cl)
                hdr_row1 = rw
        tc = None
        for coord in hd.get("Tổng", []):
            cl, rw = openpyxl.utils.cell.coordinate_from_string(coord)
            if rw == hdr_row1:
                tc = openpyxl.utils.column_index_from_string(cl)
        blk = []
        tot_row = tot_after(hdr_row1 or 0)
        regrows = reg_after(hdr_row1 or 0)
        if regrows and cat_cols and tot_row:
            r_list = [regrows[r] for r in ("Bắc", "Trung", "Nam") if r in regrows]
            c_list = sorted(cat_cols.values())
            formulas = {}
            for r in r_list:
                for c in c_list:
                    formulas[(r, c)] = val(wf.cell(r, c))
            kv["khoi_3x6_cong_thuc"] = all(is_formula(v) and "SUMIFS" in v.upper() for v in formulas.values())
            anchor = (r_list[0], c_list[0])
            base = formulas[anchor]
            same = True
            if is_formula(base):
                for (r, c), v in formulas.items():
                    if not is_formula(v):
                        same = False
                        continue
                    src = wf.cell(*anchor).coordinate
                    tgt = wf.cell(r, c).coordinate
                    exp = Translator(base, origin=src).translate_formula(tgt)
                    if exp.replace(" ", "") != str(v).replace(" ", ""):
                        same = False
            else:
                same = False
            kv["khoi_giong_nhau_khi_keo"] = same
            kv["cong_thuc_o_dau"] = base
            kv["co_dollar_hon_hop"] = bool(re.search(r"\$[A-Z]+\d+", str(base))) and bool(re.search(r"(?<!\$)[A-Z]+\$\d+", str(base)))
            kv["gia_tri_3x6"] = {f"{wf.cell(r,1).value}|{c}": wv.cell(r, c).value for r in r_list for c in c_list}
            if tc:
                kv["tong_cot_cong_thuc"] = all(is_formula(val(wf.cell(r, tc))) for r in r_list)
                kv["tong_kpi1"] = wv.cell(tot_row, tc).value
                kv["tong_kpi1_cong_thuc"] = is_formula(val(wf.cell(tot_row, tc)))
            kv["dong_tong_cong_thuc"] = all(is_formula(val(wf.cell(tot_row, c))) for c in c_list)
        # KPI2
        def find_hdr(pred):
            out = []
            for name, coords in hd.items():
                if pred(name):
                    for coord in coords:
                        cl, rw = openpyxl.utils.cell.coordinate_from_string(coord)
                        out.append((openpyxl.utils.column_index_from_string(cl), rw, name))
            return out
        so_don = find_hdr(lambda n: n.lower().startswith("số đơn"))
        aov = find_hdr(lambda n: n.upper() == "AOV")
        gm = find_hdr(lambda n: "%" in n and "margin" in n.lower())
        def rowvals(h):
            if not h:
                return None
            c, hr = h[0][0], h[0][1]
            regrows = reg_after(hr); tot_row = tot_after(hr)
            return {
                "cong_thuc": {r: val(wf.cell(rw, c)) for r, rw in regrows.items()},
                "gia_tri": {r: wv.cell(rw, c).value for r, rw in regrows.items()},
                "tong_cong_thuc": val(wf.cell(tot_row, c)) if tot_row else None,
                "tong_gia_tri": wv.cell(tot_row, c).value if tot_row else None,
                "tong_o": wf.cell(tot_row, c).coordinate if tot_row else None,
            }
        kv["so_don"], kv["aov"], kv["gm_pct"] = rowvals(so_don), rowvals(aov), rowvals(gm)
        res["kpi"] = kv

        # --- Việc 3 ---
        if "khoi_3x6_cong_thuc" not in kv:
            res["viec"]["3"] = ["Chưa đạt", "Không tìm thấy khối 3×6 (Bắc/Trung/Nam × 6 ngành) hoặc dòng Tổng"]
        else:
            probs = []
            if not kv["khoi_3x6_cong_thuc"]: probs.append("có ô gõ tay hoặc không phải SUMIFS")
            if not kv["khoi_giong_nhau_khi_keo"]: probs.append("công thức các ô không đồng nhất khi quy về một ô (không phải một công thức kéo ra)")
            if not kv["co_dollar_hon_hop"]: probs.append("không có $ hỗn hợp")
            if not kv.get("tong_cot_cong_thuc", False) or not kv["dong_tong_cong_thuc"]: probs.append("dòng/cột Tổng không là công thức")
            tt = kv.get("tong_kpi1")
            if tt != av_get(av, "Kiem_tra", "D8"):
                probs.append(f"tổng KPI1 = {tt}")
            res["viec"]["3"] = ["Đạt" if not probs else "Một phần", "; ".join(probs) or f"Một công thức SUMIFS kéo ra: {kv['cong_thuc_o_dau']}; tổng = {tt}"]
        # --- Việc 4 ---
        probs = []
        if not kv.get("so_don"): probs.append("không có cột Số đơn")
        if not kv.get("aov"): probs.append("không có cột AOV")
        if not kv.get("gm_pct"): probs.append("không có cột Gross Margin %")
        if kv.get("so_don"):
            fs = list(kv["so_don"]["cong_thuc"].values())
            if not all(is_formula(f) and "COUNTIFS" in str(f).upper() for f in fs): probs.append("Số đơn không phải COUNTIFS")
            if kv["so_don"]["tong_gia_tri"] is None: probs.append("dòng Tổng thiếu Số đơn")
        if kv.get("aov"):
            if kv["aov"]["tong_gia_tri"] is None: probs.append("dòng Tổng thiếu AOV")
            elif "AVERAGE" in str(kv["aov"]["tong_cong_thuc"]).upper(): probs.append("AOV tổng dùng AVERAGE")
        if kv.get("gm_pct"):
            if kv["gm_pct"]["tong_gia_tri"] is None: probs.append("dòng Tổng thiếu GM%")
            elif "AVERAGE" in str(kv["gm_pct"]["tong_cong_thuc"]).upper(): probs.append("GM% tổng dùng AVERAGE")
        # Quy ước: thiếu cả cột là Chưa đạt; chỉ thiếu ô ở dòng Tổng (từng miền vẫn đúng) là Một phần.
        res["viec"]["4"] = ["Đạt" if not probs else ("Chưa đạt" if any(p.startswith("không có") for p in probs) else "Một phần"), "; ".join(probs) or "đủ"]

    # ---------- Kiem_tra ----------
    ktws = av["Kiem_tra"]
    exp = {r: ktws.cell(r, 4).value for r in range(5, 13)}
    names = {r: ktws.cell(r, 1).value for r in range(5, 13)}
    sums = lambda c: sum(r[col(c)] for r in n_rows.values() if isinstance(r[col(c)], (int, float))) if c in idx else None
    got = {
        5: len(n_rows),
        6: res.get("tra_cuu", {}).get("o_loi_hoac_trong"),
        7: sums("Revenue"),
        8: kv.get("tong_kpi1"),
        9: (kv.get("so_don") or {}).get("tong_gia_tri"),
        10: sums("COGS"),
        11: (kv.get("gm_pct") or {}).get("tong_gia_tri"),
        12: (kv.get("aov") or {}).get("tong_gia_tri"),
    }
    for r in range(5, 13):
        g, e = got[r], exp[r]
        ok = g is not None and round(float(g), 6) == round(float(e), 6)
        res["kiem_tra"][names[r]] = {"hoc_vien": g, "dap_an": e, "khop": ok}

    # ---------- Việc 5: KPI_Card ----------
    card = None
    for n in nv.sheetnames:
        if n.strip().lower() == "kpi_card":
            card = nv[n]
    if card is None:
        res["viec"]["5"] = ["Chưa đạt", "Không có sheet KPI_Card"]
    else:
        fields = {}
        for row in card.iter_rows(min_row=4, values_only=True):
            if row[0]:
                fields[str(row[0]).strip()] = row[1]
        res["card"] = fields
        filled = sum(1 for v in fields.values() if v and str(v).strip())
        res["viec"]["5"] = ["Cần GV/agent đọc", f"{filled}/{len(fields)} trường có nội dung"]
    # ---------- Việc 6: Insight ----------
    ins = None
    for n in nv.sheetnames:
        if n.strip().lower() == "insight":
            ins = nv[n]
    if ins is None:
        res["viec"]["6"] = ["Chưa đạt", "Không có sheet Insight"]
    else:
        txt = " ".join(str(c.value) for row in ins.iter_rows() for c in row if c.value)
        res["insight"] = txt
        res["viec"]["6"] = ["Cần GV/agent đọc", f"{len(txt)} ký tự"]
    # ---------- dấu hiệu bất thường ----------
    sheets = set(nf.sheetnames)
    res["sheet_dap_an"] = sorted(sheets & {"Dap_an_loi", "Insight_mau"})
    try:
        core = zipfile.ZipFile(nop).read("docProps/core.xml").decode("utf8")
        res["core"] = {k: (re.search(f"<{k}>(.*?)</{k}>", core) or [None, None])[1] for k in ("dc:creator", "cp:lastModifiedBy", "dcterms:created", "dcterms:modified")}
    except Exception:
        pass
    return res

def av_get(av, sheet, coord):
    return av[sheet][coord].value

TRONG_SO = {"1": 2, "2": 2, "3": 1, "4": 1, "5": 1}  # insight (việc 6) = 3, tổng 10 điểm (thang tạm; đề và Kiem_tra chưa quy định điểm)


def diem_viec(kq, k):
    """Đạt = 100%, Một phần = 50%, Chưa đạt = 0; Không kiểm được / Cần chấm tay chưa có kết quả = 0 (tạm)."""
    return TRONG_SO[k] * {"Đạt": 1.0, "Một phần": 0.5}.get(kq, 0.0)


def r2(x):
    return float(Decimal(repr(round(x, 6))).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def ap_cham_tay(res, v5, tp):
    """v5: kết quả việc 5 (KPI_Card) do người chấm đọc; tp: số thành phần insight đạt 0-3 (việc 6)."""
    if "loi_mo_file" in res:
        return res
    res["viec"]["5"][0] = v5
    d = {k: diem_viec(res["viec"][k][0], k) for k in ("1", "2", "3", "4", "5")}
    d["6"] = float(tp)
    res["diem_viec"] = d
    res["insight_thanh_phan_dat"] = tp
    res["tong_diem"] = r2(sum(d.values()))
    return res


def similarity(results):
    items = [(r["file"], r.get("insight", "")) for r in results if "loi_mo_file" not in r]
    out = []
    for i in range(len(items)):
        for j in range(i + 1, len(items)):
            a, b = items[i][1], items[j][1]
            if a and b:
                ratio = difflib.SequenceMatcher(None, a.lower(), b.lower()).ratio()
                if ratio >= 0.6:
                    out.append({"a": items[i][0], "b": items[j][0], "ti_le_giong": round(ratio, 2)})
    return out


def main(argv):
    sys.stdout.reconfigure(encoding="utf-8")
    if len(argv) >= 4 and argv[1] == "--folder":
        files = sorted(f for f in glob.glob(os.path.join(argv[2], "**", "*.xlsx"), recursive=True) if not os.path.basename(f).startswith("~$"))
        results = [grade(f, argv[3]) for f in files]
        tay = {}
        if len(argv) >= 6 and argv[4] == "--cham-tay":
            tay = json.load(open(argv[5], encoding="utf8"))
        for r in results:
            t = tay.get(os.path.basename(r["file"]))
            if t:
                ap_cham_tay(r, t["v5"], t["insight"])
        for r in results:
            r["file"] = os.path.basename(r["file"])
        out = {"bai": results, "giong_nhau_insight": similarity(results)}
    elif len(argv) >= 3:
        out = grade(argv[1], argv[2])
    else:
        print(__doc__)
        return 2
    print(json.dumps(out, ensure_ascii=False, indent=1, default=str))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
